from __future__ import annotations

import concurrent.futures
import hashlib
import html
import json
import logging
import os
import re
import secrets
import time
from dataclasses import asdict, replace
from pathlib import Path
from typing import Any, Callable, Literal, Mapping, Union
from urllib.parse import parse_qsl, quote, urlencode, urlsplit, urlunsplit

from fastapi import APIRouter, Depends, File, Form, Header, HTTPException, Request, Response, UploadFile
from fastapi.responses import FileResponse, HTMLResponse, JSONResponse
from pydantic import BaseModel, Field, TypeAdapter, ValidationError

from backend.auth0_dependencies import (
    AuthenticatedUser,
    get_current_user,
    get_current_user_optional,
)
from backend.errors import to_http_exception
from backend.database import get_db
from backend.email_client import ConsoleEmailClient, build_default_email_client
from backend.esignature_persistence import (
    PostgresEnvelopeRepository,
    PostgresSigningTokenRepository,
)
from backend.rate_limiter.dependencies import rate_limit_for_feature
from backend.subscriptions import get_user_entitlement
from backend.upload_retention import (
    mark_upload_paths_processed,
    upload_processing_session,
)
from backend.upload import (
    UploadError,
    UploadServiceUnavailableError,
    build_uploaded_document_payload,
    build_uploaded_media_payload,
    save_pdf_edit_asset_upload,
    save_pdf_tool_upload,
)
from backend.batch_processing import (
    BATCH_UPLOAD_LIMITS_BY_PLAN,
    BatchUploadPolicy,
    BatchUploadPolicyError,
    find_duplicate_upload_content,
    require_batch_upload_entitlement,
)

from backend.src.extraction import (
    build_inline_text_payload,
    build_pdf_input_artifact_for_action,
)
from backend.src.inline_text_security import (
    AUXILIARY_PROMPT_POLICY,
    INLINE_TEXT_POLICY,
    MAX_NUMBERED_QUESTIONS,
    validate_auxiliary_prompt_text,
    validate_question_item,
)
from backend.src.processing.conversion.convert import convert_document
from backend.src.processing.data_protection.data_masking.data_mask import preview_data_mask_candidates
from backend.src.processing.data_protection.document_type_detection import detect_privacy_document_type
from backend.src.processing.data_protection.orchestration import ProtectedArtifactResult
from backend.src.processing.data_protection.redaction.redact import preview_redaction_candidates
from backend.src.processing.compliance.registry import RuleRegistryError

from backend.src.schema import (
    AddSignatureOperation,
    AnalyzerRequest,
    AnalyzerResponse,
    AnswerGenerationRequest,
    CombinePdfRequest,
    ComplianceJurisdiction,
    ComplianceRegulatoryDomain,
    ComplianceReportVariant,
    ComplianceRequest,
    ComplianceSectorPack,
    MAX_COMPLIANCE_DOCUMENT_SET_FILES,
    CompressPdfRequest,
    ConversionOutputFormat,
    ConversionRequest,
    DataMaskingRequest,
    DocumentSetPayload,
    ESignatureAction,
    ESignatureRequest,
    EditPdfRequest,
    ExplanationRequest,
    FeatureType,
    GrammarCorrectionRequest,
    MediaType,
    OutputPolicy,
    PdfCompressionLevel,
    PdfEditOperation,
    PdfJobResult,
    PdfPageRange,
    PdfSplitMode,
    QuestionGenerationRequest,
    RedactionMaskingDocumentType,
    RedactionRequest,
    SensitiveDataType,
    SignatureRepresentationType,
    SplitPdfRequest,
    StructuredDataOutputFormat,
    StructuredExtractionDocumentClass,
    StructuredExtractionRequest,
    StructuredExtractionResultShape,
    SummarizationRequest,
    SystemLanguage,
    TranscriptionRequest,
    TranslationRequest,
)
from backend.src.workflow_router import WorkflowRouter
from backend.src.esignature_service import ESignatureService, ESignatureServiceConfig
from backend.src.processing.esignature.layout import analyze_esignature_pdf
from backend.src.storage.artifacts import (
    LocalArtifactStorage,
    artifact_owner_context,
    get_artifact_owner,
    guess_content_type,
    is_artifact_expired,
)


API_V1_ANALYZER_PREFIX = "/analyzer"

router = APIRouter(prefix=API_V1_ANALYZER_PREFIX, tags=["analyzer-v1"])
logger = logging.getLogger(__name__)


# -----------------------------------------------------------------------------
# Policy / error helpers
# -----------------------------------------------------------------------------

TRANSFORMED_ACTIONS = {
    FeatureType.convert,
    FeatureType.summarize,
    FeatureType.grammar_correct,
    FeatureType.translate,
    FeatureType.transcribe,
    FeatureType.redact,
    FeatureType.data_mask,
    FeatureType.combine_pdf,
    FeatureType.split_pdf,
    FeatureType.edit_pdf,
    FeatureType.compress_pdf,
    FeatureType.e_signature,
}

GENERATED_ACTIONS = {
    FeatureType.explain,
    FeatureType.generate_questions,
    FeatureType.generate_answers,
    FeatureType.structured_extract,
    FeatureType.compliance,
}

PDF_UPLOAD_DIR = Path(os.getenv("PDF_UPLOAD_DIR", "uploads/pdf_tools"))
DEFAULT_GOOGLE_SDP_LOCATION = os.getenv("GOOGLE_SDP_LOCATION", "global")
MAX_STRUCTURED_EXTRACTION_DOCUMENT_SET_FILES = 20
MIN_PDF_COMBINE_FILES = 2
MAX_PDF_COMBINE_FILES = 25


def _policy_for_action(action: FeatureType) -> OutputPolicy:
    if action in TRANSFORMED_ACTIONS:
        return OutputPolicy(structure_preservation=True)
    if action in GENERATED_ACTIONS:
        return OutputPolicy(structure_preservation=False)
    raise ValueError(f"Unsupported action: {action}")


def _bad_request(message: str) -> HTTPException:
    return HTTPException(
        status_code=400,
        detail={
            "error": "invalid_request",
            "message": message,
        },
    )


def _service_unavailable(message: str) -> HTTPException:
    return HTTPException(
        status_code=503,
        detail={
            "error": "service_unavailable",
            "message": message,
        },
    )


def _google_sdp_project_id() -> str:
    project_id = os.getenv("GOOGLE_SDP_PROJECT_ID", "").strip()
    if not project_id:
        raise _service_unavailable(
            "Google Sensitive Data Protection is not configured. "
            "Set GOOGLE_SDP_PROJECT_ID for redaction and data masking."
        )
    return project_id


def _download_url_for_storage_key(storage_key: str | None) -> str | None:
    if not isinstance(storage_key, str) or not storage_key.strip():
        return None

    key = storage_key.strip().replace("\\", "/")
    key = key.removeprefix("/api/analyzer/artifacts/")
    key = key.removeprefix("/api/v1/analyzer/artifacts/")
    key = key.removeprefix("/artifacts/")
    key = key.removeprefix("artifacts/")
    return f"/api/v1/analyzer/artifacts/{key}"


DOWNLOAD_NAME_QUERY_PARAM = "download_name"
ANONYMOUS_ARTIFACT_COOKIE_NAME = os.getenv(
    "ANONYMOUS_ARTIFACT_COOKIE_NAME",
    "redocx_anonymous_artifact",
)
ANONYMOUS_ARTIFACT_COOKIE_MAX_AGE_SECONDS = int(
    os.getenv("ANONYMOUS_ARTIFACT_COOKIE_MAX_AGE_SECONDS", str(24 * 60 * 60))
)


def _is_production_environment() -> bool:
    environment = (
        os.getenv("APP_ENV", "").strip()
        or os.getenv("ENVIRONMENT", "").strip()
        or os.getenv("RAILWAY_ENVIRONMENT_NAME", "").strip()
    ).lower()
    return environment in {"production", "prod"}


def _valid_anonymous_artifact_token(value: str | None) -> str | None:
    token = str(value or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{32,128}", token):
        return None
    return token


def _anonymous_artifact_owner_id(
    request: Request,
    response: Response | None = None,
) -> str:
    token = _valid_anonymous_artifact_token(
        request.cookies.get(ANONYMOUS_ARTIFACT_COOKIE_NAME)
    )
    if token is None:
        if response is None:
            raise HTTPException(status_code=404, detail="Artifact not found.")
        token = secrets.token_urlsafe(32)
        response.set_cookie(
            key=ANONYMOUS_ARTIFACT_COOKIE_NAME,
            value=token,
            max_age=ANONYMOUS_ARTIFACT_COOKIE_MAX_AGE_SECONDS,
            httponly=True,
            secure=_is_production_environment(),
            samesite="lax",
            path="/",
        )
    digest = hashlib.sha256(token.encode("utf-8")).hexdigest()
    return f"anonymous:{digest}"


def _safe_download_filename(value: str | None, *, default: str = "artifact") -> str:
    raw = Path(str(value or "").replace("\\", "/")).name
    cleaned = re.sub(r"[\x00-\x1f\x7f]+", "", raw).replace('"', "").strip()
    return cleaned or default


def _uploaded_filename(upload: UploadFile | None, *, default: str = "document") -> str:
    if upload is None:
        return default
    return _safe_download_filename(upload.filename, default=default)


def _filename_stem(filename: str, *, default: str = "document") -> str:
    stem = Path(_safe_download_filename(filename, default=default)).stem.strip()
    return stem or default


def _normalized_extension(value: Any, *, default: str = "") -> str:
    raw = getattr(value, "value", value)
    text = str(raw or "").strip().lower()
    if not text:
        return default
    return text if text.startswith(".") else f".{text}"


def _conversion_artifact_extension(response: AnalyzerResponse, requested: Any) -> str:
    """Return the physical artifact extension for a conversion response."""
    filename = getattr(getattr(response, "result", None), "filename", None)
    if isinstance(filename, str) and filename.strip():
        suffix = Path(filename).suffix.strip().lower()
        if suffix:
            return suffix
    requested_extension = _normalized_extension(requested)
    return ".pdf" if requested_extension == ".pdfa" else requested_extension


def _source_output_filename(source_filename: str, *, extension: str | None = None) -> str:
    source = _safe_download_filename(source_filename, default="document")
    if extension is None:
        return source

    normalized_extension = _normalized_extension(extension)
    if not normalized_extension:
        return source

    source_suffix = Path(source).suffix
    if source_suffix.lower() == normalized_extension.lower():
        return source
    return f"{_filename_stem(source)}{normalized_extension}"


def _feature_output_filename(
    source_filename: str,
    feature_suffix: str,
    *,
    extension: str = ".pdf",
    separator: str = "_",
) -> str:
    safe_suffix = re.sub(r"[^A-Za-z0-9_-]+", "_", feature_suffix).strip("_")
    safe_separator = "." if separator == "." else "_"
    normalized_extension = _normalized_extension(extension, default=".pdf")
    return f"{_filename_stem(source_filename)}{safe_separator}{safe_suffix}{normalized_extension}"


def _source_extension(source_filename: str, *, default: str = ".txt") -> str:
    return _normalized_extension(Path(_safe_download_filename(source_filename)).suffix, default=default)


FEATURE_FILENAME_SUFFIXES: dict[FeatureType, str] = {
    FeatureType.summarize: "summarized",
    FeatureType.explain: "explained",
    FeatureType.translate: "translated",
    FeatureType.generate_questions: "generated_questions",
    FeatureType.generate_answers: "generated_answers",
    FeatureType.grammar_correct: "grammar_corrected",
    FeatureType.compliance: "compliance_report",
    FeatureType.structured_extract: "structured_extraction",
    FeatureType.redact: "redacted",
    FeatureType.data_mask: "masked",
}

PDF_FEATURE_FILENAME_SUFFIXES: dict[FeatureType, str] = {
    FeatureType.combine_pdf: "combined",
    FeatureType.compress_pdf: "compressed",
    FeatureType.edit_pdf: "edited",
    FeatureType.split_pdf: "split",
}


def _download_filename_for_action(
    action: FeatureType,
    source_filename: str,
    *,
    output_extension: str | None = None,
) -> str:
    if action == FeatureType.convert:
        return _source_output_filename(
            source_filename,
            extension=output_extension or _source_extension(source_filename),
        )
    if action == FeatureType.transcribe:
        return _source_output_filename(
            source_filename,
            extension=output_extension or ".pdf",
        )
    if action in PDF_FEATURE_FILENAME_SUFFIXES:
        return _feature_output_filename(
            source_filename,
            PDF_FEATURE_FILENAME_SUFFIXES[action],
            extension=output_extension or ".pdf",
            separator=".",
        )

    feature_suffix = FEATURE_FILENAME_SUFFIXES.get(action)
    if feature_suffix:
        return _feature_output_filename(
            source_filename,
            feature_suffix,
            extension=output_extension or _source_extension(source_filename),
        )

    return _source_output_filename(source_filename, extension=output_extension)


def _with_download_filename(url: str | None, filename: str | None) -> str | None:
    if not isinstance(url, str) or not url.strip():
        return url

    parts = urlsplit(url)
    normalized_path = parts.path if parts.path.startswith("/") else f"/{parts.path}"
    artifact_prefixes = (
        "/api/analyzer/artifacts/",
        "/api/v1/analyzer/artifacts/",
        "/artifacts/",
    )
    if not normalized_path.startswith(artifact_prefixes):
        return url

    safe_filename = _safe_download_filename(filename)
    query = dict(parse_qsl(parts.query, keep_blank_values=True))
    query[DOWNLOAD_NAME_QUERY_PARAM] = safe_filename
    return urlunsplit(
        (
            parts.scheme,
            parts.netloc,
            parts.path,
            urlencode(query),
            parts.fragment,
        )
    )



def _object_field_names(value: Any) -> list[str]:
    model_fields = getattr(type(value), "model_fields", None)
    if isinstance(model_fields, dict):
        return list(model_fields)
    instance_fields = getattr(value, "__dict__", None)
    if isinstance(instance_fields, dict):
        return list(instance_fields)
    return []


def _apply_download_filename(
    value: Any,
    filename: str,
    *,
    _seen: set[int] | None = None,
) -> Any:
    """Attach an authoritative download name to every artifact URL in a result.

    Only URL-bearing artifact nodes are changed. Input/source metadata remains
    untouched so response semantics do not change beyond download behavior.
    """
    if value is None or isinstance(value, (str, bytes, int, float, bool)):
        return value

    seen = _seen if _seen is not None else set()
    value_id = id(value)
    if value_id in seen:
        return value
    seen.add(value_id)

    if isinstance(value, (list, tuple)):
        for item in value:
            _apply_download_filename(item, filename, _seen=seen)
        return value

    if isinstance(value, dict):
        storage_key = value.get("storage_key") or value.get("storageKey")
        download_key = (
            "download_url"
            if "download_url" in value
            else "downloadUrl"
            if "downloadUrl" in value
            else None
        )
        download_url = value.get(download_key) if download_key else None
        if not download_url and storage_key:
            download_url = _download_url_for_storage_key(str(storage_key))
        if download_url:
            value[download_key or "download_url"] = _with_download_filename(
                str(download_url),
                filename,
            )
            for filename_key in (
                "file_name",
                "original_artifact_name",
                "artifact_name",
                "output_filename",
            ):
                if filename_key in value:
                    value[filename_key] = filename
            value["filename"] = filename

        for nested in value.values():
            _apply_download_filename(nested, filename, _seen=seen)
        return value

    storage_key = getattr(value, "storage_key", None) or getattr(
        value,
        "storageKey",
        None,
    )
    download_attr = (
        "download_url"
        if hasattr(value, "download_url")
        else "downloadUrl"
        if hasattr(value, "downloadUrl")
        else None
    )
    download_url = getattr(value, download_attr, None) if download_attr else None
    if not download_url and storage_key:
        download_url = _download_url_for_storage_key(str(storage_key))
    if download_url and download_attr:
        try:
            setattr(
                value,
                download_attr,
                _with_download_filename(str(download_url), filename),
            )
        except (AttributeError, TypeError, ValueError):
            pass
        for filename_attr in (
            "filename",
            "file_name",
            "original_artifact_name",
            "artifact_name",
            "output_filename",
        ):
            if not hasattr(value, filename_attr):
                continue
            try:
                setattr(value, filename_attr, filename)
            except (AttributeError, TypeError, ValueError):
                continue

    for field_name in _object_field_names(value):
        try:
            nested = getattr(value, field_name)
        except (AttributeError, TypeError, ValueError):
            continue
        _apply_download_filename(nested, filename, _seen=seen)
    return value


def _compliance_download_extension(
    report_variant: ComplianceReportVariant,
    *,
    file_count: int,
) -> str:
    if report_variant == ComplianceReportVariant.machine_readable_report:
        return ".json"
    if report_variant == ComplianceReportVariant.annotated_source_output and file_count > 1:
        return ".zip"
    return ".pdf"


def _apply_split_download_filenames(
    response: AnalyzerResponse,
    source_filename: str,
) -> AnalyzerResponse:
    """Name split artifacts deterministically without creating collisions."""
    result = response.result
    output_files = (
        result.get("output_files")
        if isinstance(result, dict)
        else getattr(result, "output_files", None)
    )
    output_files = list(output_files or [])
    base_filename = _download_filename_for_action(
        FeatureType.split_pdf,
        source_filename,
    )
    base_stem = _filename_stem(base_filename)

    for index, output_file in enumerate(output_files, start=1):
        output_filename = (
            base_filename
            if len(output_files) == 1
            else f"{base_stem}.{index}.pdf"
        )
        _apply_download_filename(output_file, output_filename)

    archive_file = (
        result.get("archive_file")
        if isinstance(result, dict)
        else getattr(result, "archive_file", None)
    )
    if archive_file is not None:
        _apply_download_filename(
            archive_file,
            _download_filename_for_action(
                FeatureType.split_pdf,
                source_filename,
                output_extension=".zip",
            ),
        )

    return response


def _apply_esignature_download_filenames(
    response: AnalyzerResponse,
    source_filename: str,
) -> AnalyzerResponse:
    """Give signed, preview, and certificate artifacts distinct safe names."""
    result = response.result
    stem = _filename_stem(source_filename)
    signed_pdf = getattr(result, "signed_pdf", None)
    if signed_pdf is not None:
        _apply_download_filename(signed_pdf, f"{stem}-signed.pdf")

    certificate = getattr(result, "audit_certificate", None)
    if certificate is not None:
        _apply_download_filename(certificate, f"{stem}-certificate.pdf")

    for index, preview in enumerate(getattr(result, "previews", ()) or (), start=1):
        preview_pdf = getattr(preview, "preview_pdf", None)
        if preview_pdf is not None:
            _apply_download_filename(preview_pdf, f"{stem}-preview-{index}.pdf")
    return response


workflow_router = WorkflowRouter(download_url_builder=_download_url_for_storage_key)


def _run_request(
    request: Union[AnalyzerRequest, Mapping[str, Any]],
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    workflow_router_override: WorkflowRouter | None = None,
    **context: Any,
) -> AnalyzerResponse:
    try:
        feature = getattr(getattr(request, "action", None), "value", None)
        with artifact_owner_context(
            artifact_owner_user_id,
            organization_id=artifact_owner_organization_id,
            feature=feature,
        ), upload_processing_session(request):
            dispatcher = workflow_router_override or workflow_router
            return dispatcher.handle(
                request,
                artifact_owner_user_id=artifact_owner_user_id,
                artifact_owner_organization_id=artifact_owner_organization_id,
                **context,
            )
    except HTTPException:
        raise
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValidationError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        raise _bad_request(str(exc)) from exc
    except TypeError as exc:
        raise _bad_request(str(exc)) from exc
    except RuntimeError as exc:
        raise _service_unavailable(str(exc)) from exc



def _run_workflow_execution(
    request: AnalyzerRequest,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    **context: Any,
):
    try:
        with artifact_owner_context(
            artifact_owner_user_id,
            organization_id=artifact_owner_organization_id,
            feature=request.action.value,
        ), upload_processing_session(request):
            return workflow_router.execute(
                request,
                artifact_owner_user_id=artifact_owner_user_id,
                artifact_owner_organization_id=artifact_owner_organization_id,
                **context,
            )
    except HTTPException:
        raise
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except RuleRegistryError as exc:
        raise _bad_request(str(exc)) from exc
    except ValidationError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        raise _bad_request(str(exc)) from exc
    except TypeError as exc:
        raise _bad_request(str(exc)) from exc
    except RuntimeError as exc:
        raise _service_unavailable(str(exc)) from exc



# -----------------------------------------------------------------------------
# Input builders
# -----------------------------------------------------------------------------


def _build_document_input(
    *,
    action: FeatureType,
    file: UploadFile | None,
    text: str | None,
):
    has_file = file is not None
    has_text = text is not None and text.strip() != ""

    if has_file and has_text:
        raise _bad_request("Provide either file or text, not both.")
    if not has_file and not has_text:
        raise _bad_request("Either file or text is required.")

    try:
        if has_file:
            return build_uploaded_document_payload(action=action, upload=file)  # type: ignore[arg-type]
        return build_inline_text_payload(text=text)  # type: ignore[arg-type]
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        raise _bad_request(str(exc)) from exc


def _save_upload_to_disk(upload: UploadFile, *, subdir: str, default_name: str) -> Path:
    try:
        return save_pdf_tool_upload(
            upload,
            subdir=subdir,
            default_name=default_name,
            base_dir=PDF_UPLOAD_DIR,
        )
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc


def _build_single_pdf_input(action: FeatureType, file: UploadFile):
    saved_path = _save_upload_to_disk(
        file,
        subdir=action.value,
        default_name="document.pdf",
    )
    try:
        return build_pdf_input_artifact_for_action(
            action=action,
            file_path=saved_path,
            storage_key=str(saved_path),
            mime_type=file.content_type or "application/pdf",
        )
    except ValueError as exc:
        Path(saved_path).unlink(missing_ok=True)
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        Path(saved_path).unlink(missing_ok=True)
        raise _bad_request(str(exc)) from exc


def _build_pdf_set_input(action: FeatureType, files: list[UploadFile]):
    saved_paths = [
        _save_upload_to_disk(
            file,
            subdir=action.value,
            default_name=f"document-{index}.pdf",
        )
        for index, file in enumerate(files, start=1)
    ]
    try:
        return build_pdf_input_artifact_for_action(
            action=action,
            file_paths=saved_paths,
            storage_keys=[str(path) for path in saved_paths],
            mime_types=[file.content_type or "application/pdf" for file in files],
        )
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        raise _bad_request(str(exc)) from exc


def _deduplicate_uploads(files: list[UploadFile] | None) -> list[UploadFile]:
    upload_list = list(files or [])
    duplicate_of = find_duplicate_upload_content(upload_list)
    if not duplicate_of:
        return upload_list
    return [
        upload
        for index, upload in enumerate(upload_list, start=1)
        if index not in duplicate_of
    ]


def _validate_pdf_combine_files(files: list[UploadFile]) -> None:
    file_count = len(files or [])
    if file_count < MIN_PDF_COMBINE_FILES:
        raise _bad_request(
            f"PDF combine requires at least {MIN_PDF_COMBINE_FILES} PDF files."
        )
    if file_count > MAX_PDF_COMBINE_FILES:
        raise _bad_request(
            f"PDF combine accepts at most {MAX_PDF_COMBINE_FILES} PDF files per request."
        )


def _loads_json(value: str | None, *, default: Any = None) -> Any:
    if value is None or not str(value).strip():
        return default
    try:
        return json.loads(value)
    except json.JSONDecodeError as exc:
        raise _bad_request(f"Invalid JSON: {exc.msg}") from exc


def _parse_int_list(value: str | None) -> list[int]:
    if value is None or not value.strip():
        return []
    raw = value.strip()
    loaded = _loads_json(raw, default=None) if raw.startswith("[") else None
    if loaded is not None:
        if not isinstance(loaded, list):
            raise _bad_request("selected_pages must be a JSON array or comma-separated integers.")
        values = loaded
    else:
        values = [item.strip() for item in raw.split(",") if item.strip()]

    try:
        pages: list[int] = []
        for item in values:
            if isinstance(item, bool):
                raise ValueError("Boolean values are not page numbers.")
            if isinstance(item, int):
                page = item
            elif isinstance(item, str) and re.fullmatch(r"[0-9]+", item.strip()):
                page = int(item.strip())
            else:
                raise ValueError(f"Invalid page number: {item!r}.")
            if page < 1:
                raise ValueError("Page numbers must be greater than or equal to 1.")
            pages.append(page)
        return pages
    except (TypeError, ValueError) as exc:
        raise _bad_request(
            "selected_pages must contain only positive whole-number page numbers."
        ) from exc


def _parse_page_ranges(value: str | None) -> list[PdfPageRange]:
    if value is None or not value.strip():
        return []

    raw = value.strip()
    if raw.startswith("["):
        loaded = _loads_json(raw, default=[])
        if not isinstance(loaded, list):
            raise _bad_request("page_ranges must be a JSON array or a comma-separated range string.")
        try:
            return [PdfPageRange.model_validate(item) for item in loaded]
        except (ValidationError, TypeError, ValueError) as exc:
            raise _bad_request(
                "page_ranges must contain valid positive page ranges where the end page is not before the start page."
            ) from exc

    ranges: list[PdfPageRange] = []
    try:
        for item in raw.split(","):
            text = item.strip()
            if not text:
                continue
            match = re.fullmatch(r"([0-9]+)(?:\s*-\s*([0-9]+))?", text)
            if match is None:
                raise ValueError(f"Invalid page range: {text!r}.")
            start_page = int(match.group(1))
            end_page = int(match.group(2) or match.group(1))
            ranges.append(PdfPageRange(start_page=start_page, end_page=end_page))
        return ranges
    except (ValidationError, TypeError, ValueError) as exc:
        raise _bad_request(
            "page_ranges must use positive page numbers or inclusive ranges such as 1-3,5,8-10."
        ) from exc


def _save_pdf_edit_assets(files: list[UploadFile]) -> dict[str, str]:
    if len(files) > 25:
        raise _bad_request("PDF edit accepts at most 25 image or signature assets per request.")

    saved: dict[str, str] = {}
    try:
        for upload in files:
            filename = Path(upload.filename or "").name
            reference = Path(filename).stem
            if not re.fullmatch(r"op_[A-Za-z0-9_-]+", reference):
                raise _bad_request("A PDF edit image has an invalid operation reference.")
            key = f"asset:{reference}"
            if key in saved:
                raise _bad_request("The same PDF edit asset reference was uploaded more than once.")
            saved[key] = str(save_pdf_edit_asset_upload(upload))
    except UploadServiceUnavailableError as exc:
        for saved_path in saved.values():
            Path(saved_path).unlink(missing_ok=True)
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        for saved_path in saved.values():
            Path(saved_path).unlink(missing_ok=True)
        raise _bad_request(str(exc)) from exc
    except Exception:
        for saved_path in saved.values():
            Path(saved_path).unlink(missing_ok=True)
        raise
    return saved


def _parse_edit_operations(
    operations_json: str,
    *,
    asset_paths: Mapping[str, str] | None = None,
) -> list[PdfEditOperation]:
    loaded = _loads_json(operations_json, default=[])
    if not isinstance(loaded, list):
        raise _bad_request("operations_json must be a JSON array.")

    resolved_assets = dict(asset_paths or {})
    asset_fields = {
        "add_image": "image_storage_key",
        "drawn": "signature_svg_storage_key",
        "uploaded_image": "signature_image_storage_key",
    }
    used_assets: set[str] = set()

    for item in loaded:
        if not isinstance(item, dict):
            raise _bad_request("Every PDF edit operation must be a JSON object.")

        operation = str(item.get("operation") or "")
        if operation == "draw" and str(item.get("strokes_storage_key") or "").strip():
            raise _bad_request(
                "Draw operations submitted through the PDF edit endpoint must use path_svg."
            )
        field = None
        if operation == "add_image":
            field = asset_fields["add_image"]
        elif operation == "add_signature":
            field = asset_fields.get(str(item.get("signature_type") or ""))

        if field is None:
            continue

        reference = str(item.get(field) or "")
        if not reference.startswith("asset:") or reference not in resolved_assets:
            raise _bad_request(
                "Every added image or drawn/uploaded signature must include its uploaded image file."
            )
        item[field] = resolved_assets[reference]
        used_assets.add(reference)

    if set(resolved_assets) != used_assets:
        raise _bad_request("One or more uploaded PDF edit images are not used by an operation.")
    try:
        return TypeAdapter(list[PdfEditOperation]).validate_python(loaded)
    except (ValidationError, TypeError, ValueError) as exc:
        raise _bad_request(f"Invalid PDF edit operations: {exc}") from exc


def _parse_esignature_request(payload_json: str) -> ESignatureRequest:
    loaded = _loads_json(payload_json, default={})
    if not isinstance(loaded, dict):
        raise _bad_request("payload_json must be a JSON object.")
    loaded.setdefault("feature", FeatureType.e_signature.value)
    return ESignatureRequest.model_validate(loaded)


def _resolve_esignature_signature_assets(
    payload: ESignatureRequest,
    *,
    asset_paths: Mapping[str, str],
) -> ESignatureRequest:
    """Replace the owner's untrusted asset reference with its validated upload path."""
    signature = payload.self_signer.signature if payload.self_signer else None
    if signature is None:
        if asset_paths:
            raise _bad_request("A signature image was uploaded but no owner signature uses it.")
        return payload

    signature_type = signature.signature_type.value
    storage_field = (
        "signature_image_storage_key"
        if signature_type in {"drawn", "uploaded_image"}
        else None
    )
    if storage_field is None:
        if asset_paths:
            raise _bad_request("Typed signatures must not include a signature image upload.")
        return payload

    reference = str(getattr(signature, storage_field, None) or "").strip()
    if not reference.startswith("asset:") or reference not in asset_paths:
        raise _bad_request(
            "Drawn and uploaded-image signatures must include their matching image upload."
        )
    if set(asset_paths) != {reference}:
        raise _bad_request("Exactly one matching owner signature image is allowed.")

    resolved_signature = signature.model_copy(
        update={storage_field: asset_paths[reference]}
    )
    assert payload.self_signer is not None
    resolved_self_signer = payload.self_signer.model_copy(
        update={"signature": resolved_signature}
    )
    return payload.model_copy(update={"self_signer": resolved_self_signer})


def _parse_optional_signature(value: str | None) -> AddSignatureOperation | None:
    if value is None or not value.strip():
        return None
    loaded = _loads_json(value, default=None)
    return TypeAdapter(AddSignatureOperation).validate_python(loaded)


def _client_ip(request: Request) -> str | None:
    return request.client.host[:128] if request.client else None


def _user_agent(request: Request) -> str | None:
    value = request.headers.get("user-agent")
    return value[:512] if value else None


def _user_email(user: AuthenticatedUser | None) -> str | None:
    if user is None:
        return None
    for attr in ("email", "user_email", "sub"):
        value = getattr(user, attr, None)
        if isinstance(value, str) and "@" in value and value.strip():
            return value.strip()
    claims = getattr(user, "claims", None)
    if isinstance(claims, Mapping):
        value = claims.get("email")
        if isinstance(value, str) and "@" in value and value.strip():
            return value.strip()
    return None


def _user_name(user: AuthenticatedUser | None) -> str | None:
    if user is None:
        return None
    for attr in ("name", "display_name"):
        value = getattr(user, attr, None)
        if isinstance(value, str) and value.strip():
            return value.strip()
    claims = getattr(user, "claims", None)
    if isinstance(claims, Mapping):
        for key in ("name", "nickname", "email"):
            value = claims.get(key)
            if isinstance(value, str) and value.strip():
                return value.strip()
    return None


def _user_organization_id(user: AuthenticatedUser | None) -> str | None:
    if user is None:
        return None
    for attr in ("organization_id", "org_id"):
        value = getattr(user, attr, None)
        if value is not None and str(value).strip():
            return str(value).strip()
    return None


def _artifact_owner_kwargs(
    user: AuthenticatedUser | None,
    *,
    request: Request | None = None,
    response: Response | None = None,
) -> dict[str, str | None]:
    if user is not None:
        owner_user_id = str(user.user_id)
    else:
        if request is None:
            raise RuntimeError("Anonymous artifact ownership requires the current request.")
        owner_user_id = _anonymous_artifact_owner_id(request, response)

    return {
        "artifact_owner_user_id": owner_user_id,
        "artifact_owner_organization_id": _user_organization_id(user),
    }


# -----------------------------------------------------------------------------
# Download URL / artifact serialization helpers
# -----------------------------------------------------------------------------


def _ensure_download_url(
    response: AnalyzerResponse,
    *,
    download_filename: str | None = None,
) -> AnalyzerResponse:
    result = response.result

    storage_key = getattr(result, "storage_key", None)
    download_url = getattr(result, "download_url", None)
    if storage_key and not download_url and hasattr(result, "download_url"):
        result.download_url = _download_url_for_storage_key(storage_key)

    pdf_artifact = getattr(result, "pdf_artifact", None)
    if pdf_artifact is not None:
        pdf_storage_key = getattr(pdf_artifact, "storage_key", None)
        pdf_download_url = getattr(pdf_artifact, "download_url", None)
        if pdf_storage_key and not pdf_download_url and hasattr(pdf_artifact, "download_url"):
            pdf_artifact.download_url = _download_url_for_storage_key(pdf_storage_key)

    if download_filename:
        _apply_download_filename(result, _safe_download_filename(download_filename))

    return response


def _artifact_storage_download_candidates() -> list[LocalArtifactStorage]:
    configured_root = Path(
        os.getenv("ARTIFACT_STORAGE_DIR", "artifacts")
    ).expanduser()
    candidate_base_dirs = [
        configured_root,
        configured_root / "ai_documents",
        configured_root / "compliance",
        configured_root / "structured_extraction",
        configured_root / "pdf_tools" / "combine",
        configured_root / "pdf_tools" / "split",
        configured_root / "pdf_tools" / "edit",
        configured_root / "pdf_tools" / "compress",
        configured_root / "pdf_tools" / "preview",
        configured_root / "pdf_tools" / "preview" / "pages",
        configured_root / "esignature",
        configured_root / "esignature" / "sources",
        configured_root / "esignature" / "signed",
        configured_root / "esignature" / "previews",
        configured_root / "esignature" / "certificates",
        Path(os.getenv("ESIGNATURE_ARTIFACT_STORAGE_DIR", "artifacts/esignature")),
    ]

    storages: list[LocalArtifactStorage] = []
    seen: set[str] = set()
    for base_dir in candidate_base_dirs:
        storage = LocalArtifactStorage(base_dir=str(base_dir))
        resolved_base_dir = str(storage.base_dir.resolve())
        if resolved_base_dir in seen:
            continue
        seen.add(resolved_base_dir)
        storages.append(storage)
    return storages



def _privacy_source_path(input_payload: Any) -> str:
    filename = getattr(input_payload, "filename", None)
    if not isinstance(filename, str) or not filename.strip():
        raise _bad_request("Uploaded privacy document is missing its persisted file path.")
    return filename.strip()


def _run_privacy_request(
    request: AnalyzerRequest,
    *,
    source_path: str,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    custom_redactions: list[str] | None = None,
) -> ProtectedArtifactResult:
    try:
        with artifact_owner_context(
            artifact_owner_user_id,
            organization_id=artifact_owner_organization_id,
            feature=request.action.value,
        ), upload_processing_session(request):
            execution = workflow_router.execute(
                request,
                privacy_source_path=source_path,
                custom_redactions=custom_redactions,
                artifact_owner_user_id=artifact_owner_user_id,
                artifact_owner_organization_id=artifact_owner_organization_id,
            )
        if execution.protected_artifact is None:
            raise RuntimeError("Privacy workflow did not return a protected artifact.")
        return execution.protected_artifact
    except HTTPException as exc:
        raise to_http_exception(exc) from exc
    except RuntimeError as exc:
        raise _service_unavailable(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        raise _bad_request(str(exc)) from exc


def _build_docx_preview_artifact(
    processed: ProtectedArtifactResult,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
) -> dict[str, Any] | None:
    original_name = processed.artifact.original_artifact_name.lower()
    if not original_name.endswith(".docx"):
        return None
    with artifact_owner_context(
        artifact_owner_user_id,
        organization_id=artifact_owner_organization_id,
        feature="privacy_preview",
    ):
        preview = convert_document(
            input_format="docx",
            output_format="pdf",
            source_reference=processed.artifact.stored_path,
            source_name_hint=processed.artifact.original_artifact_name,
        )
    return {
        "filename": preview.file_name,
        "storage_key": preview.storage_key,
        "download_url": preview.download_url or _download_url_for_storage_key(preview.storage_key),
        "content_type": "application/pdf",
    }


def _build_review_preview_artifact(
    processed: ProtectedArtifactResult,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
) -> dict[str, Any]:
    docx_preview = _build_docx_preview_artifact(
        processed,
        artifact_owner_user_id=artifact_owner_user_id,
        artifact_owner_organization_id=artifact_owner_organization_id,
    )
    if docx_preview is not None:
        return docx_preview
    artifact = processed.artifact
    return {
        "filename": artifact.original_artifact_name,
        "storage_key": artifact.storage_key,
        "download_url": _download_url_for_storage_key(artifact.storage_key),
        "content_type": artifact.content_type,
    }


def _serialize_review_preview_artifact(
    processed: ProtectedArtifactResult,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    download_filename: str | None = None,
) -> dict[str, Any]:
    del download_filename
    preview = _build_review_preview_artifact(
        processed,
        artifact_owner_user_id=artifact_owner_user_id,
        artifact_owner_organization_id=artifact_owner_organization_id,
    )
    return {
        **preview,
        "artifact_name": preview["filename"],
        "original_artifact_name": preview["filename"],
        "preview_only": True,
    }


def _serialize_processed_result(
    processed: ProtectedArtifactResult,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    download_filename: str | None = None,
) -> dict[str, Any]:
    analyzer_response = _ensure_download_url(
        processed.analyzer_response,
        download_filename=download_filename,
    )
    artifact = asdict(processed.artifact)
    artifact["download_url"] = artifact.get("download_url") or _download_url_for_storage_key(
        artifact.get("storage_key")
    )
    if download_filename:
        _apply_download_filename(
            artifact,
            _safe_download_filename(download_filename),
        )
    return {
        "analyzer_response": analyzer_response.model_dump(mode="python"),
        "artifact": artifact,
        "preview_artifact": _build_docx_preview_artifact(
            processed,
            artifact_owner_user_id=artifact_owner_user_id,
            artifact_owner_organization_id=artifact_owner_organization_id,
        ),
    }


def _run_structured_extraction_request_with_preview(
    request: AnalyzerRequest,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    download_filename: str | None = None,
) -> dict[str, Any]:
    execution = _run_workflow_execution(
        request,
        artifact_owner_user_id=artifact_owner_user_id,
        artifact_owner_organization_id=artifact_owner_organization_id,
        structured_preview=True,
        structured_preview_rows_limit=200,
    )
    response = _ensure_download_url(
        execution.response,
        download_filename=download_filename,
    )
    return {
        "analyzer_response": response.model_dump(mode="json"),
        "preview_payload": execution.preview_payload,
        "preview_rows": execution.preview_rows or [],
        "preview_truncated": execution.preview_truncated,
    }


def _run_standalone_feature_request(
    request: AnalyzerRequest,
    *,
    artifact_owner_user_id: str,
    artifact_owner_organization_id: str | None = None,
    download_filename: str | None = None,
) -> AnalyzerResponse:
    return _ensure_download_url(
        _run_request(
            request,
            artifact_owner_user_id=artifact_owner_user_id,
            artifact_owner_organization_id=artifact_owner_organization_id,
        ),
        download_filename=download_filename,
    )


def _validate_numbered_questions(questions: list[str]) -> list[str]:
    cleaned = [str(item).strip() for item in questions if str(item).strip()]
    if not cleaned:
        raise _bad_request("At least one generated question is required.")
    if len(cleaned) > MAX_NUMBERED_QUESTIONS:
        raise _bad_request(
            f"At most {MAX_NUMBERED_QUESTIONS} numbered questions are allowed."
        )

    validated: list[str] = []
    for index, question in enumerate(cleaned, start=1):
        if not question.lstrip().startswith(f"{index}."):
            raise _bad_request("Questions must be sequentially numbered starting at 1.")
        try:
            validated.append(validate_question_item(question))
        except (TypeError, ValueError) as exc:
            raise _bad_request(str(exc)) from exc

    try:
        validate_auxiliary_prompt_text(
            "\n".join(validated),
            field_name="Generated questions",
        )
    except (TypeError, ValueError) as exc:
        raise _bad_request(str(exc)) from exc

    return validated


def _numbered_questions_from_plain_text(raw: str) -> list[str]:
    matches = re.findall(r"(?:^|\n)\s*(\d+)\.\s+([\s\S]*?)(?=\n\s*\d+\.\s+|$)", raw)
    return [f"{index}. {body.strip()}" for index, (_, body) in enumerate(matches, start=1) if body.strip()]


def _coerce_numbered_questions(value: Any) -> list[str]:
    if isinstance(value, dict):
        value = value.get("questions")

    if isinstance(value, list):
        return _validate_numbered_questions([str(item).strip() for item in value if str(item).strip()])

    if isinstance(value, str):
        raw = value.strip()
        if not raw:
            raise _bad_request("At least one generated question is required.")
        try:
            raw = validate_auxiliary_prompt_text(
                raw,
                field_name="questions_json",
            )
        except (TypeError, ValueError) as exc:
            raise _bad_request(str(exc)) from exc
        if raw.startswith("[") or raw.startswith("{"):
            return _coerce_numbered_questions(_loads_json(raw, default=None))
        return _validate_numbered_questions(_numbered_questions_from_plain_text(raw))

    raise _bad_request("questions_json must be a JSON array, numbered text, or an object with a questions array.")


def _parse_numbered_questions(value: str | None) -> list[str]:
    """
    Parse generated questions supplied by the frontend for the follow-on
    generate_answers action.

    Accepted forms:
    - JSON array: ["1. ...", "2. ..."]
    - JSON object with a questions array: {"questions": [...]}
    - Plain numbered text: "1. ...\n2. ..."
    """
    if value is None or not str(value).strip():
        raise _bad_request("questions_json is required to generate answers.")
    return _coerce_numbered_questions(str(value).strip())


def _parse_batch_questions_by_index(value: str | None, *, file_count: int) -> dict[int, list[str]]:
    """Parse per-file questions for batch answer generation.

    Accepted forms:
    - shared JSON array / numbered text: applied to every file;
    - {"questions": [...]}: shared questions applied to every file;
    - {"questions_by_index": {"1": [...], "2": [...]}}: per-file questions;
    - {"1": [...], "2": [...]}: compact per-file shape.
    """
    if value is None or not str(value).strip():
        raise _bad_request("questions_json is required to generate answers.")

    raw = str(value).strip()
    parsed = _loads_json(raw, default=None) if raw.startswith("{") or raw.startswith("[") else None

    if isinstance(parsed, dict):
        per_file = parsed.get("questions_by_index") or parsed.get("questionsByIndex") or parsed.get("by_index")
        if per_file is None and all(str(key).isdigit() for key in parsed.keys()):
            per_file = parsed

        if isinstance(per_file, dict):
            questions_by_index: dict[int, list[str]] = {}
            for index in range(1, file_count + 1):
                raw_questions = per_file.get(str(index), per_file.get(index))
                if raw_questions is not None:
                    questions_by_index[index] = _coerce_numbered_questions(raw_questions)
            return questions_by_index

    shared_questions = _coerce_numbered_questions(parsed if parsed is not None else raw)
    return {index: shared_questions for index in range(1, file_count + 1)}

def _artifact_path_from_result(result: Any) -> Path | None:
    storage_key = getattr(result, "storage_key", None)
    filename = getattr(result, "filename", None)
    candidates: list[Path] = []

    if isinstance(storage_key, str) and storage_key.strip():
        raw_key = storage_key.strip().replace("\\", "/")
        candidates.append(Path(raw_key))

        key = raw_key
        key = key.removeprefix("/api/analyzer/artifacts/")
        key = key.removeprefix("/api/v1/analyzer/artifacts/")
        key = key.removeprefix("/artifacts/")
        key = key.removeprefix("artifacts/")

        for storage in _artifact_storage_download_candidates():
            candidates.append(storage.base_dir / key)

    if isinstance(filename, str) and filename.strip():
        for storage in _artifact_storage_download_candidates():
            candidates.append(storage.base_dir / Path(filename).name)

    seen: set[str] = set()
    for candidate in candidates:
        try:
            resolved = candidate.expanduser().resolve()
        except OSError:
            continue
        marker = str(resolved)
        if marker in seen:
            continue
        seen.add(marker)
        if resolved.exists() and resolved.is_file():
            return resolved

    return None


def _read_text_from_artifact(path: Path) -> str:
    suffix = path.suffix.lower()

    if suffix == ".txt":
        return path.read_text(encoding="utf-8", errors="replace").strip()

    if suffix == ".docx":
        try:
            import docx  # python-docx
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("python-docx is required to read generated DOCX questions.") from exc
        document = docx.Document(str(path))
        return "\n".join(paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()).strip()

    if suffix == ".pdf":
        try:
            import fitz  # PyMuPDF
        except ImportError as exc:  # pragma: no cover
            raise RuntimeError("PyMuPDF is required to read generated PDF questions.") from exc
        with fitz.open(path) as pdf:
            return "\n".join(page.get_text("text") for page in pdf).strip()

    return ""


def _generated_questions_text_from_response(response: AnalyzerResponse) -> str | None:
    result = response.result
    inline_content = getattr(result, "content", None)
    if isinstance(inline_content, str) and inline_content.strip():
        return inline_content.strip()

    generated_content = getattr(result, "generated_questions_text", None)
    if isinstance(generated_content, str) and generated_content.strip():
        return generated_content.strip()

    path = _artifact_path_from_result(result)
    if path is None:
        return None

    try:
        text = _read_text_from_artifact(path)
    except Exception:
        return None

    return text or None


def _clean_repeated_strings(values: list[str] | None) -> list[str]:
    if not values:
        return []
    cleaned: list[str] = []
    seen: set[str] = set()
    for value in values:
        text = str(value).strip()
        if not text or text in seen:
            continue
        cleaned.append(text)
        seen.add(text)
    return cleaned


def _serialize_candidates(candidates: list[Any]) -> list[dict[str, Any]]:
    serialized: list[dict[str, Any]] = []
    for index, candidate in enumerate(candidates):
        serialized.append(
            {
                "id": f"{candidate.label}|{candidate.source}|{candidate.quote}|{index}",
                "label": candidate.label,
                "quote": candidate.quote,
                "occurrences": candidate.occurrences,
                "source": candidate.source,
            }
        )
    return serialized


# -----------------------------------------------------------------------------
# Request builders for privacy/compliance
# -----------------------------------------------------------------------------


def _normalize_compliance_sector_packs(
    sector_packs: list[ComplianceSectorPack] | None,
) -> list[ComplianceSectorPack]:
    core_pack = ComplianceSectorPack.core_control_library
    legacy_core_pack = getattr(ComplianceSectorPack, "nigeria_core_control_library", None)

    resolved: list[ComplianceSectorPack] = []
    seen: set[ComplianceSectorPack] = set()
    for pack in list(sector_packs or []):
        if legacy_core_pack is not None and pack == legacy_core_pack:
            pack = core_pack
        if pack not in seen:
            resolved.append(pack)
            seen.add(pack)

    if core_pack not in seen:
        resolved.insert(0, core_pack)
    return resolved


def _default_compliance_sector_packs(
    jurisdiction: ComplianceJurisdiction,
) -> list[ComplianceSectorPack]:
    return [ComplianceSectorPack.core_control_library]


def _compliance_uploads(
    *,
    file: UploadFile | None,
    files: list[UploadFile] | None,
) -> list[UploadFile]:
    """
    Normalize Compliance uploads without enabling batch processing.

    This is one Compliance document-set request, not the paid-plan batch feature.
    The maximum is fixed at MAX_COMPLIANCE_DOCUMENT_SET_FILES.
    """
    upload_list = list(files or [])

    if file is not None and upload_list:
        raise _bad_request(
            "Submit compliance uploads under either the legacy 'file' field or the multi-document 'files' field, not both."
        )

    if file is not None:
        upload_list = [file]

    if not upload_list:
        raise _bad_request("At least one file is required for compliance.")

    upload_list = _deduplicate_uploads(upload_list)

    if len(upload_list) > MAX_COMPLIANCE_DOCUMENT_SET_FILES:
        raise _bad_request(
            f"Compliance accepts a maximum of {MAX_COMPLIANCE_DOCUMENT_SET_FILES} unique files per document-set request."
        )

    return upload_list


def _build_compliance_input_payload(
    *,
    file: UploadFile | None,
    files: list[UploadFile] | None,
):
    uploads = _compliance_uploads(file=file, files=files)

    try:
        documents = [
            build_uploaded_document_payload(
                action=FeatureType.compliance,
                upload=upload,
            )
            for upload in uploads
        ]
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except FileNotFoundError as exc:
        raise _bad_request(str(exc)) from exc

    if len(documents) == 1:
        return documents[0]

    return DocumentSetPayload(documents=documents)


def _build_compliance_request(
    *,
    file: UploadFile | None = None,
    files: list[UploadFile] | None = None,
    jurisdiction: ComplianceJurisdiction,
    sector_packs: list[ComplianceSectorPack] | None,
    regulatory_domains: list[ComplianceRegulatoryDomain] | None,
    report_variant: ComplianceReportVariant,
    system_language: SystemLanguage,
) -> AnalyzerRequest:
    input_payload = _build_compliance_input_payload(file=file, files=files)

    payload = ComplianceRequest(
        feature=FeatureType.compliance,
        jurisdiction=jurisdiction,
        sector_packs=_normalize_compliance_sector_packs(
            sector_packs or _default_compliance_sector_packs(jurisdiction)
        ),
        regulatory_domains=regulatory_domains or [],
        report_variant=report_variant,
        require_human_review=True,
    )

    return AnalyzerRequest(
        action=FeatureType.compliance,
        input=input_payload,
        payload=payload,
        policy=_policy_for_action(FeatureType.compliance),
        system_language=system_language,
    )



def _privacy_payload_kwargs(
    *,
    feature: FeatureType,
    document_type: RedactionMaskingDocumentType | None,
    target_data: list[SensitiveDataType] | None,
    review_exclusions: list[str] | None,
) -> dict[str, Any]:
    payload_kwargs: dict[str, Any] = {"feature": feature}
    if document_type is not None:
        payload_kwargs["document_type"] = document_type
    if target_data:
        payload_kwargs["target_data"] = target_data
    if review_exclusions:
        payload_kwargs["review_exclusions"] = [item.strip() for item in review_exclusions if item and item.strip()]
    return payload_kwargs


def _build_privacy_request(
    *,
    action: FeatureType,
    file: UploadFile,
    document_type: RedactionMaskingDocumentType | None,
    target_data: list[SensitiveDataType] | None,
    review_exclusions: list[str] | None,
    system_language: SystemLanguage,
) -> tuple[Any, AnalyzerRequest]:
    try:
        input_payload = build_uploaded_document_payload(action=action, upload=file)
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc

    payload_kwargs = _privacy_payload_kwargs(
        feature=action,
        document_type=document_type,
        target_data=target_data,
        review_exclusions=review_exclusions,
    )
    payload = RedactionRequest(**payload_kwargs) if action == FeatureType.redact else DataMaskingRequest(**payload_kwargs)

    request = AnalyzerRequest(
        action=action,
        input=input_payload,
        payload=payload,
        policy=_policy_for_action(action),
        system_language=system_language,
    )
    return input_payload, request


def _detect_privacy_document_type_upload(
    *,
    action: FeatureType,
    file: UploadFile,
) -> dict[str, object]:
    if action not in {FeatureType.redact, FeatureType.data_mask}:
        raise ValueError("Privacy document-type detection only supports redact and data_mask.")

    original_filename = _uploaded_filename(file)
    try:
        input_payload = build_uploaded_document_payload(action=action, upload=file)
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc

    # The detection upload is short-lived and never becomes an artifact. Reuse the
    # existing retention session so the saved source is marked for prompt cleanup.
    with upload_processing_session(input_payload):
        detection = detect_privacy_document_type(
            text=input_payload.text,
            filename=original_filename,
        )
    return detection.as_dict()


# -----------------------------------------------------------------------------
# Paid-plan batch upload helpers
# -----------------------------------------------------------------------------


def _batch_policy_exception(exc: BatchUploadPolicyError) -> HTTPException:
    return HTTPException(
        status_code=getattr(exc, "status_code", 400),
        detail={
            "error": getattr(exc, "error_code", "invalid_batch_upload"),
            "message": str(exc),
            "limits": BATCH_UPLOAD_LIMITS_BY_PLAN,
        },
    )


def _require_batch_upload_policy(
    *,
    current_user: AuthenticatedUser,
    action: FeatureType,
    files: list[UploadFile],
) -> BatchUploadPolicy:
    try:
        # Use the same authoritative subscription entitlement as the feature
        # access/rate-limit dependency. AuthenticatedUser represents identity
        # and token claims; it is not the billing source of truth.
        entitlement = get_user_entitlement(current_user.user_id)
        return require_batch_upload_entitlement(
            entitlement,
            feature=action.value,
            files=files,
        )
    except BatchUploadPolicyError as exc:
        raise _batch_policy_exception(exc) from exc


def _serialize_batch_result(
    value: Any,
    *,
    download_filename: str | None = None,
) -> Any:
    if isinstance(value, AnalyzerResponse):
        return _ensure_download_url(
            value,
            download_filename=download_filename,
        ).model_dump(mode="json")
    if hasattr(value, "model_dump"):
        serialized = value.model_dump(mode="json")
    else:
        serialized = value
    if download_filename:
        _apply_download_filename(serialized, _safe_download_filename(download_filename))
    return serialized


def _batch_error_payload(status_code: int, error: str, message: str) -> dict[str, Any]:
    return {
        "status_code": status_code,
        "error": error,
        "message": message,
    }


def _http_error_detail(exc: HTTPException) -> dict[str, Any]:
    detail = exc.detail
    if isinstance(detail, dict):
        return _batch_error_payload(
            int(exc.status_code),
            str(detail.get("error") or "request_failed"),
            str(detail.get("message") or detail.get("detail") or "Request failed."),
        )
    return _batch_error_payload(int(exc.status_code), "request_failed", str(detail or "Request failed."))



def _batch_download_filename(
    *,
    action: FeatureType,
    original_filename: str,
    policy: BatchUploadPolicy,
    output_extension: str | None = None,
) -> str:
    return _download_filename_for_action(
        action,
        original_filename,
        output_extension=output_extension or policy.extension,
    )


def _batch_item_from_upload(
    *,
    action: FeatureType,
    upload: UploadFile,
    index: int,
    policy: BatchUploadPolicy,
    operation: Callable[[UploadFile], Any],
    output_extension: str | None = None,
) -> dict[str, Any]:
    original_filename = (upload.filename or f"upload-{index}{policy.extension}").strip()
    item_started = time.perf_counter()

    try:
        result = operation(upload)
        effective_output_extension = output_extension
        if action == FeatureType.convert and isinstance(result, AnalyzerResponse):
            effective_output_extension = _conversion_artifact_extension(
                result, output_extension or policy.extension
            )
        download_filename = _batch_download_filename(
            action=action,
            original_filename=original_filename,
            policy=policy,
            output_extension=effective_output_extension,
        )
        item = {
            "index": index,
            "filename": original_filename,
            "success": True,
            "response": _serialize_batch_result(
                result,
                download_filename=download_filename,
            ),
        }
    except HTTPException as exc:
        item = {
            "index": index,
            "filename": original_filename,
            "success": False,
            "error": _http_error_detail(exc),
        }
    except UploadServiceUnavailableError as exc:
        item = {
            "index": index,
            "filename": original_filename,
            "success": False,
            "error": _batch_error_payload(503, "service_unavailable", str(exc)),
        }
    except (UploadError, ValidationError, ValueError, FileNotFoundError, TypeError) as exc:
        item = {
            "index": index,
            "filename": original_filename,
            "success": False,
            "error": _batch_error_payload(400, "invalid_request", str(exc)),
        }
    except RuntimeError as exc:
        item = {
            "index": index,
            "filename": original_filename,
            "success": False,
            "error": _batch_error_payload(503, "service_unavailable", str(exc)),
        }
    except Exception as exc:  # pragma: no cover - defensive isolation per file.
        item = {
            "index": index,
            "filename": original_filename,
            "success": False,
            "error": _batch_error_payload(500, "batch_item_failed", str(exc)),
        }

    item["elapsed_ms"] = round((time.perf_counter() - item_started) * 1000)
    return item


def _run_batch_uploads(
    *,
    action: FeatureType,
    files: list[UploadFile],
    policy: BatchUploadPolicy,
    operation: Callable[[UploadFile], Any],
    output_extension: str | None = None,
) -> JSONResponse:
    """Run unique batch items concurrently and reject duplicates per item."""
    batch_started = time.perf_counter()
    indexed_uploads = list(enumerate(files, start=1))
    duplicate_of = find_duplicate_upload_content(files)
    unique_uploads = [
        (index, upload)
        for index, upload in indexed_uploads
        if index not in duplicate_of
    ]
    worker_count = max(
        1,
        min(
            int(getattr(policy, "max_concurrency", 1) or 1),
            len(unique_uploads) or 1,
        ),
    )
    items_by_index: dict[int, dict[str, Any]] = {}

    for duplicate_index, original_index in duplicate_of.items():
        duplicate_upload = files[duplicate_index - 1]
        original_upload = files[original_index - 1]
        items_by_index[duplicate_index] = {
            "index": duplicate_index,
            "filename": (
                duplicate_upload.filename
                or f"upload-{duplicate_index}{policy.extension}"
            ).strip(),
            "success": False,
            "error": _batch_error_payload(
                400,
                "duplicate_batch_upload",
                (
                    f"Duplicate file rejected: "
                    f"{_uploaded_filename(duplicate_upload)!r} has the same content as "
                    f"{_uploaded_filename(original_upload)!r}."
                ),
            ),
            "duplicate_of_index": original_index,
            "elapsed_ms": 0,
        }

    if worker_count == 1:
        for index, upload in unique_uploads:
            items_by_index[index] = _batch_item_from_upload(
                action=action,
                upload=upload,
                index=index,
                policy=policy,
                operation=operation,
                output_extension=output_extension,
            )
    else:
        with concurrent.futures.ThreadPoolExecutor(
            max_workers=worker_count,
            thread_name_prefix=f"redocx-{action.value}-batch",
        ) as executor:
            future_by_index = {
                executor.submit(
                    _batch_item_from_upload,
                    action=action,
                    upload=upload,
                    index=index,
                    policy=policy,
                    operation=operation,
                    output_extension=output_extension,
                ): index
                for index, upload in unique_uploads
            }

            for future in concurrent.futures.as_completed(future_by_index):
                index = future_by_index[future]
                try:
                    items_by_index[index] = future.result()
                except Exception as exc:  # pragma: no cover - worker envelope fallback.
                    upload = files[index - 1]
                    items_by_index[index] = {
                        "index": index,
                        "filename": (
                            upload.filename or f"upload-{index}{policy.extension}"
                        ).strip(),
                        "success": False,
                        "error": _batch_error_payload(
                            500,
                            "batch_worker_failed",
                            str(exc),
                        ),
                        "elapsed_ms": 0,
                    }

    items = [
        items_by_index[index]
        for index, _ in indexed_uploads
        if index in items_by_index
    ]
    succeeded = sum(1 for item in items if item.get("success") is True)
    failed = len(items) - succeeded
    elapsed_ms = round((time.perf_counter() - batch_started) * 1000)

    payload = {
        "success": failed == 0,
        "feature": action.value,
        "batch": {
            "plan": policy.plan,
            "limit": policy.max_uploads,
            "extension": policy.extension,
            "file_count": policy.file_count,
            "unique_file_count": len(unique_uploads),
            "duplicate_file_count": len(duplicate_of),
            "succeeded": succeeded,
            "failed": failed,
            "concurrency": worker_count,
            "processing_mode": "concurrent" if worker_count > 1 else "sequential",
            "elapsed_ms": elapsed_ms,
        },
        "items": items,
    }
    return JSONResponse(
        status_code=200 if failed == 0 else 207,
        content=payload,
    )

BATCH_CONVERSION_OUTPUTS_BY_INPUT_EXTENSION: dict[str, set[str]] = {
    ".pdf": {"docx", "jpg", "pptx", "xlsx", "pdfa"},
    ".docx": {"pdf"},
    ".xlsx": {"pdf"},
    ".html": {"pdf"},
    ".htm": {"pdf"},
    ".pptx": {"pdf"},
    ".jpg": {"pdf", "docx"},
    ".jpeg": {"pdf", "docx"},
    ".png": {"jpg", "jpeg"},
}

BATCH_TRANSCRIBE_MEDIA_TYPE_BY_EXTENSION: dict[str, str] = {
    ".mp3": "audio",
    ".mp4": "video",
    ".mkv": "video",
    ".mov": "video",
}


def _form_value(value: Any) -> str:
    return str(getattr(value, "value", value) or "").strip().lower()


def _require_batch_extension(
    *,
    policy: BatchUploadPolicy,
    allowed_extensions: set[str],
    feature_label: str,
) -> None:
    normalized_allowed = {item if item.startswith(".") else f".{item}" for item in allowed_extensions}
    if policy.extension not in normalized_allowed:
        raise _bad_request(
            f"{feature_label} batch processing only accepts {', '.join(sorted(normalized_allowed))} files. "
            f"Received {policy.extension}."
        )


def _require_batch_conversion_action(
    *,
    policy: BatchUploadPolicy,
    output_format: ConversionOutputFormat,
) -> None:
    output = _form_value(output_format)
    allowed_outputs = BATCH_CONVERSION_OUTPUTS_BY_INPUT_EXTENSION.get(policy.extension, set())
    if not allowed_outputs:
        raise _bad_request(
            f"Batch conversion does not support {policy.extension} uploads."
        )
    if output not in allowed_outputs:
        raise _bad_request(
            "All files in a conversion batch must follow one valid conversion action. "
            f"For {policy.extension} input batches, allowed output formats are: "
            f"{', '.join(sorted(allowed_outputs))}. Received: {output or 'unknown'}."
        )


def _require_batch_transcription_action(
    *,
    policy: BatchUploadPolicy,
    media_type: MediaType,
) -> None:
    expected_media_type = BATCH_TRANSCRIBE_MEDIA_TYPE_BY_EXTENSION.get(policy.extension)
    requested_media_type = _form_value(media_type)
    if expected_media_type is None:
        raise _bad_request(
            f"Speech-to-text batch processing does not support {policy.extension} uploads."
        )
    if requested_media_type != expected_media_type:
        raise _bad_request(
            "All files in a speech-to-text batch must follow one valid media action. "
            f"{policy.extension} batches must be submitted as {expected_media_type}, "
            f"not {requested_media_type or 'unknown'}."
        )


# -----------------------------------------------------------------------------
# Paid-plan batch processing routes
# -----------------------------------------------------------------------------


@router.post("/batch/convert", dependencies=[Depends(rate_limit_for_feature(FeatureType.convert))])
def batch_convert_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    output_format: ConversionOutputFormat = Form(...),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.convert, files=files)
    _require_batch_conversion_action(policy=policy, output_format=output_format)

    def operation(upload: UploadFile) -> AnalyzerResponse:
        input_payload = build_uploaded_document_payload(action=FeatureType.convert, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.convert,
            input=input_payload,
            payload=ConversionRequest(feature=FeatureType.convert, output_format=output_format),
            policy=_policy_for_action(FeatureType.convert),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(
        action=FeatureType.convert,
        files=files,
        policy=policy,
        operation=operation,
        output_extension=(
            ".pdf" if _normalized_extension(output_format) == ".pdfa"
            else _normalized_extension(output_format)
        ),
    )


@router.post("/batch/summarize", dependencies=[Depends(rate_limit_for_feature(FeatureType.summarize))])
def batch_summarize_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.summarize, files=files)

    def operation(upload: UploadFile) -> AnalyzerResponse:
        input_payload = build_uploaded_document_payload(action=FeatureType.summarize, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.summarize,
            input=input_payload,
            payload=SummarizationRequest(feature=FeatureType.summarize),
            policy=_policy_for_action(FeatureType.summarize),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(action=FeatureType.summarize, files=files, policy=policy, operation=operation)


@router.post("/batch/grammar-correct", dependencies=[Depends(rate_limit_for_feature(FeatureType.grammar_correct))])
def batch_grammar_correct_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.grammar_correct, files=files)

    def operation(upload: UploadFile) -> AnalyzerResponse:
        input_payload = build_uploaded_document_payload(action=FeatureType.grammar_correct, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.grammar_correct,
            input=input_payload,
            payload=GrammarCorrectionRequest(feature=FeatureType.grammar_correct),
            policy=_policy_for_action(FeatureType.grammar_correct),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(action=FeatureType.grammar_correct, files=files, policy=policy, operation=operation)


@router.post("/batch/translate", dependencies=[Depends(rate_limit_for_feature(FeatureType.translate))])
def batch_translate_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    target_language: str = Form(...),
    source_language: str = Form("auto"),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.translate, files=files)

    def operation(upload: UploadFile) -> AnalyzerResponse:
        input_payload = build_uploaded_document_payload(action=FeatureType.translate, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.translate,
            input=input_payload,
            payload=TranslationRequest(
                feature=FeatureType.translate,
                source_language=source_language,
                target_language=target_language,
            ),
            policy=_policy_for_action(FeatureType.translate),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(action=FeatureType.translate, files=files, policy=policy, operation=operation)


@router.post("/batch/explain", dependencies=[Depends(rate_limit_for_feature(FeatureType.explain))])
def batch_explain_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    allow_external_knowledge: bool = Form(False),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.explain, files=files)

    def operation(upload: UploadFile) -> AnalyzerResponse:
        input_payload = build_uploaded_document_payload(action=FeatureType.explain, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.explain,
            input=input_payload,
            payload=ExplanationRequest(
                feature=FeatureType.explain,
                allow_external_knowledge=allow_external_knowledge,
            ),
            policy=_policy_for_action(FeatureType.explain),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(action=FeatureType.explain, files=files, policy=policy, operation=operation)


@router.post("/batch/generate-questions", dependencies=[Depends(rate_limit_for_feature(FeatureType.generate_questions))])
def batch_generate_questions_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.generate_questions, files=files)

    def operation(upload: UploadFile) -> dict[str, Any]:
        input_payload = build_uploaded_document_payload(action=FeatureType.generate_questions, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.generate_questions,
            input=input_payload,
            payload=QuestionGenerationRequest(feature=FeatureType.generate_questions),
            policy=_policy_for_action(FeatureType.generate_questions),
            system_language=system_language,
        )
        response = _run_request(request, **_artifact_owner_kwargs(current_user))
        body = response.model_dump(mode="json")
        generated_questions_text = _generated_questions_text_from_response(response)
        if generated_questions_text:
            body["generated_questions_text"] = generated_questions_text
        return body

    return _run_batch_uploads(action=FeatureType.generate_questions, files=files, policy=policy, operation=operation)


@router.post("/batch/generate-answers", dependencies=[Depends(rate_limit_for_feature(FeatureType.generate_answers))])
def batch_generate_answers_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    questions_json: str = Form(..., max_length=AUXILIARY_PROMPT_POLICY.max_chars),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.generate_answers, files=files)
    questions_by_index = _parse_batch_questions_by_index(questions_json, file_count=len(files))
    index_by_upload_id = {id(upload): index for index, upload in enumerate(files, start=1)}

    def operation(upload: UploadFile) -> AnalyzerResponse:
        upload_index = index_by_upload_id[id(upload)]
        questions = questions_by_index.get(upload_index)
        if not questions:
            raise _bad_request("No generated questions were available for this file.")
        input_payload = build_uploaded_document_payload(action=FeatureType.generate_answers, upload=upload)
        request = AnalyzerRequest(
            action=FeatureType.generate_answers,
            input=input_payload,
            payload=AnswerGenerationRequest(
                feature=FeatureType.generate_answers,
                questions=questions,
            ),
            policy=_policy_for_action(FeatureType.generate_answers),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(action=FeatureType.generate_answers, files=files, policy=policy, operation=operation)


@router.post("/batch/pdf/compress", dependencies=[Depends(rate_limit_for_feature(FeatureType.compress_pdf))])
def batch_compress_pdf_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    compression_level: PdfCompressionLevel = Form(PdfCompressionLevel.balanced),
    output_filename: str = Form("compressed-document.pdf"),
    async_processing: bool = Form(True),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.compress_pdf, files=files)
    _require_batch_extension(policy=policy, allowed_extensions={".pdf"}, feature_label="PDF compression")

    del output_filename

    def operation(upload: UploadFile) -> AnalyzerResponse:
        source_filename = _uploaded_filename(upload)
        input_payload = _build_single_pdf_input(FeatureType.compress_pdf, upload)
        request = AnalyzerRequest(
            action=FeatureType.compress_pdf,
            input=input_payload,
            payload=CompressPdfRequest(
                feature=FeatureType.compress_pdf,
                compression_level=compression_level,
                output_filename=_download_filename_for_action(FeatureType.compress_pdf, source_filename),
                async_processing=async_processing,
            ),
            policy=_policy_for_action(FeatureType.compress_pdf),
            system_language=system_language,
        )
        return _run_request(
            request,
            pdf_job_owner_id=str(current_user.user_id),
            **_artifact_owner_kwargs(current_user),
        )

    return _run_batch_uploads(action=FeatureType.compress_pdf, files=files, policy=policy, operation=operation)


@router.post("/batch/transcribe", dependencies=[Depends(rate_limit_for_feature(FeatureType.transcribe))])
def batch_transcribe_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    media_type: MediaType = Form(...),
    duration_seconds: list[int] = Form(...),
    preserve_filler_words: bool = Form(True),
    remove_background_noise: bool = Form(False),
    diarize_speakers: bool = Form(True),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    policy = _require_batch_upload_policy(current_user=current_user, action=FeatureType.transcribe, files=files)
    _require_batch_transcription_action(policy=policy, media_type=media_type)
    durations = [int(item) for item in duration_seconds]
    if len(durations) == 1 and len(files) > 1:
        durations = durations * len(files)
    if len(durations) != len(files):
        raise _bad_request("duration_seconds must be supplied once per uploaded media file.")
    duration_by_filename = {id(upload): durations[index] for index, upload in enumerate(files)}

    def operation(upload: UploadFile) -> AnalyzerResponse:
        try:
            input_payload = build_uploaded_media_payload(
                upload=upload,
                media_type=media_type,
                duration_seconds=duration_by_filename[id(upload)],
            )
        except UploadServiceUnavailableError as exc:
            raise _service_unavailable(str(exc)) from exc
        except UploadError as exc:
            raise _bad_request(str(exc)) from exc
        except ValueError as exc:
            raise _bad_request(str(exc)) from exc

        request = AnalyzerRequest(
            action=FeatureType.transcribe,
            input=input_payload,
            payload=TranscriptionRequest(
                feature=FeatureType.transcribe,
                preserve_filler_words=preserve_filler_words,
                remove_background_noise=remove_background_noise,
                diarize_speakers=diarize_speakers,
            ),
            policy=_policy_for_action(FeatureType.transcribe),
            system_language=system_language,
        )
        return _run_request(request, **_artifact_owner_kwargs(current_user))

    return _run_batch_uploads(
        action=FeatureType.transcribe,
        files=files,
        policy=policy,
        operation=operation,
        output_extension=".pdf",
    )

# -----------------------------------------------------------------------------
# Existing AI/document routes
# -----------------------------------------------------------------------------


@router.post("/convert", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.convert))])
def convert_route(
    http_request: Request,
    http_response: Response,
    current_user: AuthenticatedUser | None = Depends(get_current_user_optional),
    file: UploadFile = File(...),
    output_format: ConversionOutputFormat = Form(...),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    source_filename = _uploaded_filename(file)
    try:
        input_payload = build_uploaded_document_payload(action=FeatureType.convert, upload=file)
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc

    request = AnalyzerRequest(
        action=FeatureType.convert,
        input=input_payload,
        payload=ConversionRequest(feature=FeatureType.convert, output_format=output_format),
        policy=_policy_for_action(FeatureType.convert),
        system_language=system_language,
    )
    response = _run_request(
        request,
        **_artifact_owner_kwargs(
            current_user,
            request=http_request,
            response=http_response,
        ),
    )
    return _ensure_download_url(
        response,
        download_filename=_download_filename_for_action(
            FeatureType.convert,
            source_filename,
            output_extension=_conversion_artifact_extension(response, output_format),
        ),
    )


@router.post("/summarize", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.summarize))])
def summarize_route(
    http_request: Request,
    http_response: Response,
    current_user: AuthenticatedUser | None = Depends(get_current_user_optional),
    file: UploadFile | None = File(default=None),
    text: str | None = Form(default=None, max_length=INLINE_TEXT_POLICY.max_chars),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    input_payload = _build_document_input(action=FeatureType.summarize, file=file, text=text)
    request = AnalyzerRequest(
        action=FeatureType.summarize,
        input=input_payload,
        payload=SummarizationRequest(feature=FeatureType.summarize),
        policy=_policy_for_action(FeatureType.summarize),
        system_language=system_language,
    )
    response = _run_request(
        request,
        **_artifact_owner_kwargs(
            current_user,
            request=http_request,
            response=http_response,
        ),
    )
    return _ensure_download_url(
        response,
        download_filename=(
            _download_filename_for_action(FeatureType.summarize, _uploaded_filename(file))
            if file is not None
            else None
        ),
    )


@router.post("/grammar-correct", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.grammar_correct))])
def grammar_correct_route(
    http_request: Request,
    http_response: Response,
    current_user: AuthenticatedUser | None = Depends(get_current_user_optional),
    file: UploadFile | None = File(default=None),
    text: str | None = Form(default=None, max_length=INLINE_TEXT_POLICY.max_chars),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    input_payload = _build_document_input(action=FeatureType.grammar_correct, file=file, text=text)
    request = AnalyzerRequest(
        action=FeatureType.grammar_correct,
        input=input_payload,
        payload=GrammarCorrectionRequest(feature=FeatureType.grammar_correct),
        policy=_policy_for_action(FeatureType.grammar_correct),
        system_language=system_language,
    )
    response = _run_request(
        request,
        **_artifact_owner_kwargs(
            current_user,
            request=http_request,
            response=http_response,
        ),
    )
    return _ensure_download_url(
        response,
        download_filename=(
            _download_filename_for_action(FeatureType.grammar_correct, _uploaded_filename(file))
            if file is not None
            else None
        ),
    )


@router.post("/translate", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.translate))])
def translate_route(
    http_request: Request,
    http_response: Response,
    current_user: AuthenticatedUser | None = Depends(get_current_user_optional),
    target_language: str = Form(...),
    source_language: str = Form("auto"),
    file: UploadFile | None = File(default=None),
    text: str | None = Form(default=None, max_length=INLINE_TEXT_POLICY.max_chars),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    input_payload = _build_document_input(action=FeatureType.translate, file=file, text=text)
    request = AnalyzerRequest(
        action=FeatureType.translate,
        input=input_payload,
        payload=TranslationRequest(
            feature=FeatureType.translate,
            source_language=source_language,
            target_language=target_language,
        ),
        policy=_policy_for_action(FeatureType.translate),
        system_language=system_language,
    )
    response = _run_request(
        request,
        **_artifact_owner_kwargs(
            current_user,
            request=http_request,
            response=http_response,
        ),
    )
    return _ensure_download_url(
        response,
        download_filename=(
            _download_filename_for_action(FeatureType.translate, _uploaded_filename(file))
            if file is not None
            else None
        ),
    )


@router.post("/transcribe", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.transcribe))])
def transcribe_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    media_type: MediaType = Form(...),
    duration_seconds: int = Form(...),
    preserve_filler_words: bool = Form(True),
    remove_background_noise: bool = Form(False),
    diarize_speakers: bool = Form(True),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    source_filename = _uploaded_filename(file)
    try:
        input_payload = build_uploaded_media_payload(upload=file, media_type=media_type, duration_seconds=duration_seconds)
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc

    request = AnalyzerRequest(
        action=FeatureType.transcribe,
        input=input_payload,
        payload=TranscriptionRequest(
            feature=FeatureType.transcribe,
            preserve_filler_words=preserve_filler_words,
            remove_background_noise=remove_background_noise,
            diarize_speakers=diarize_speakers,
        ),
        policy=_policy_for_action(FeatureType.transcribe),
        system_language=system_language,
    )
    return _ensure_download_url(
        _run_request(request, **_artifact_owner_kwargs(current_user)),
        download_filename=_download_filename_for_action(
            FeatureType.transcribe,
            source_filename,
            output_extension=".pdf",
        ),
    )


@router.post("/explain", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.explain))])
def explain_route(
    http_request: Request,
    http_response: Response,
    current_user: AuthenticatedUser | None = Depends(get_current_user_optional),
    file: UploadFile | None = File(default=None),
    text: str | None = Form(default=None, max_length=INLINE_TEXT_POLICY.max_chars),
    allow_external_knowledge: bool = Form(False),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    input_payload = _build_document_input(action=FeatureType.explain, file=file, text=text)
    request = AnalyzerRequest(
        action=FeatureType.explain,
        input=input_payload,
        payload=ExplanationRequest(
            feature=FeatureType.explain,
            allow_external_knowledge=allow_external_knowledge,
        ),
        policy=_policy_for_action(FeatureType.explain),
        system_language=system_language,
    )
    response = _run_request(
        request,
        **_artifact_owner_kwargs(
            current_user,
            request=http_request,
            response=http_response,
        ),
    )
    return _ensure_download_url(
        response,
        download_filename=(
            _download_filename_for_action(FeatureType.explain, _uploaded_filename(file))
            if file is not None
            else None
        ),
    )


@router.post("/generate-questions", dependencies=[Depends(rate_limit_for_feature(FeatureType.generate_questions))])
def generate_questions_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile | None = File(default=None),
    text: str | None = Form(default=None, max_length=INLINE_TEXT_POLICY.max_chars),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    input_payload = _build_document_input(action=FeatureType.generate_questions, file=file, text=text)
    request = AnalyzerRequest(
        action=FeatureType.generate_questions,
        input=input_payload,
        payload=QuestionGenerationRequest(feature=FeatureType.generate_questions),
        policy=_policy_for_action(FeatureType.generate_questions),
        system_language=system_language,
    )
    response = _ensure_download_url(
        _run_request(request, **_artifact_owner_kwargs(current_user)),
        download_filename=(
            _download_filename_for_action(
                FeatureType.generate_questions,
                _uploaded_filename(file),
            )
            if file is not None
            else None
        ),
    )
    body = response.model_dump(mode="json")
    generated_questions_text = _generated_questions_text_from_response(response)
    if generated_questions_text:
        body["generated_questions_text"] = generated_questions_text
    return body


@router.post("/generate-answers", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.generate_answers))])
def generate_answers_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile | None = File(default=None),
    text: str | None = Form(default=None, max_length=INLINE_TEXT_POLICY.max_chars),
    questions_json: str = Form(..., max_length=AUXILIARY_PROMPT_POLICY.max_chars),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    input_payload = _build_document_input(action=FeatureType.generate_answers, file=file, text=text)
    questions = _parse_numbered_questions(questions_json)
    request = AnalyzerRequest(
        action=FeatureType.generate_answers,
        input=input_payload,
        payload=AnswerGenerationRequest(
            feature=FeatureType.generate_answers,
            questions=questions,
        ),
        policy=_policy_for_action(FeatureType.generate_answers),
        system_language=system_language,
    )
    response = _run_request(request, **_artifact_owner_kwargs(current_user))
    return _ensure_download_url(
        response,
        download_filename=(
            _download_filename_for_action(
                FeatureType.generate_answers,
                _uploaded_filename(file),
            )
            if file is not None
            else None
        ),
    )


# -----------------------------------------------------------------------------
# Privacy routes
# -----------------------------------------------------------------------------


@router.post("/redact/detect-document-type")
def redact_detect_document_type_route(
    _current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
) -> dict[str, object]:
    return _detect_privacy_document_type_upload(action=FeatureType.redact, file=file)


@router.post("/data-mask/detect-document-type")
def data_mask_detect_document_type_route(
    _current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
) -> dict[str, object]:
    return _detect_privacy_document_type_upload(action=FeatureType.data_mask, file=file)


@router.post("/redact/review", dependencies=[Depends(rate_limit_for_feature(FeatureType.redact))])
def redact_review_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    document_type: RedactionMaskingDocumentType | None = Form(default=None),
    target_data: list[SensitiveDataType] | None = Form(default=None),
    review_exclusions: list[str] | None = Form(default=None),
    custom_redactions: list[str] | None = Form(default=None),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    source_filename = _uploaded_filename(file)
    input_payload, request = _build_privacy_request(
        action=FeatureType.redact,
        file=file,
        document_type=document_type,
        target_data=target_data,
        review_exclusions=review_exclusions,
        system_language=system_language,
    )
    cleaned_custom_redactions = _clean_repeated_strings(custom_redactions)
    processed = _run_privacy_request(
        request,
        source_path=_privacy_source_path(input_payload),
        custom_redactions=cleaned_custom_redactions,
        **_artifact_owner_kwargs(current_user),
    )
    candidates = preview_redaction_candidates(
        request,
        project_id=_google_sdp_project_id(),
        location=DEFAULT_GOOGLE_SDP_LOCATION,
        custom_redactions=cleaned_custom_redactions,
    )
    # The redacted bytes returned here are a draft preview; /redact is the only
    # endpoint that publishes the final artifact after the user's review. The
    # artifact alias preserves the current page contract without exposing the
    # final analyzer response or a generated server path.
    preview_artifact = _serialize_review_preview_artifact(
        processed,
        download_filename=_download_filename_for_action(FeatureType.redact, source_filename),
        **_artifact_owner_kwargs(current_user),
    )
    return {
        "candidates": _serialize_candidates(candidates),
        "preview_artifact": preview_artifact,
        "artifact": preview_artifact,
    }


@router.post("/data-mask/review", dependencies=[Depends(rate_limit_for_feature(FeatureType.data_mask))])
def data_mask_review_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    document_type: RedactionMaskingDocumentType | None = Form(default=None),
    target_data: list[SensitiveDataType] | None = Form(default=None),
    review_exclusions: list[str] | None = Form(default=None),
    custom_redactions: list[str] | None = Form(default=None),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    source_filename = _uploaded_filename(file)
    input_payload, request = _build_privacy_request(
        action=FeatureType.data_mask,
        file=file,
        document_type=document_type,
        target_data=target_data,
        review_exclusions=review_exclusions,
        system_language=system_language,
    )
    cleaned_custom_redactions = _clean_repeated_strings(custom_redactions)
    processed = _run_privacy_request(
        request,
        source_path=_privacy_source_path(input_payload),
        custom_redactions=cleaned_custom_redactions,
        **_artifact_owner_kwargs(current_user),
    )
    candidates = preview_data_mask_candidates(
        request,
        project_id=_google_sdp_project_id(),
        location=DEFAULT_GOOGLE_SDP_LOCATION,
        custom_redactions=cleaned_custom_redactions,
    )
    return {
        **_serialize_processed_result(
            processed,
            download_filename=_download_filename_for_action(FeatureType.data_mask, source_filename),
            **_artifact_owner_kwargs(current_user),
        ),
        "candidates": _serialize_candidates(candidates),
    }


@router.post("/redact", dependencies=[Depends(rate_limit_for_feature(FeatureType.redact))])
def redact_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    document_type: RedactionMaskingDocumentType | None = Form(default=None),
    target_data: list[SensitiveDataType] | None = Form(default=None),
    review_exclusions: list[str] | None = Form(default=None),
    custom_redactions: list[str] | None = Form(default=None),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    source_filename = _uploaded_filename(file)
    input_payload, request = _build_privacy_request(
        action=FeatureType.redact,
        file=file,
        document_type=document_type,
        target_data=target_data,
        review_exclusions=review_exclusions,
        system_language=system_language,
    )
    processed = _run_privacy_request(
        request,
        source_path=_privacy_source_path(input_payload),
        custom_redactions=_clean_repeated_strings(custom_redactions),
        **_artifact_owner_kwargs(current_user),
    )
    return _serialize_processed_result(
        processed,
        download_filename=_download_filename_for_action(FeatureType.redact, source_filename),
        **_artifact_owner_kwargs(current_user),
    )


@router.post("/data-mask", dependencies=[Depends(rate_limit_for_feature(FeatureType.data_mask))])
def data_mask_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    document_type: RedactionMaskingDocumentType | None = Form(default=None),
    target_data: list[SensitiveDataType] | None = Form(default=None),
    review_exclusions: list[str] | None = Form(default=None),
    custom_redactions: list[str] | None = Form(default=None),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    source_filename = _uploaded_filename(file)
    input_payload, request = _build_privacy_request(
        action=FeatureType.data_mask,
        file=file,
        document_type=document_type,
        target_data=target_data,
        review_exclusions=review_exclusions,
        system_language=system_language,
    )
    processed = _run_privacy_request(
        request,
        source_path=_privacy_source_path(input_payload),
        custom_redactions=_clean_repeated_strings(custom_redactions),
        **_artifact_owner_kwargs(current_user),
    )
    return _serialize_processed_result(
        processed,
        download_filename=_download_filename_for_action(FeatureType.data_mask, source_filename),
        **_artifact_owner_kwargs(current_user),
    )


# -----------------------------------------------------------------------------
# Structured extraction / compliance routes
# -----------------------------------------------------------------------------


def _structured_extraction_uploads(
    *,
    file: UploadFile | None,
    files: list[UploadFile] | None,
) -> list[UploadFile]:
    file_list = list(files or [])

    if file is not None and file_list:
        raise _bad_request(
            "Submit structured extraction uploads under either 'files' or the legacy 'file' field, not both."
        )

    if file is not None:
        file_list = [file]

    if not file_list:
        raise _bad_request("At least one file is required for structured extraction.")

    file_list = _deduplicate_uploads(file_list)

    if len(file_list) > MAX_STRUCTURED_EXTRACTION_DOCUMENT_SET_FILES:
        raise _bad_request(
            f"Structured extraction accepts a maximum of {MAX_STRUCTURED_EXTRACTION_DOCUMENT_SET_FILES} unique files per document-set request."
        )

    return file_list


def _build_structured_extraction_input_payload(
    *,
    file: UploadFile | None,
    files: list[UploadFile] | None,
):
    uploads = _structured_extraction_uploads(file=file, files=files)

    try:
        documents = [
            build_uploaded_document_payload(
                action=FeatureType.structured_extract,
                upload=upload,
            )
            for upload in uploads
        ]
    except UploadServiceUnavailableError as exc:
        raise _service_unavailable(str(exc)) from exc
    except UploadError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc

    if len(documents) == 1:
        return documents[0]

    return DocumentSetPayload(documents=documents)



@router.post("/structured-extraction", dependencies=[Depends(rate_limit_for_feature(FeatureType.structured_extract))])
def structured_extraction_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] | None = File(default=None),
    file: UploadFile | None = File(default=None),
    document_classes: list[StructuredExtractionDocumentClass] | None = Form(default=None),
    selected_fields: list[str] | None = Form(default=None),
    output_format: StructuredDataOutputFormat = Form(StructuredDataOutputFormat.json),
    result_shape: StructuredExtractionResultShape = Form(StructuredExtractionResultShape.machine_readable),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    source_uploads = _structured_extraction_uploads(file=file, files=files)
    source_filenames = [_uploaded_filename(upload) for upload in source_uploads]
    input_payload = _build_structured_extraction_input_payload(
        file=None,
        files=source_uploads,
    )

    payload = StructuredExtractionRequest(
        feature=FeatureType.structured_extract,
        document_classes=document_classes or [],
        selected_fields=_clean_repeated_strings(selected_fields),
        output_format=output_format,
        result_shape=result_shape,
        allow_external_knowledge=False,
        require_human_review=True,
    )
    request = AnalyzerRequest(
        action=FeatureType.structured_extract,
        input=input_payload,
        payload=payload,
        policy=_policy_for_action(FeatureType.structured_extract),
        system_language=system_language,
    )
    output_extension = _normalized_extension(output_format, default=".json")
    download_filename = (
        f"structured_extraction{output_extension}"
        if len(source_filenames) > 1
        else _download_filename_for_action(
            FeatureType.structured_extract,
            source_filenames[0],
            output_extension=output_extension,
        )
    )
    return _run_structured_extraction_request_with_preview(
        request,
        download_filename=download_filename,
        **_artifact_owner_kwargs(current_user),
    )


@router.post("/compliance", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.compliance))])
def compliance_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile | None = File(default=None),
    files: list[UploadFile] | None = File(default=None),
    jurisdiction: ComplianceJurisdiction = Form(ComplianceJurisdiction.nigeria),
    sector_packs: list[ComplianceSectorPack] | None = Form(default=None),
    regulatory_domains: list[ComplianceRegulatoryDomain] | None = Form(default=None),
    report_variant: ComplianceReportVariant = Form(ComplianceReportVariant.human_readable_report),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    source_uploads = _compliance_uploads(file=file, files=files)
    source_filenames = [_uploaded_filename(upload) for upload in source_uploads]
    request = _build_compliance_request(
        file=None,
        files=source_uploads,
        jurisdiction=jurisdiction,
        sector_packs=sector_packs,
        regulatory_domains=regulatory_domains,
        report_variant=report_variant,
        system_language=system_language,
    )
    output_extension = _compliance_download_extension(
        report_variant,
        file_count=len(source_filenames),
    )
    download_filename = (
        f"compliance_report{output_extension}"
        if len(source_filenames) > 1
        else _download_filename_for_action(
            FeatureType.compliance,
            source_filenames[0],
            output_extension=output_extension,
        )
    )
    return _run_standalone_feature_request(
        request,
        download_filename=download_filename,
        **_artifact_owner_kwargs(current_user),
    )


@router.post("/compliance/set", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.compliance))])
def compliance_document_set_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    jurisdiction: ComplianceJurisdiction = Form(ComplianceJurisdiction.nigeria),
    sector_packs: list[ComplianceSectorPack] | None = Form(default=None),
    regulatory_domains: list[ComplianceRegulatoryDomain] | None = Form(default=None),
    report_variant: ComplianceReportVariant = Form(ComplianceReportVariant.human_readable_report),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    """
    Multi-document Compliance route.

    This is not the paid-plan batch-processing feature. It creates one
    DocumentSetPayload, capped at MAX_COMPLIANCE_DOCUMENT_SET_FILES files.
    """
    request = _build_compliance_request(
        file=None,
        files=files,
        jurisdiction=jurisdiction,
        sector_packs=sector_packs,
        regulatory_domains=regulatory_domains,
        report_variant=report_variant,
        system_language=system_language,
    )
    output_extension = _compliance_download_extension(
        report_variant,
        file_count=len(files),
    )
    return _run_standalone_feature_request(
        request,
        download_filename=f"compliance_report{output_extension}",
        **_artifact_owner_kwargs(current_user),
    )


@router.post("/compliance/preview", dependencies=[Depends(rate_limit_for_feature(FeatureType.compliance))])
def compliance_preview_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile | None = File(default=None),
    files: list[UploadFile] | None = File(default=None),
    jurisdiction: ComplianceJurisdiction = Form(ComplianceJurisdiction.nigeria),
    sector_packs: list[ComplianceSectorPack] | None = Form(default=None),
    regulatory_domains: list[ComplianceRegulatoryDomain] | None = Form(default=None),
    report_variant: ComplianceReportVariant = Form(ComplianceReportVariant.human_readable_report),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> dict[str, Any]:
    request = _build_compliance_request(
        file=file,
        files=files,
        jurisdiction=jurisdiction,
        sector_packs=sector_packs,
        regulatory_domains=regulatory_domains,
        report_variant=report_variant,
        system_language=system_language,
    )
    try:
        with artifact_owner_context(
            str(current_user.user_id),
            organization_id=_user_organization_id(current_user),
            feature=FeatureType.compliance.value,
        ), upload_processing_session(request):
            preview = workflow_router.preview_compliance(request)
        report = preview.report.model_dump(mode="json")
        return {
            "preview_markdown": preview.preview_markdown,
            "report": report,
            "counts": report.get("counts"),
            "rule_results": report.get("rule_results", []),
            "human_review": preview.human_review.model_dump(mode="json"),
        }
    except HTTPException:
        raise
    except RuleRegistryError as exc:
        raise _bad_request(str(exc)) from exc
    except ValueError as exc:
        raise _bad_request(str(exc)) from exc
    except RuntimeError as exc:
        raise _service_unavailable(str(exc)) from exc


# -----------------------------------------------------------------------------
# PDF tools routes
# -----------------------------------------------------------------------------


@router.post("/pdf/combine", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.combine_pdf))])
def combine_pdf_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    files: list[UploadFile] = File(...),
    output_filename: str = Form("combined-document.pdf"),
    preserve_bookmarks: bool = Form(True),
    preserve_metadata: bool = Form(False),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    """Combine multiple PDFs into one output document.

    This is one PDF-tool operation whose input inherently contains multiple
    documents. It is not the paid-plan batch-processing feature.
    """
    del output_filename
    files = _deduplicate_uploads(files)
    _validate_pdf_combine_files(files)
    resolved_output_filename = _download_filename_for_action(
        FeatureType.combine_pdf,
        _uploaded_filename(files[0]),
    )
    input_payload = _build_pdf_set_input(FeatureType.combine_pdf, files)
    request = AnalyzerRequest(
        action=FeatureType.combine_pdf,
        input=input_payload,
        payload=CombinePdfRequest(
            feature=FeatureType.combine_pdf,
            output_filename=resolved_output_filename,
            preserve_bookmarks=preserve_bookmarks,
            preserve_metadata=preserve_metadata,
        ),
        policy=_policy_for_action(FeatureType.combine_pdf),
        system_language=system_language,
    )
    return _ensure_download_url(
        _run_request(request, **_artifact_owner_kwargs(current_user)),
        download_filename=resolved_output_filename,
    )


@router.post("/pdf/split", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.split_pdf))])
def split_pdf_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    mode: PdfSplitMode = Form(...),
    selected_pages: str | None = Form(default=None),
    page_ranges: str | None = Form(default=None),
    output_basename: str = Form("split-document"),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    del output_basename
    source_filename = _uploaded_filename(file)
    resolved_output_basename = f"{_filename_stem(source_filename)}.split"
    try:
        payload = SplitPdfRequest(
            feature=FeatureType.split_pdf,
            mode=mode,
            selected_pages=_parse_int_list(selected_pages),
            page_ranges=_parse_page_ranges(page_ranges),
            output_basename=resolved_output_basename,
        )
    except HTTPException:
        raise
    except (ValidationError, TypeError, ValueError) as exc:
        raise _bad_request(f"Invalid PDF split request: {exc}") from exc

    input_payload = _build_single_pdf_input(FeatureType.split_pdf, file)
    try:
        request = AnalyzerRequest(
            action=FeatureType.split_pdf,
            input=input_payload,
            payload=payload,
            policy=_policy_for_action(FeatureType.split_pdf),
            system_language=system_language,
        )
    except (ValidationError, TypeError, ValueError) as exc:
        raise _bad_request(f"Invalid PDF split request: {exc}") from exc
    return _apply_split_download_filenames(
        _run_request(request, **_artifact_owner_kwargs(current_user)),
        source_filename,
    )


@router.post("/pdf/edit", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.edit_pdf))])
def edit_pdf_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    edit_assets: list[UploadFile] = File(default=[]),
    operations_json: str = Form(...),
    output_filename: str = Form("edited-document.pdf"),
    generate_preview: bool = Form(True),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    del output_filename
    resolved_output_filename = _download_filename_for_action(
        FeatureType.edit_pdf,
        _uploaded_filename(file),
    )
    input_payload = _build_single_pdf_input(FeatureType.edit_pdf, file)
    source_path_value = (
        getattr(input_payload, "storage_key", None)
        or getattr(input_payload, "filename", None)
    )
    source_path = Path(source_path_value).expanduser().resolve() if source_path_value else None
    asset_paths: dict[str, str] = {}
    operation_succeeded = False
    try:
        asset_paths = _save_pdf_edit_assets(edit_assets)
        payload = EditPdfRequest(
            feature=FeatureType.edit_pdf,
            operations=_parse_edit_operations(operations_json, asset_paths=asset_paths),
            output_filename=resolved_output_filename,
            generate_preview=generate_preview,
        )
        request = AnalyzerRequest(
            action=FeatureType.edit_pdf,
            input=input_payload,
            payload=payload,
            policy=_policy_for_action(FeatureType.edit_pdf),
            system_language=system_language,
        )
        response = _run_request(
            request,
            artifact_owner_user_id=str(current_user.user_id),
            artifact_owner_organization_id=_user_organization_id(current_user),
        )
        preview = getattr(response.result, "preview", None)
        preview_filename = getattr(preview, "filename", None)
        response = _ensure_download_url(
            response,
            download_filename=resolved_output_filename,
        )
        if preview is not None and preview_filename:
            # The generic filename helper walks nested artifacts. Restore the
            # preview's distinct name so downloading it cannot overwrite or be
            # confused with the final edited PDF.
            _apply_download_filename(
                preview,
                _safe_download_filename(preview_filename),
            )
        operation_succeeded = True
        return response
    except HTTPException:
        raise
    except (ValidationError, TypeError, ValueError) as exc:
        raise _bad_request(f"Invalid PDF edit request: {exc}") from exc
    finally:
        del source_path
        mark_upload_paths_processed(
            asset_paths.values(),
            success=operation_succeeded,
        )


@router.post("/pdf/compress", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.compress_pdf))])
def compress_pdf_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    compression_level: PdfCompressionLevel = Form(PdfCompressionLevel.balanced),
    output_filename: str = Form("compressed-document.pdf"),
    async_processing: bool = Form(True),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    del output_filename
    resolved_output_filename = _download_filename_for_action(
        FeatureType.compress_pdf,
        _uploaded_filename(file),
    )
    input_payload = _build_single_pdf_input(FeatureType.compress_pdf, file)
    request = AnalyzerRequest(
        action=FeatureType.compress_pdf,
        input=input_payload,
        payload=CompressPdfRequest(
            feature=FeatureType.compress_pdf,
            compression_level=compression_level,
            output_filename=resolved_output_filename,
            async_processing=async_processing,
        ),
        policy=_policy_for_action(FeatureType.compress_pdf),
        system_language=system_language,
    )
    response = _run_request(
        request,
        pdf_job_owner_id=str(current_user.user_id),
        **_artifact_owner_kwargs(current_user),
    )
    return _ensure_download_url(
        response,
        download_filename=resolved_output_filename,
    )


@router.get("/pdf/compress/jobs/{job_id}", response_model=PdfJobResult)
def compress_pdf_job_status_route(
    job_id: str,
    current_user: AuthenticatedUser = Depends(get_current_user),
) -> PdfJobResult:
    """Return an authenticated user's queued PDF compression job status."""
    try:
        return workflow_router.get_pdf_compression_job(
            job_id=job_id,
            owner_id=str(current_user.user_id),
        )
    except LookupError as exc:
        raise HTTPException(
            status_code=404,
            detail={
                "error": "compression_job_not_found",
                "message": "Compression job was not found.",
            },
        ) from exc


# -----------------------------------------------------------------------------
# E-signature route
# -----------------------------------------------------------------------------


class RecipientSigningSubmission(BaseModel):
    signature: AddSignatureOperation
    field_values: dict[str, str] = Field(default_factory=dict)


def _esignature_public_base_url() -> str:
    value = (
        os.getenv("ESIGN_SIGNING_BASE_URL", "").strip()
        or os.getenv("APP_BASE_URL", "").strip()
        or os.getenv("FRONTEND_URL", "").strip()
        or os.getenv("NEXT_PUBLIC_APP_URL", "").strip()
    ).rstrip("/")
    if not value.startswith(("http://", "https://")):
        raise RuntimeError(
            "Set ESIGN_SIGNING_BASE_URL (or APP_BASE_URL) to the public ReDOCX URL."
        )
    if _is_production_environment() and not value.startswith("https://"):
        raise RuntimeError("Production e-signature links require an HTTPS public URL.")
    return value


def _build_esignature_service(
    conn: Any,
    *,
    require_email: bool,
    require_token_access: bool = False,
) -> ESignatureService:
    if getattr(conn, "autocommit", False):
        raise RuntimeError(
            "E-signature persistence requires a transactional PostgreSQL connection."
        )
    token_secret = os.getenv("ESIGN_TOKEN_PEPPER", "").strip() or None
    if (require_email or require_token_access) and token_secret is None:
        raise RuntimeError("ESIGN_TOKEN_PEPPER is required for recipient signing links.")
    if token_secret is not None and len(token_secret) < 32:
        raise RuntimeError("ESIGN_TOKEN_PEPPER must contain at least 32 characters.")

    signing_base_url = _esignature_public_base_url() if require_email else None
    email_client = build_default_email_client() if require_email else None
    if (
        require_email
        and _is_production_environment()
        and isinstance(email_client, ConsoleEmailClient)
    ):
        raise RuntimeError(
            "Production e-signature email delivery is not configured. "
            "Set EMAIL_PROVIDER to zeptomail or smtp and configure that provider."
        )

    artifact_dir = Path(
        os.getenv("ESIGNATURE_ARTIFACT_STORAGE_DIR", "artifacts/esignature")
    ).expanduser()

    # E-signature source PDFs must remain available for the entire lifetime of a
    # valid signing link. The schema permits envelopes to live for up to 180 days,
    # while the generic artifact store defaults to only 30 minutes. Fail closed if
    # deployment configuration would make a still-valid signing link lose its PDF.
    try:
        esignature_retention_days = int(
            os.getenv("ESIGNATURE_ARTIFACT_RETENTION_DAYS", "180")
        )
    except ValueError as exc:
        raise RuntimeError(
            "ESIGNATURE_ARTIFACT_RETENTION_DAYS must be an integer."
        ) from exc
    if esignature_retention_days < 180:
        raise RuntimeError(
            "ESIGNATURE_ARTIFACT_RETENTION_DAYS must be at least 180 so artifacts "
            "outlive every valid e-signature envelope."
        )

    storage = LocalArtifactStorage(
        base_dir=str(artifact_dir),
        retention_minutes=esignature_retention_days * 24 * 60,
    )

    def resolve_esignature_source(payload):
        # Initial requests carry the server-side upload path in storage_key. After
        # ESignatureService persists the source, storage_key becomes a relative
        # LocalArtifactStorage key. Support both forms so the same resolver works
        # before and after durable persistence.
        raw_key = str(getattr(payload, "storage_key", "") or "").strip()
        if raw_key:
            candidate = Path(raw_key).expanduser()
            if candidate.exists() and candidate.is_file():
                return candidate.resolve()
            return storage.resolve_storage_key(raw_key)
        raise FileNotFoundError("The e-signature PDF source could not be resolved.")

    envelope_repository = PostgresEnvelopeRepository(conn)
    token_repository = PostgresSigningTokenRepository(conn)

    return ESignatureService(
        config=ESignatureServiceConfig(
            algorithm_version="esignature-service-v1.2.0",
            signed_artifacts_dir=str(artifact_dir / "work" / "signed"),
            preview_artifacts_dir=str(artifact_dir / "work" / "previews"),
            certificate_artifacts_dir=str(artifact_dir / "work" / "certificates"),
            signing_base_url=signing_base_url,
            token_secret=token_secret,
            send_completion_emails=True,
        ),
        storage_backend=storage,
        source_path_resolver=resolve_esignature_source,
        asset_path_resolver=storage.resolve_storage_key,
        email_client=email_client,
        envelope_repository=envelope_repository,
        token_repository=token_repository,
        download_url_builder=_download_url_for_storage_key,
    )


def _recipient_token(value: str | None) -> str:
    token = str(value or "").strip()
    if not re.fullmatch(r"[A-Za-z0-9_-]{43,256}", token):
        raise HTTPException(
            status_code=404,
            detail={
                "error": "signing_link_invalid",
                "message": "This signing link is invalid or no longer available.",
            },
        )
    return token


def _recipient_http_exception(exc: Exception) -> HTTPException:
    message = str(exc).lower()
    if isinstance(exc, KeyError) or any(
        token in message for token in ("expired", "revoked", "already been used")
    ):
        return HTTPException(
            status_code=410,
            detail={
                "error": "signing_link_expired",
                "message": "This signing link has expired or has already been used.",
            },
        )
    if isinstance(exc, ValueError):
        return HTTPException(
            status_code=409,
            detail={"error": "signing_not_available", "message": str(exc)},
        )
    return HTTPException(
        status_code=503,
        detail={
            "error": "signing_service_unavailable",
            "message": "The signing service is temporarily unavailable.",
        },
    )


def _esignature_owner_state_for_token(
    service: ESignatureService,
    raw_token: str,
    *,
    for_update: bool = False,
):
    """Resolve artifact ownership before touching owner-scoped storage."""
    if service.token_repository is None or service.envelope_repository is None:
        raise RuntimeError("E-signature persistence is not configured.")
    token = service.token_repository.get_valid_for_raw_token(
        raw_token,
        secret=service.config.token_secret,
        for_update=for_update,
    )
    state = (
        service.envelope_repository.get_for_update(token.envelope_id)
        if for_update
        else service.envelope_repository.get(token.envelope_id)
    )
    if not state.owner_user_id:
        raise RuntimeError("The envelope is missing its artifact owner identity.")
    return state


@router.post(
    "/e-signature/layout",
    dependencies=[Depends(rate_limit_for_feature(FeatureType.e_signature))],
)
def esignature_layout_route(
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    fields_json: str = Form("[]"),
    signers_json: str = Form("[]"),
    page_number: int = Form(1),
    add_signature_page: bool = Form(False),
    detect_signature_lines: bool = Form(False),
):
    """Analyze one sender-selected PDF page and validate all proposed fields.

    The route is authenticated and stateless: the uploaded PDF is security-scanned,
    analyzed locally, and removed after the response is built. Suggestions are
    advisory; the main e-signature service repeats collision validation before send.
    """
    if len(fields_json) > 512_000 or len(signers_json) > 128_000:
        raise _bad_request("E-signature layout metadata is too large.")

    fields = _loads_json(fields_json, default=[])
    signers = _loads_json(signers_json, default=[])
    if not isinstance(fields, list) or len(fields) > 250:
        raise _bad_request("fields_json must be a JSON array with at most 250 fields.")
    if not isinstance(signers, list) or len(signers) > 25:
        raise _bad_request("signers_json must be a JSON array with at most 25 signers.")
    if any(not isinstance(item, dict) for item in fields):
        raise _bad_request("Every layout field must be a JSON object.")
    if any(not isinstance(item, dict) for item in signers):
        raise _bad_request("Every layout signer must be a JSON object.")

    source_path = _save_upload_to_disk(
        file,
        subdir="e-signature-layout",
        default_name="document.pdf",
    )
    succeeded = False
    try:
        result = analyze_esignature_pdf(
            source_path,
            fields=fields,
            preview_page_number=page_number,
            add_signature_page=add_signature_page,
            signers=signers,
            detect_signature_lines=detect_signature_lines,
        )
        succeeded = True
        return JSONResponse(
            result,
            headers={
                "Cache-Control": "private, no-store, max-age=0",
                "Pragma": "no-cache",
                "Referrer-Policy": "no-referrer",
                "X-Content-Type-Options": "nosniff",
            },
        )
    except HTTPException:
        raise
    except (TypeError, ValueError, FileNotFoundError) as exc:
        raise _bad_request(f"Could not analyze e-signature layout: {exc}") from exc
    except RuntimeError as exc:
        raise _service_unavailable(str(exc)) from exc
    finally:
        mark_upload_paths_processed([source_path], success=succeeded)
        Path(source_path).unlink(missing_ok=True)


@router.post("/e-signature", response_model=AnalyzerResponse, dependencies=[Depends(rate_limit_for_feature(FeatureType.e_signature))])
def esignature_route(
    http_request: Request,
    current_user: AuthenticatedUser = Depends(get_current_user),
    file: UploadFile = File(...),
    signature_assets: list[UploadFile] = File(default=[]),
    payload_json: str = Form(...),
    signer_email: str | None = Form(default=None),
    signer_signature_json: str | None = Form(default=None),
    send_emails: bool = Form(True),
    system_language: SystemLanguage = Form(SystemLanguage.english),
) -> AnalyzerResponse:
    source_filename = _uploaded_filename(file)
    input_payload = _build_single_pdf_input(FeatureType.e_signature, file)
    asset_paths: dict[str, str] = {}
    operation_succeeded = False
    try:
        asset_paths = _save_pdf_edit_assets(signature_assets)
        payload = _resolve_esignature_signature_assets(
            _parse_esignature_request(payload_json),
            asset_paths=asset_paths,
        )
        authenticated_email = _user_email(current_user)
        if authenticated_email is None:
            raise _bad_request(
                "Your authenticated account must provide an email address for e-signature."
            )
        if (
            payload.self_signer is not None
            and payload.self_signer.email.strip().lower()
            != authenticated_email.strip().lower()
        ):
            raise _bad_request(
                "The self-signer email must match your authenticated account email."
            )
        if payload.workflow.value == "self_sign":
            allowed_actions = {
                ESignatureAction.create_draft,
                ESignatureAction.complete_signing,
            }
        else:
            allowed_actions = {
                ESignatureAction.create_draft,
                ESignatureAction.send,
            }
        if payload.action not in allowed_actions:
            raise _bad_request(
                "External recipients must sign through their token-protected email link."
            )
        if (
            signer_email
            and signer_email.strip().lower() != authenticated_email.strip().lower()
        ):
            raise _bad_request(
                "The submitted signer email must match your authenticated account email."
            )
        supplied_signature = _parse_optional_signature(signer_signature_json)
        signer_signature = (
            payload.self_signer.signature
            if payload.self_signer is not None and payload.self_signer.signature is not None
            else supplied_signature
        )

        request = AnalyzerRequest(
            action=FeatureType.e_signature,
            input=input_payload,
            payload=payload,
            policy=_policy_for_action(FeatureType.e_signature),
            system_language=system_language,
        )
        requires_email = payload.action == ESignatureAction.send and bool(payload.recipients)
        if requires_email and not send_emails:
            raise _bad_request(
                "Recipient workflows require invitation email delivery."
            )

        with get_db() as conn:
            esignature_service = _build_esignature_service(
                conn,
                require_email=requires_email,
            )
            dispatcher = WorkflowRouter(
                esignature_service=esignature_service,
                download_url_builder=_download_url_for_storage_key,
            )
            response = _run_request(
                request,
                workflow_router_override=dispatcher,
                signer_email=signer_email,
                signer_signature=signer_signature,
                sender_email=authenticated_email,
                sender_name=_user_name(current_user),
                send_emails=send_emails,
                ip_address=_client_ip(http_request),
                user_agent=_user_agent(http_request),
                **_artifact_owner_kwargs(current_user),
            )
            persisted_state = esignature_service.envelope_repository.get(
                response.result.envelope_id
            )
            esignature_service.envelope_repository.save(
                replace(
                    persisted_state,
                    owner_user_id=str(current_user.user_id),
                    owner_organization_id=_user_organization_id(current_user),
                )
            )
        operation_succeeded = True
    except HTTPException:
        raise
    except (ValidationError, TypeError, ValueError) as exc:
        raise _bad_request(f"Invalid e-signature request: {exc}") from exc
    except RuntimeError as exc:
        raise _service_unavailable(str(exc)) from exc
    finally:
        mark_upload_paths_processed(
            asset_paths.values(),
            success=operation_succeeded,
        )
    return _apply_esignature_download_filenames(
        _ensure_download_url(response),
        source_filename,
    )


@router.get("/e-signature/recipient")
def esignature_recipient_context_route(
    http_request: Request,
    x_redocx_signing_token: str | None = Header(
        default=None,
        alias="X-ReDOCX-Signing-Token",
    ),
):
    raw_token = _recipient_token(x_redocx_signing_token)
    try:
        with get_db() as conn:
            service = _build_esignature_service(
                conn,
                require_email=False,
                require_token_access=True,
            )
            owner_state = _esignature_owner_state_for_token(service, raw_token)
            with artifact_owner_context(
                owner_state.owner_user_id,
                organization_id=owner_state.owner_organization_id,
                feature=FeatureType.e_signature.value,
            ):
                session = service.get_recipient_session(
                    raw_token,
                    mark_viewed_event=True,
                    ip_address=_client_ip(http_request),
                    user_agent=_user_agent(http_request),
                )
            source_request = session.state.source_request
            filename = (
                source_request.input.filename
                if source_request is not None
                and hasattr(source_request.input, "filename")
                else "document.pdf"
            )
            return JSONResponse(
                {
                    "envelope_id": session.state.envelope_id,
                    "workflow": session.state.workflow.value,
                    "status": session.state.status.value,
                    "document_filename": filename,
                    "signer": session.signer.model_dump(mode="json"),
                    "fields": [
                        field.model_dump(mode="json")
                        for field in session.fields
                    ],
                },
                headers={
                    "Cache-Control": "private, no-store, max-age=0",
                    "Pragma": "no-cache",
                    "Referrer-Policy": "no-referrer",
                    "X-Content-Type-Options": "nosniff",
                },
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise _recipient_http_exception(exc) from exc


@router.get("/e-signature/recipient/document")
def esignature_recipient_document_route(
    http_request: Request,
    x_redocx_signing_token: str | None = Header(
        default=None,
        alias="X-ReDOCX-Signing-Token",
    ),
):
    raw_token = _recipient_token(x_redocx_signing_token)
    try:
        with get_db() as conn:
            service = _build_esignature_service(
                conn,
                require_email=False,
                require_token_access=True,
            )
            owner_state = _esignature_owner_state_for_token(service, raw_token)
            with artifact_owner_context(
                owner_state.owner_user_id,
                organization_id=owner_state.owner_organization_id,
                feature=FeatureType.e_signature.value,
            ):
                session = service.get_recipient_session(
                    raw_token,
                    mark_viewed_event=False,
                    ip_address=_client_ip(http_request),
                    user_agent=_user_agent(http_request),
                )
            response = FileResponse(
                path=session.current_pdf_path,
                media_type="application/pdf",
                filename="document-to-sign.pdf",
            )
            response.headers["Content-Disposition"] = (
                'inline; filename="document-to-sign.pdf"'
            )
            response.headers["Cache-Control"] = "private, no-store, max-age=0"
            response.headers["Pragma"] = "no-cache"
            response.headers["X-Content-Type-Options"] = "nosniff"
            response.headers["Content-Security-Policy"] = "sandbox"
            response.headers["Referrer-Policy"] = "no-referrer"
            return response
    except HTTPException:
        raise
    except Exception as exc:
        raise _recipient_http_exception(exc) from exc


@router.post("/e-signature/recipient", response_model=AnalyzerResponse)
def esignature_recipient_sign_route(
    submission: RecipientSigningSubmission,
    http_request: Request,
    x_redocx_signing_token: str | None = Header(
        default=None,
        alias="X-ReDOCX-Signing-Token",
    ),
) -> AnalyzerResponse:
    raw_token = _recipient_token(x_redocx_signing_token)
    if submission.signature.signature_type != SignatureRepresentationType.typed:
        raise _bad_request("Recipient signing currently accepts typed signatures only.")
    if len(submission.field_values) > 250:
        raise _bad_request("At most 250 e-signature field values are allowed.")
    for key, value in submission.field_values.items():
        if len(str(key)) > 256 or len(str(value)) > 10_000:
            raise _bad_request("An e-signature field value exceeds the supported size limit.")

    try:
        with get_db() as conn:
            service = _build_esignature_service(
                conn,
                require_email=True,
                require_token_access=True,
            )
            locked_state = _esignature_owner_state_for_token(
                service,
                raw_token,
                for_update=True,
            )
            with artifact_owner_context(
                locked_state.owner_user_id,
                organization_id=locked_state.owner_organization_id,
                feature=FeatureType.e_signature.value,
            ):
                response = service.sign_recipient(
                    raw_token,
                    signature=submission.signature,
                    field_values=submission.field_values,
                    ip_address=_client_ip(http_request),
                    user_agent=_user_agent(http_request),
                )
            return _ensure_download_url(response)
    except HTTPException:
        raise
    except Exception as exc:
        raise _recipient_http_exception(exc) from exc


@router.get("/e-signature/completed")
def esignature_completed_context_route(
    x_redocx_signing_token: str | None = Header(
        default=None,
        alias="X-ReDOCX-Signing-Token",
    ),
):
    raw_token = _recipient_token(x_redocx_signing_token)
    try:
        with get_db() as conn:
            service = _build_esignature_service(
                conn,
                require_email=False,
                require_token_access=True,
            )
            owner_state = _esignature_owner_state_for_token(service, raw_token)
            with artifact_owner_context(
                owner_state.owner_user_id,
                organization_id=owner_state.owner_organization_id,
                feature=FeatureType.e_signature.value,
            ):
                session = service.get_completed_session(raw_token)
            return JSONResponse(
                {
                    "envelope_id": session.state.envelope_id,
                    "status": session.state.status.value,
                    "document_filename": session.document_filename,
                },
                headers={
                    "Cache-Control": "private, no-store, max-age=0",
                    "Pragma": "no-cache",
                    "Referrer-Policy": "no-referrer",
                    "X-Content-Type-Options": "nosniff",
                },
            )
    except HTTPException:
        raise
    except Exception as exc:
        raise _recipient_http_exception(exc) from exc


@router.get("/e-signature/completed/document")
def esignature_completed_document_route(
    artifact: Literal["signed_pdf", "certificate"] = "signed_pdf",
    x_redocx_signing_token: str | None = Header(
        default=None,
        alias="X-ReDOCX-Signing-Token",
    ),
):
    raw_token = _recipient_token(x_redocx_signing_token)
    try:
        with get_db() as conn:
            service = _build_esignature_service(
                conn,
                require_email=False,
                require_token_access=True,
            )
            owner_state = _esignature_owner_state_for_token(service, raw_token)
            with artifact_owner_context(
                owner_state.owner_user_id,
                organization_id=owner_state.owner_organization_id,
                feature=FeatureType.e_signature.value,
            ):
                session = service.get_completed_session(raw_token)
            if artifact == "certificate":
                path = session.certificate_path
                filename = f"{Path(session.document_filename).stem}-certificate.pdf"
            else:
                path = session.signed_pdf_path
                filename = f"{Path(session.document_filename).stem}-signed.pdf"

            response = FileResponse(
                path=path,
                media_type="application/pdf",
                filename=_safe_download_filename(filename, default="signed-document.pdf"),
            )
            response.headers["Cache-Control"] = "private, no-store, max-age=0"
            response.headers["Pragma"] = "no-cache"
            response.headers["X-Content-Type-Options"] = "nosniff"
            response.headers["Content-Security-Policy"] = "sandbox"
            response.headers["Referrer-Policy"] = "no-referrer"
            return response
    except HTTPException:
        raise
    except Exception as exc:
        raise _recipient_http_exception(exc) from exc


# -----------------------------------------------------------------------------
# Artifact download route
# -----------------------------------------------------------------------------

OFFICE_PRINT_PREVIEW_EXTENSIONS = {
    ".docx",
    ".xlsx",
    ".xls",
    ".ods",
    ".pptx",
}
MAX_SPREADSHEET_PRINT_CELLS = 250_000


def _print_preview_headers() -> dict[str, str]:
    return {
        "Cache-Control": "private, no-store, max-age=0",
        "Pragma": "no-cache",
        "X-Content-Type-Options": "nosniff",
        "Content-Security-Policy": "sandbox",
        "Cross-Origin-Resource-Policy": "same-origin",
        "Referrer-Policy": "no-referrer",
    }


def _print_preview_html_document(
    *,
    title: str,
    body: str,
    landscape: bool = False,
) -> str:
    page_size = "A4 landscape" if landscape else "A4 portrait"
    return f"""<!doctype html>
<html lang="en">
  <head>
    <meta charset="utf-8">
    <meta name="viewport" content="width=device-width, initial-scale=1">
    <title>{html.escape(title)}</title>
    <style>
      @page {{ size: {page_size}; margin: 14mm; }}
      * {{ box-sizing: border-box; }}
      body {{
        margin: 0;
        color: #111827;
        background: #ffffff;
        font: 11pt/1.45 Arial, Helvetica, sans-serif;
        overflow-wrap: anywhere;
      }}
      h1, h2, h3, h4, h5, h6 {{ break-after: avoid; line-height: 1.2; }}
      p {{ margin: 0 0 0.65em; white-space: pre-wrap; }}
      table {{ width: 100%; border-collapse: collapse; table-layout: auto; }}
      thead {{ display: table-header-group; }}
      tr {{ break-inside: avoid; }}
      th, td {{
        border: 1px solid #9ca3af;
        padding: 5px 7px;
        text-align: left;
        vertical-align: top;
        white-space: pre-wrap;
      }}
      th {{ background: #f3f4f6; font-weight: 700; }}
      .document-title {{ margin: 0 0 12mm; font-size: 16pt; }}
      .sheet {{ break-before: page; }}
      .sheet:first-of-type {{ break-before: auto; }}
      .sheet-title {{ margin: 0 0 5mm; font-size: 14pt; }}
      .empty {{ color: #6b7280; font-style: italic; }}
    </style>
  </head>
  <body>
    {body}
  </body>
</html>"""


def _docx_print_fallback(path: Path, display_name: str) -> str:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as exc:  # pragma: no cover - deployment dependency
        raise RuntimeError(
            "DOCX print preview requires LibreOffice or python-docx."
        ) from exc

    document = docx.Document(str(path))
    fragments = [
        f'<h1 class="document-title">{html.escape(display_name)}</h1>'
    ]

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            paragraph = Paragraph(child, document)
            text = html.escape(paragraph.text).replace("\n", "<br>")
            style_name = str(getattr(paragraph.style, "name", "") or "").lower()
            heading_match = re.match(r"heading\s+([1-6])", style_name)
            if heading_match and text:
                level = heading_match.group(1)
                fragments.append(f"<h{level}>{text}</h{level}>")
            elif text:
                fragments.append(f"<p>{text}</p>")
        elif child.tag.endswith("}tbl"):
            table = Table(child, document)
            rows = []
            for row_index, row in enumerate(table.rows):
                cell_tag = "th" if row_index == 0 else "td"
                cells = "".join(
                    f"<{cell_tag}>"
                    f"{html.escape(cell.text).replace(chr(10), '<br>')}"
                    f"</{cell_tag}>"
                    for cell in row.cells
                )
                rows.append(f"<tr>{cells}</tr>")
            if rows:
                fragments.append(
                    "<table><thead>"
                    f"{rows[0]}"
                    "</thead><tbody>"
                    f"{''.join(rows[1:])}"
                    "</tbody></table>"
                )

    if len(fragments) == 1:
        fragments.append('<p class="empty">This document has no printable text.</p>')
    return _print_preview_html_document(
        title=display_name,
        body="\n".join(fragments),
    )


def _spreadsheet_print_fallback(path: Path, display_name: str) -> str:
    try:
        import openpyxl
    except ImportError as exc:  # pragma: no cover - deployment dependency
        raise RuntimeError(
            "Spreadsheet print preview requires LibreOffice or openpyxl."
        ) from exc

    workbook = openpyxl.load_workbook(
        filename=str(path),
        read_only=True,
        data_only=True,
    )
    fragments: list[str] = []
    rendered_cells = 0

    try:
        for sheet in workbook.worksheets:
            rows = []
            for row in sheet.iter_rows(values_only=True):
                values = list(row)
                while values and values[-1] is None:
                    values.pop()
                if not values:
                    continue

                rendered_cells += len(values)
                if rendered_cells > MAX_SPREADSHEET_PRINT_CELLS:
                    raise RuntimeError(
                        "This workbook is too large to prepare safely for browser printing."
                    )

                cells = "".join(
                    "<td>"
                    f"{html.escape(str(value) if value is not None else '').replace(chr(10), '<br>')}"
                    "</td>"
                    for value in values
                )
                rows.append(f"<tr>{cells}</tr>")

            sheet_title = html.escape(str(sheet.title or "Sheet"))
            if rows:
                first_row = rows[0].replace("<td>", "<th>").replace("</td>", "</th>")
                table = (
                    f'<section class="sheet"><h2 class="sheet-title">{sheet_title}</h2>'
                    f"<table><thead>{first_row}</thead>"
                    f"<tbody>{''.join(rows[1:])}</tbody></table></section>"
                )
            else:
                table = (
                    f'<section class="sheet"><h2 class="sheet-title">{sheet_title}</h2>'
                    '<p class="empty">This worksheet is empty.</p></section>'
                )
            fragments.append(table)
    finally:
        workbook.close()

    return _print_preview_html_document(
        title=display_name,
        body="\n".join(fragments),
        landscape=True,
    )


def _office_print_fallback(path: Path, display_name: str) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_print_fallback(path, display_name)
    if suffix == ".xlsx":
        return _spreadsheet_print_fallback(path, display_name)
    raise RuntimeError(
        f"{suffix or 'This Office format'} requires the document conversion service "
        "for print preview."
    )


def _build_office_print_preview(
    *,
    path: Path,
    display_name: str,
    owner_user_id: str,
    owner_organization_id: str | None,
) -> tuple[Path | None, str]:
    suffix = path.suffix.lower()
    try:
        with artifact_owner_context(
            owner_user_id,
            organization_id=owner_organization_id,
            feature="print_preview",
        ):
            preview = convert_document(
                input_format=suffix.removeprefix("."),
                output_format="pdf",
                source_reference=str(path),
                source_name_hint=display_name,
            )
        preview_path = _artifact_path_from_result(preview)
        if preview_path is None:
            raise RuntimeError("The generated print preview could not be resolved.")
        if preview_path.suffix.lower() != ".pdf":
            raise RuntimeError("The document conversion service did not return a PDF.")
        preview_name = f"{_filename_stem(display_name, default='document')}.pdf"
        return preview_path, preview_name
    except Exception:
        logger.exception(
            "Office-to-PDF print preview failed; using the safe HTML fallback.",
            extra={"artifact_suffix": suffix},
        )

    return None, _office_print_fallback(path, display_name)


@router.api_route("/artifacts/{storage_key:path}", methods=["GET", "HEAD"])
def download_artifact(
    request: Request,
    storage_key: str,
    disposition: str = "attachment",
    download_name: str | None = None,
    print_preview: bool = False,
    current_user: AuthenticatedUser | None = Depends(get_current_user_optional),
):
    requested_disposition = disposition.strip().lower()
    if requested_disposition not in {"attachment", "inline"}:
        raise HTTPException(
            status_code=400,
            detail="Artifact disposition must be either 'attachment' or 'inline'.",
        )

    normalized_requested_key = storage_key.strip().replace("\\", "/")
    if Path(normalized_requested_key).name.startswith(".artifact_owners"):
        raise HTTPException(status_code=404, detail="Artifact not found.")

    requesting_user_id = (
        str(current_user.user_id)
        if current_user is not None
        else _anonymous_artifact_owner_id(request)
    )
    inline_content_types = {"application/pdf", "image/jpeg", "image/png"}

    def file_response(path: Path, filename: str):
        content_type = guess_content_type(str(path)) or "application/octet-stream"
        normalized_content_type = content_type.split(";", 1)[0].strip().lower()
        response_disposition = (
            "inline"
            if requested_disposition == "inline"
            and normalized_content_type in inline_content_types
            else "attachment"
        )
        safe_filename = _safe_download_filename(filename, default=f"artifact{path.suffix}")
        if path.suffix and Path(safe_filename).suffix.lower() != path.suffix.lower():
            safe_filename = f"{_filename_stem(safe_filename, default='artifact')}{path.suffix}"
        encoded_filename = quote(safe_filename)
        ascii_filename = safe_filename.encode("ascii", "ignore").decode("ascii").strip()
        if not ascii_filename:
            ascii_filename = f"artifact{path.suffix}"

        response = FileResponse(path=str(path), media_type=content_type)
        response.headers["Content-Disposition"] = (
            f'{response_disposition}; filename="{ascii_filename}"; '
            f"filename*=UTF-8''{encoded_filename}"
        )
        response.headers["X-Content-Type-Options"] = "nosniff"
        response.headers["Content-Security-Policy"] = "sandbox"
        response.headers["Cache-Control"] = "private, no-store, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Cross-Origin-Resource-Policy"] = "same-origin"
        response.headers["Referrer-Policy"] = "no-referrer"
        return response

    for storage in _artifact_storage_download_candidates():
        try:
            path = storage.resolve_storage_key(normalized_requested_key)
        except ValueError:
            continue
        if not path.exists() or not path.is_file():
            continue

        owner = get_artifact_owner(
            normalized_requested_key,
            base_dir=str(storage.base_dir),
        )
        # Mandatory fail-closed ownership: a file without metadata is never served.
        if owner is None or owner.owner_user_id != requesting_user_id:
            raise HTTPException(status_code=404, detail="Artifact not found.")
        if is_artifact_expired(owner):
            path.unlink(missing_ok=True)
            raise HTTPException(status_code=404, detail="Artifact not found.")

        default_display_name = owner.original_artifact_name or re.sub(
            r"^[0-9a-fA-F]{16}-",
            "",
            path.name,
        )
        display_name = _safe_download_filename(
            download_name,
            default=default_display_name,
        )

        if (
            print_preview
            and request.method == "GET"
            and path.suffix.lower() in OFFICE_PRINT_PREVIEW_EXTENSIONS
        ):
            try:
                preview_path, preview_payload = _build_office_print_preview(
                    path=path,
                    display_name=display_name,
                    owner_user_id=requesting_user_id,
                    owner_organization_id=(
                        _user_organization_id(current_user)
                        if current_user is not None
                        else None
                    ),
                )
            except RuntimeError as exc:
                raise HTTPException(
                    status_code=422,
                    detail={
                        "error": "print_preview_unavailable",
                        "message": str(exc),
                    },
                ) from exc

            if preview_path is not None:
                return file_response(preview_path, preview_payload)

            response = HTMLResponse(
                content=preview_payload,
                headers=_print_preview_headers(),
            )
            safe_preview_name = (
                f"{_filename_stem(display_name, default='document')}.html"
            )
            response.headers["Content-Disposition"] = (
                f"inline; filename={quote(safe_preview_name)}"
            )
            return response

        return file_response(path, display_name)

    # The former direct-filesystem fallback was intentionally removed. Every
    # downloadable file must resolve through an owner-scoped storage backend.
    raise HTTPException(status_code=404, detail="Artifact not found.")


__all__ = ["router", "API_V1_ANALYZER_PREFIX"]
