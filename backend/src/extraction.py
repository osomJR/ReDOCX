from __future__ import annotations

import hashlib
import io
import math
import mimetypes
from os import PathLike
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional, Sequence, Union, overload

import cv2
import docx  # python-docx
from docx.text.paragraph import Paragraph
import fitz  # PyMuPDF
import numpy as np
import pytesseract
from PIL import Image

from .inline_text_security import (
    validate_inline_text,
    validate_text_to_speech_inline_text,
)
from .schema import (
    DocumentInputFormat,
    DocumentMetadata,
    DocumentPayload,
    DocumentSetPayload,
    FeatureType,
    InputArtifact,
    MAX_COMBINE_PDF_FILES,
    MAX_FILE_SIZE_MB,
    MAX_PDF_TOOL_FILE_SIZE_MB,
    PdfDocumentMetadata,
    PdfFilePayload,
    PdfFileSetPayload,
    TEXT_AI_DOC_ACTIONS_REQUIRING_TEXT_AND_WORDCOUNT,
    VaultFilePayload,
    VaultItemReferencePayload,
    VaultOperation,
    VaultQueryPayload,
    classify_word_count,
)

Pathish = Union[str, Path, PathLike[str]]

# OCR configuration
OCR_CONFIG = "--oem 3 --psm 6"
OCR_SPARSE_CONFIG = "--oem 3 --psm 11"

# Common BCP-47 / language-name aliases -> Tesseract traineddata codes.
# This is only used when callers supply explicit OCR language hints.
_TESSERACT_LANG_ALIASES: dict[str, str] = {
    "ar": "ara",
    "arabic": "ara",
    "de": "deu",
    "german": "deu",
    "en": "eng",
    "english": "eng",
    "es": "spa",
    "spanish": "spa",
    "fa": "fas",
    "farsi": "fas",
    "fr": "fra",
    "french": "fra",
    "ha": "hau",
    "hausa": "hau",
    "hi": "hin",
    "hindi": "hin",
    "ig": "ibo",
    "igbo": "ibo",
    "it": "ita",
    "italian": "ita",
    "ja": "jpn",
    "japanese": "jpn",
    "ko": "kor",
    "korean": "kor",
    "nl": "nld",
    "dutch": "nld",
    "pl": "pol",
    "polish": "pol",
    "pt": "por",
    "pt-br": "por",
    "portuguese": "por",
    "ru": "rus",
    "russian": "rus",
    "sw": "swa",
    "swahili": "swa",
    "tr": "tur",
    "turkish": "tur",
    "uk": "ukr",
    "ukrainian": "ukr",
    "yo": "yor",
    "yoruba": "yor",
    "zh": "chi_sim",
    "zh-cn": "chi_sim",
    "zh-hans": "chi_sim",
    "zh-tw": "chi_tra",
    "zh-hant": "chi_tra",
}

# ----------------------------
# Action groups aligned with schema.py / validation.py
# ----------------------------
CONVERSION_ACTIONS = {FeatureType.convert}

TEXT_AI_DOC_INPUT_FORMATS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.txt,
}

TEXT_TO_SPEECH_INPUT_FORMATS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.txt,
}

REDACTION_MASKING_INPUT_FORMATS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.jpg,
    DocumentInputFormat.jpeg,
    DocumentInputFormat.png,
}

STRUCTURED_EXTRACTION_INPUT_FORMATS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.jpg,
    DocumentInputFormat.jpeg,
    DocumentInputFormat.png,
}

COMPLIANCE_INPUT_FORMATS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.jpg,
    DocumentInputFormat.jpeg,
    DocumentInputFormat.png,
}

CONVERSION_INPUT_FORMATS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.jpg,
    DocumentInputFormat.jpeg,
    DocumentInputFormat.png,
}

OPTIONAL_TEXT_DOCUMENT_ACTIONS = {
    FeatureType.redact,
    FeatureType.data_mask,
    FeatureType.structured_extract,
    FeatureType.compliance,
}

DOCUMENT_SET_ACTIONS = {
    FeatureType.structured_extract,
    FeatureType.compliance,
}

PDF_DOCUMENT_ACTIONS = {
    FeatureType.combine_pdf,
    FeatureType.split_pdf,
    FeatureType.edit_pdf,
    FeatureType.compress_pdf,
    FeatureType.lock_pdf,
    FeatureType.e_signature,
}

PDF_SINGLE_FILE_ACTIONS = {
    FeatureType.split_pdf,
    FeatureType.edit_pdf,
    FeatureType.compress_pdf,
    FeatureType.lock_pdf,
    FeatureType.e_signature,
}

_ALLOWED_INPUT_FORMATS_BY_ACTION: dict[FeatureType, set[DocumentInputFormat]] = {
    FeatureType.convert: CONVERSION_INPUT_FORMATS,
    FeatureType.summarize: TEXT_AI_DOC_INPUT_FORMATS,
    FeatureType.grammar_correct: TEXT_AI_DOC_INPUT_FORMATS,
    FeatureType.translate: TEXT_AI_DOC_INPUT_FORMATS,
    FeatureType.explain: TEXT_AI_DOC_INPUT_FORMATS,
    FeatureType.generate_questions: TEXT_AI_DOC_INPUT_FORMATS,
    FeatureType.generate_answers: TEXT_AI_DOC_INPUT_FORMATS,
    FeatureType.text_to_speech: TEXT_TO_SPEECH_INPUT_FORMATS,
    FeatureType.redact: REDACTION_MASKING_INPUT_FORMATS,
    FeatureType.data_mask: REDACTION_MASKING_INPUT_FORMATS,
    FeatureType.structured_extract: STRUCTURED_EXTRACTION_INPUT_FORMATS,
    FeatureType.compliance: COMPLIANCE_INPUT_FORMATS,
}

_PDF_MIME_TYPES = {"application/pdf", "application/x-pdf"}


# ----------------------------
# OCR helpers
# ----------------------------
def get_available_tesseract_languages() -> list[str]:
    """Return installed Tesseract language codes, excluding utility packs."""
    languages = pytesseract.get_languages(config="")
    return sorted(lang for lang in languages if lang not in {"osd", "equ"})


def _normalize_ocr_language_token(token: str) -> str:
    normalized = token.strip().lower().replace("_", "-")
    if not normalized:
        raise ValueError("OCR language token cannot be empty.")
    return _TESSERACT_LANG_ALIASES.get(normalized, normalized)


def resolve_ocr_lang(ocr_languages: Optional[Sequence[str]] = None) -> str:
    """
    Resolve the Tesseract language bundle.

    Behavior:
    - if caller supplies OCR language hints, normalize and validate them
    - otherwise, use all installed OCR languages to maximize scanned-document coverage
    - include osd when present for orientation/script detection support
    """
    available = set(get_available_tesseract_languages())
    raw_available = set(pytesseract.get_languages(config=""))

    if ocr_languages:
        requested: list[str] = []
        for item in ocr_languages:
            code = _normalize_ocr_language_token(item)
            if code not in available:
                raise ValueError(
                    f"Requested OCR language '{item}' resolved to '{code}', "
                    "but that traineddata is not installed in Tesseract."
                )
            if code not in requested:
                requested.append(code)
        selected = requested
    else:
        selected = sorted(available)

    if not selected:
        raise ValueError(
            "No OCR languages are installed in Tesseract. Install traineddata files "
            "for the languages you need before processing scanned documents."
        )

    if "osd" in raw_available:
        selected = [*selected, "osd"]

    return "+".join(selected)


# ----------------------------
# File helpers
# ----------------------------
def _as_existing_file(file_path: Pathish) -> Path:
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"File not found: {path}")
    if not path.is_file():
        raise ValueError(f"Path is not a file: {path}")
    return path


def preprocess_for_ocr(image: Image.Image) -> Image.Image:
    """
    Improve OCR accuracy by cleaning the image.
    Steps:
    - convert to grayscale
    - denoise
    - adaptive threshold
    """
    img = np.array(image)
    gray = cv2.cvtColor(img, cv2.COLOR_RGB2GRAY)
    gray = cv2.medianBlur(gray, 3)
    thresh = cv2.adaptiveThreshold(
        gray,
        255,
        cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
        cv2.THRESH_BINARY,
        31,
        2,
    )
    return Image.fromarray(thresh)


def _ocr_data_score(data: Mapping[str, Sequence[object]]) -> tuple[float, float, int]:
    """Rank OCR passes by confident, useful text rather than raw token count."""
    weighted_confidence = 0.0
    alphanumeric_count = 0
    token_count = 0

    texts = data.get("text", ())
    confidences = data.get("conf", ())
    for index, raw in enumerate(texts):
        token = str(raw or "").strip()
        if not token:
            continue

        useful_length = sum(char.isalnum() for char in token)
        if useful_length <= 0:
            continue

        try:
            confidence = float(confidences[index])
        except (IndexError, TypeError, ValueError):
            confidence = 0.0

        weighted_confidence += max(0.0, confidence) * useful_length
        alphanumeric_count += useful_length
        token_count += 1

    average_confidence = (
        weighted_confidence / alphanumeric_count
        if alphanumeric_count
        else 0.0
    )
    quality = average_confidence * math.log1p(alphanumeric_count)
    return quality, average_confidence, token_count


def extract_ocr_data(
    image: Image.Image,
    *,
    ocr_lang: str,
    config: str = OCR_CONFIG,
) -> Mapping[str, Sequence[object]]:
    """Return the stronger of the original-image and thresholded OCR passes.

    Adaptive thresholding helps ordinary scans, but it can destroy characters
    printed over guilloches and other security backgrounds. Running both passes
    and selecting by Tesseract confidence keeps scan cleanup without sacrificing
    identity documents.
    """
    original = image.convert("RGB")
    processed = preprocess_for_ocr(original)
    candidates = (
        (original, OCR_SPARSE_CONFIG),
        (processed, config),
    )
    results = [
        pytesseract.image_to_data(
            candidate,
            lang=ocr_lang,
            config=candidate_config,
            output_type=pytesseract.Output.DICT,
        )
        for candidate, candidate_config in candidates
    ]
    return max(results, key=_ocr_data_score)


def ocr_text_from_data(data: Mapping[str, Sequence[object]]) -> str:
    """Reconstruct OCR text while preserving line boundaries."""
    parts: list[str] = []
    previous_line: tuple[object, object, object] | None = None

    texts = data.get("text", ())
    page_numbers = data.get("page_num", ())
    block_numbers = data.get("block_num", ())
    line_numbers = data.get("line_num", ())

    for index, raw in enumerate(texts):
        token = str(raw or "").strip()
        if not token:
            continue

        def item(values: Sequence[object]) -> object:
            return values[index] if index < len(values) else 0

        current_line = (
            item(page_numbers),
            item(block_numbers),
            item(line_numbers),
        )
        if parts:
            parts.append("\n" if current_line != previous_line else " ")
        parts.append(token)
        previous_line = current_line

    return "".join(parts).strip()


def get_file_size_mb(file_path: Pathish, *, max_size_mb: float = MAX_FILE_SIZE_MB) -> float:
    """
    Return size in MB and enforce the provided size limit.

    The default limit is the strict AI-document limit. PDF tools should call this
    with max_size_mb=MAX_PDF_TOOL_FILE_SIZE_MB through get_pdf_tool_file_size_mb().
    """
    path = _as_existing_file(file_path)
    size_bytes = path.stat().st_size
    size_mb = size_bytes / (1024 * 1024)

    if size_mb <= 0:
        raise ValueError("File is empty.")
    if size_mb > max_size_mb:
        raise ValueError(f"File exceeds maximum allowed size of {max_size_mb:g} MB.")

    return round(size_mb, 4)


def get_pdf_tool_file_size_mb(file_path: Pathish) -> float:
    """Return file size using the larger PDF Tool / Lock PDF / E-Signature upload limit."""
    return get_file_size_mb(file_path, max_size_mb=MAX_PDF_TOOL_FILE_SIZE_MB)


def detect_format(file_path: Pathish) -> DocumentInputFormat:
    path = Path(file_path)
    suffix = path.suffix.lower().lstrip(".")
    try:
        return DocumentInputFormat(suffix)
    except ValueError as exc:
        raise ValueError(f"Unsupported file format: {suffix}") from exc


def guess_mime_type(file_path: Pathish, *, fallback: str = "application/octet-stream") -> str:
    guessed, _encoding = mimetypes.guess_type(str(file_path))
    return guessed or fallback


def compute_sha256_hex(file_path: Pathish, *, chunk_size: int = 1024 * 1024) -> str:
    """Compute SHA-256 for storage identity, audit evidence, and duplicate-upload guards."""
    path = _as_existing_file(file_path)
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(chunk_size), b""):
            digest.update(chunk)
    return digest.hexdigest()


def build_vault_file_payload(
    file_path: Pathish,
    *,
    storage_key: Optional[str] = None,
    upload_id: Optional[str] = None,
    content_type: Optional[str] = None,
    checksum_sha256: Optional[str] = None,
    client_encrypted: bool = False,
) -> VaultFilePayload:
    """
    Build the binary-safe input artifact for a Vault store operation.

    Vault accepts arbitrary file types, so this path intentionally performs no
    document-format detection, text extraction, OCR, or document-size capping.
    The upload layer must supply at least one persisted-file reference.
    """
    path = _as_existing_file(file_path)
    resolved_content_type = content_type or guess_mime_type(path)

    return VaultFilePayload(
        kind="vault_file",
        filename=path.name,
        content_type=resolved_content_type,
        file_size_bytes=path.stat().st_size,
        checksum_sha256=checksum_sha256 or compute_sha256_hex(path),
        storage_key=storage_key,
        upload_id=upload_id,
        client_encrypted=client_encrypted,
    )


def build_vault_item_reference_payload(item_id: str) -> VaultItemReferencePayload:
    """Build the owner-scoped input artifact for Vault retrieve or delete."""
    return VaultItemReferencePayload(kind="vault_item_reference", item_id=item_id)


def build_vault_query_payload(
    *,
    limit: int = 50,
    cursor: Optional[str] = None,
    filename_contains: Optional[str] = None,
    content_type: Optional[str] = None,
) -> VaultQueryPayload:
    """Build the pagination/filter input artifact for a Vault list operation."""
    return VaultQueryPayload(
        kind="vault_query",
        limit=limit,
        cursor=cursor,
        filename_contains=filename_contains,
        content_type=content_type,
    )


def build_vault_input_artifact(
    *,
    operation: VaultOperation,
    file_path: Optional[Pathish] = None,
    item_id: Optional[str] = None,
    storage_key: Optional[str] = None,
    upload_id: Optional[str] = None,
    content_type: Optional[str] = None,
    checksum_sha256: Optional[str] = None,
    client_encrypted: bool = False,
    limit: int = 50,
    cursor: Optional[str] = None,
    filename_contains: Optional[str] = None,
) -> Union[VaultFilePayload, VaultItemReferencePayload, VaultQueryPayload]:
    """Build the schema-required Vault input type for the selected operation."""
    if operation == VaultOperation.store:
        if file_path is None:
            raise ValueError("Vault store requires file_path.")
        if item_id is not None:
            raise ValueError("Vault store does not accept item_id.")
        return build_vault_file_payload(
            file_path,
            storage_key=storage_key,
            upload_id=upload_id,
            content_type=content_type,
            checksum_sha256=checksum_sha256,
            client_encrypted=client_encrypted,
        )

    if operation in {VaultOperation.retrieve, VaultOperation.delete}:
        if file_path is not None:
            raise ValueError(f"Vault {operation.value} does not accept file_path.")
        if item_id is None:
            raise ValueError(f"Vault {operation.value} requires item_id.")
        return build_vault_item_reference_payload(item_id)

    if operation == VaultOperation.list:
        if file_path is not None or item_id is not None:
            raise ValueError("Vault list does not accept file_path or item_id.")
        return build_vault_query_payload(
            limit=limit,
            cursor=cursor,
            filename_contains=filename_contains,
            content_type=content_type,
        )

    raise ValueError(f"Unsupported Vault operation: {operation!r}")


def _validate_pdf_path_and_mime(path: Path, mime_type: str) -> None:
    if path.suffix.lower() != ".pdf":
        raise ValueError("PDF tools, Lock PDF, and e-signature require a .pdf file.")
    if mime_type.lower() not in _PDF_MIME_TYPES:
        raise ValueError(
            "PDF tools, Lock PDF, and e-signature inputs must use "
            "application/pdf or application/x-pdf."
        )


# ----------------------------
# PDF metadata helpers for PDF Tools + Lock PDF + E-Signature
# ----------------------------
def inspect_pdf_metadata(file_path: Pathish, *, password: Optional[str] = None) -> dict[str, Optional[Union[int, bool]]]:
    """
    Inspect PDF metadata needed by PdfDocumentMetadata.

    Returned keys:
    - page_count: int | None
    - encrypted: bool
    - password_protected: bool

    The metadata reports whether the original file is encrypted/password-protected.
    It does not silently mark encrypted files as processable. validation.py rejects
    encrypted/password-protected PDFs until you add a dedicated unlock workflow.
    """
    path = _as_existing_file(file_path)
    try:
        with fitz.open(path) as pdf:
            encrypted = bool(getattr(pdf, "is_encrypted", False))
            needs_pass = bool(getattr(pdf, "needs_pass", False))
            authenticated = False

            if needs_pass and password:
                authenticated = bool(pdf.authenticate(password))
                if not authenticated:
                    raise ValueError("PDF password authentication failed.")

            page_count: Optional[int]
            if encrypted and not authenticated:
                page_count = None
            else:
                page_count = int(pdf.page_count)

            return {
                "page_count": page_count,
                "encrypted": encrypted,
                "password_protected": needs_pass,
            }
    except ValueError:
        raise
    except Exception as exc:  # PyMuPDF raises several low-level exceptions for invalid PDFs.
        raise ValueError(f"Invalid or unreadable PDF file: {path.name}") from exc


def build_pdf_document_metadata(
    file_path: Pathish,
    *,
    password: Optional[str] = None,
    checksum_sha256: Optional[str] = None,
) -> PdfDocumentMetadata:
    """Build PdfDocumentMetadata for PDF Tools, Lock PDF, and E-Signature inputs."""
    path = _as_existing_file(file_path)
    file_size_mb = get_pdf_tool_file_size_mb(path)
    pdf_info = inspect_pdf_metadata(path, password=password)
    checksum = checksum_sha256 or compute_sha256_hex(path)

    return PdfDocumentMetadata(
        input_format=DocumentInputFormat.pdf,
        file_size_mb=file_size_mb,
        page_count=pdf_info["page_count"],
        encrypted=bool(pdf_info["encrypted"]),
        password_protected=bool(pdf_info["password_protected"]),
        checksum_sha256=checksum,
    )


def build_pdf_file_payload(
    file_path: Pathish,
    *,
    storage_key: Optional[str] = None,
    upload_id: Optional[str] = None,
    mime_type: Optional[str] = None,
    password: Optional[str] = None,
    checksum_sha256: Optional[str] = None,
) -> PdfFilePayload:
    """
    Build a PdfFilePayload for Split, Edit, Compress, Lock PDF, and E-Signature workflows.

    This intentionally does not extract text or run OCR. PDF tools use structural
    PDF metadata, a file identity/hash, and the persisted file reference.
    """
    path = _as_existing_file(file_path)
    resolved_mime = mime_type or guess_mime_type(path, fallback="application/pdf")
    _validate_pdf_path_and_mime(path, resolved_mime)

    return PdfFilePayload(
        kind="pdf_file",
        metadata=build_pdf_document_metadata(
            path,
            password=password,
            checksum_sha256=checksum_sha256,
        ),
        filename=path.name,
        mime_type=resolved_mime,
        storage_key=storage_key,
        upload_id=upload_id,
    )


def _sequence_item(
    value: Optional[Sequence[Optional[str]]],
    index: int,
    *,
    field_name: str,
    expected_length: int,
) -> Optional[str]:
    if value is None:
        return None
    if len(value) != expected_length:
        raise ValueError(f"{field_name} length must match file_paths length.")
    return value[index]


def build_pdf_file_set_payload(
    file_paths: Sequence[Pathish],
    *,
    storage_keys: Optional[Sequence[Optional[str]]] = None,
    upload_ids: Optional[Sequence[Optional[str]]] = None,
    mime_types: Optional[Sequence[Optional[str]]] = None,
    passwords: Optional[Sequence[Optional[str]]] = None,
    checksums_sha256: Optional[Sequence[Optional[str]]] = None,
) -> PdfFileSetPayload:
    """
    Build a PdfFileSetPayload for Combine PDF.

    ReDOCX supports 2..10 PDFs per combine request. The order of file_paths is
    preserved so frontend move-up/move-down order is the backend combine order.
    """
    if not file_paths:
        raise ValueError("file_paths cannot be empty.")
    if len(file_paths) < 2:
        raise ValueError("Combine PDF requires at least 2 PDF files.")
    if len(file_paths) > MAX_COMBINE_PDF_FILES:
        raise ValueError(f"Combine PDF supports at most {MAX_COMBINE_PDF_FILES} PDF files.")

    expected_length = len(file_paths)
    documents = [
        build_pdf_file_payload(
            file_path,
            storage_key=_sequence_item(storage_keys, index, field_name="storage_keys", expected_length=expected_length),
            upload_id=_sequence_item(upload_ids, index, field_name="upload_ids", expected_length=expected_length),
            mime_type=_sequence_item(mime_types, index, field_name="mime_types", expected_length=expected_length),
            password=_sequence_item(passwords, index, field_name="passwords", expected_length=expected_length),
            checksum_sha256=_sequence_item(
                checksums_sha256,
                index,
                field_name="checksums_sha256",
                expected_length=expected_length,
            ),
        )
        for index, file_path in enumerate(file_paths)
    ]

    # Prevent accidentally combining the exact same uploaded object twice. Different
    # files may share a display filename, so checksum/storage identity is preferred.
    identities = [
        document.storage_key
        or document.upload_id
        or (document.metadata.checksum_sha256.lower() if document.metadata.checksum_sha256 else None)
        or document.filename
        for document in documents
    ]
    if len(set(identities)) != len(identities):
        raise ValueError("Combine PDF input documents must not contain duplicate uploaded files.")

    return PdfFileSetPayload(kind="pdf_file_set", documents=documents)


# ----------------------------
# Text extraction
# ----------------------------
def extract_text_from_txt(file_path: Pathish) -> str:
    path = _as_existing_file(file_path)
    try:
        return path.read_text(encoding="utf-8").strip()
    except UnicodeDecodeError as exc:
        raise ValueError("TXT file must be valid UTF-8.") from exc


def _iter_docx_table_paragraphs(table: Any):
    for row in table.rows:
        for cell in row.cells:
            yield from cell.paragraphs
            for nested_table in cell.tables:
                yield from _iter_docx_table_paragraphs(nested_table)


def _iter_docx_paragraphs(document: Any):
    """Yield body, table, header, footer, and text-box paragraphs once."""
    # Retain the XML elements. Tracking only id(element) is unsafe because
    # lxml proxy objects can be released during iteration and Python may reuse
    # the same numeric id for a different paragraph.
    seen: set[Any] = set()

    def emit(paragraphs: Iterable[Any]):
        for paragraph in paragraphs:
            element = paragraph._p
            if element in seen:
                continue
            seen.add(element)
            yield paragraph
            for nested_element in paragraph._p.xpath(".//w:txbxContent//w:p"):
                if nested_element in seen:
                    continue
                seen.add(nested_element)
                yield Paragraph(nested_element, paragraph._parent)

    yield from emit(document.paragraphs)
    for table in document.tables:
        yield from emit(_iter_docx_table_paragraphs(table))

    for section in document.sections:
        for story in (section.header, section.footer):
            yield from emit(story.paragraphs)
            for table in story.tables:
                yield from emit(_iter_docx_table_paragraphs(table))


def extract_text_from_docx(file_path: Pathish) -> str:
    path = _as_existing_file(file_path)
    document = docx.Document(path)
    return "\n".join(
        paragraph.text
        for paragraph in _iter_docx_paragraphs(document)
        if paragraph.text and paragraph.text.strip()
    ).strip()


def extract_text_from_image(file_path: Pathish, *, ocr_lang: str) -> str:
    path = _as_existing_file(file_path)
    image = Image.open(path).convert("RGB")
    return ocr_text_from_data(
        extract_ocr_data(image, ocr_lang=ocr_lang),
    )


def extract_text_from_pdf_text(file_path: Pathish) -> str:
    path = _as_existing_file(file_path)
    chunks: list[str] = []
    with fitz.open(path) as pdf:
        if bool(getattr(pdf, "needs_pass", False)):
            raise ValueError("Cannot extract text from password-protected PDF without an unlock workflow.")
        for page in pdf:
            chunks.append(page.get_text())
    return "\n".join(chunks).strip()


def _pixmap_to_pil(pix: fitz.Pixmap) -> Image.Image:
    mode = "RGBA" if pix.alpha else "RGB"
    return Image.frombytes(mode, (pix.width, pix.height), pix.samples)


def _embedded_ocr_images(pdf: fitz.Document, page: fitz.Page) -> list[Image.Image]:
    """Extract meaningful image placements from an otherwise textless PDF page."""
    page_area = max(1.0, float(page.rect.width * page.rect.height))
    images: list[Image.Image] = []
    seen_xrefs: set[int] = set()

    for item in page.get_images(full=True):
        xref = int(item[0])
        width = int(item[2])
        height = int(item[3])
        if xref <= 0 or xref in seen_xrefs or width < 96 or height < 48:
            continue

        placements = page.get_image_rects(xref)
        if not placements:
            continue
        if max(float(rect.width * rect.height) for rect in placements) / page_area < 0.01:
            continue

        try:
            extracted = pdf.extract_image(xref)
            image_bytes = extracted.get("image")
            if not image_bytes:
                continue
            with Image.open(io.BytesIO(image_bytes)) as source:
                images.append(source.convert("RGB"))
            seen_xrefs.add(xref)
        except Exception:
            # Rendering the whole page below remains the safe fallback for
            # unsupported image encodings or malformed image objects.
            continue

    return images


def extract_text_from_pdf_ocr(file_path: Pathish, *, ocr_lang: str, zoom: float = 4.0) -> str:
    path = _as_existing_file(file_path)
    matrix = fitz.Matrix(zoom, zoom)
    chunks: list[str] = []

    with fitz.open(path) as pdf:
        if bool(getattr(pdf, "needs_pass", False)):
            raise ValueError("Cannot OCR password-protected PDF without an unlock workflow.")
        for page in pdf:
            embedded_images = _embedded_ocr_images(pdf, page)
            if embedded_images:
                chunks.extend(
                    ocr_text_from_data(extract_ocr_data(image, ocr_lang=ocr_lang))
                    for image in embedded_images
                )
                continue

            pix = page.get_pixmap(matrix=matrix, alpha=False, colorspace=fitz.csRGB)
            image = _pixmap_to_pil(pix).convert("RGB")
            chunks.append(
                ocr_text_from_data(extract_ocr_data(image, ocr_lang=ocr_lang))
            )

    return "\n".join(chunks).strip()


def extract_text_by_format(
    file_path: Pathish,
    fmt: DocumentInputFormat,
    *,
    ocr_languages: Optional[Sequence[str]] = None,
) -> tuple[str, bool]:
    """
    Return (text, ocr_used).

    Contract alignment:
    - txt/docx => ocr_used=False
    - pdf => OCR fallback only if native extraction is empty
    - images => OCR always used when text extraction is attempted
    """
    if fmt == DocumentInputFormat.txt:
        return extract_text_from_txt(file_path), False

    if fmt == DocumentInputFormat.docx:
        return extract_text_from_docx(file_path), False

    if fmt == DocumentInputFormat.pdf:
        text = extract_text_from_pdf_text(file_path)
        if text:
            return text, False
        ocr_lang = resolve_ocr_lang(ocr_languages)
        return extract_text_from_pdf_ocr(file_path, ocr_lang=ocr_lang), True

    if fmt in (DocumentInputFormat.jpg, DocumentInputFormat.jpeg, DocumentInputFormat.png):
        ocr_lang = resolve_ocr_lang(ocr_languages)
        return extract_text_from_image(file_path, ocr_lang=ocr_lang), True

    raise ValueError(f"Unsupported file format: {fmt.value}")


# ----------------------------
# Word-count helpers
# ----------------------------
def count_words(text: str) -> int:
    return len(text.split())


def enforce_text_ai_word_contract(word_count: int) -> None:
    if word_count < 1:
        raise ValueError("Document contains no words.")
    classify_word_count(word_count)


# ----------------------------
# DocumentPayload builders for existing AI/document features
# ----------------------------
def build_inline_text_payload(
    text: str,
    *,
    input_format: DocumentInputFormat = DocumentInputFormat.txt,
) -> DocumentPayload:
    """
    Normalize inline text into a schema-compliant DocumentPayload.

    Notes:
    - Inline text is represented as txt input.
    - detected_language is intentionally omitted because request-side validation
      forbids client-supplied detected_language; analyzer/server handles detection.
    - The strict 1..1000 word-count contract is enforced only for text AI actions.
    """
    normalized = validate_inline_text(text)

    if input_format != DocumentInputFormat.txt:
        raise ValueError("Inline text payload must use input_format='txt'.")

    word_count = count_words(normalized)
    enforce_text_ai_word_contract(word_count)

    metadata = DocumentMetadata(
        input_format=DocumentInputFormat.txt,
        file_size_mb=0.0,
        extracted_word_count=word_count,
        ocr_used=False,
    )
    return DocumentPayload(text=normalized, metadata=metadata, mime_type="text/plain")


def build_text_to_speech_inline_payload(text: str) -> DocumentPayload:
    """
    Build an inline TXT payload for Text-to-Speech without the text-AI word cap.

    The normalized text stored in this payload is the exact string the speech
    provider must synthesize and the response character count must reference.
    """
    normalized = validate_text_to_speech_inline_text(text)

    metadata = DocumentMetadata(
        input_format=DocumentInputFormat.txt,
        file_size_mb=0.0,
        extracted_word_count=count_words(normalized),
        ocr_used=False,
    )
    return DocumentPayload(text=normalized, metadata=metadata, mime_type="text/plain")


def _build_document_payload(
    file_path: Pathish,
    *,
    allowed_formats: set[DocumentInputFormat],
    require_text: bool,
    enforce_text_ai_range: bool,
    extract_optional_text: bool = True,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    path = _as_existing_file(file_path)

    fmt = detect_format(path)
    if fmt not in allowed_formats:
        allowed = ", ".join(item.value for item in sorted(allowed_formats, key=lambda x: x.value))
        raise ValueError(f"Unsupported input format for this action. Allowed formats: {allowed}.")

    file_size_mb = get_file_size_mb(path)
    mime_type = guess_mime_type(path)

    text: Optional[str] = None
    extracted_word_count: Optional[int] = None
    ocr_used = False

    should_extract_text = require_text or enforce_text_ai_range or extract_optional_text
    if should_extract_text and fmt in {
        DocumentInputFormat.txt,
        DocumentInputFormat.docx,
        DocumentInputFormat.pdf,
        DocumentInputFormat.jpg,
        DocumentInputFormat.jpeg,
        DocumentInputFormat.png,
    }:
        extracted_text, ocr_used = extract_text_by_format(path, fmt, ocr_languages=ocr_languages)
        normalized = extracted_text.strip()
        if normalized:
            text = normalized
            extracted_word_count = count_words(normalized)

    if require_text and not text:
        raise ValueError("Document text could not be extracted.")

    if text and enforce_text_ai_range:
        enforce_text_ai_word_contract(extracted_word_count or 0)

    metadata = DocumentMetadata(
        input_format=fmt,
        file_size_mb=file_size_mb,
        extracted_word_count=extracted_word_count,
        ocr_used=ocr_used,
    )
    return DocumentPayload(text=text, metadata=metadata, filename=path.name, mime_type=mime_type)


def build_conversion_document_payload(
    file_path: Pathish,
    *,
    extract_optional_text: bool = False,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    """
    Build a DocumentPayload for the convert action.

    Contract alignment:
    - convert accepts pdf/docx/jpg/jpeg/png
    - text is optional
    - extracted_word_count is optional

    By default this avoids OCR/native extraction because conversion engines normally
    only need the persisted file. Set extract_optional_text=True if your pipeline
    wants pre-extracted text for logging or downstream heuristics.
    """
    return _build_document_payload(
        file_path,
        allowed_formats=CONVERSION_INPUT_FORMATS,
        require_text=False,
        enforce_text_ai_range=False,
        extract_optional_text=extract_optional_text,
        ocr_languages=ocr_languages,
    )


def build_text_ai_document_payload(
    file_path: Pathish,
    *,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    """
    Build a DocumentPayload for text AI document actions:
    summarize, grammar_correct, translate, explain, generate_questions, generate_answers.

    Contract alignment:
    - only pdf/docx/txt are allowed
    - extracted text is required
    - extracted_word_count is required and must fit the 1..1000 contract range
    """
    return _build_document_payload(
        file_path,
        allowed_formats=TEXT_AI_DOC_INPUT_FORMATS,
        require_text=True,
        enforce_text_ai_range=True,
        extract_optional_text=True,
        ocr_languages=ocr_languages,
    )


def build_text_to_speech_document_payload(
    file_path: Pathish,
    *,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    """
    Build the PDF, DOCX, or TXT input required by Text-to-Speech.

    Extracted text is mandatory, but the legacy 1..1000-word text-AI limit is
    intentionally not applied because neither schema.py nor validation.py applies
    that cap to Text-to-Speech.
    """
    return _build_document_payload(
        file_path,
        allowed_formats=TEXT_TO_SPEECH_INPUT_FORMATS,
        require_text=True,
        enforce_text_ai_range=False,
        extract_optional_text=True,
        ocr_languages=ocr_languages,
    )


def build_redaction_or_masking_document_payload(
    file_path: Pathish,
    *,
    extract_optional_text: bool = True,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    """
    Build a DocumentPayload for redact/data_mask.

    Contract alignment:
    - accepts pdf/docx/jpg/jpeg/png
    - text is optional in schema
    - extracted_word_count may be present, but is not capped at 1000 here
    """
    return _build_document_payload(
        file_path,
        allowed_formats=REDACTION_MASKING_INPUT_FORMATS,
        require_text=False,
        enforce_text_ai_range=False,
        extract_optional_text=extract_optional_text,
        ocr_languages=ocr_languages,
    )


def build_structured_extraction_or_compliance_document_payload(
    file_path: Pathish,
    *,
    action: FeatureType,
    extract_optional_text: bool = True,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    """
    Build a single-document payload for structured_extract or compliance.

    Contract alignment:
    - accepts pdf/docx/jpg/jpeg/png
    - text is optional in schema
    - extracted_word_count may be present, but is not capped at 1000 here
    """
    if action not in DOCUMENT_SET_ACTIONS:
        raise ValueError("This builder only supports structured_extract and compliance actions.")

    allowed_formats = _ALLOWED_INPUT_FORMATS_BY_ACTION[action]
    return _build_document_payload(
        file_path,
        allowed_formats=allowed_formats,
        require_text=False,
        enforce_text_ai_range=False,
        extract_optional_text=extract_optional_text,
        ocr_languages=ocr_languages,
    )


def build_document_set_payload(
    file_paths: Sequence[Pathish],
    *,
    action: FeatureType,
    extract_optional_text: bool = True,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentSetPayload:
    """Build a DocumentSetPayload for structured_extract or compliance."""
    if action not in DOCUMENT_SET_ACTIONS:
        raise ValueError("Document sets are only supported for structured_extract and compliance.")
    if not file_paths:
        raise ValueError("file_paths cannot be empty.")

    documents = [
        build_structured_extraction_or_compliance_document_payload(
            file_path,
            action=action,
            extract_optional_text=extract_optional_text,
            ocr_languages=ocr_languages,
        )
        for file_path in file_paths
    ]
    return DocumentSetPayload(documents=documents)


def build_document_payload_for_action(
    *,
    action: FeatureType,
    file_path: Optional[Pathish] = None,
    inline_text: Optional[str] = None,
    extract_optional_text: bool = True,
    ocr_languages: Optional[Sequence[str]] = None,
) -> DocumentPayload:
    """
    Runtime-aware normalization entrypoint for single-document actions.

    Rules:
    - inline_text => txt payload for text AI document actions or Text-to-Speech
    - convert => text optional
    - text AI actions => extracted text + extracted_word_count required
    - Text-to-Speech => exact extracted text required, without the text-AI word cap
    - redact/data_mask/structured_extract/compliance => extracted text optional
    - Vault uses Vault-specific artifacts and is not routed through this builder
    - PDF tools/e-signature are not routed through this builder; use PdfFilePayload builders
    """
    if action == FeatureType.vault:
        raise ValueError(
            "vault uses VaultFilePayload/VaultItemReferencePayload/VaultQueryPayload. "
            "Use build_vault_input_artifact or build_input_artifact_for_action."
        )

    if action in PDF_DOCUMENT_ACTIONS:
        raise ValueError(
            f"{action.value} uses PdfFilePayload/PdfFileSetPayload. "
            "Use build_pdf_file_payload, build_pdf_file_set_payload, or build_input_artifact_for_action."
        )

    if inline_text is not None:
        if file_path is not None:
            raise ValueError("Provide either file_path or inline_text, not both.")
        if action == FeatureType.text_to_speech:
            return build_text_to_speech_inline_payload(inline_text)
        if action not in TEXT_AI_DOC_ACTIONS_REQUIRING_TEXT_AND_WORDCOUNT:
            raise ValueError(f"{action.value} does not support inline text through this extraction path.")
        return build_inline_text_payload(inline_text)

    if file_path is None:
        raise ValueError("Either file_path or inline_text must be provided.")

    if action in CONVERSION_ACTIONS:
        return build_conversion_document_payload(
            file_path,
            extract_optional_text=extract_optional_text,
            ocr_languages=ocr_languages,
        )

    if action in TEXT_AI_DOC_ACTIONS_REQUIRING_TEXT_AND_WORDCOUNT:
        return build_text_ai_document_payload(file_path, ocr_languages=ocr_languages)

    if action == FeatureType.text_to_speech:
        return build_text_to_speech_document_payload(
            file_path,
            ocr_languages=ocr_languages,
        )

    if action in {FeatureType.redact, FeatureType.data_mask}:
        return build_redaction_or_masking_document_payload(
            file_path,
            extract_optional_text=extract_optional_text,
            ocr_languages=ocr_languages,
        )

    if action in DOCUMENT_SET_ACTIONS:
        return build_structured_extraction_or_compliance_document_payload(
            file_path,
            action=action,
            extract_optional_text=extract_optional_text,
            ocr_languages=ocr_languages,
        )

    raise ValueError(f"Unsupported document action for extraction: {action.value}")


# ----------------------------
# Unified input-artifact builder aligned with schema.InputArtifact
# ----------------------------
def build_input_artifact_for_action(
    *,
    action: FeatureType,
    file_path: Optional[Pathish] = None,
    file_paths: Optional[Sequence[Pathish]] = None,
    inline_text: Optional[str] = None,
    storage_key: Optional[str] = None,
    upload_id: Optional[str] = None,
    storage_keys: Optional[Sequence[Optional[str]]] = None,
    upload_ids: Optional[Sequence[Optional[str]]] = None,
    mime_type: Optional[str] = None,
    mime_types: Optional[Sequence[Optional[str]]] = None,
    password: Optional[str] = None,
    passwords: Optional[Sequence[Optional[str]]] = None,
    checksum_sha256: Optional[str] = None,
    checksums_sha256: Optional[Sequence[Optional[str]]] = None,
    vault_operation: Optional[VaultOperation] = None,
    vault_item_id: Optional[str] = None,
    vault_content_type: Optional[str] = None,
    vault_client_encrypted: bool = False,
    vault_limit: int = 50,
    vault_cursor: Optional[str] = None,
    vault_filename_contains: Optional[str] = None,
    extract_optional_text: bool = True,
    ocr_languages: Optional[Sequence[str]] = None,
) -> InputArtifact:
    """
    Runtime-aware normalization entrypoint aligned with schema.InputArtifact rules.

    - vault returns the operation-specific Vault input artifact
    - combine_pdf returns PdfFileSetPayload
    - split_pdf/edit_pdf/compress_pdf/lock_pdf/e_signature return PdfFilePayload
    - structured_extract/compliance may return DocumentSetPayload when file_paths is supplied
    - existing single-document actions return DocumentPayload
    - inline_text is supported for text AI actions and Text-to-Speech
    """
    if action == FeatureType.vault:
        if file_paths is not None or inline_text is not None:
            raise ValueError("vault does not accept file_paths or inline_text.")
        if vault_operation is None:
            raise ValueError("vault requires vault_operation.")
        return build_vault_input_artifact(
            operation=vault_operation,
            file_path=file_path,
            item_id=vault_item_id,
            storage_key=storage_key,
            upload_id=upload_id,
            content_type=vault_content_type or mime_type,
            checksum_sha256=checksum_sha256,
            client_encrypted=vault_client_encrypted,
            limit=vault_limit,
            cursor=vault_cursor,
            filename_contains=vault_filename_contains,
        )

    provided = sum(value is not None for value in (file_path, file_paths, inline_text))
    if provided != 1:
        raise ValueError("Provide exactly one of file_path, file_paths, or inline_text.")

    if action == FeatureType.combine_pdf:
        if file_paths is None:
            raise ValueError("combine_pdf requires file_paths.")
        return build_pdf_file_set_payload(
            file_paths,
            storage_keys=storage_keys,
            upload_ids=upload_ids,
            mime_types=mime_types,
            passwords=passwords,
            checksums_sha256=checksums_sha256,
        )

    if action in PDF_SINGLE_FILE_ACTIONS:
        if file_path is None:
            raise ValueError(f"{action.value} requires file_path.")
        return build_pdf_file_payload(
            file_path,
            storage_key=storage_key,
            upload_id=upload_id,
            mime_type=mime_type,
            password=password,
            checksum_sha256=checksum_sha256,
        )

    if inline_text is not None:
        return build_document_payload_for_action(action=action, inline_text=inline_text)

    if file_paths is not None:
        return build_document_set_payload(
            file_paths,
            action=action,
            extract_optional_text=extract_optional_text,
            ocr_languages=ocr_languages,
        )

    return build_document_payload_for_action(
        action=action,
        file_path=file_path,
        extract_optional_text=extract_optional_text,
        ocr_languages=ocr_languages,
    )


# Optional convenience alias for backend upload routers.
def build_pdf_input_artifact_for_action(
    *,
    action: FeatureType,
    file_path: Optional[Pathish] = None,
    file_paths: Optional[Sequence[Pathish]] = None,
    storage_key: Optional[str] = None,
    upload_id: Optional[str] = None,
    storage_keys: Optional[Sequence[Optional[str]]] = None,
    upload_ids: Optional[Sequence[Optional[str]]] = None,
    mime_type: Optional[str] = None,
    mime_types: Optional[Sequence[Optional[str]]] = None,
    password: Optional[str] = None,
    passwords: Optional[Sequence[Optional[str]]] = None,
    checksum_sha256: Optional[str] = None,
    checksums_sha256: Optional[Sequence[Optional[str]]] = None,
) -> Union[PdfFilePayload, PdfFileSetPayload]:
    """Build only PDF-tool, Lock PDF, and e-signature input artifacts."""
    if action not in PDF_DOCUMENT_ACTIONS:
        raise ValueError("This helper only supports PDF tool and e-signature actions.")
    artifact = build_input_artifact_for_action(
        action=action,
        file_path=file_path,
        file_paths=file_paths,
        storage_key=storage_key,
        upload_id=upload_id,
        storage_keys=storage_keys,
        upload_ids=upload_ids,
        mime_type=mime_type,
        mime_types=mime_types,
        password=password,
        passwords=passwords,
        checksum_sha256=checksum_sha256,
        checksums_sha256=checksums_sha256,
    )
    if not isinstance(artifact, (PdfFilePayload, PdfFileSetPayload)):
        raise TypeError("Expected PdfFilePayload or PdfFileSetPayload.")
    return artifact


__all__ = [
    "OCR_CONFIG",
    "OCR_SPARSE_CONFIG",
    "CONVERSION_ACTIONS",
    "TEXT_AI_DOC_INPUT_FORMATS",
    "TEXT_TO_SPEECH_INPUT_FORMATS",
    "REDACTION_MASKING_INPUT_FORMATS",
    "STRUCTURED_EXTRACTION_INPUT_FORMATS",
    "COMPLIANCE_INPUT_FORMATS",
    "CONVERSION_INPUT_FORMATS",
    "OPTIONAL_TEXT_DOCUMENT_ACTIONS",
    "DOCUMENT_SET_ACTIONS",
    "PDF_DOCUMENT_ACTIONS",
    "PDF_SINGLE_FILE_ACTIONS",
    "get_available_tesseract_languages",
    "resolve_ocr_lang",
    "preprocess_for_ocr",
    "extract_ocr_data",
    "ocr_text_from_data",
    "get_file_size_mb",
    "get_pdf_tool_file_size_mb",
    "detect_format",
    "guess_mime_type",
    "compute_sha256_hex",
    "build_vault_file_payload",
    "build_vault_item_reference_payload",
    "build_vault_query_payload",
    "build_vault_input_artifact",
    "inspect_pdf_metadata",
    "build_pdf_document_metadata",
    "build_pdf_file_payload",
    "build_pdf_file_set_payload",
    "extract_text_from_txt",
    "extract_text_from_docx",
    "extract_text_from_image",
    "extract_text_from_pdf_text",
    "extract_text_from_pdf_ocr",
    "extract_text_by_format",
    "count_words",
    "enforce_text_ai_word_contract",
    "build_inline_text_payload",
    "build_text_to_speech_inline_payload",
    "build_conversion_document_payload",
    "build_text_ai_document_payload",
    "build_text_to_speech_document_payload",
    "build_redaction_or_masking_document_payload",
    "build_structured_extraction_or_compliance_document_payload",
    "build_document_set_payload",
    "build_document_payload_for_action",
    "build_input_artifact_for_action",
    "build_pdf_input_artifact_for_action",
]
