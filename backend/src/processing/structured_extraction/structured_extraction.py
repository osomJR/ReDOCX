from __future__ import annotations

"""
Production-grade structured extraction engine for ReDOCX v1.

Purpose:
- own request validation, deterministic extraction, artifact generation, and response construction
  for the structured_extract feature outside analyzer.py
- remain aligned with schema.py, validation.py, extraction.py, and Product Contract v1
- support both single-document and document-set inputs
- provide document-class-aware extraction strategies and strict result-shape enforcement

Design principles:
- stateless and deterministic
- source-grounded: every extracted value is derived from the provided document text only
- no external knowledge or enrichment
- explicit human-review requirement before reliance/final export
- clean writer layer for JSON, CSV, and XLSX artifacts
"""

from dataclasses import dataclass, field
from collections import defaultdict, deque
from datetime import datetime, timezone
import csv
import hashlib
import json
import re
import secrets
import unicodedata
from pathlib import Path
from typing import Any, Iterable, Mapping, Optional, Protocol, Sequence, Union

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.src.extraction import extract_text_by_format
from backend.src.schema import (
    AnalyzerRequest,
    AnalyzerResponse,
    DocumentPayload,
    DocumentSetPayload,
    FeatureType,
    HumanReviewRequirement,
    OutputPolicy,
    StructuredDataOutputFormat,
    StructuredExtractionDocumentClass,
    StructuredExtractionRequest,
    StructuredExtractionResultShape,
)
from backend.src.validation import (
    build_structured_extraction_file_result,
    validate_analyzer_request,
    validate_analyzer_response,
)
try:
    from backend.src.storage.artifacts import LocalArtifactStorage, guess_content_type
except ImportError:  # pragma: no cover
    from backend.src.storage.artifacts import LocalArtifactStorage, guess_content_type


STRUCTURED_EXTRACTION_RULES = """
TASK: STRUCTURED EXTRACTION
RULES:
- Extract predefined or user-selected fields strictly from the provided document or document set
- Preserve source meaning without interpretation or invention
- No external knowledge may be introduced
- Output may be key-value fields, tables, row-based records, or machine-readable structured data
- Human review remains required before reliance or final export
""".strip()

MAX_EXTRACTED_FIELDS_PER_DOCUMENT = 5_000
MAX_TABULAR_ROWS_PER_REQUEST = 100_000
MAX_EXPORT_CELLS_PER_REQUEST = 5_000_000
MAX_PREVIEW_BYTES = 1_048_576
GENERIC_KEY_VALUE_PATTERN = re.compile(
    r"^\s*(?P<key>[^:|–—]{1,81}?)\s*(?::|[–—]|\||\s-\s)\s*(?P<value>\S.*?)\s*$",
    re.UNICODE,
)


# =========================
# Artifact and config models
# =========================


@dataclass(frozen=True)
class PersistedArtifact:
    file_name: str
    file_extension: str
    file_size_mb: float
    file_path: str
    storage_key: Optional[str] = None
    download_url: Optional[str] = None


@dataclass(frozen=True)
class StructuredExtractionConfig:
    algorithm_version: Optional[str] = "structured-extraction-v1.1.1"
    artifact_base_dir: str = "artifacts/structured_extraction"
    include_empty_selected_fields: bool = True
    max_context_excerpt_chars: int = 240
    max_searchable_characters_per_document: int = 10_000_000


@dataclass(frozen=True)
class SourceEvidence:
    source_document_index: int
    field_name: str
    value: str
    line_number: Optional[int] = None
    excerpt: Optional[str] = None

    def to_dict(self) -> dict[str, Any]:
        return {
            "source_document_index": self.source_document_index,
            "field_name": self.field_name,
            "value": self.value,
            "line_number": self.line_number,
            "excerpt": self.excerpt,
        }
@dataclass(frozen=True)
class StructuredExtractionExecution:
    response: AnalyzerResponse
    preview_payload: dict[str, Any]
    preview_rows: list[dict[str, Any]]

@dataclass(frozen=True)
class ExtractedField:
    name: str
    value: Any
    source_document_index: int
    confidence: float = 1.0
    evidence: list[SourceEvidence] = field(default_factory=list)
    source_kind: str = "field"

    def to_dict(self) -> dict[str, Any]:
        return {
            "name": self.name,
            "normalized_name": _normalize_field_name(self.name),
            "value": self.value,
            "source_document_index": self.source_document_index,
            "confidence": self.confidence,
            "confidence_basis": "deterministic_pattern_strength_not_probability",
            "source_kind": self.source_kind,
            "evidence": [item.to_dict() for item in self.evidence],
        }


@dataclass(frozen=True)
class ExtractedTable:
    table_index: int
    source_document_index: int
    columns: list[str]
    rows: list[dict[str, Any]]
    source_line_numbers: list[int] = field(default_factory=list)
    row_line_numbers: list[int] = field(default_factory=list)

    def to_dict(self) -> dict[str, Any]:
        return {
            "table_index": self.table_index,
            "source_document_index": self.source_document_index,
            "columns": list(self.columns),
            "rows": [dict(row) for row in self.rows],
            "source_line_numbers": list(self.source_line_numbers),
            "row_line_numbers": list(self.row_line_numbers),
        }


@dataclass(frozen=True)
class DocumentExtraction:
    source_document_index: int
    filename: Optional[str]
    input_format: str
    ocr_used: bool
    checksum_sha256: str
    extracted_character_count: int
    document_classes: list[str]
    fields: list[ExtractedField]
    tables: list[ExtractedTable]
    row_records: list[dict[str, Any]]
    requested_field_count: int = 0
    populated_requested_field_count: int = 0
    extraction_quality_score: float = 0.0
    warnings: list[str] = field(default_factory=list)

    def fields_as_mapping(self) -> dict[str, Any]:
        output: dict[str, Any] = {}
        for item in self.fields:
            output[item.name] = item.value
        return output

    def to_machine_readable(self) -> dict[str, Any]:
        return {
            "source_document_index": self.source_document_index,
            "filename": self.filename,
            "input_format": self.input_format,
            "ocr_used": self.ocr_used,
            "checksum_sha256": self.checksum_sha256,
            "extracted_character_count": self.extracted_character_count,
            "document_classes": list(self.document_classes),
            "fields": [field_item.to_dict() for field_item in self.fields],
            "tables": [table.to_dict() for table in self.tables],
            "row_based_records": [dict(row) for row in self.row_records],
            "quality": {
                "requested_field_count": self.requested_field_count,
                "populated_requested_field_count": self.populated_requested_field_count,
                "score": self.extraction_quality_score,
                "human_review_required": True,
            },
            "warnings": list(self.warnings),
        }


@dataclass(frozen=True)
class StructuredExtractionOutput:
    payload: dict[str, Any]
    tabular_rows: list[dict[str, Any]]
    documents: list[DocumentExtraction]


# =========================
# Strategy interfaces
# =========================


class DocumentClassExtractionStrategy(Protocol):
    document_class: StructuredExtractionDocumentClass

    def extract(
        self,
        *,
        text: str,
        lines: Sequence[str],
        source_document_index: int,
        selected_fields: Sequence[str],
        config: StructuredExtractionConfig,
    ) -> tuple[list[ExtractedField], list[dict[str, Any]], list[str]]:
        """Return fields, class-specific row records, warnings."""
        ...


class StructuredExtractionBackend(Protocol):
    def extract(
        self,
        *,
        documents: Sequence[DocumentPayload],
        result_shape: StructuredExtractionResultShape,
        selected_fields: Sequence[str],
        document_classes: Sequence[StructuredExtractionDocumentClass],
        config: StructuredExtractionConfig,
    ) -> StructuredExtractionOutput:
        ...


# =========================
# Deterministic backend
# =========================


class DeterministicStructuredExtractionBackend:
    """
    Rule-based, document-class-aware extraction backend.

    This backend is intentionally deterministic. It does not infer facts outside
    the document text. It extracts explicit key-value pairs, tabular structures,
    and class-specific fields/records using conservative regex patterns.
    """

    def __init__(self, strategies: Optional[Sequence[DocumentClassExtractionStrategy]] = None) -> None:
        provided = list(strategies or _default_strategies())
        self._strategies: dict[StructuredExtractionDocumentClass, DocumentClassExtractionStrategy] = {
            strategy.document_class: strategy for strategy in provided
        }

    def extract(
        self,
        *,
        documents: Sequence[DocumentPayload],
        result_shape: StructuredExtractionResultShape,
        selected_fields: Sequence[str],
        document_classes: Sequence[StructuredExtractionDocumentClass],
        config: StructuredExtractionConfig,
    ) -> StructuredExtractionOutput:
        if not documents:
            raise ValueError("structured_extract requires at least one document.")

        normalized_selected = _normalize_selected_fields(selected_fields)
        explicit_classes = _dedupe_document_classes(document_classes)

        document_outputs: list[DocumentExtraction] = []
        request_row_count = 0
        request_cell_count = 0
        for index, document in enumerate(documents):
            text, actual_ocr_used, source_checksum = _ensure_document_text(
                document, config=config
            )
            lines = list(_iter_lines(text))
            warnings = _source_extraction_warnings(lines)
            raw_tables = _extract_generic_tables(
                lines=lines, source_document_index=index, warnings=warnings
            )
            table_line_numbers = {
                number for table in raw_tables for number in table.source_line_numbers
            }
            # Keep positions for evidence, but never interpret table headers or
            # cells as standalone labels or scalar values.
            field_lines = [
                "" if number in table_line_numbers else line
                for number, line in enumerate(lines, start=1)
            ]
            resolved_classes = explicit_classes or _infer_document_classes(
                text=text,
                lines=lines,
                registered_classes=set(self._strategies),
            )
            auto_detected_classes = not explicit_classes

            generic_fields = _extract_generic_key_values(
                lines=field_lines,
                source_document_index=index,
                selected_fields=normalized_selected,
                config=config,
            )
            generic_fields.extend(
                _extract_requested_fields_across_lines(
                    lines=field_lines,
                    source_document_index=index,
                    selected_fields=normalized_selected,
                    existing_fields=generic_fields,
                    config=config,
                )
            )
            generic_record_rows = _extract_repeated_key_value_records(
                lines=field_lines,
                source_document_index=index,
                selected_fields=normalized_selected,
            )

            class_fields: list[ExtractedField] = []
            class_rows: list[dict[str, Any]] = []
            if auto_detected_classes:
                if resolved_classes == [StructuredExtractionDocumentClass.form]:
                    warnings.append(
                        "No specialized document type was identified; generic extraction was used."
                    )
                else:
                    warnings.append(
                        "Auto-detected document classes: "
                        + ", ".join(item.value for item in resolved_classes)
                        + ". Confirm this classification during review."
                    )

            for document_class in resolved_classes:
                strategy = self._strategies.get(document_class)
                if strategy is None:
                    warnings.append(f"No specialized extractor registered for document_class='{document_class.value}'.")
                    continue
                fields, rows, strategy_warnings = strategy.extract(
                    text=text,
                    lines=field_lines,
                    source_document_index=index,
                    selected_fields=normalized_selected,
                    config=config,
                )
                class_fields.extend(fields)
                class_rows.extend(rows)
                warnings.extend(strategy_warnings)

            fields = _merge_fields(
                [
                    *generic_fields,
                    *class_fields,
                    *_selected_table_fields(raw_tables, normalized_selected, lines, config),
                ],
                selected_fields=normalized_selected,
                include_empty_selected_fields=config.include_empty_selected_fields,
                source_document_index=index,
            )
            populated_requested = sum(
                1
                for field_item in fields
                if _normalize_field_name(field_item.name) in normalized_selected
                and _clean_value(field_item.value)
            )
            if normalized_selected:
                quality_score = populated_requested / len(normalized_selected)
                if populated_requested < len(normalized_selected):
                    warnings.append(
                        f"{len(normalized_selected) - populated_requested} of "
                        f"{len(normalized_selected)} requested fields were not found."
                    )
            else:
                extracted_items = len([item for item in fields if _clean_value(item.value)])
                quality_score = 1.0 if extracted_items or raw_tables or class_rows or generic_record_rows else 0.0
                if quality_score == 0.0:
                    warnings.append(
                        "No reliable key-value fields, tables, or records were detected."
                    )
            if actual_ocr_used:
                warnings.append(
                    "OCR was used; confirm names, identifiers, dates, and amounts against the source."
                )
            row_records = _derive_row_records(
                source_document_index=index,
                fields=fields,
                tables=raw_tables,
                class_rows=[*generic_record_rows, *class_rows],
            )
            row_records = _project_record_rows(row_records, normalized_selected)
            if not row_records:
                row_records = _derive_row_records(
                    source_document_index=index, fields=fields, tables=[], class_rows=[]
                )
            row_records = _with_document_fields(row_records, fields)
            generic_tables = _project_tables(raw_tables, normalized_selected)
            request_row_count += len(row_records)
            if request_row_count > MAX_TABULAR_ROWS_PER_REQUEST:
                raise ValueError(
                    "Structured extraction exceeds the 100,000-row request limit; "
                    "no partial export was created. Narrow the document set."
                )
            display_name = Path(document.filename or f"document-{index + 1}").name
            row_records = _attach_row_provenance(
                row_records,
                filename=display_name,
                input_format=document.metadata.input_format.value,
                checksum_sha256=source_checksum,
                source_document_index=index,
            )
            request_cell_count += sum(len(row) for row in row_records)
            if request_cell_count > MAX_EXPORT_CELLS_PER_REQUEST:
                raise ValueError('Structured records exceed the request cell budget; no partial export was created. Narrow the request.')
            document_outputs.append(
                DocumentExtraction(
                    source_document_index=index,
                    filename=display_name,
                    input_format=document.metadata.input_format.value,
                    ocr_used=actual_ocr_used,
                    checksum_sha256=source_checksum,
                    extracted_character_count=len(text),
                    document_classes=[item.value for item in resolved_classes],
                    fields=fields,
                    tables=generic_tables,
                    row_records=row_records,
                    requested_field_count=len(normalized_selected),
                    populated_requested_field_count=populated_requested,
                    extraction_quality_score=round(quality_score, 4),
                    warnings=list(dict.fromkeys(warnings)),
                )
            )

        shaped_payload = ShapeEnforcer.enforce(
            result_shape=result_shape,
            selected_fields=list(selected_fields),
            documents=document_outputs,
        )
        tabular_rows = ShapeEnforcer.to_tabular_rows(
            result_shape=result_shape,
            selected_fields=list(selected_fields),
            documents=document_outputs,
        )
        if len(tabular_rows) > MAX_TABULAR_ROWS_PER_REQUEST:
            raise ValueError(
                f"Structured extraction produced more than {MAX_TABULAR_ROWS_PER_REQUEST:,} "
                "rows. Narrow the document set or requested result."
            )
        return StructuredExtractionOutput(
            payload=shaped_payload,
            tabular_rows=tabular_rows,
            documents=document_outputs,
        )

def _rx(pattern: str) -> re.Pattern[str]:
    return re.compile(pattern, re.IGNORECASE)


AUTO_DETECT_CLASS_PRIORITY: tuple[StructuredExtractionDocumentClass, ...] = (
    StructuredExtractionDocumentClass.invoice,
    StructuredExtractionDocumentClass.receipt,
    StructuredExtractionDocumentClass.bank_statement,
    StructuredExtractionDocumentClass.kyc_document,
    StructuredExtractionDocumentClass.id_document,
    StructuredExtractionDocumentClass.contract,
    StructuredExtractionDocumentClass.legal_record,
    StructuredExtractionDocumentClass.procurement_document,
    StructuredExtractionDocumentClass.insurance_document,
    StructuredExtractionDocumentClass.medical_record,
    StructuredExtractionDocumentClass.hr_record,
    StructuredExtractionDocumentClass.onboarding_document,
    StructuredExtractionDocumentClass.incident_report,
    StructuredExtractionDocumentClass.technical_report,
    StructuredExtractionDocumentClass.ticket,
    StructuredExtractionDocumentClass.memo,
    StructuredExtractionDocumentClass.form,
)

AUTO_DETECT_MAX_MATCHED_CLASSES = 2


def _dedupe_document_classes(
    document_classes: Sequence[StructuredExtractionDocumentClass],
) -> list[StructuredExtractionDocumentClass]:
    output: list[StructuredExtractionDocumentClass] = []
    for item in document_classes:
        if item not in output:
            output.append(item)
    return output


def _infer_document_classes(
    *,
    text: str,
    lines: Sequence[str],
    registered_classes: set[StructuredExtractionDocumentClass],
) -> list[StructuredExtractionDocumentClass]:
    """Infer document classes when the request leaves document_classes empty.

    Empty document_classes means conservative auto-detection. A specialized
    extractor runs only when multiple class signals or a distinctive document
    signature is present. Ambiguous prose falls back to the generic form parser;
    it must never trigger every specialized extractor.
    """
    del lines
    normalized_text = text.lower()
    signal_patterns: Mapping[StructuredExtractionDocumentClass, Sequence[re.Pattern[str]]] = {
        StructuredExtractionDocumentClass.invoice: (
            _rx(r"\binvoice\b"),
            _rx(r"\binv(?:oice)?\s*(?:no\.?|number|#)"),
            _rx(r"\bamount\s+due\b"),
            _rx(r"\bdue\s+date\b"),
        ),
        StructuredExtractionDocumentClass.receipt: (
            _rx(r"\breceipt\b"),
            _rx(r"\btransaction\s+date\b"),
            _rx(r"\bamount\s+paid\b"),
            _rx(r"\bmerchant\b"),
        ),
        StructuredExtractionDocumentClass.bank_statement: (
            _rx(r"\bbank\s+statement\b"),
            _rx(r"\baccount\s+(?:no\.?|number|name)\b"),
            _rx(r"\bopening\s+balance\b"),
            _rx(r"\bclosing\s+balance\b"),
            _rx(r"\bstatement\s+period\b"),
        ),
        StructuredExtractionDocumentClass.kyc_document: (
            _rx(r"\bkyc\b"),
            _rx(r"\bknow\s+your\s+customer\b"),
            _rx(r"\bnational\s+id\b"),
            _rx(r"\bnin\b"),
            _rx(r"\bdate\s+of\s+birth\b"),
        ),
        StructuredExtractionDocumentClass.id_document: (
            _rx(r"\bidentity\s+(?:card|document)\b"),
            _rx(r"\bpassport\b"),
            _rx(r"\bdriver'?s\s+licen[cs]e\b"),
            _rx(r"\bid\s+(?:no\.?|number)\b"),
        ),
        StructuredExtractionDocumentClass.contract: (
            _rx(r"\bagreement\b"),
            _rx(r"\bcontract\b"),
            _rx(r"\bgoverning\s+law\b"),
            _rx(r"\beffective\s+date\b"),
            _rx(r"\bparty\s+[ab]\b"),
        ),
        StructuredExtractionDocumentClass.legal_record: (
            _rx(r"\blegal\b"),
            _rx(r"\bcourt\b"),
            _rx(r"\bcase\s+(?:no\.?|number)\b"),
            _rx(r"\bclause\b"),
        ),
        StructuredExtractionDocumentClass.procurement_document: (
            _rx(r"\bpurchase\s+order\b"),
            _rx(r"\bpo\s*(?:no\.?|number|#)\b"),
            _rx(r"\bvendor\b"),
            _rx(r"\bsupplier\b"),
            _rx(r"\bdelivery\s+date\b"),
        ),
        StructuredExtractionDocumentClass.insurance_document: (
            _rx(r"\bpolicy\s*(?:no\.?|number|#)\b"),
            _rx(r"\binsured\b"),
            _rx(r"\bpremium\b"),
            _rx(r"\bcoverage\s+period\b"),
            _rx(r"\bclaim\s*(?:no\.?|number|#)\b"),
        ),
        StructuredExtractionDocumentClass.medical_record: (
            _rx(r"\bpatient\b"),
            _rx(r"\bdiagnosis\b"),
            _rx(r"\bmedical\s+record\b"),
            _rx(r"\bphysician\b"),
        ),
        StructuredExtractionDocumentClass.hr_record: (
            _rx(r"\bemployee\b"),
            _rx(r"\bstaff\s+id\b"),
            _rx(r"\bdepartment\b"),
            _rx(r"\bjob\s+title\b"),
        ),
        StructuredExtractionDocumentClass.onboarding_document: (
            _rx(r"\bonboarding\b"),
            _rx(r"\bnew\s+hire\b"),
            _rx(r"\bemployment\s+date\b"),
        ),
        StructuredExtractionDocumentClass.incident_report: (
            _rx(r"\bincident\b"),
            _rx(r"\bseverity\b"),
            _rx(r"\bincident\s+(?:date|location)\b"),
        ),
        StructuredExtractionDocumentClass.technical_report: (
            _rx(r"\btechnical\s+report\b"),
            _rx(r"\breport\s+title\b"),
            _rx(r"\bprepared\s+by\b"),
            _rx(r"\bexecutive\s+summary\b"),
        ),
        StructuredExtractionDocumentClass.ticket: (
            _rx(r"\bticket\b"),
            _rx(r"\bpriority\b"),
            _rx(r"\bassignee\b"),
            _rx(r"\bstatus\b"),
        ),
        StructuredExtractionDocumentClass.memo: (
            _rx(r"(?m)^\s*to\s*[:\-]"),
            _rx(r"(?m)^\s*from\s*[:\-]"),
            _rx(r"(?m)^\s*(?:subject|re)\s*[:\-]"),
        ),
    }

    distinctive_patterns: Mapping[
        StructuredExtractionDocumentClass, re.Pattern[str]
    ] = {
        StructuredExtractionDocumentClass.invoice: _rx(r"\binvoice\b"),
        StructuredExtractionDocumentClass.receipt: _rx(r"\breceipt\b"),
        StructuredExtractionDocumentClass.bank_statement: _rx(r"\bbank\s+statement\b"),
        StructuredExtractionDocumentClass.kyc_document: _rx(r"\b(?:kyc|know\s+your\s+customer)\b"),
        StructuredExtractionDocumentClass.id_document: _rx(
            r"\b(?:identity\s+(?:card|document)|passport|driver'?s\s+licen[cs]e)\b"
        ),
        StructuredExtractionDocumentClass.contract: _rx(r"\b(?:agreement|contract)\b"),
        StructuredExtractionDocumentClass.procurement_document: _rx(r"\bpurchase\s+order\b"),
        StructuredExtractionDocumentClass.insurance_document: _rx(
            r"\b(?:insurance\s+policy|policy\s*(?:no\.?|number|#))\b"
        ),
        StructuredExtractionDocumentClass.technical_report: _rx(r"\btechnical\s+report\b"),
        StructuredExtractionDocumentClass.incident_report: _rx(r"\bincident\s+report\b"),
        StructuredExtractionDocumentClass.ticket: _rx(r"\b(?:support\s+ticket|ticket\s*(?:id|no\.?|number|#))\b"),
    }

    scores: dict[StructuredExtractionDocumentClass, int] = {}
    for document_class, patterns in signal_patterns.items():
        if document_class not in registered_classes:
            continue
        score = sum(1 for pattern in patterns if pattern.search(normalized_text))
        distinctive = distinctive_patterns.get(document_class)
        if score >= 2 or (distinctive is not None and distinctive.search(normalized_text)):
            scores[document_class] = score

    if scores:
        ranked = sorted(
            scores,
            key=lambda item: (-scores[item], AUTO_DETECT_CLASS_PRIORITY.index(item)),
        )
        return ranked[:AUTO_DETECT_MAX_MATCHED_CLASSES]

    if StructuredExtractionDocumentClass.form in registered_classes:
        return [StructuredExtractionDocumentClass.form]
    return []


def _money_rx(label_pattern: str) -> re.Pattern[str]:
    # Match a complete labeled amount. A word boundary after an optional
    # two-digit fraction used to accept only the integer prefix of 1234.5.
    currency = r"(?:NGN|₦|USD|\$|EUR|€|GBP|£)"
    number = r"[+-]?[0-9](?:[0-9,.\u00a0\u202f]*[0-9])?"
    money = rf"(?:{currency}\s*)?(?:{number}|\({number}\))(?:\s*{currency})?"
    return _rx(rf"^\s*(?:{label_pattern})\s*[:\-]?\s*(?P<value>{money})\s*$")

# =========================
# Document-class strategies
# =========================

class BaseRegexStrategy:
    document_class: StructuredExtractionDocumentClass
    field_patterns: Mapping[str, Sequence[re.Pattern[str]]] = {}

    def extract(
        self,
        *,
        text: str,
        lines: Sequence[str],
        source_document_index: int,
        selected_fields: Sequence[str],
        config: StructuredExtractionConfig,
    ) -> tuple[list[ExtractedField], list[dict[str, Any]], list[str]]:
        del text
        fields: list[ExtractedField] = []
        warnings: list[str] = []
        for field_name, patterns in self.field_patterns.items():
            if selected_fields and _normalize_field_name(field_name) not in selected_fields:
                continue
            match = _first_regex_match(lines, patterns, field_name=field_name)
            if match is None:
                continue
            value, line_number, excerpt = match
            fields.append(
                ExtractedField(
                    name=field_name,
                    value=value,
                    source_document_index=source_document_index,
                    confidence=0.92,
                    evidence=[
                        SourceEvidence(
                            source_document_index=source_document_index,
                            field_name=field_name,
                            value=value,
                            line_number=line_number,
                            excerpt=_truncate(excerpt, config.max_context_excerpt_chars),
                        )
                    ],
                )
            )
        return fields, [], warnings


class InvoiceStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.invoice
    field_patterns = {
        "invoice_number": [_rx(r"\b(?:invoice|inv)\s*(?:no\.?|number|#)\s*[:#\-]?\s*(?P<value>[A-Z0-9][A-Z0-9\-\/]+)\b")],
        "invoice_date": [
            _rx(r"^\s*invoice\s+date\s*[:\-]?\s*(?P<value>.+?)\s*$"),
            _rx(r"^\s*date\s*[:\-]?\s*(?P<value>.+?)\s*$"),
        ],
        "due_date": [_rx(r"\bdue\s+date\s*[:\-]?\s*(?P<value>\d{1,2}[\/\-.]\d{1,2}[\/\-.]\d{2,4}|[A-Za-z]+\s+\d{1,2},?\s+\d{4})")],
        "subtotal": [_money_rx("subtotal")],
        "tax": [_money_rx("tax|vat")],
        "total": [_money_rx(r"total|amount\s+due|balance\s+due")],
    }

    def extract(self, **kwargs: Any) -> tuple[list[ExtractedField], list[dict[str, Any]], list[str]]:
        fields, rows, warnings = super().extract(**kwargs)
        rows.extend(_extract_line_item_rows(kwargs["lines"], kwargs["source_document_index"]))
        return fields, rows, warnings


class ReceiptStrategy(InvoiceStrategy):
    document_class = StructuredExtractionDocumentClass.receipt
    field_patterns = {
        "receipt_number": [_rx(r"\b(?:receipt|rcpt)\s*(?:no\.?|number|#)\s*[:#\-]?\s*(?P<value>[A-Z0-9][A-Z0-9\-\/]+)\b")],
        "transaction_date": [_rx(r"\b(?:date|transaction\s+date)\s*[:\-]?\s*(?P<value>\d{1,2}[\/\-.]\d{1,2}[\/\-.]\d{2,4}|[A-Za-z]+\s+\d{1,2},?\s+\d{4})")],
        "merchant": [_rx(r"\b(?:merchant|vendor|seller)\s*[:\-]?\s*(?P<value>.+)$")],
        "subtotal": [_money_rx("subtotal")],
        "tax": [_money_rx("tax|vat")],
        "total": [_money_rx(r"total|amount\s+paid")],
    }


class BankStatementStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.bank_statement
    field_patterns = {
        "account_name": [_rx(r"\baccount\s+name\s*[:\-]?\s*(?P<value>.+)$")],
        "account_number": [_rx(r"\baccount\s+(?:no\.?|number)\s*[:#\-]?\s*(?P<value>[0-9Xx*\-\s]{5,})")],
        "statement_period": [_rx(r"\b(?:statement\s+period|period)\s*[:\-]?\s*(?P<value>.+)$")],
        "opening_balance": [_money_rx(r"opening\s+balance")],
        "closing_balance": [_money_rx(r"closing\s+balance")],
    }

    def extract(self, **kwargs: Any) -> tuple[list[ExtractedField], list[dict[str, Any]], list[str]]:
        fields, rows, warnings = super().extract(**kwargs)
        rows.extend(_extract_transaction_rows(kwargs["lines"], kwargs["source_document_index"]))
        return fields, rows, warnings


class KycDocumentStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.kyc_document
    field_patterns = {
        "full_name": [_rx(r"\b(?:full\s+name|name)\s*[:\-]?\s*(?P<value>[A-Za-z][A-Za-z\s.'-]{2,})$")],
        "date_of_birth": [_rx(r"\b(?:date\s+of\s+birth|dob)\s*[:\-]?\s*(?P<value>\d{1,2}[\/\-.]\d{1,2}[\/\-.]\d{2,4})")],
        "national_id": [_rx(r"\b(?:national\s+id|nin|id\s+number)\s*[:#\-]?\s*(?P<value>[A-Z0-9\-]{5,})")],
        "phone_number": [_rx(r"\b(?:phone|mobile|telephone)\s*[:\-]?\s*(?P<value>\+?[0-9][0-9\s()\-]{7,})")],
        "email_address": [_rx(r"\b(?:email|e-mail)\s*[:\-]?\s*(?P<value>[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})")],
        "address": [_rx(r"\b(?:address|residential\s+address)\s*[:\-]?\s*(?P<value>.+)$")],
    }


class IdDocumentStrategy(KycDocumentStrategy):
    document_class = StructuredExtractionDocumentClass.id_document


class ContractStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.contract
    field_patterns = {
        "effective_date": [_rx(r"\b(?:effective\s+date|commencement\s+date)\s*[:\-]?\s*(?P<value>.+)$")],
        "termination_date": [_rx(r"\b(?:termination\s+date|expiry\s+date|expiration\s+date)\s*[:\-]?\s*(?P<value>.+)$")],
        "governing_law": [_rx(r"\bgoverning\s+law\s*[:\-]?\s*(?P<value>.+)$")],
        "party_a": [_rx(r"\b(?:party\s+a|first\s+party|between)\s*[:\-]?\s*(?P<value>.+)$")],
        "party_b": [_rx(r"\b(?:party\s+b|second\s+party|and)\s*[:\-]?\s*(?P<value>.+)$")],
    }

    def extract(self, **kwargs: Any) -> tuple[list[ExtractedField], list[dict[str, Any]], list[str]]:
        fields, rows, warnings = super().extract(**kwargs)
        rows.extend(_extract_clause_rows(kwargs["lines"], kwargs["source_document_index"]))
        return fields, rows, warnings


class LegalRecordStrategy(ContractStrategy):
    document_class = StructuredExtractionDocumentClass.legal_record


class MedicalRecordStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.medical_record
    field_patterns = {
        "patient_name": [_rx(r"\b(?:patient\s+name|name)\s*[:\-]?\s*(?P<value>[A-Za-z][A-Za-z\s.'-]{2,})$")],
        "patient_id": [_rx(r"\b(?:patient\s+id|hospital\s+number|medical\s+record\s+number)\s*[:#\-]?\s*(?P<value>[A-Z0-9\-\/]{3,})")],
        "date_of_birth": [_rx(r"\b(?:date\s+of\s+birth|dob)\s*[:\-]?\s*(?P<value>\d{1,2}[\/\-.]\d{1,2}[\/\-.]\d{2,4})")],
        "diagnosis": [_rx(r"\bdiagnosis\s*[:\-]?\s*(?P<value>.+)$")],
        "provider": [_rx(r"\b(?:doctor|physician|provider|consultant)\s*[:\-]?\s*(?P<value>.+)$")],
    }


class ProcurementStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.procurement_document
    field_patterns = {
        "purchase_order_number": [_rx(r"\b(?:purchase\s+order|po)\s*(?:no\.?|number|#)?\s*[:#\-]?\s*(?P<value>[A-Z0-9\-\/]+)")],
        "vendor": [_rx(r"\b(?:vendor|supplier)\s*[:\-]?\s*(?P<value>.+)$")],
        "delivery_date": [_rx(r"\b(?:delivery\s+date|required\s+by)\s*[:\-]?\s*(?P<value>.+)$")],
        "total": [_money_rx(r"total|grand\s+total|amount")],
    }

    def extract(self, **kwargs: Any) -> tuple[list[ExtractedField], list[dict[str, Any]], list[str]]:
        fields, rows, warnings = super().extract(**kwargs)
        rows.extend(_extract_line_item_rows(kwargs["lines"], kwargs["source_document_index"]))
        return fields, rows, warnings


class InsuranceStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.insurance_document
    field_patterns = {
        "policy_number": [_rx(r"\bpolicy\s*(?:no\.?|number|#)\s*[:#\-]?\s*(?P<value>[A-Z0-9\-\/]+)")],
        "insured_name": [_rx(r"\b(?:insured|policyholder)\s*(?:name)?\s*[:\-]?\s*(?P<value>.+)$")],
        "premium": [_money_rx("premium")],
        "coverage_period": [_rx(r"\b(?:coverage\s+period|policy\s+period)\s*[:\-]?\s*(?P<value>.+)$")],
        "claim_number": [_rx(r"\bclaim\s*(?:no\.?|number|#)\s*[:#\-]?\s*(?P<value>[A-Z0-9\-\/]+)")],
    }


class HRRecordStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.hr_record
    field_patterns = {
        "employee_name": [_rx(r"\b(?:employee\s+name|staff\s+name|name)\s*[:\-]?\s*(?P<value>[A-Za-z][A-Za-z\s.'-]{2,})$")],
        "employee_id": [_rx(r"\b(?:employee\s+id|staff\s+id)\s*[:#\-]?\s*(?P<value>[A-Z0-9\-\/]+)")],
        "department": [_rx(r"\bdepartment\s*[:\-]?\s*(?P<value>.+)$")],
        "job_title": [_rx(r"\b(?:job\s+title|role|position)\s*[:\-]?\s*(?P<value>.+)$")],
        "start_date": [_rx(r"\b(?:start\s+date|employment\s+date|hire\s+date)\s*[:\-]?\s*(?P<value>.+)$")],
    }


class OnboardingStrategy(HRRecordStrategy):
    document_class = StructuredExtractionDocumentClass.onboarding_document


class TicketStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.ticket
    field_patterns = {
        "ticket_id": [_rx(r"\b(?:ticket|case|issue)\s*(?:id|no\.?|number|#)?\s*[:#\-]?\s*(?P<value>[A-Z0-9\-\/]+)")],
        "status": [_rx(r"\bstatus\s*[:\-]?\s*(?P<value>.+)$")],
        "priority": [_rx(r"\bpriority\s*[:\-]?\s*(?P<value>.+)$")],
        "assignee": [_rx(r"\b(?:assignee|assigned\s+to|owner)\s*[:\-]?\s*(?P<value>.+)$")],
        "created_date": [_rx(r"\b(?:created|opened|reported)\s*(?:date)?\s*[:\-]?\s*(?P<value>.+)$")],
    }


class ReportStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.technical_report
    field_patterns = {
        "report_title": [_rx(r"\b(?:report\s+title|title)\s*[:\-]?\s*(?P<value>.+)$")],
        "report_date": [_rx(r"\b(?:report\s+date|date)\s*[:\-]?\s*(?P<value>.+)$")],
        "author": [_rx(r"\b(?:author|prepared\s+by|reported\s+by)\s*[:\-]?\s*(?P<value>.+)$")],
        "summary": [_rx(r"\b(?:summary|executive\s+summary)\s*[:\-]?\s*(?P<value>.+)$")],
    }


class IncidentReportStrategy(ReportStrategy):
    document_class = StructuredExtractionDocumentClass.incident_report
    field_patterns = {
        **ReportStrategy.field_patterns,
        "incident_date": [_rx(r"\bincident\s+date\s*[:\-]?\s*(?P<value>.+)$")],
        "incident_location": [_rx(r"\b(?:incident\s+location|location)\s*[:\-]?\s*(?P<value>.+)$")],
        "severity": [_rx(r"\bseverity\s*[:\-]?\s*(?P<value>.+)$")],
    }


class GenericFormStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.form


class MemoStrategy(BaseRegexStrategy):
    document_class = StructuredExtractionDocumentClass.memo
    field_patterns = {
        "to": [_rx(r"^to\s*[:\-]?\s*(?P<value>.+)$")],
        "from": [_rx(r"^from\s*[:\-]?\s*(?P<value>.+)$")],
        "date": [_rx(r"^date\s*[:\-]?\s*(?P<value>.+)$")],
        "subject": [_rx(r"^(?:subject|re)\s*[:\-]?\s*(?P<value>.+)$")],
    }


# =========================
# Shape enforcement
# =========================


class ShapeEnforcer:
    @staticmethod
    def enforce(
        *,
        result_shape: StructuredExtractionResultShape,
        selected_fields: Sequence[str],
        documents: Sequence[DocumentExtraction],
    ) -> dict[str, Any]:
        base = {
            "contract_version": "v1",
            "feature": FeatureType.structured_extract.value,
            "result_shape": result_shape.value,
            "selected_fields": list(selected_fields),
            "selected_field_status": _selected_field_status(documents, selected_fields),
            "generated_at": _timestamp_iso(),
            "human_review_required": True,
            "extraction_method": "deterministic_source_grounded",
            "source_documents": [
                {
                    "source_document_index": doc.source_document_index,
                    "filename": doc.filename,
                    "input_format": doc.input_format,
                    "checksum_sha256": doc.checksum_sha256,
                    "ocr_used": doc.ocr_used,
                    "extracted_character_count": doc.extracted_character_count,
                }
                for doc in documents
            ],
            "quality_summary": _quality_summary(documents),
            "document_warnings": [
                {
                    "source_document_index": doc.source_document_index,
                    "filename": doc.filename,
                    "warnings": list(doc.warnings),
                }
                for doc in documents
                if doc.warnings
            ],
            "reliance_notice": (
                "Extracted values are source-grounded candidates, not verified facts. "
                "A person must compare important fields with the source before reliance."
            ),
        }

        if result_shape in {StructuredExtractionResultShape.row_based_records, StructuredExtractionResultShape.tables}:
            base["document_fields"] = [
                {"source_document_index": doc.source_document_index, "filename": doc.filename,
                 "fields": [item.to_dict() for item in doc.fields]}
                for doc in documents
            ]
        if result_shape in {StructuredExtractionResultShape.row_based_records, StructuredExtractionResultShape.key_value_fields}:
            base["extracted_tables"] = [table.to_dict() for doc in documents for table in doc.tables]

        if result_shape == StructuredExtractionResultShape.key_value_fields:
            base["documents"] = [
                {
                    "source_document_index": doc.source_document_index,
                    "filename": doc.filename,
                    "fields": doc.fields_as_mapping(),
                    "evidence": [evidence.to_dict() for field_item in doc.fields for evidence in field_item.evidence],
                    "warnings": list(doc.warnings),
                }
                for doc in documents
            ]
            return base

        if result_shape == StructuredExtractionResultShape.tables:
            base["documents"] = [
                {
                    "source_document_index": doc.source_document_index,
                    "filename": doc.filename,
                    "tables": [table.to_dict() for table in doc.tables],
                    "warnings": list(doc.warnings),
                }
                for doc in documents
            ]
            return base

        if result_shape == StructuredExtractionResultShape.row_based_records:
            rows: list[dict[str, Any]] = []
            for doc in documents:
                rows.extend([dict(row) for row in doc.row_records])
            base["rows"] = rows
            return base

        if result_shape == StructuredExtractionResultShape.machine_readable:
            base["documents"] = [doc.to_machine_readable() for doc in documents]
            base["aggregate"] = _aggregate_documents(documents)
            return base

        raise ValueError(f"Unsupported result_shape: {result_shape.value}")

    @staticmethod
    def to_tabular_rows(
        *,
        result_shape: StructuredExtractionResultShape,
        selected_fields: Sequence[str],
        documents: Sequence[DocumentExtraction],
    ) -> list[dict[str, Any]]:
        if result_shape == StructuredExtractionResultShape.key_value_fields:
            return [_document_fields_to_row(doc, selected_fields=selected_fields) for doc in documents]

        if result_shape == StructuredExtractionResultShape.tables:
            rows: list[dict[str, Any]] = []
            for doc in documents:
                if not doc.tables:
                    rows.append(
                        {
                            "source_document_index": doc.source_document_index,
                            "filename": doc.filename,
                            "input_format": doc.input_format,
                            "source_checksum_sha256": doc.checksum_sha256,
                        }
                    )
                    continue
                for table in doc.tables:
                    for row in table.rows:
                        normalized = {
                            "source_document_index": doc.source_document_index,
                            "filename": doc.filename,
                            "input_format": doc.input_format,
                            "source_checksum_sha256": doc.checksum_sha256,
                            "table_index": table.table_index,
                        }
                        normalized["record_type"] = "table_row"
                        _add_business_values(normalized, row)
                        rows.append(normalized)
            return rows

        if result_shape == StructuredExtractionResultShape.row_based_records:
            rows = []
            for doc in documents:
                rows.extend([dict(row) for row in doc.row_records])
            return rows or [{"source_document_index": 0}]

        # Full machine-readable spreadsheet exports include actual records,
        # not merely counts. Document fields/tables/evidence are also retained
        # in their existing payload or companion workbook sheets.
        return [dict(row) for doc in documents for row in doc.row_records]


# =========================
# Writer layer
# =========================


class StructuredArtifactWriter(Protocol):
    def write(
        self,
        *,
        output_format: StructuredDataOutputFormat,
        payload: Mapping[str, Any],
        rows: Sequence[Mapping[str, Any]],
        stem: str,
    ) -> PersistedArtifact:
        ...


class LocalStructuredArtifactWriter:
    def __init__(self, base_dir: str) -> None:
        self.base_dir = Path(base_dir)
        self.base_dir.mkdir(parents=True, exist_ok=True)

    def write(
        self,
        *,
        output_format: StructuredDataOutputFormat,
        payload: Mapping[str, Any],
        rows: Sequence[Mapping[str, Any]],
        stem: str,
    ) -> PersistedArtifact:
        if output_format == StructuredDataOutputFormat.json:
            return self._write_json(payload=payload, file_name=f"{stem}.json")
        if output_format == StructuredDataOutputFormat.csv:
            return self._write_csv(rows=rows, file_name=f"{stem}.csv", payload=payload)
        if output_format == StructuredDataOutputFormat.xlsx:
            return self._write_xlsx(payload=payload, rows=rows, file_name=f"{stem}.xlsx")
        raise ValueError(f"Unsupported structured extraction output format: {output_format.value}")

    def _write_json(self, *, payload: Mapping[str, Any], file_name: str) -> PersistedArtifact:
        target = self.base_dir / _safe_filename(file_name)
        temporary = _temporary_artifact_path(target)
        try:
            temporary.write_text(
                json.dumps(payload, ensure_ascii=False, indent=2, default=str),
                encoding="utf-8",
            )
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)
        return _artifact_from_path(target)

    def _write_csv(self, *, rows: Sequence[Mapping[str, Any]], file_name: str,
                   payload: Optional[Mapping[str, Any]] = None) -> PersistedArtifact:
        target = self.base_dir / _safe_filename(file_name)
        export_rows = _complete_csv_rows(payload or {}, rows)
        headers = _check_export_rows(export_rows)
        temporary = _temporary_artifact_path(target)
        try:
            with temporary.open('w', encoding='utf-8', newline='') as handle:
                writer = csv.writer(handle)
                writer.writerow([_safe_spreadsheet_cell(header) for header in headers])
                for row in export_rows:
                    rendered = _stringify_row(row)
                    writer.writerow([rendered.get(header, '') for header in headers])
            temporary.replace(target)
        finally:
            temporary.unlink(missing_ok=True)
        return _artifact_from_path(target)

    def _write_xlsx(
        self, *, payload: Mapping[str, Any], rows: Sequence[Mapping[str, Any]], file_name: str,
    ) -> PersistedArtifact:
        target = self.base_dir / _safe_filename(file_name)
        sheets = [('records', rows), ('fields', _field_rows_from_payload(payload)),
                  ('tables', _table_rows_from_payload(payload)),
                  ('evidence', _evidence_rows_from_payload(payload))]
        cell_count = 0
        for _, data in sheets:
            headers = _check_export_rows(data, xlsx=True)
            cell_count += (len(data) + 1) * len(headers)
        if cell_count > MAX_EXPORT_CELLS_PER_REQUEST:
            raise ValueError('The complete workbook exceeds the export cell budget. Narrow the request or use JSON; no partial export was created.')
        workbook = Workbook(write_only=True)
        temporary = _temporary_artifact_path(target)
        try:
            _write_summary_sheet(workbook.create_sheet('summary'), payload)
            for name, data in sheets:
                _write_rows_sheet(workbook.create_sheet(name), data)
            workbook.save(temporary)
            temporary.replace(target)
        finally:
            workbook.close()
            temporary.unlink(missing_ok=True)
        return _artifact_from_path(target)


# Backward-compatible alias from the earlier scaffold.
LocalStructuredExtractionArtifactStore = LocalStructuredArtifactWriter


# =========================
# Engine facade
# =========================

class StructuredExtractionEngine:
    def __init__(
        self,
        backend: Optional[StructuredExtractionBackend] = None,
        writer: Optional[StructuredArtifactWriter] = None,
        config: Optional[StructuredExtractionConfig] = None,
    ) -> None:
        self.config = config or StructuredExtractionConfig()
        self.backend = backend or DeterministicStructuredExtractionBackend()
        self.writer = writer or LocalStructuredArtifactWriter(self.config.artifact_base_dir)

    def _validate_engine_scope(self, request: AnalyzerRequest) -> None:
        if request.action != FeatureType.structured_extract:
            raise ValueError(
                "StructuredExtractionEngine only handles the structured_extract action."
            )

        if not isinstance(request.payload, StructuredExtractionRequest):
            raise ValueError(
                "structured_extract requires StructuredExtractionRequest payload."
            )

        if not isinstance(request.policy, OutputPolicy):
            raise ValueError("structured_extract requires OutputPolicy.")

        if not isinstance(request.input, (DocumentPayload, DocumentSetPayload)):
            raise ValueError(
                "structured_extract requires DocumentPayload or DocumentSetPayload input."
            )

    def execute(
        self,
        request: Union[AnalyzerRequest, Mapping[str, Any]],
    ) -> StructuredExtractionExecution:
        req = validate_analyzer_request(request)
        self._validate_engine_scope(req)

        payload = req.payload
        assert isinstance(payload, StructuredExtractionRequest)

        documents = _iter_request_documents(req)

        extraction_output = self.backend.extract(
            documents=documents,
            result_shape=payload.result_shape,
            selected_fields=payload.selected_fields,
            document_classes=payload.document_classes,
            config=self.config,
        )

        artifact = self.writer.write(
            output_format=payload.output_format,
            payload=extraction_output.payload,
            rows=extraction_output.tabular_rows,
            stem=(
                f"structured_extract_{payload.result_shape.value}_"
                f"{_timestamp_slug()}_{_source_digest_slug(extraction_output.documents)}"
            ),
        )

        response = AnalyzerResponse(
            action=FeatureType.structured_extract,
            input_format=_expected_response_input_format(req),
            policy=req.policy,
            system_language=req.system_language,
            detected_language=None,
            output_language=None,
            result=build_structured_extraction_file_result(
                filename=artifact.file_name,
                output_format=payload.output_format,
                file_size_mb=artifact.file_size_mb,
                result_shape=payload.result_shape,
                selected_fields=list(payload.selected_fields),
                storage_key=artifact.storage_key,
                download_url=artifact.download_url,
                algorithm_version=self.config.algorithm_version,
            ),
            human_review=HumanReviewRequirement(),
        )

        validated_response = validate_analyzer_response(response, request=req)

        return StructuredExtractionExecution(
            response=validated_response,
            preview_payload=extraction_output.payload,
            preview_rows=extraction_output.tabular_rows,
        )

    def run(
        self,
        request: Union[AnalyzerRequest, Mapping[str, Any]],
    ) -> AnalyzerResponse:
        return self.execute(request).response


def run_structured_extraction(request: Union[AnalyzerRequest, Mapping[str, Any]]) -> AnalyzerResponse:
    return StructuredExtractionEngine().run(request)


def run_structured_extraction_with_preview(
    request: Union[AnalyzerRequest, Mapping[str, Any]],
) -> StructuredExtractionExecution:
    return StructuredExtractionEngine().execute(request)


# =========================
# Extraction helpers
# =========================


def _default_strategies() -> list[DocumentClassExtractionStrategy]:
    return [
        GenericFormStrategy(),
        MemoStrategy(),
        InvoiceStrategy(),
        ReceiptStrategy(),
        BankStatementStrategy(),
        KycDocumentStrategy(),
        IdDocumentStrategy(),
        ContractStrategy(),
        LegalRecordStrategy(),
        MedicalRecordStrategy(),
        ProcurementStrategy(),
        ReportStrategy(),
        IncidentReportStrategy(),
        InsuranceStrategy(),
        HRRecordStrategy(),
        OnboardingStrategy(),
        TicketStrategy(),
    ]


def _iter_request_documents(request: AnalyzerRequest) -> list[DocumentPayload]:
    if isinstance(request.input, DocumentPayload):
        return [request.input]
    if isinstance(request.input, DocumentSetPayload):
        return list(request.input.documents)
    raise ValueError("structured_extract requires DocumentPayload or DocumentSetPayload input.")


def _ensure_document_text(
    document: DocumentPayload,
    *,
    config: StructuredExtractionConfig,
) -> tuple[str, bool, str]:
    normalized = document.text if document.text and document.text.strip() else ""
    actual_ocr_used = bool(document.metadata.ocr_used)
    filename = (document.filename or "").strip()
    source_path = Path(filename) if filename else None
    if source_path is not None and not source_path.is_file():
        source_path = None

    if not normalized and source_path is not None:
        try:
            extracted_text, fallback_ocr_used = extract_text_by_format(
                source_path, document.metadata.input_format
            )
        except Exception as exc:
            raise ValueError(
                "Structured extraction could not read reliable text from an uploaded document."
            ) from exc
        normalized = extracted_text
        actual_ocr_used = actual_ocr_used or bool(fallback_ocr_used)

    if not normalized:
        raise ValueError(
            "Structured extraction stopped because an uploaded document contains no "
            "reliable searchable text. Upload a clearer file or a searchable PDF."
        )
    if len(normalized) > config.max_searchable_characters_per_document:
        raise ValueError(
            "An uploaded document exceeds the structured extraction searchable-text limit."
        )

    digest = hashlib.sha256()
    if source_path is not None:
        with source_path.open("rb") as source:
            for chunk in iter(lambda: source.read(1024 * 1024), b""):
                digest.update(chunk)
    else:
        digest.update(normalized.encode("utf-8"))
    return normalized, actual_ocr_used, digest.hexdigest()


def _extract_generic_key_values(
    *,
    lines: Sequence[str],
    source_document_index: int,
    selected_fields: Sequence[str],
    config: StructuredExtractionConfig,
) -> list[ExtractedField]:
    fields: list[ExtractedField] = []
    seen: set[str] = set()
    for line_number, raw_line in enumerate(lines, start=1):
        parsed = _parse_generic_key_value_line(raw_line)
        if parsed is None:
            continue
        key, value = parsed
        normalized_key = _normalize_field_name(key)
        if selected_fields and normalized_key not in selected_fields:
            continue
        if normalized_key in seen:
            continue
        seen.add(normalized_key)
        fields.append(
            ExtractedField(
                name=key,
                value=value,
                source_document_index=source_document_index,
                confidence=0.88,
                evidence=[
                    SourceEvidence(
                        source_document_index=source_document_index,
                        field_name=key,
                        value=value,
                        line_number=line_number,
                        excerpt=_truncate(raw_line, config.max_context_excerpt_chars),
                    )
                ],
            )
        )
        if len(fields) > MAX_EXTRACTED_FIELDS_PER_DOCUMENT:
            raise ValueError("A document exceeds the 5,000-field extraction limit; no partial export was created. Narrow the document or selected fields.")
    return fields


def _parse_generic_key_value_line(raw_line: str) -> Optional[tuple[str, str]]:
    if _is_structure_boundary(raw_line):
        return None
    match = GENERIC_KEY_VALUE_PATTERN.match(raw_line)
    if not match:
        return None
    key = _clean_key(match.group("key"))
    value = _clean_value(match.group("value"))
    if (
        not key
        or not value
        or not any(character.isalnum() for character in key)
        or len(key) > 70
        or key.count(" ") > 10
    ):
        return None
    return key, value


def _extract_repeated_key_value_records(
    *, lines: Sequence[str], source_document_index: int,
    selected_fields: Sequence[str],
) -> list[dict[str, Any]]:
    """Preserve explicit group boundaries, including sparse selected records."""
    records: list[dict[str, Any]] = []
    current: dict[str, Any] = {}
    seen: set[str] = set()
    selected = set(selected_fields)

    def flush() -> None:
        nonlocal current, seen
        if current:
            row = {"source_document_index": source_document_index,
                   "record_type": "key_value_record", "record_index": len(records) + 1}
            _add_business_values(row, current)
            records.append(row)
            if len(records) > MAX_TABULAR_ROWS_PER_REQUEST:
                raise ValueError("Repeated key-value records exceed the output row limit.")
        current, seen = {}, set()

    for line in lines:
        if not line.strip() or _is_structure_boundary(line):
            flush()
            continue
        parsed = _parse_generic_key_value_line(line)
        if parsed is None:
            if re.fullmatch(r"[A-Z][A-Z0-9 _/\-]{1,80}", line.strip()):
                flush()
            continue
        key, value = parsed
        normalized = _normalize_field_name(key)
        if normalized in seen:
            flush()
        seen.add(normalized)
        if not selected or normalized in selected:
            current[key] = value
    flush()
    return records if len(records) >= 2 else []


def _extract_requested_fields_across_lines(
    *, lines: Sequence[str], source_document_index: int,
    selected_fields: Sequence[str], existing_fields: Sequence[ExtractedField],
    config: StructuredExtractionConfig,
) -> list[ExtractedField]:
    """Recover explicitly labeled values, never the following field's label."""
    if not selected_fields:
        return []
    existing = {_normalize_field_name(item.name) for item in existing_fields}
    labels = set(selected_fields)
    for line in lines:
        parsed = _parse_generic_key_value_line(line)
        if parsed:
            labels.add(_normalize_field_name(parsed[0]))
    for strategy in _default_strategies():
        labels.update(getattr(strategy, 'field_patterns', {}))
    labels.update({'phone', 'email', 'name', 'address', 'date', 'amount', 'total'})
    output: list[ExtractedField] = []
    for selected in selected_fields:
        if selected in existing:
            continue
        expression = r"\s+".join(re.escape(word) for word in selected.split('_') if word)
        if not expression:
            continue
        for index, line in enumerate(lines):
            if not line.strip() or _is_structure_boundary(line):
                continue
            value, excerpt = '', line
            if _normalize_field_name(line) == selected and index + 1 < len(lines):
                candidate = lines[index + 1]
                if (candidate.strip() and not _is_structure_boundary(candidate)
                        and _parse_generic_key_value_line(candidate) is None
                        and _normalize_field_name(candidate) not in labels
                        and not candidate.rstrip().endswith(':')):
                    value = _clean_value(candidate)
                    excerpt = f'{line} | {candidate}'
            else:
                match = re.fullmatch(rf"\s*{expression}\s+(?P<value>\S.*?)\s*", line, re.IGNORECASE)
                if match:
                    candidate = match.group('value')
                    if (_parse_generic_key_value_line(candidate) is None
                            and _normalize_field_name(candidate) not in labels):
                        value = _clean_value(candidate)
            if not value or len(value) > 10_000 or not _valid_requested_value(selected, value):
                continue
            output.append(ExtractedField(name=selected, value=value,
                source_document_index=source_document_index, confidence=0.78,
                evidence=[SourceEvidence(source_document_index=source_document_index,
                    field_name=selected, value=value, line_number=index + 1,
                    excerpt=_truncate(excerpt, config.max_context_excerpt_chars))]))
            break
    return output


def _extract_generic_tables(
    *, lines: Sequence[str], source_document_index: int,
    warnings: Optional[list[str]] = None,
) -> list[ExtractedTable]:
    tables: list[ExtractedTable] = []
    total_rows = 0
    for group in _numbered_tabular_groups(lines):
        parsed = [(number, _split_tabular_line(line)) for number, line in group]
        if not _reliable_tabular_group(parsed):
            continue
        has_header = _is_likely_table_header(parsed[0][1])
        width = max(len(cells) for _, cells in parsed)
        columns = (_unique_table_columns(parsed[0][1]) if has_header
                   else [f'column_{index}' for index in range(1, width + 1)])
        for number in range(len(columns) + 1, width + 1):
            columns.append(_next_unique_column_name(columns, f'column_{number}'))
        rows, row_lines = [], []
        for number, cells in parsed[1 if has_header else 0:]:
            if _is_table_separator_row(cells):
                continue
            if has_header and _is_repeated_table_header(cells, columns):
                continue
            padded = [*cells, *([''] * (len(columns) - len(cells)))]
            rows.append(dict(zip(columns, padded)))
            row_lines.append(number)
            total_rows += 1
            if total_rows > MAX_TABULAR_ROWS_PER_REQUEST:
                raise ValueError('A document exceeds the 100,000-row table limit; no partial export was created.')
        if rows:
            if not has_header and warnings is not None:
                warnings.append('A table had no reliable header; generated column names were used and the first row was retained.')
            tables.append(ExtractedTable(table_index=len(tables) + 1,
                source_document_index=source_document_index, columns=columns, rows=rows,
                source_line_numbers=[number for number, _ in parsed], row_line_numbers=row_lines))
    return tables


def _extract_line_item_rows(lines: Sequence[str], source_document_index: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    money = r"(?:NGN|₦|USD|\$|EUR|€|GBP|£)?\s?[0-9][0-9,]*(?:\.\d{2})?"
    pattern = re.compile(
        rf"^(?P<description>[A-Za-z][A-Za-z0-9 .,'\-/()]+?)\s+(?P<quantity>\d+(?:\.\d+)?)\s+(?P<unit_price>{money})\s+(?P<amount>{money})$",
        re.IGNORECASE,
    )
    for line in lines:
        match = pattern.match(line.strip())
        if not match:
            continue
        row = {"source_document_index": source_document_index, "record_type": "line_item"}
        row.update({key: _clean_value(value) for key, value in match.groupdict().items()})
        if len(rows) >= MAX_TABULAR_ROWS_PER_REQUEST:
            raise ValueError("Specialized records exceed the output row limit; no partial export was created.")
        rows.append(row)
    return rows


def _extract_transaction_rows(lines: Sequence[str], source_document_index: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    date = r"\d{1,2}[\/\-.]\d{1,2}[\/\-.]\d{2,4}"
    money = r"(?:NGN|₦|USD|\$|EUR|€|GBP|£)?\s?[0-9][0-9,]*(?:\.\d{2})?"
    pattern = re.compile(
        rf"^(?P<date>{date})\s+(?P<description>.+?)\s+(?P<amount>-?{money})(?:\s+(?P<balance>{money}))?$",
        re.IGNORECASE,
    )
    for line in lines:
        match = pattern.match(line.strip())
        if not match:
            continue
        row = {"source_document_index": source_document_index, "record_type": "transaction"}
        row.update({key: _clean_value(value or "") for key, value in match.groupdict().items()})
        if len(rows) >= MAX_TABULAR_ROWS_PER_REQUEST:
            raise ValueError("Specialized records exceed the output row limit; no partial export was created.")
        rows.append(row)
    return rows


def _extract_clause_rows(lines: Sequence[str], source_document_index: int) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    pattern = re.compile(r"^(?P<section>\d+(?:\.\d+)*)\s*[.)]?\s+(?P<title>[A-Z][A-Za-z0-9 ,&/()\-]{2,80})")
    for line in lines:
        match = pattern.match(line.strip())
        if not match:
            continue
        if len(rows) >= MAX_TABULAR_ROWS_PER_REQUEST:
            raise ValueError("Specialized records exceed the output row limit; no partial export was created.")
        rows.append(
            {
                "source_document_index": source_document_index,
                "record_type": "clause",
                "section": match.group("section"),
                "title": match.group("title").strip(),
            }
        )
    return rows


def _derive_row_records(
    *, source_document_index: int, fields: Sequence[ExtractedField],
    tables: Sequence[ExtractedTable], class_rows: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    specialized = [dict(row, source_document_index=source_document_index) for row in class_rows]
    table_rows = []
    for table in tables:
        for row in table.rows:
            result = {'source_document_index': source_document_index,
                      'record_type': 'table_row', 'table_index': table.table_index}
            _add_business_values(result, row)
            table_rows.append(result)
    if table_rows:
        return _merge_specialized_rows_into_table_rows(table_rows=table_rows, specialized_rows=specialized)
    if specialized:
        return specialized
    if fields:
        result = {'source_document_index': source_document_index, 'record_type': 'field_set'}
        _add_business_values(result, {item.name: item.value for item in fields})
        return [result]
    return [{'source_document_index': source_document_index, 'record_type': 'empty'}]


def _attach_row_provenance(
    rows: Sequence[Mapping[str, Any]], *, filename: str, input_format: str,
    checksum_sha256: str, source_document_index: Optional[int] = None,
) -> list[dict[str, Any]]:
    output = []
    for row in rows:
        result = dict(row)
        result['filename'] = filename
        result['input_format'] = input_format
        result['source_checksum_sha256'] = checksum_sha256
        if source_document_index is not None:
            result['source_document_index'] = source_document_index
        output.append(result)
    return output


_ROW_METADATA_KEYS = {
    "source_document_index",
    "filename",
    "input_format",
    "ocr_used",
    "record_type",
    "table_index",
    "record_index",
    "source_checksum_sha256",
    "source_field_names",
    "extraction_quality_score",
}


def _row_business_values(row: Mapping[str, Any]) -> dict[str, str]:
    aliases = row.get('source_field_names', {})
    result = {}
    for key, value in row.items():
        if _normalize_field_name(str(key)) in _ROW_METADATA_KEYS:
            continue
        original = aliases.get(key, key) if isinstance(aliases, Mapping) else key
        normalized = _normalize_field_name(str(original))
        if normalized in _ROW_METADATA_KEYS:
            normalized = 'field_' + normalized
        rendered = _clean_value(value).casefold()
        if normalized and rendered:
            result[normalized] = rendered
    return result


def _row_match_score(left: Mapping[str, Any], right: Mapping[str, Any]) -> int:
    left_fingerprint, right_fingerprint = _row_fingerprint(left), _row_fingerprint(right)
    if left_fingerprint is None or left_fingerprint != right_fingerprint:
        return 0
    return len(left_fingerprint)


def _merge_specialized_rows_into_table_rows(
    *, table_rows: Sequence[Mapping[str, Any]], specialized_rows: Sequence[Mapping[str, Any]],
) -> list[dict[str, Any]]:
    # Exact canonical records with an identity field may represent the same
    # source row twice. Quantities/prices alone never establish identity.
    # Queues consume matches one-for-one, preserving duplicate occurrences.
    merged = [dict(row) for row in table_rows]
    indexed = defaultdict(deque)
    for index, row in enumerate(merged):
        fingerprint = _row_fingerprint(row)
        if fingerprint is not None:
            indexed[fingerprint].append(index)
    for row in specialized_rows:
        fingerprint = _row_fingerprint(row)
        candidates = indexed.get(fingerprint) if fingerprint is not None else None
        if candidates:
            candidates.popleft()
            continue
        merged.append(dict(row))
        if len(merged) > MAX_TABULAR_ROWS_PER_REQUEST:
            raise ValueError('Structured extraction exceeds the 100,000-row limit; no partial export was created.')
    return merged


def _merge_fields(
    fields: Sequence[ExtractedField],
    *,
    selected_fields: Sequence[str],
    include_empty_selected_fields: bool,
    source_document_index: int,
) -> list[ExtractedField]:
    by_normalized: dict[str, ExtractedField] = {}
    for field_item in fields:
        normalized = _normalize_field_name(field_item.name)
        if selected_fields and normalized not in selected_fields:
            continue
        existing = by_normalized.get(normalized)
        if (existing is None or not _clean_value(existing.value)
                or (_clean_value(existing.value) == _clean_value(field_item.value)
                    and field_item.confidence > existing.confidence)):
            by_normalized[normalized] = field_item

    if include_empty_selected_fields:
        for raw_selected in selected_fields:
            if raw_selected not in by_normalized:
                by_normalized[raw_selected] = ExtractedField(
                    name=raw_selected,
                    value="",
                    source_document_index=source_document_index,
                    confidence=0.0,
                    evidence=[],
                )

    return sorted(by_normalized.values(), key=lambda item: _normalize_field_name(item.name))


def _first_regex_match(
    lines: Sequence[str], patterns: Sequence[re.Pattern[str]],
    *, field_name: Optional[str] = None,
) -> Optional[tuple[str, int, str]]:
    # Pattern order expresses precedence: an explicit invoice date wins over
    # the generic Date fallback, regardless of their order in the source.
    for pattern in patterns:
        for number, line in enumerate(lines, start=1):
            stripped = line.strip()
            if not stripped or _is_structure_boundary(stripped):
                continue
            match = pattern.search(stripped)
            if not match or match.start() != 0:
                continue
            value = _clean_value(match.group('value'))
            prefix = stripped[:match.start('value')]
            if field_name and field_name.endswith(('_id', '_number')):
                if not any(char.isdigit() for char in value) and not any(char in prefix for char in ':#'):
                    continue
            if value:
                return value, number, line
    return None


# =========================
# Generic utilities
# =========================

def _iter_lines(text: str) -> Iterable[str]:
    # Empty lines separate records; leading/trailing tabs encode empty cells.
    for line in text.splitlines():
        yield line.strip(' \r\v\f')


def _normalize_selected_fields(fields: Sequence[str]) -> list[str]:
    output: list[str] = []
    for item in fields:
        normalized = _normalize_field_name(item)
        if normalized and normalized not in output:
            output.append(normalized)
    return output


def _normalize_field_name(value: str) -> str:
    normalized = unicodedata.normalize("NFKC", str(value)).strip().casefold()
    return re.sub(r"[^\w]+", "_", normalized, flags=re.UNICODE).strip("_")


def _clean_key(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip().strip(":-–—|"))


def _clean_value(value: Any) -> str:
    return re.sub(r'\s+', ' ', str(value).strip())


def _truncate(value: Optional[str], max_chars: int) -> Optional[str]:
    if value is None:
        return None
    text = str(value)
    if len(text) <= max_chars:
        return text
    return f"{text[: max_chars - 1]}…"


def _detect_tabular_delimiter(line: str) -> Optional[str]:
    stripped = line.strip(' \r\n')
    if '\t' in stripped:
        return '\t'
    if stripped.strip('|').count('|') >= 1:
        return '|'
    # Labeled amounts/addresses are not CSV, even when their values contain
    # commas or semicolons. A standalone thousands-formatted number is not a row.
    if _parse_generic_key_value_line(stripped) is not None:
        return None
    if re.fullmatch(r'(?:[A-Z]{3}|[$€£₦])?\s*[+-]?\d{1,3}(?:,\d{3})+(?:\.\d+)?', stripped):
        return None
    for delimiter in (';', ','):
        try:
            if len(next(csv.reader([stripped], delimiter=delimiter))) >= 2:
                return delimiter
        except (csv.Error, StopIteration):
            pass
    return None


def _split_tabular_line(line: str) -> list[str]:
    source = line.strip(' \r\n')
    delimiter = _detect_tabular_delimiter(source)
    if delimiter is None:
        return [source]
    if delimiter == '|':
        # Remove one framing pipe, never all empty leading/trailing cells.
        if source.startswith('|'):
            source = source[1:]
        if source.endswith('|'):
            source = source[:-1]
    try:
        cells = next(csv.reader([source], delimiter=delimiter, skipinitialspace=True))
    except (csv.Error, StopIteration):
        cells = source.split(delimiter)
    return [_clean_value(cell) for cell in cells]


def _looks_tabular(line: str) -> bool:
    return _detect_tabular_delimiter(line) is not None


def _group_tabular_lines(lines: Sequence[str]) -> list[list[str]]:
    return [[line for _, line in group] for group in _numbered_tabular_groups(lines)]


def _next_unique_column_name(columns: Sequence[str], preferred: str) -> str:
    existing = {_normalize_field_name(column) for column in columns}
    candidate = preferred
    suffix = 2
    while _normalize_field_name(candidate) in existing:
        candidate = f"{preferred}_{suffix}"
        suffix += 1
    return candidate


def _unique_table_columns(raw_columns: Sequence[str]) -> list[str]:
    columns: list[str] = []
    occupied: set[str] = set()
    next_suffix: dict[str, int] = {}
    for index, raw_column in enumerate(raw_columns, start=1):
        preferred = _clean_key(raw_column) or f"column_{index}"
        candidate = preferred
        normalized = _normalize_field_name(candidate)
        suffix = next_suffix.get(normalized, 2)
        while _normalize_field_name(candidate) in occupied:
            candidate = f"{preferred}_{suffix}"
            suffix += 1
        next_suffix[normalized] = suffix
        occupied.add(_normalize_field_name(candidate))
        columns.append(candidate)
    return columns


def _is_table_separator_row(cells: Sequence[str]) -> bool:
    return bool(cells) and all(
        not cell or re.fullmatch(r":?[-=_]{2,}:?", cell) is not None
        for cell in cells
    )


def _is_repeated_table_header(cells: Sequence[str], columns: Sequence[str]) -> bool:
    if len(cells) != len(columns):
        return False
    return all(
        _normalize_field_name(cell) == _normalize_field_name(column)
        for cell, column in zip(cells, columns)
    )


def _aggregate_documents(documents: Sequence[DocumentExtraction]) -> dict[str, Any]:
    return {
        "document_count": len(documents),
        "field_count": sum(len(doc.fields) for doc in documents),
        "table_count": sum(len(doc.tables) for doc in documents),
        "row_record_count": sum(len(doc.row_records) for doc in documents),
        "warnings_count": sum(len(doc.warnings) for doc in documents),
    }


def _quality_summary(documents: Sequence[DocumentExtraction]) -> dict[str, Any]:
    requested = sum(doc.requested_field_count for doc in documents)
    populated = sum(doc.populated_requested_field_count for doc in documents)
    if requested:
        score = populated / requested
        coverage_mode = "requested_fields"
    else:
        score = None
        coverage_mode = "not_applicable_without_requested_fields"
    return {
        "score": round(score, 4) if score is not None else None,
        "coverage_mode": coverage_mode,
        "requested_field_count": requested,
        "populated_requested_field_count": populated,
        "documents_requiring_attention": [
            doc.source_document_index
            for doc in documents
            if doc.extraction_quality_score < 1.0 or doc.warnings
        ],
        "human_review_required": True,
    }


def _document_fields_to_row(doc: DocumentExtraction, *, selected_fields: Sequence[str]) -> dict[str, Any]:
    row = {'source_document_index': doc.source_document_index, 'filename': doc.filename or '',
           'input_format': doc.input_format, 'ocr_used': doc.ocr_used,
           'source_checksum_sha256': doc.checksum_sha256,
           'extraction_quality_score': doc.extraction_quality_score}
    mapping = doc.fields_as_mapping()
    if selected_fields:
        normalized = {_normalize_field_name(key): value for key, value in mapping.items()}
        values = {name: normalized.get(_normalize_field_name(name), '') for name in selected_fields}
    else:
        values = mapping
    _add_business_values(row, values)
    return row


def _headers_for_rows(rows: Sequence[Mapping[str, Any]]) -> list[str]:
    preferred = ['source_document_index', 'filename', 'input_format', 'ocr_used',
                 'source_checksum_sha256', 'extraction_quality_score', 'record_type',
                 'record_index', 'table_index', 'source_field_names']
    discovered = dict.fromkeys(str(key) for row in rows for key in row)
    headers = [key for key in preferred if key in discovered]
    present = set(headers)
    headers.extend(key for key in discovered if key not in present)
    return headers or ['source_document_index']


def _stringify_row(row: Mapping[str, Any], *, max_chars: Optional[int] = None) -> dict[str, str]:
    output = {}
    for key, value in row.items():
        rendered = (json.dumps(value, ensure_ascii=False, default=str)
                    if isinstance(value, (dict, list)) else '' if value is None else str(value))
        output[str(key)] = _safe_spreadsheet_cell(rendered, max_chars=max_chars)
    return output


MAX_SPREADSHEET_CELL_CHARACTERS = 32_767


def _safe_spreadsheet_cell(value: Any, *, max_chars: Optional[int] = None) -> str:
    rendered = str(value if value is not None else '')
    if rendered.lstrip().startswith(('=', '+', '-', '@')):
        rendered = "'" + rendered
    if max_chars is not None and len(rendered) > max_chars:
        raise ValueError(
            'An extracted value exceeds the 32,767-character Excel cell limit. '
            'No shortened workbook was exported. Choose JSON or CSV to preserve the complete value.'
        )
    return rendered


def _write_summary_sheet(sheet: Worksheet, payload: Mapping[str, Any]) -> None:
    rows = [(key, payload.get(key, '')) for key in [
        'contract_version', 'feature', 'result_shape', 'generated_at', 'human_review_required']]
    rows.append(('selected_fields', ', '.join(payload.get('selected_fields', []) or [])))
    for key, value in (payload.get('aggregate', {}) or {}).items():
        rows.append((f'aggregate.{key}', value))
    for key, value in (payload.get('quality_summary', {}) or {}).items():
        rows.append((f'quality.{key}', json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value))
    for document in payload.get('document_warnings', []) or []:
        for warning in document.get('warnings', []) or []:
            rows.append((f"warning.document_{document.get('source_document_index', '')}", warning))
    sheet.append(['property', 'value'])
    for row in rows:
        sheet.append([_safe_spreadsheet_cell(item, max_chars=MAX_SPREADSHEET_CELL_CHARACTERS) for item in row])


def _write_rows_sheet(sheet: Worksheet, rows: Sequence[Mapping[str, Any]]) -> None:
    headers = _headers_for_rows(rows)
    sheet.append([_safe_spreadsheet_cell(header, max_chars=MAX_SPREADSHEET_CELL_CHARACTERS) for header in headers])
    for row in rows:
        rendered = _stringify_row(row, max_chars=MAX_SPREADSHEET_CELL_CHARACTERS)
        sheet.append([rendered.get(header, '') for header in headers])


def _evidence_rows_from_payload(payload: Mapping[str, Any]) -> list[dict[str, Any]]:
    output = []
    documents = payload.get('document_fields') or payload.get('documents') or []
    for document in documents:
        if not isinstance(document, Mapping):
            continue
        evidence = list(document.get('evidence', []) or [])
        if isinstance(document.get('fields'), list):
            evidence.extend(item for field in document['fields'] if isinstance(field, Mapping)
                            for item in field.get('evidence', []) or [])
        for item in evidence:
            if isinstance(item, Mapping):
                row = dict(item)
                row['source_document_index'] = document.get('source_document_index')
                row['filename'] = document.get('filename', '')
                row['record_type'] = 'field_evidence'
                output.append(row)
    return output


def _timestamp_slug() -> str:
    return datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%S%fZ")


def _timestamp_iso() -> str:
    return datetime.now(timezone.utc).isoformat().replace("+00:00", "Z")


def _artifact_from_path(path: Path) -> PersistedArtifact:
    storage = LocalArtifactStorage()

    stored = storage.persist(
        source_file_path=str(path),
        artifact_name=path.name,
        content_type=guess_content_type(str(path)),
    )

    stored_path = Path(stored.stored_path)
    size_mb = round(stored_path.stat().st_size / (1024 * 1024), 4)

    return PersistedArtifact(
        file_name=stored.original_artifact_name,
        file_extension=stored_path.suffix.lstrip("."),
        file_size_mb=size_mb,
        file_path=stored.stored_path,
        storage_key=stored.storage_key,
        download_url=stored.download_url,
    )


def _temporary_artifact_path(target: Path) -> Path:
    return target.with_name(
        f".{target.stem}.{secrets.token_hex(8)}.tmp{target.suffix}"
    )


def _source_digest_slug(documents: Sequence[DocumentExtraction]) -> str:
    identity = "|".join(document.checksum_sha256 for document in documents)
    return hashlib.sha256(identity.encode("ascii")).hexdigest()[:12]


def _safe_filename(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", name).strip("._")
    if not cleaned:
        raise ValueError("Artifact filename cannot be empty.")
    return cleaned


def _expected_response_input_format(request: AnalyzerRequest) -> Union[str, Any]:
    if isinstance(request.input, DocumentSetPayload):
        return "document_set"
    assert isinstance(request.input, DocumentPayload)
    return request.input.metadata.input_format


_HEADER_WORDS = {
    "name", "full_name", "first_name", "last_name", "city", "age", "email",
    "email_address", "phone", "address", "country", "state", "date", "id",
    "item", "description", "product", "quantity", "qty", "price", "unit_price",
    "amount", "total", "balance", "debit", "credit", "currency", "reference",
    "account", "account_number", "customer", "customer_id", "order", "order_id",
    "salary", "department", "status", "category", "code", "sku", "value",
    "x", "y", "z", "nom", "prénom", "âge", "ville", "montant", "prix",
    "quantité", "libellé", "référence", "statut", "filename", "input_format",
    "source_document_index", "record_type", "source_checksum_sha256",
}


def _is_structure_boundary(line: str) -> bool:
    return line.strip().startswith(("[STRUCTURED ", "[/STRUCTURED "))


def _source_extraction_warnings(lines: Sequence[str]) -> list[str]:
    prefix = "[STRUCTURED EXTRACTION WARNING] "
    return [line[len(prefix):] for line in lines if line.startswith(prefix)]


def _is_likely_table_header(cells: Sequence[str]) -> bool:
    nonempty = [_normalize_field_name(cell) for cell in cells if cell.strip()]
    if not nonempty or any(re.search(r"\d|[@.!?]", cell) for cell in cells):
        return False
    if any(len(cell.split()) > 5 for cell in cells):
        return False
    return any(value in _HEADER_WORDS or value.endswith(("_id", "_name", "_date", "_amount"))
               for value in nonempty)


def _numbered_tabular_groups(lines: Sequence[str]) -> list[list[tuple[int, str]]]:
    groups: list[list[tuple[int, str]]] = []
    current: list[tuple[int, str]] = []
    delimiter: Optional[str] = None

    def flush() -> None:
        nonlocal current, delimiter
        if len(current) >= 2:
            groups.append(current)
        current, delimiter = [], None

    for number, line in enumerate(lines, start=1):
        if not line.strip() or _is_structure_boundary(line):
            flush()
            continue
        found = _detect_tabular_delimiter(line)
        if (found is None and current and delimiter == '\t'
                and _is_likely_table_header(_split_tabular_line(current[0][1]))
                and _parse_generic_key_value_line(line) is None
                and re.fullmatch(r"[^\s:;|,.!?]{1,120}", line.strip())):
            # DocumentPayload strips trailing whitespace at the outer text
            # boundary. Retain an otherwise atomic last row whose empty tail
            # cells consequently lost their trailing tabs.
            found = delimiter
        if found is None:
            flush()
            continue
        cells = _split_tabular_line(line)
        if current and (found != delimiter or (
                _is_likely_table_header(cells)
                and _is_likely_table_header(_split_tabular_line(current[0][1]))
                and not _is_repeated_table_header(cells, _split_tabular_line(current[0][1])))):
            flush()
        if not current:
            delimiter = found
        current.append((number, line))
    flush()
    return groups


def _reliable_tabular_group(group: Sequence[tuple[int, list[str]]]) -> bool:
    if len(group) < 2 or len(group[0][1]) < 2:
        return False
    if _is_likely_table_header(group[0][1]):
        return True
    # Headerless atomic records are retained with generated column names.
    # Prose containing commas/semicolons is not promoted to a table.
    widths = {len(cells) for _, cells in group}
    return len(widths) == 1 and min(widths) >= 2 and all(
        len(cell.split()) <= 3 and not re.search(r"[.!?]\s*$", cell)
        for _, cells in group for cell in cells
    )


def _valid_requested_value(name: str, value: str) -> bool:
    if "email" in name.split("_"):
        return re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", value) is not None
    if name in {"phone", "phone_number", "mobile", "telephone"}:
        return (re.fullmatch(r"\+?[0-9 ()\-.]{6,}", value) is not None
                and sum(char.isdigit() for char in value) >= 6)
    return True


def _selected_table_fields(
    tables: Sequence[ExtractedTable], selected: Sequence[str],
    lines: Sequence[str], config: StructuredExtractionConfig,
) -> list[ExtractedField]:
    wanted, found = set(selected), set()
    output: list[ExtractedField] = []
    for table in tables:
        for column in table.columns:
            name = _normalize_field_name(column)
            if name not in wanted or name in found:
                continue
            for index, row in enumerate(table.rows):
                value = _clean_value(row.get(column, ''))
                if not value:
                    continue
                number = table.row_line_numbers[index] if index < len(table.row_line_numbers) else None
                excerpt = lines[number - 1] if number and number <= len(lines) else value
                output.append(ExtractedField(name=column, value=value,
                    source_document_index=table.source_document_index, confidence=0.8,
                    source_kind='table', evidence=[SourceEvidence(
                        source_document_index=table.source_document_index, field_name=column,
                        value=value, line_number=number,
                        excerpt=_truncate(f"{column}: {value}", config.max_context_excerpt_chars))]))
                found.add(name)
                break
    return output


def _project_tables(tables: Sequence[ExtractedTable], selected: Sequence[str]) -> list[ExtractedTable]:
    if not selected:
        return list(tables)
    wanted = set(selected)
    output = []
    for table in tables:
        columns = [column for column in table.columns if _normalize_field_name(column) in wanted]
        if columns:
            output.append(ExtractedTable(table_index=table.table_index,
                source_document_index=table.source_document_index, columns=columns,
                rows=[{column: row.get(column, '') for column in columns} for row in table.rows],
                source_line_numbers=table.source_line_numbers, row_line_numbers=table.row_line_numbers))
    return output


def _add_business_values(target: dict[str, Any], values: Mapping[str, Any]) -> None:
    """Keep untrusted column names separate from trusted provenance names."""
    aliases = dict(target.get('source_field_names', {}))
    occupied = None
    for raw_key, value in values.items():
        key = str(raw_key)
        if _normalize_field_name(key) in _ROW_METADATA_KEYS:
            if occupied is None:
                occupied = [*target, *map(str, values)]
            key = _next_unique_column_name(occupied, 'field_' + _normalize_field_name(key))
            aliases[key] = str(raw_key)
            occupied.append(key)
        target[key] = value
    if aliases:
        target['source_field_names'] = aliases


def _project_record_rows(rows: Sequence[Mapping[str, Any]], selected: Sequence[str]) -> list[dict[str, Any]]:
    if not selected:
        return [dict(row) for row in rows]
    wanted = set(selected)
    output = []
    for row in rows:
        aliases = row.get('source_field_names', {})
        result = {key: value for key, value in row.items()
                  if key in _ROW_METADATA_KEYS and key != 'source_field_names'}
        retained = {}
        for key, value in row.items():
            if key in _ROW_METADATA_KEYS:
                continue
            original = aliases.get(key, key) if isinstance(aliases, Mapping) else key
            if _normalize_field_name(original) in wanted:
                result[key] = value
                if original != key:
                    retained[key] = original
        if any(key not in _ROW_METADATA_KEYS for key in result):
            if retained:
                result['source_field_names'] = retained
            output.append(result)
    return output


def _with_document_fields(
    rows: Sequence[Mapping[str, Any]], fields: Sequence[ExtractedField],
) -> list[dict[str, Any]]:
    output = []
    cells = 0
    for row in rows:
        result = dict(row)
        if result.get('record_type') == 'key_value_record':
            # Explicit groups represent separate entities, not document-wide
            # header fields that should be copied into every other entity.
            cells += len(result)
            if cells > MAX_EXPORT_CELLS_PER_REQUEST:
                raise ValueError('Structured records exceed the cell budget; no partial export was created. Narrow the request.')
            output.append(result)
            continue
        originals = result.get('source_field_names', {})
        existing = {_normalize_field_name(originals.get(key, key)): key
                    for key in result if key not in _ROW_METADATA_KEYS}
        for item in fields:
            if item.source_kind == 'table':
                continue
            normalized = _normalize_field_name(item.name)
            if normalized in existing:
                key = existing[normalized]
                if _clean_value(result[key]) == _clean_value(item.value):
                    continue
                key = _next_unique_column_name(list(result), 'document_' + normalized)
                result[key] = item.value
                aliases = dict(result.get('source_field_names', {}))
                aliases[key] = item.name
                result['source_field_names'] = aliases
            else:
                _add_business_values(result, {item.name: item.value})
        cells += len(result)
        if cells > MAX_EXPORT_CELLS_PER_REQUEST:
            raise ValueError('Structured records exceed the cell budget; no partial export was created. Narrow the request.')
        output.append(result)
    return output


_ROW_FIELD_ALIASES = {
    'item': 'description', 'product': 'description', 'item_description': 'description',
    'qty': 'quantity', 'price': 'unit_price', 'unitprice': 'unit_price',
    'line_total': 'amount', 'transaction_date': 'date',
}
_ROW_IDENTITY_FIELDS = {
    'description', 'name', 'full_name', 'customer_id', 'employee_id', 'item_id',
    'sku', 'code', 'reference', 'transaction_id', 'section', 'ticket_id',
}


def _row_fingerprint(row: Mapping[str, Any]) -> Optional[tuple[tuple[str, str], ...]]:
    values = {}
    for key, value in _row_business_values(row).items():
        canonical = _ROW_FIELD_ALIASES.get(key, key)
        if canonical in values and values[canonical] != value:
            return None
        values[canonical] = value
    if not any(key in _ROW_IDENTITY_FIELDS and value for key, value in values.items()):
        return None
    return tuple(sorted(values.items())) if len(values) >= 2 else None


def _selected_field_status(
    documents: Sequence[DocumentExtraction], selected: Sequence[str],
) -> list[dict[str, Any]]:
    if not selected:
        return []
    best: dict[str, ExtractedField] = {}
    for document in documents:
        for item in document.fields:
            name = _normalize_field_name(item.name)
            current = best.get(name)
            if current is None or (bool(item.value), item.confidence) > (bool(current.value), current.confidence):
                best[name] = item
    output = []
    for name in selected:
        normalized = _normalize_field_name(name)
        item = best.get(normalized)
        value = _clean_value(item.value) if item is not None else ''
        evidence = item.evidence[0].to_dict() if item is not None and item.evidence else None
        if evidence:
            evidence = {**evidence, 'value': _truncate(evidence['value'], 240),
                        'excerpt': _truncate(evidence.get('excerpt'), 120)}
        output.append({'name': name, 'normalized_name': normalized,
            'value': _truncate(value, 240), 'value_truncated': len(value) > 240,
            'confidence': item.confidence if item is not None else 0.0,
            'found': bool(value), 'evidence': evidence})
    return output


def _payload_sources(payload: Mapping[str, Any]) -> dict[Any, Mapping[str, Any]]:
    return {item.get('source_document_index'): item for item in payload.get('source_documents', [])
            if isinstance(item, Mapping)}


def _field_rows_from_payload(payload: Mapping[str, Any]) -> list[dict[str, Any]]:
    output = []
    sources = _payload_sources(payload)
    for document in payload.get('document_fields') or payload.get('documents') or []:
        if not isinstance(document, Mapping):
            continue
        fields = document.get('fields')
        if isinstance(fields, list):
            values = {item['name']: item.get('value', '') for item in fields if isinstance(item, Mapping) and 'name' in item}
        elif isinstance(fields, Mapping):
            values = dict(fields)
        else:
            continue
        if not values:
            continue
        index = document.get('source_document_index')
        source = sources.get(index, {})
        row = {'source_document_index': index, 'filename': source.get('filename', document.get('filename', '')),
               'input_format': source.get('input_format', ''),
               'source_checksum_sha256': source.get('checksum_sha256', ''), 'record_type': 'field_set'}
        _add_business_values(row, values)
        output.append(row)
    return output


def _table_rows_from_payload(payload: Mapping[str, Any]) -> list[dict[str, Any]]:
    sources = _payload_sources(payload)
    tables = payload.get('extracted_tables')
    if tables is None:
        tables = [table for document in payload.get('documents', []) if isinstance(document, Mapping)
                  for table in document.get('tables', []) or []]
    output = []
    for table in tables:
        index = table.get('source_document_index')
        source = sources.get(index, {})
        for values in table.get('rows', []) or []:
            row = {'source_document_index': index, 'filename': source.get('filename', ''),
                   'input_format': source.get('input_format', ''),
                   'source_checksum_sha256': source.get('checksum_sha256', ''),
                   'table_index': table.get('table_index'), 'record_type': 'table_row'}
            _add_business_values(row, values)
            output.append(row)
    return output


def _complete_csv_rows(payload: Mapping[str, Any], rows: Sequence[Mapping[str, Any]]) -> list[dict[str, Any]]:
    shape = payload.get('result_shape')
    output = [dict(row) for row in rows]
    if shape == StructuredExtractionResultShape.key_value_fields.value:
        for row in output:
            row['record_type'] = 'field_set'
        output.extend(_table_rows_from_payload(payload))
    elif shape == StructuredExtractionResultShape.tables.value:
        output.extend(_field_rows_from_payload(payload))
    output.extend(_evidence_rows_from_payload(payload))
    for document in payload.get('document_warnings', []) or []:
        for warning in document.get('warnings', []) or []:
            output.append({'source_document_index': document.get('source_document_index'),
                           'filename': document.get('filename', ''),
                           'record_type': 'extraction_warning', 'warning': warning})
    return output


def _check_export_rows(rows: Sequence[Mapping[str, Any]], *, xlsx: bool = False) -> list[str]:
    if len(rows) > MAX_TABULAR_ROWS_PER_REQUEST:
        raise ValueError('The complete export exceeds 100,000 rows including companion data. Narrow the request or use JSON; no partial export was created.')
    headers = _headers_for_rows(rows)
    if (len(rows) + 1) * len(headers) > MAX_EXPORT_CELLS_PER_REQUEST:
        raise ValueError('The complete export exceeds the cell budget. Narrow the request or use JSON; no partial export was created.')
    if xlsx:
        if len(headers) > 16_384:
            raise ValueError('The export exceeds the Excel column limit. Use JSON; no partial export was created.')
        for header in headers:
            _safe_spreadsheet_cell(header, max_chars=MAX_SPREADSHEET_CELL_CHARACTERS)
        for row in rows:
            _stringify_row(row, max_chars=MAX_SPREADSHEET_CELL_CHARACTERS)
    return headers


def limit_structured_extraction_preview(
    payload: Mapping[str, Any], rows: Sequence[Mapping[str, Any]], *, rows_limit: int = 200,
    max_bytes: int = MAX_PREVIEW_BYTES,
) -> tuple[dict[str, Any], list[dict[str, Any]], bool]:
    """Bound every preview container; exported artifacts always retain full data."""
    limit = max(1, min(int(rows_limit), 200))
    max_bytes = max(1024, min(int(max_bytes), MAX_PREVIEW_BYTES))
    shortened = False
    omitted = object()
    preferred = ('selected_field_status', 'quality_summary', 'source_documents', 'document_warnings',
                 'contract_version', 'feature', 'result_shape', 'selected_fields', 'generated_at',
                 'human_review_required', 'extraction_method', 'reliance_notice')

    def bounded(value: Any, budget: list[int], key: str = '') -> Any:
        nonlocal shortened
        if budget[0] <= 16:
            shortened = True
            return omitted
        if isinstance(value, Mapping):
            result = {}
            keys = ([name for name in preferred if name in value] +
                    [name for name in value if name not in preferred]) if key == 'root' else value
            for index, name in enumerate(keys):
                if index >= 200 or budget[0] <= 16:
                    shortened = True
                    break
                name = str(name)
                if len(name) > budget[0] // 6:
                    shortened = True
                    continue
                cost = len(json.dumps(name, ensure_ascii=False).encode('utf-8')) + 2
                budget[0] -= cost
                item = bounded(value[name], budget, name)
                if item is omitted:
                    break
                result[name] = item
            return result
        if isinstance(value, (list, tuple)):
            result = []
            collection_limit = 200 if key in {'selected_field_status', 'selected_fields'} else limit
            for index, item in enumerate(value):
                if index >= collection_limit or budget[0] <= 16:
                    shortened = True
                    break
                budget[0] -= 2
                sampled = bounded(item, budget, key)
                if sampled is omitted:
                    break
                result.append(sampled)
            return result
        if isinstance(value, str):
            available = max(0, (budget[0] - 16) // 6)
            if len(value) > available:
                value = value[:max(0, available - 1)] + '…'
                shortened = True
            budget[0] -= len(json.dumps(value, ensure_ascii=False).encode('utf-8'))
            return value
        budget[0] -= len(json.dumps(value, default=str)) + 1
        return value

    preview = bounded(payload, [max_bytes * 3 // 4], 'root')
    preview_rows = bounded(rows, [max_bytes // 4], 'preview_rows')
    preview['preview_truncated'] = shortened
    preview['preview_limits'] = {'rows_per_collection': limit, 'max_bytes': max_bytes}
    return preview, preview_rows, shortened


__all__ = [
    "STRUCTURED_EXTRACTION_RULES",
    "PersistedArtifact",
    "StructuredExtractionConfig",
    "SourceEvidence",
    "ExtractedField",
    "ExtractedTable",
    "DocumentExtraction",
    "StructuredExtractionOutput",
    "DocumentClassExtractionStrategy",
    "StructuredExtractionBackend",
    "DeterministicStructuredExtractionBackend",
    "ShapeEnforcer",
    "StructuredArtifactWriter",
    "LocalStructuredArtifactWriter",
    "LocalStructuredExtractionArtifactStore",
    "StructuredExtractionEngine",
    "run_structured_extraction",
]
