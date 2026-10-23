from __future__ import annotations

from copy import deepcopy
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import re
from typing import Any, Mapping, Optional, Sequence
import uuid

import docx
import fitz  # PyMuPDF
import pytesseract
from PIL import Image, ImageDraw, ImageOps, ImageFont
import math
import json
import hashlib
import sys
import os
import shutil
import subprocess
import tempfile
from functools import lru_cache
from lxml import etree
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from docx.shared import Pt

from backend.src.extraction import (
    OCR_CONFIG, preprocess_for_ocr, resolve_ocr_lang, extract_redaction_ocr_passes,
    _redaction_docx_parts, redaction_pdf_surfaces,
)
from backend.src.redaction_analysis import analysis_active, analysis_cached, finding_context, image_fingerprint
from backend.src.processing.data_protection.client import (
    DEFAULT_DLP_LOCATION,
    DEFAULT_MIN_LIKELIHOOD,
    DetectionCandidate,
    ImageFinding,
    GoogleSDPClient,
    TextFinding,
    build_google_sdp_client,
    id_document_field_findings,
    inspect_sensitive_image,
    inspect_sensitive_text,
    is_id_document_payload,
    merge_overlapping_findings,
)
from backend.src.schema import (
    AnalyzerRequest,
    AnalyzerResponse,
    DocumentInputFormat,
    DocumentPayload,
    FeatureType,
    HumanReviewRequirement,
    RedactionRequest,
)
from backend.src.validation import build_document_file_result, validate_analyzer_request, validate_analyzer_response

DEFAULT_PDF_RENDER_SCALE = 2.0
ID_DOCUMENT_FALLBACK_RENDER_SCALE = 4.0
ID_DOCUMENT_OCR_CONFIG = "--oem 3 --psm 11"
ID_DOCUMENT_OCR_TARGET_WIDTH = 1000
DEFAULT_BLACK = (0, 0, 0)
LANCZOS_RESAMPLING = getattr(Image, "Resampling", Image).LANCZOS
PDF_VISIBILITY_RENDER_SCALE = 2.0
PDF_VISIBLE_PIXEL_THRESHOLD = 245
VISUAL_SIGNATURE_QUOTE = "visual signature"
_PDF_REVIEWABLE_METADATA_KEYS = ("author", "title", "subject", "keywords")
_GENERIC_METADATA_AUTHORS = {"admin", "administrator", "anonymous", "unknown", "user"}

_SIGNATURE_ANCHOR_RE = re.compile(
    r"(?i)\b(?:authorized\s+signature|customer\s+signature|applicant\s+signature|"
    r"holder(?:'s)?\s+signature|signature|signed\s+by|signatory|"
    r"signature\s+autoris[eé]e|sign[eé]\s+par)\b"
)
_SIGNATURE_ROLE_RE = re.compile(
    r"(?i)\b(?:registrar|registar|istrar|piatrer|vice\s+chancellor|chancellor|authorized\s+officer|"
    r"approving\s+officer|director)\b"
)

_IMAGE_OUTPUT_MAP: dict[DocumentInputFormat, str] = {
    DocumentInputFormat.jpg: "JPEG",
    DocumentInputFormat.jpeg: "JPEG",
    DocumentInputFormat.png: "PNG",
}

_FILE_OUTPUT_MAP = {
    DocumentInputFormat.pdf: "pdf",
    DocumentInputFormat.docx: "docx",
    DocumentInputFormat.jpg: "jpg",
    DocumentInputFormat.jpeg: "jpeg",
    DocumentInputFormat.png: "png",
}

_SUPPORTED_INPUTS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.jpg,
    DocumentInputFormat.jpeg,
    DocumentInputFormat.png,
}


@dataclass(frozen=True)
class OCRWord:
    text: str
    bbox: tuple[int, int, int, int]
    start: int
    end: int
    line_id: tuple[int, int, int]
    character_boxes: tuple[tuple[int, int, int, int], ...] = ()


@dataclass(frozen=True)
class RunSpan:
    start: int
    end: int
    text: str
    rpr_xml: Any


@dataclass(frozen=True)
class PDFWord:
    start: int
    end: int
    text: str
    rect: fitz.Rect
    block_no: int
    line_no: int
    word_no: int


@dataclass(frozen=True)
class OCRImageSurface:
    image: Image.Image
    page_rect: fitz.Rect


def redact_text_with_findings(text: str, findings: Sequence[TextFinding]) -> str:
    result = text
    for finding in sorted(findings, key=lambda item: item.start, reverse=True):
        replacement = "█" * max(1, finding.end - finding.start)
        result = result[:finding.start] + replacement + result[finding.end:]
    return result


def _ocr_words_from_image(
    image: Image.Image,
    *,
    ocr_lang: Optional[str] = None,
    sparse_text: bool = False,
    include_character_boxes: bool = False,
    ocr_pass: tuple[Any, Image.Image, str] | None = None,
) -> tuple[str, list[OCRWord]]:
    source = image.convert("RGB")
    coordinate_scale = 1.0

    if sparse_text:
        coordinate_scale = min(4.0, max(1.0, ID_DOCUMENT_OCR_TARGET_WIDTH / max(1, source.width)))
        if coordinate_scale > 1:
            source = source.resize((round(source.width * coordinate_scale), round(source.height * coordinate_scale)), LANCZOS_RESAMPLING)
    if ocr_pass is None:
        ocr_pass = extract_redaction_ocr_passes(source, ocr_lang=ocr_lang or resolve_ocr_lang())[0]
    data, processed, ocr_config = ocr_pass
    ocr_lang = str(data.get("_redaction_ocr_lang") or ocr_lang or resolve_ocr_lang())

    words: list[OCRWord] = []
    text_parts: list[str] = []
    cursor = 0
    texts = data.get("text", [])
    lefts = data.get("left", [])
    tops = data.get("top", [])
    widths = data.get("width", [])
    heights = data.get("height", [])
    block_numbers = data.get("block_num", [])
    paragraph_numbers = data.get("par_num", [])
    line_numbers = data.get("line_num", [])
    previous_line_id: tuple[int, int, int] | None = None

    for idx, raw in enumerate(texts):
        token = (raw or "").strip()
        if not token:
            continue
        line_id = (
            int(block_numbers[idx]) if idx < len(block_numbers) else 0,
            int(paragraph_numbers[idx]) if idx < len(paragraph_numbers) else 0,
            int(line_numbers[idx]) if idx < len(line_numbers) else 0,
        )
        if text_parts:
            separator = "\n" if previous_line_id != line_id else " "
            text_parts.append(separator)
            cursor += len(separator)
        start = cursor
        text_parts.append(token)
        cursor += len(token)
        end = cursor
        bbox = (
            round(int(lefts[idx]) / coordinate_scale),
            round(int(tops[idx]) / coordinate_scale),
            round((int(lefts[idx]) + int(widths[idx])) / coordinate_scale),
            round((int(tops[idx]) + int(heights[idx])) / coordinate_scale),
        )
        raw_bbox = (int(lefts[idx]), int(tops[idx]), int(lefts[idx])+int(widths[idx]), int(tops[idx])+int(heights[idx]))
        symbols = data.get("_redaction_hocr_words", {}).get(raw_bbox, [])
        exact_boxes = ()
        if "".join(value for value, _box in symbols).casefold() == token.casefold():
            exact_boxes = tuple(tuple(round(value/coordinate_scale) for value in box) for _char,box in symbols)
        words.append(OCRWord(text=token, bbox=bbox, start=start, end=end,
                             line_id=line_id, character_boxes=exact_boxes))
        previous_line_id = line_id
    if include_character_boxes and words:
        words = _attach_ocr_character_boxes(
            image=processed,
            words=words,
            ocr_lang=ocr_lang or resolve_ocr_lang(),
            ocr_config=ocr_config,
            coordinate_scale=coordinate_scale,
        )
    return "".join(text_parts), words


def _attach_ocr_character_boxes(
    *,
    image: Image.Image,
    words: Sequence[OCRWord],
    ocr_lang: str,
    ocr_config: str,
    coordinate_scale: float = 1.0,
) -> list[OCRWord]:
    """Attach exact Tesseract symbol boxes when they align with an OCR word."""
    try:
        raw_boxes = analysis_cached(
            "symbols", (image_fingerprint(image), ocr_lang, ocr_config),
            lambda: pytesseract.image_to_boxes(image, lang=ocr_lang, config=ocr_config),
        )
    except Exception:
        return list(words)

    symbols: list[tuple[str, tuple[int, int, int, int]]] = []
    for raw_line in str(raw_boxes or "").splitlines():
        fields = raw_line.rsplit(maxsplit=5)
        if len(fields) < 6:
            continue
        symbol = fields[0]
        try:
            left, bottom, right, top = (int(value) for value in fields[1:5])
        except ValueError:
            continue
        box = (
            round(left / coordinate_scale),
            round((image.height - top) / coordinate_scale),
            round(right / coordinate_scale),
            round((image.height - bottom) / coordinate_scale),
        )
        if symbol and box[2] > box[0] and box[3] > box[1]:
            symbols.append((symbol, box))

    attached: list[OCRWord] = []
    for word in words:
        x0, y0, x1, y1 = word.bbox
        candidates = [
            (symbol, box)
            for symbol, box in symbols
            if x0 - 3 <= (box[0] + box[2]) / 2 <= x1 + 3
            and y0 - 3 <= (box[1] + box[3]) / 2 <= y1 + 3
        ]
        candidates.sort(key=lambda item: (item[1][0], item[1][1]))
        character_boxes = (
            tuple(box for _symbol, box in candidates)
            if len(candidates) == len(word.text)
            and "".join(symbol for symbol, _box in candidates).casefold() == word.text.casefold()
            else ()
        )
        attached.append(
            OCRWord(
                text=word.text,
                bbox=word.bbox,
                start=word.start,
                end=word.end,
                line_id=word.line_id,
                character_boxes=character_boxes,
            )
        )
    return attached


def _normalize_text_for_compare(value: str) -> str:
    import re

    return re.sub(r"\s+", " ", value.strip()).casefold()


def _normalize_manual_redactions(values: Optional[Sequence[str]] = None) -> list[str]:
    cleaned: list[str] = []
    seen: set[str] = set()

    for value in values or []:
        term = str(value).strip()
        if not term:
            continue
        key = _normalize_text_for_compare(term)
        if key in seen:
            continue
        cleaned.append(term)
        seen.add(key)

    return cleaned


_MANUAL_MATCHES: ContextVar[Any] = ContextVar("redocx_redaction_manual_matches",default=None)


def _record_manual_match(quote: str) -> None:
    state = _MANUAL_MATCHES.get()
    if state is None:
        return
    for term in state["required"]:
        if _literal_text_findings(quote,custom_redactions=[term]):
            state["seen"].add(_normalize_text_for_compare(term))


@contextmanager
def _manual_redaction_scope(custom_redactions, review_exclusions, output_path: Path):
    exclusions = {_normalize_text_for_compare(term) for term in review_exclusions}
    required = [term for term in _normalize_manual_redactions(custom_redactions)
                if _normalize_text_for_compare(term) not in exclusions]
    state = {"required":required,"seen":set()}
    token = _MANUAL_MATCHES.set(state)
    try:
        yield
        if any(_normalize_text_for_compare(term) not in state["seen"] for term in required):
            raise ValueError("Some requested redaction text could not be located reliably. "
                             "Correct the text or provide a clearer document; no output was released.")
    except Exception:
        output_path.unlink(missing_ok=True)
        raise
    finally:
        _MANUAL_MATCHES.reset(token)


def _literal_text_findings(
    text: str,
    *,
    custom_redactions: Optional[Sequence[str]] = None,
    review_exclusions: Sequence[str] = (),
) -> list[TextFinding]:
    import re

    terms = _normalize_manual_redactions(custom_redactions)
    if not text or not terms:
        return []

    excluded = {_normalize_text_for_compare(item) for item in review_exclusions if item and item.strip()}
    findings: list[TextFinding] = []

    for term in terms:
        if _normalize_text_for_compare(term) in excluded:
            continue

        literal = r"\s+".join(re.escape(piece) for piece in re.split(r"\s+", term))
        pattern = re.compile(rf"(?=({literal}))", re.IGNORECASE)
        for match in pattern.finditer(text):
            start, end = match.span(1)
            quote = text[start:end]
            if not quote:
                continue
            findings.append(
                TextFinding(
                    start=start,
                    end=end,
                    quote=quote,
                    label="custom_redaction",
                    source="manual",
                )
            )

    return findings


def _manual_candidates_from_text(
    text: str,
    *,
    custom_redactions: Optional[Sequence[str]] = None,
    review_exclusions: Sequence[str] = (),
) -> list[DetectionCandidate]:
    terms = _normalize_manual_redactions(custom_redactions)
    if not text or not terms:
        return []

    excluded = {_normalize_text_for_compare(item) for item in review_exclusions if item and item.strip()}
    candidates: list[DetectionCandidate] = []

    for term in terms:
        if _normalize_text_for_compare(term) in excluded:
            continue

        findings = _literal_text_findings(
            text,
            custom_redactions=[term],
            review_exclusions=review_exclusions,
        )
        if not findings:
            continue

        candidates.append(
            DetectionCandidate(
                label="custom_redaction",
                quote=term,
                occurrences=len(findings),
                source="manual",
            )
        )

    return candidates


def _is_id_document_payload(payload: RedactionRequest) -> bool:
    return is_id_document_payload(payload)


def _id_document_field_findings(
    text: str,
    *,
    payload: RedactionRequest,
) -> list[TextFinding]:
    return id_document_field_findings(text, payload=payload)


def _uncached_redaction_findings(
    *,
    sdp: GoogleSDPClient,
    text: str,
    payload: RedactionRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[TextFinding]:
    detected = inspect_sensitive_text(
        sdp=sdp,
        text=text,
        targets=payload.target_data,
        review_exclusions=payload.review_exclusions,
    )
    if _is_id_document_payload(payload):
        # Explicit ID fields supplement valid generic findings. Missing field
        # labels must never erase an already detected name or address.
        detected.extend(_id_document_field_findings(text, payload=payload))
    manual = _literal_text_findings(
        text,
        custom_redactions=custom_redactions,
        review_exclusions=payload.review_exclusions,
    )
    for finding in manual:
        _record_manual_match(finding.quote)
    return merge_overlapping_findings([*detected, *manual], original_text=text)


def _redaction_findings(
    *, sdp: GoogleSDPClient, text: str, payload: RedactionRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[TextFinding]:
    if not analysis_active():
        return _uncached_redaction_findings(
            sdp=sdp, text=text, payload=payload, custom_redactions=custom_redactions,
        )
    findings = analysis_cached(
        "text-findings", (text, finding_context(sdp, payload, custom_redactions)),
        lambda: _uncached_redaction_findings(
            sdp=sdp, text=text, payload=payload, custom_redactions=custom_redactions,
        ),
    )
    # Cached detection must still satisfy the active manual-removal scope.
    # A generic finding can subsume a manual one during overlap merging.
    for finding in _literal_text_findings(
        text, custom_redactions=custom_redactions, review_exclusions=payload.review_exclusions,
    ):
        _record_manual_match(finding.quote)
    return list(findings)


def _ocr_word_overlap_box(
    word: OCRWord,
    finding: TextFinding,
) -> tuple[int, int, int, int] | None:
    overlap_start = max(word.start, finding.start)
    overlap_end = min(word.end, finding.end)
    if overlap_end <= overlap_start:
        return None

    x0, y0, x1, y1 = word.bbox
    character_count = max(1, word.end - word.start)
    if overlap_start == word.start and overlap_end == word.end:
        return word.bbox

    local_start = overlap_start - word.start
    local_end = overlap_end - word.start
    if len(word.character_boxes) == character_count:
        selected = word.character_boxes[local_start:local_end]
        if selected:
            return (
                min(box[0] for box in selected),
                min(box[1] for box in selected),
                max(box[2] for box in selected),
                max(box[3] for box in selected),
            )

    raise RuntimeError(
        "A partial scanned-word redaction could not be mapped to reliable character boxes. "
        "Redact the complete word or provide a clearer scan."
    )


def _fallback_ocr_words_for_quote(
    quote: str,
    words: Sequence[OCRWord],
) -> list[OCRWord]:
    needle = _normalize_text_for_compare(quote)
    if not needle:
        return []

    # The window scales with the requested literal/name instead of imposing a
    # fixed word limit. This keeps extremely long names mappable while still
    # terminating as soon as a candidate grows beyond the needle.
    maximum_window = max(1, len(needle.split()) + 2)
    for start_index in range(len(words)):
        buffered: list[OCRWord] = []
        for end_index in range(start_index, min(len(words), start_index + maximum_window)):
            buffered.append(words[end_index])
            candidate = _normalize_text_for_compare(
                " ".join(item.text for item in buffered)
            )
            if candidate == needle:
                return buffered
            if len(candidate) > len(needle):
                break
    return []


def _boxes_for_text_spans(findings: Sequence[TextFinding], words: Sequence[OCRWord]) -> list[tuple[int, int, int, int]]:
    boxes: list[tuple[int, int, int, int]] = []
    for finding in findings:
        matched = [w for w in words if not (w.end <= finding.start or w.start >= finding.end)]
        use_full_word_boxes = False
        if not matched:
            matched = _fallback_ocr_words_for_quote(finding.quote, words)
            use_full_word_boxes = True
        if not matched:
            continue

        # Never bridge unrelated OCR lines with one large rectangle. Addresses
        # and other multi-line findings are redacted line-by-line.
        by_line: dict[tuple[int, int, int], list[tuple[int, int, int, int]]] = {}
        for word in matched:
            box = word.bbox if use_full_word_boxes else _ocr_word_overlap_box(word, finding)
            if box is not None:
                by_line.setdefault(word.line_id, []).append(box)

        for line_boxes in by_line.values():
            x0 = min(box[0] for box in line_boxes)
            y0 = min(box[1] for box in line_boxes)
            x1 = max(box[2] for box in line_boxes)
            y1 = max(box[3] for box in line_boxes)
            if x1 > x0 and y1 > y0:
                boxes.append((x0, y0, x1, y1))

    unique: list[tuple[int, int, int, int]] = []
    seen: set[tuple[int, int, int, int]] = set()
    for box in boxes:
        if box not in seen:
            unique.append(box)
            seen.add(box)
    return unique


def draw_redaction_boxes(image: Image.Image, boxes: Sequence[tuple[int, int, int, int]]) -> Image.Image:
    out = image.convert("RGBA") if "transparency" in image.info else (image.copy() if image.mode in {"RGB", "RGBA", "L", "LA", "I;16", "I"} else image.convert("RGB"))
    draw = ImageDraw.Draw(out)
    for box in boxes:
        fill = (0, 0, 0, 255) if out.mode == "RGBA" else ((0, 255) if out.mode == "LA" else (0 if out.mode in {"L", "I;16", "I"} else DEFAULT_BLACK))
        draw.rectangle(box, fill=fill)
    return out


def _signature_target_selected(payload: RedactionRequest) -> bool:
    return any(
        str(getattr(target, "value", target)) == "signature"
        for target in payload.target_data
    )


def _visual_signature_is_excluded(payload: RedactionRequest) -> bool:
    excluded = {
        _normalize_text_for_compare(item)
        for item in payload.review_exclusions
        if item and item.strip()
    }
    return _normalize_text_for_compare(VISUAL_SIGNATURE_QUOTE) in excluded


def _clip_pixel_box(
    box: tuple[int, int, int, int],
    *,
    width: int,
    height: int,
    padding: int = 0,
) -> tuple[int, int, int, int] | None:
    x0, y0, x1, y1 = box
    clipped = (
        max(0, min(width, x0 - padding)),
        max(0, min(height, y0 - padding)),
        max(0, min(width, x1 + padding)),
        max(0, min(height, y1 + padding)),
    )
    return clipped if clipped[2] > clipped[0] and clipped[3] > clipped[1] else None


def _dark_content_bbox(
    image: Image.Image,
    search_box: tuple[int, int, int, int],
) -> tuple[int, int, int, int] | None:
    clipped = _clip_pixel_box(search_box, width=image.width, height=image.height)
    if clipped is None:
        return None

    grayscale = image.crop(clipped).convert("L")
    ink = grayscale.point(lambda value: 255 if value < 232 else 0)
    local = ink.getbbox()
    if local is None:
        return None

    left, top, right, bottom = local
    histogram = ink.histogram()
    ink_pixels = sum(histogram[1:])
    if right - left < 5 or bottom - top < 3 or ink_pixels < 16:
        return None
    return (
        clipped[0] + left,
        clipped[1] + top,
        clipped[0] + right,
        clipped[1] + bottom,
    )


def _ocr_signature_boxes(
    image: Image.Image,
    words: Sequence[OCRWord],
) -> list[tuple[int, int, int, int]]:
    lines: dict[tuple[int, int, int], list[OCRWord]] = {}
    for word in words:
        lines.setdefault(word.line_id, []).append(word)

    boxes: list[tuple[int, int, int, int]] = []
    for line_words in lines.values():
        ordered = sorted(line_words, key=lambda word: (word.bbox[0], word.start))
        line_text = " ".join(word.text for word in ordered)
        explicit_anchor = bool(_SIGNATURE_ANCHOR_RE.search(line_text))
        role_anchor = bool(
            _SIGNATURE_ROLE_RE.search(line_text)
            and min(word.bbox[1] for word in ordered) >= image.height * 0.45
        )
        if not explicit_anchor and not role_anchor:
            continue

        anchor = (
            min(word.bbox[0] for word in ordered),
            min(word.bbox[1] for word in ordered),
            max(word.bbox[2] for word in ordered),
            max(word.bbox[3] for word in ordered),
        )
        anchor_height = max(8, anchor[3] - anchor[1])
        above = (
            max(0, anchor[0] - anchor_height * 2),
            max(0, anchor[1] - max(90, anchor_height * 8)),
            min(image.width, anchor[2] + anchor_height * 3),
            max(0, anchor[1] - 2),
        )
        search_boxes = [above]
        if explicit_anchor:
            search_boxes.extend(
                [
                    (
                        anchor[2] + 2,
                        max(0, anchor[1] - anchor_height),
                        min(image.width, anchor[2] + max(180, anchor_height * 18)),
                        min(image.height, anchor[3] + max(45, anchor_height * 3)),
                    ),
                    (
                        max(0, anchor[0] - anchor_height),
                        anchor[3] + 2,
                        min(image.width, anchor[0] + max(220, anchor_height * 22)),
                        min(image.height, anchor[3] + max(70, anchor_height * 6)),
                    ),
                ]
            )

        for search_box in search_boxes:
            content_box = _dark_content_bbox(image, search_box)
            if content_box is not None:
                boxes.append(content_box)

    unique: list[tuple[int, int, int, int]] = []
    seen: set[tuple[int, int, int, int]] = set()
    for box in boxes:
        clipped = _clip_pixel_box(
            box,
            width=image.width,
            height=image.height,
            padding=3,
        )
        if clipped is not None and clipped not in seen:
            unique.append(clipped)
            seen.add(clipped)
    return unique


def _google_signature_boxes(
    *,
    image: Image.Image,
    sdp: GoogleSDPClient,
    payload: RedactionRequest,
) -> list[tuple[int, int, int, int]]:
    if not _signature_target_selected(payload) or _visual_signature_is_excluded(payload):
        return []

    encoded = BytesIO()
    image.convert("RGB").save(encoded, format="JPEG", quality=90)
    findings = inspect_sensitive_image(
        sdp=sdp,
        image_bytes=encoded.getvalue(),
        image_type="IMAGE_JPEG",
        targets=[target for target in payload.target_data if str(getattr(target, "value", target)) == "signature"],
        review_exclusions=payload.review_exclusions,
        min_likelihood=getattr(sdp, "min_likelihood", DEFAULT_MIN_LIKELIHOOD),
    )
    return [finding.bbox for finding in findings if finding.label == "signature"]


def _signature_boxes_for_image(
    *,
    image: Image.Image,
    words: Sequence[OCRWord],
    sdp: GoogleSDPClient,
    payload: RedactionRequest,
) -> list[tuple[int, int, int, int]]:
    if not _signature_target_selected(payload) or _visual_signature_is_excluded(payload):
        return []
    boxes = _google_signature_boxes(image=image, sdp=sdp, payload=payload)
    unique: list[tuple[int, int, int, int]] = []
    seen: set[tuple[int, int, int, int]] = set()
    for box in boxes:
        clipped = _clip_pixel_box(box, width=image.width, height=image.height, padding=2)
        if clipped is not None and clipped not in seen:
            unique.append(clipped)
            seen.add(clipped)
    return unique


def _render_page_surface_for_signatures(page: fitz.Page) -> OCRImageSurface:
    scale = min(
        4.0,
        max(2.5, ID_DOCUMENT_OCR_TARGET_WIDTH / max(1.0, page.rect.width)),
    )
    pix = page.get_pixmap(
        matrix=fitz.Matrix(scale, scale),
        alpha=False,
        colorspace=fitz.csRGB,
        annots=False,
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    return OCRImageSurface(image=image, page_rect=fitz.Rect(page.rect))


def _visual_signature_occurrence_count(
    *,
    source_path: Path,
    sdp: GoogleSDPClient,
    payload: RedactionRequest,
    ocr_lang: str,
) -> int:
    if not _signature_target_selected(payload) or _visual_signature_is_excluded(payload):
        return 0

    suffix = source_path.suffix.lower()
    if suffix in {".jpg", ".jpeg", ".png"}:
        image = Image.open(source_path).convert("RGB")
        _text, words = _ocr_words_from_image(image, ocr_lang=ocr_lang, sparse_text=True)
        return len(_signature_boxes_for_image(image=image, words=words, sdp=sdp, payload=payload))

    if suffix != ".pdf":
        return 0

    count = 0
    with fitz.open(source_path) as document:
        for page in document:
            surface = _render_page_surface_for_signatures(page)
            _text, words = _ocr_words_from_image(
                surface.image,
                ocr_lang=ocr_lang,
                sparse_text=True,
            )
            count += len(
                _signature_boxes_for_image(
                    image=surface.image,
                    words=words,
                    sdp=sdp,
                    payload=payload,
                )
            )
    return count


def _primary_page_image_surface(doc: fitz.Document, page: fitz.Page) -> OCRImageSurface | None:
    candidates: list[tuple[int, float, int, int, fitz.Rect]] = []
    seen: set[tuple[int, float, float, float, float]] = set()

    for image_info in page.get_images(full=True):
        xref = int(image_info[0])
        width = int(image_info[2])
        height = int(image_info[3])
        if xref <= 0 or width < 160 or height < 90:
            continue
        for placement in page.get_image_rects(xref):
            rect = fitz.Rect(placement) & page.rect
            key = (
                xref,
                round(rect.x0, 3),
                round(rect.y0, 3),
                round(rect.x1, 3),
                round(rect.y1, 3),
            )
            if key in seen or rect.is_empty or rect.is_infinite:
                continue
            seen.add(key)
            candidates.append((width * height, rect.get_area(), width, height, rect))

    if not candidates:
        return None

    _pixel_area, _page_area, source_width, source_height, rect = max(
        candidates,
        key=lambda item: (item[0], item[1]),
    )
    del doc
    del source_width, source_height
    render_scale = 2.0
    pix = page.get_pixmap(
        matrix=fitz.Matrix(render_scale, render_scale),
        clip=rect,
        alpha=False,
        colorspace=fitz.csRGB,
        annots=False,
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    return OCRImageSurface(image=image, page_rect=fitz.Rect(rect))


def _id_document_text_from_source(
    source_path: str | Path,
    *,
    ocr_lang: Optional[str] = None,
) -> str:
    source = Path(source_path)
    suffix = source.suffix.lower()
    chunks: list[str] = []

    if suffix == ".pdf":
        with fitz.open(source) as doc:
            for page in doc:
                native_text, _native_words = _page_words_with_offsets(page)
                native_text = native_text.strip()
                if native_text:
                    chunks.append(native_text)
                    continue
                surface = _primary_page_image_surface(doc, page)
                if surface is not None:
                    text, _words = _ocr_words_from_image(
                        surface.image,
                        ocr_lang=ocr_lang,
                        sparse_text=True,
                    )
                else:
                    pix = page.get_pixmap(
                        matrix=fitz.Matrix(
                            ID_DOCUMENT_FALLBACK_RENDER_SCALE,
                            ID_DOCUMENT_FALLBACK_RENDER_SCALE,
                        )
                    )
                    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    text, _words = _ocr_words_from_image(
                        image,
                        ocr_lang=ocr_lang,
                        sparse_text=True,
                    )
                if text.strip():
                    chunks.append(text)
        return "\n".join(chunks).strip()

    if suffix in {".jpg", ".jpeg", ".png"}:
        image = Image.open(source).convert("RGB")
        text, _words = _ocr_words_from_image(
            image,
            ocr_lang=ocr_lang,
            sparse_text=True,
        )
        return text.strip()

    return ""


def _image_box_to_page_rect(
    box: tuple[int, int, int, int],
    *,
    image_size: tuple[int, int],
    page_rect: fitz.Rect,
) -> fitz.Rect | None:
    image_width, image_height = image_size
    x0, y0, x1, y1 = box
    if image_width <= 0 or image_height <= 0 or x1 <= x0 or y1 <= y0:
        return None

    box_area_ratio = ((x1 - x0) * (y1 - y0)) / (image_width * image_height)
    if box_area_ratio > 0.25:
        return None

    x0 = max(0, min(image_width, x0))
    x1 = max(0, min(image_width, x1))
    y0 = max(0, min(image_height, y0))
    y1 = max(0, min(image_height, y1))
    if x1 <= x0 or y1 <= y0:
        return None

    return fitz.Rect(
        page_rect.x0 + (x0 / image_width) * page_rect.width,
        page_rect.y0 + (y0 / image_height) * page_rect.height,
        page_rect.x0 + (x1 / image_width) * page_rect.width,
        page_rect.y0 + (y1 / image_height) * page_rect.height,
    )


def _iter_table_paragraphs(table: Any):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_document_paragraphs(document: docx.Document, *, stories=None):
    # Do not materialize unused header/footer variants or process linked parts twice.
    if stories is None:
        stories = list(_redaction_docx_parts(document))
    for part, root in stories:
        for element in root.iter(qn("w:p")):
            paragraph = Paragraph(element, part if hasattr(part, "get_style") else document.part)
            paragraph._redaction_owner_part = part
            yield paragraph


def _docx_story_surfaces(document: Any, *, stories=None):
    for part,root in (stories if stories is not None else _redaction_docx_parts(document)):
        paragraphs=list(_iter_document_paragraphs(document,stories=[(part,root)]))
        parts=[];offsets=[];cursor=0
        for paragraph in paragraphs:
            text,_spans=_paragraph_run_spans(paragraph)
            if parts:
                cursor+=1
            offsets.append(cursor);parts.append(text);cursor+=len(text)
        yield paragraphs,offsets,parts,"\n".join(parts)


def _ensure_redaction_request(request: AnalyzerRequest) -> tuple[DocumentPayload, RedactionRequest]:
    if request.action != FeatureType.redact:
        raise ValueError("redact.py only handles action='redact'.")
    if not isinstance(request.input, DocumentPayload):
        raise ValueError("redact requires DocumentPayload input.")
    if not isinstance(request.payload, RedactionRequest):
        raise ValueError("redact requires RedactionRequest payload.")
    if request.input.metadata.input_format not in _SUPPORTED_INPUTS:
        raise ValueError("redact only supports pdf, docx, jpg, jpeg, png.")
    if request.policy.structure_preservation is not True:
        raise ValueError("structure_preservation must be True for redact.")
    return request.input, request.payload


def _output_path(source_path: Path, output_dir: Path, suffix: str = "_redacted") -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / f"{source_path.stem}{suffix}{source_path.suffix.lower()}"


def _build_redaction_response(*, request: AnalyzerRequest, output_path: Path) -> AnalyzerResponse:
    input_payload, _ = _ensure_redaction_request(request)
    result = build_document_file_result(
        filename=output_path.name,
        output_format=_FILE_OUTPUT_MAP[input_payload.metadata.input_format],
        file_size_mb=round(output_path.stat().st_size / (1024 * 1024), 4),
        algorithm_version="google-sdp-redact-v3",
    )
    response = AnalyzerResponse(
        action=FeatureType.redact,
        input_format=input_payload.metadata.input_format,
        policy=request.policy,
        system_language=request.system_language,
        result=result,
        human_review=HumanReviewRequirement(),
    )
    return validate_analyzer_response(response, request=request)


def _resolve_sdp(
    *,
    sdp: GoogleSDPClient | None = None,
    project_id: str | None = None,
    location: str = DEFAULT_DLP_LOCATION,
    min_likelihood: str = DEFAULT_MIN_LIKELIHOOD,
    client: Any | None = None,
) -> GoogleSDPClient:
    if sdp is not None:
        return sdp
    if not project_id:
        raise ValueError("Provide either sdp or project_id.")
    return build_google_sdp_client(
        project_id=project_id,
        location=location,
        min_likelihood=min_likelihood,
        client=client,
    )


def preview_redaction_candidates(
    request: AnalyzerRequest | Mapping[str, Any], *, sdp: GoogleSDPClient | None = None,
    project_id: str | None = None, location: str = DEFAULT_DLP_LOCATION,
    min_likelihood: str = DEFAULT_MIN_LIKELIHOOD, client: Any | None = None,
    custom_redactions: Optional[Sequence[str]] = None,
    ocr_languages: Optional[Sequence[str]] = None,
) -> list[DetectionCandidate]:
    req = validate_analyzer_request(request)
    input_payload, payload = _ensure_redaction_request(req)
    resolved = _resolve_sdp(sdp=sdp, project_id=project_id, location=location,
                            min_likelihood=min_likelihood, client=client)
    grouped = {}
    def add_surface(texts, image=None):
        # Complementary OCR passes describe the same surface. Use the greatest
        # occurrence count per quote, rather than counting each pass as a copy.
        counts = {}
        sources = {}
        quotes = {}
        for text in texts:
            current = {}
            for finding in _redaction_findings(sdp=resolved, text=text, payload=payload, custom_redactions=custom_redactions):
                key = (finding.label, _normalize_text_for_compare(finding.quote))
                current[key] = current.get(key, 0) + 1
                sources[key] = finding.source; quotes[key] = finding.quote
            for key, count in current.items():
                counts[key] = max(counts.get(key, 0), count)
        if image is not None:
            current = {}
            for finding in _image_findings(image, sdp=resolved, payload=payload, custom_redactions=custom_redactions):
                quote = VISUAL_SIGNATURE_QUOTE if finding.label == "signature" else finding.quote
                if not quote:
                    # A quote-less image detection still needs a reviewable item;
                    # excluding it by label is not silently inferred from OCR.
                    quote = "visual " + finding.label.replace("_", " ")
                key = (finding.label, _normalize_text_for_compare(quote))
                current[key] = current.get(key, 0) + 1
                sources[key] = "visual_signature" if finding.label == "signature" else finding.source
                quotes[key] = quote
            for key, count in current.items():
                counts[key] = max(counts.get(key, 0), count)
        for key, count in counts.items():
            label, _normalized = key
            existing = grouped.get(key)
            grouped[key] = DetectionCandidate(label=label, quote=quotes[key],
                occurrences=count + (existing.occurrences if existing else 0), source=sources[key])
    source = Path(str(input_payload.filename or ""))
    if not source.is_file():
        if not input_payload.text:
            raise ValueError("Redaction review requires a source document or extracted text.")
        add_surface([input_payload.text])
    elif source.suffix.lower() == ".docx":
        document = docx.Document(source)
        for _paragraphs,_offsets,_parts,text in _docx_story_surfaces(document):
            add_surface([text])
        for part in document.part.package.parts:
            if part.content_type.startswith("image/"):
                with Image.open(BytesIO(part.blob)) as raw:
                    image = ImageOps.exif_transpose(raw)
                    add_surface([text for text,_words in _image_text_passes(image, ocr_lang=resolve_ocr_lang(ocr_languages), partial=False)], image)
    elif source.suffix.lower() == ".pdf":
        with fitz.open(source) as document:
            document.bake(annots=True, widgets=True)
            for page in document:
                text, _chars = _pdf_character_records(page)
                texts = [text]
                image = None
                if page.get_image_info() or not text.strip():
                    add_surface([text])
                    review_surfaces = list(redaction_pdf_surfaces(page, original_source=True))
                    for image, _rect in review_surfaces:
                        add_surface([text for text,_words in _image_text_passes(image, ocr_lang=resolve_ocr_lang(ocr_languages), partial=False)], image)
                else:
                    if _signature_target_selected(payload):
                        pix = page.get_pixmap(matrix=fitz.Matrix(DEFAULT_PDF_RENDER_SCALE, DEFAULT_PDF_RENDER_SCALE), alpha=False, colorspace=fitz.csRGB)
                        image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    add_surface(texts, image)
            for _key, value, _findings in _pdf_metadata_findings(document=document, sdp=resolved, payload=payload, custom_redactions=custom_redactions):
                add_surface([value])
    else:
        with Image.open(source) as raw:
            image = ImageOps.exif_transpose(raw)
            add_surface([text for text,_words in _image_text_passes(image, ocr_lang=resolve_ocr_lang(ocr_languages), partial=False)], image)
    return [grouped[key] for key in sorted(grouped)]


def _paragraph_run_spans(paragraph: Any) -> tuple[str, list[RunSpan]]:
    parts = []
    spans = []
    cursor = 0
    for run_element in paragraph._p.iter(qn("w:r")):
        owner = run_element.getparent()
        while owner is not None and owner.tag != qn("w:p"):
            owner = owner.getparent()
        if owner is not paragraph._p:
            continue
        run = Run(run_element, paragraph)
        text = "".join((node.text or "") if node.tag in {qn("w:t"), qn("w:delText"), qn("w:instrText")} else
                       ("\t" if node.tag == qn("w:tab") else "\n")
                       for node in run_element if node.tag in {qn("w:t"), qn("w:delText"), qn("w:instrText"), qn("w:tab"), qn("w:br"), qn("w:cr")})
        if text:
            spans.append(RunSpan(cursor, cursor + len(text), text, deepcopy(run_element.rPr)))
            parts.append(text)
            cursor += len(text)
    return "".join(parts), spans


def _copy_rpr(dst_run: Any, rpr_xml: Any) -> None:
    if dst_run._r.rPr is not None:
        dst_run._r.remove(dst_run._r.rPr)
    if rpr_xml is not None:
        dst_run._r.insert(0, deepcopy(rpr_xml))


def _rewrite_paragraph_from_fragments(paragraph: Any, fragments: list[tuple[Any, str]]) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    for rpr_xml, text in fragments:
        if not text:
            continue
        new_run = paragraph.add_run()
        _copy_rpr(new_run, rpr_xml)
        new_run.text = text


def _run_span_for_offset(spans: Sequence[RunSpan], offset: int) -> RunSpan:
    for span in spans:
        if span.start <= offset < span.end:
            return span
    return spans[-1]


def redact_paragraph_runs(paragraph: Any, findings: Sequence[TextFinding]) -> None:
    full_text, _spans = _paragraph_run_spans(paragraph)
    if not full_text or not findings:
        return
    selected = merge_overlapping_findings(findings, original_text=full_text)
    cursor = 0
    for run_element in list(paragraph._p.iter(qn("w:r"))):
        owner = run_element.getparent()
        while owner is not None and owner.tag != qn("w:p"):
            owner = owner.getparent()
        if owner is not paragraph._p:
            continue
        run = Run(run_element, paragraph)
        for node in list(run_element):
            if node.tag in {qn("w:tab"), qn("w:br"), qn("w:cr")}:
                cursor += 1
                continue
            if node.tag not in {qn("w:t"), qn("w:delText"), qn("w:instrText")}:
                continue
            text = node.text or ""
            cuts = {0, len(text)}
            for finding in selected:
                if finding.start < cursor + len(text) and finding.end > cursor:
                    cuts.update((max(0, finding.start - cursor), min(len(text), finding.end - cursor)))
            cuts = sorted(cuts)
            if len(cuts) == 2 and not any(f.start <= cursor and f.end >= cursor + len(text) for f in selected):
                cursor += len(text)
                continue
            if node.tag != qn("w:t"):
                # Hidden field instructions / deleted revisions must not retain PII.
                node.text = _remove_metadata_findings(text, [TextFinding(
                    start=max(0, f.start-cursor), end=min(len(text), f.end-cursor),
                    quote="", label=f.label, source=f.source,
                ) for f in selected if f.start < cursor+len(text) and f.end > cursor])
                cursor += len(text)
                continue
            index = run_element.index(node)
            # Split only this text node. Existing images, breaks, field objects,
            # hyperlinks and their relative ordering are retained untouched.
            for a, b in zip(cuts, cuts[1:]):
                fragment = text[a:b]
                covered = any(f.start <= cursor+a and f.end >= cursor+b for f in selected)
                if covered:
                    for piece in re.findall(r"\s+|\S+", fragment):
                        if piece.isspace():
                            child = OxmlElement("w:t"); child.text = piece
                            child.set(qn("xml:space"), "preserve")
                        else:
                            child = _docx_mask_drawing(run, piece)
                        run_element.insert(index, child); index += 1
                else:
                    child = OxmlElement("w:t"); child.text = fragment
                    child.set(qn("xml:space"), "preserve")
                    run_element.insert(index, child); index += 1
            run_element.remove(node)
            cursor += len(text)


def _page_words_with_offsets(page: fitz.Page) -> tuple[str, list[PDFWord]]:
    raw = page.get_text("words", sort=True)
    words: list[PDFWord] = []
    parts: list[str] = []
    cursor = 0
    prev_line: tuple[int, int] | None = None

    for x0, y0, x1, y1, token, block_no, line_no, word_no in raw:
        token = (token or "").strip()
        if not token:
            continue

        current_line = (block_no, line_no)
        if parts:
            separator = "\n" if prev_line != current_line else " "
            parts.append(separator)
            cursor += len(separator)

        start = cursor
        parts.append(token)
        cursor += len(token)
        end = cursor

        words.append(
            PDFWord(
                start=start,
                end=end,
                text=token,
                rect=fitz.Rect(x0, y0, x1, y1),
                block_no=block_no,
                line_no=line_no,
                word_no=word_no,
            )
        )
        prev_line = current_line

    return "".join(parts), words


def _pdf_rects_for_findings(
    findings: Sequence[TextFinding],
    words: Sequence[PDFWord],
    *,
    original_text: str,
) -> list[fitz.Rect]:
    rects: list[fitz.Rect] = []
    merged = merge_overlapping_findings(findings, original_text=original_text)

    for finding in merged:
        matched = [w for w in words if not (w.end <= finding.start or w.start >= finding.end)]
        if not matched:
            continue

        by_line: dict[tuple[int, int], list[fitz.Rect]] = {}
        for word in matched:
            overlap_start = max(word.start, finding.start)
            overlap_end = min(word.end, finding.end)
            if overlap_end <= overlap_start:
                continue

            rect = fitz.Rect(word.rect)
            if overlap_start != word.start or overlap_end != word.end:
                character_count = max(1, word.end - word.start)
                original_x0 = rect.x0
                width = rect.width
                rect.x0 = original_x0 + width * (
                    (overlap_start - word.start) / character_count
                )
                rect.x1 = original_x0 + width * (
                    (overlap_end - word.start) / character_count
                )
            by_line.setdefault((word.block_no, word.line_no), []).append(rect)

        for line_rects in by_line.values():
            rect = line_rects[0]
            for word_rect in line_rects[1:]:
                rect = rect | word_rect
            rects.append(rect)

    unique: list[fitz.Rect] = []
    seen: set[tuple[float, float, float, float]] = set()
    for rect in rects:
        key = (round(rect.x0, 3), round(rect.y0, 3), round(rect.x1, 3), round(rect.y1, 3))
        if key not in seen:
            seen.add(key)
            unique.append(rect)
    return unique


def _pdf_word_subrect(
    word: PDFWord,
    start_index: int,
    end_index: int,
) -> fitz.Rect:
    rect = fitz.Rect(word.rect)
    character_count = max(1, len(word.text))
    original_x0 = rect.x0
    width = rect.width
    rect.x0 = original_x0 + width * (start_index / character_count)
    rect.x1 = original_x0 + width * (end_index / character_count)
    return rect


def _native_redaction_word_specs(
    findings: Sequence[TextFinding],
    words: Sequence[PDFWord],
    *,
    original_text: str,
) -> list[
    tuple[
        fitz.Rect,
        list[tuple[fitz.Rect, str]],
        list[fitz.Rect],
    ]
]:
    merged = merge_overlapping_findings(findings, original_text=original_text)
    specs: list[
        tuple[fitz.Rect, list[tuple[fitz.Rect, str]], list[fitz.Rect]]
    ] = []

    for word in words:
        covered = [False] * len(word.text)
        for finding in merged:
            overlap_start = max(word.start, finding.start)
            overlap_end = min(word.end, finding.end)
            for offset in range(overlap_start, overlap_end):
                index = offset - word.start
                if 0 <= index < len(covered):
                    covered[index] = True
        if not any(covered):
            continue

        safe_fragments: list[tuple[fitz.Rect, str]] = []
        sensitive_rects: list[fitz.Rect] = []
        run_start = 0
        run_value = covered[0]
        for index in range(1, len(covered) + 1):
            value = covered[index] if index < len(covered) else not run_value
            if value == run_value:
                continue
            rect = _pdf_word_subrect(word, run_start, index)
            if run_value:
                sensitive_rects.append(rect)
            else:
                safe_fragments.append((rect, word.text[run_start:index]))
            if index < len(covered):
                run_start = index
                run_value = value

        specs.append((fitz.Rect(word.rect), safe_fragments, sensitive_rects))

    return specs


def _fit_pdf_fragment_font_size(text: str, rect: fitz.Rect) -> float:
    font_size = max(3.5, min(10.0, rect.height * 0.78))
    if not text or rect.width <= 0:
        return font_size
    try:
        width = fitz.get_text_length(text, fontname="helv", fontsize=font_size)
    except Exception:
        return font_size
    if width > rect.width and width > 0:
        font_size = max(3.0, font_size * (rect.width / width) * 0.96)
    return font_size


def _restore_native_redaction_word(
    page: fitz.Page,
    *,
    safe_fragments: Sequence[tuple[fitz.Rect, str]],
    sensitive_rects: Sequence[fitz.Rect],
) -> None:
    for rect, text in safe_fragments:
        if not text:
            continue
        page.insert_text(
            fitz.Point(rect.x0, rect.y1 - max(0.5, rect.height * 0.15)),
            text,
            fontname="helv",
            fontsize=_fit_pdf_fragment_font_size(text, rect),
            color=(0, 0, 0),
            overlay=True,
        )
    for rect in sensitive_rects:
        page.draw_rect(
            rect,
            color=DEFAULT_BLACK,
            fill=DEFAULT_BLACK,
            width=0,
            overlay=True,
        )


def _pdf_rect_has_visible_content(page: fitz.Page, rect: fitz.Rect) -> bool:
    """Reject coordinates belonging only to invisible/covered PDF text.

    Some PDFs retain a hidden text layer from a different page. Redacting those
    coordinates creates bars in visually blank areas. A small rendered clip is
    the source of truth: if no visible ink exists there, no annotation is added.
    """
    candidate = fitz.Rect(rect)
    candidate_area = max(1.0, candidate.get_area())
    try:
        paint_log = list(page.get_bboxlog())
    except Exception:
        paint_log = []
    text_indexes = [
        index
        for index, item in enumerate(paint_log)
        if str(item[0]).endswith("text")
        and not (fitz.Rect(item[1]) & candidate).is_empty
    ]
    if text_indexes:
        text_index = max(text_indexes)
        for item in paint_log[text_index + 1:]:
            if item[0] != "fill-image":
                continue
            covered = fitz.Rect(item[1]) & candidate
            if not covered.is_empty and covered.get_area() / candidate_area >= 0.95:
                return False

    clipped = fitz.Rect(
        candidate.x0 - 0.5,
        candidate.y0 - 0.5,
        candidate.x1 + 0.5,
        candidate.y1 + 0.5,
    ) * page.rotation_matrix
    clipped &= page.rect
    if clipped.is_empty:
        return False
    pix = page.get_pixmap(
        matrix=fitz.Matrix(PDF_VISIBILITY_RENDER_SCALE, PDF_VISIBILITY_RENDER_SCALE),
        clip=clipped,
        alpha=False,
        colorspace=fitz.csRGB,
        annots=False,
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("L")
    histogram = image.histogram()
    visible_pixels = sum(histogram[:PDF_VISIBLE_PIXEL_THRESHOLD])
    minimum_pixels = max(2, round(image.width * image.height * 0.002))
    return visible_pixels >= minimum_pixels


def _apply_page_redactions(
    page: fitz.Page,
    *,
    image_mode: str,
    remove_graphics: bool = False,
) -> None:
    kwargs: dict[str, Any] = {}

    if image_mode == "none" and hasattr(fitz, "PDF_REDACT_IMAGE_NONE"):
        kwargs["images"] = fitz.PDF_REDACT_IMAGE_NONE
    elif image_mode == "pixels" and hasattr(fitz, "PDF_REDACT_IMAGE_PIXELS"):
        kwargs["images"] = fitz.PDF_REDACT_IMAGE_PIXELS

    if remove_graphics and hasattr(fitz, "PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED"):
        kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED
    elif hasattr(fitz, "PDF_REDACT_LINE_ART_NONE"):
        kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_NONE
    if hasattr(fitz, "PDF_REDACT_TEXT_REMOVE"):
        kwargs["text"] = fitz.PDF_REDACT_TEXT_REMOVE

    page.apply_redactions(**kwargs)


def _page_has_rendered_content(page: fitz.Page) -> bool:
    pixmap = page.get_pixmap(
        matrix=fitz.Matrix(0.5, 0.5),
        alpha=False,
        colorspace=fitz.csRGB,
        annots=False,
    )
    return min(pixmap.samples, default=255) < 250


def _save_pdf_checked(
    document: fitz.Document,
    output_path: Path,
    *,
    temporary_paths: list[Path],
    rasterized_pages: Optional[set[int]] = None,
) -> None:
    """Atomically publish a readable PDF without accepting blank-page loss."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    expected_content = [
        _page_has_rendered_content(page)
        for page in document
    ]
    expected_pixels = [page.get_pixmap(alpha=False, annots=False) for page in document]
    raster_signatures = {index: _raster_page_signature(document[index])
                         for index in (rasterized_pages or set())}
    last_error: Exception | None = None

    for _attempt in range(2):
        temporary_path = output_path.with_name(
            f".{output_path.stem}.{uuid.uuid4().hex}.tmp{output_path.suffix}"
        )
        temporary_paths.append(temporary_path)
        try:
            # Explicitly mark the catalog key dirty on each write. In supported
            # MuPDF versions, del_xml_metadata() alone can reappear on a second
            # save of the same document object.
            document.xref_set_key(document.pdf_catalog(), "Metadata", "null")
            # The first file is an unserved verification draft. Delay xref
            # compaction until the final write: MuPDF caches can otherwise keep
            # old colour-profile xrefs when a recovery image is inserted.
            document.save(temporary_path, garbage=0 if _attempt == 0 else 4,
                          deflate=True, clean=False)
            if not temporary_path.exists() or temporary_path.stat().st_size <= 0:
                raise RuntimeError("PDF writer produced an empty file.")

            with fitz.open(temporary_path) as verification:
                if len(verification) != len(expected_content):
                    raise RuntimeError("PDF writer changed the page count.")
                if verification.get_xml_metadata() or verification.embfile_names():
                    raise RuntimeError("PDF writer retained uninspected metadata or attachments.")
                roundtrip_failures = []
                for page_index, expected in enumerate(expected_content):
                    if verification.xref_get_key(verification[page_index].xref, "Thumb")[0] != "null":
                        raise RuntimeError("PDF writer retained an unredacted page thumbnail.")
                    if page_index in raster_signatures:
                        if _raster_page_signature(verification[page_index]) != raster_signatures[page_index]:
                            raise RuntimeError("PDF writer changed verified raster content.")
                    else:
                        try:
                            _verify_pdf_untouched_pixels(verification[page_index], expected_pixels[page_index], [])
                        except RuntimeError:
                            roundtrip_failures.append(page_index)
                    if expected and not _page_has_rendered_content(verification[page_index]):
                        raise RuntimeError(
                            f"PDF writer lost visible page content on page {page_index + 1}."
                        )
                if roundtrip_failures:
                    if _attempt:
                        raise RuntimeError("PDF writer could not preserve verified page appearances.")
                    # Some image/ICC combinations sample differently after a
                    # write. Rebuild only those pages from the already verified,
                    # redacted appearance, then verify their lossless pixels on
                    # the second write. Do not repeat an identical failed save.
                    for page_index in roundtrip_failures:
                        page = document[page_index]
                        rotation = page.rotation
                        page.set_rotation(0)
                        scene = page.get_displaylist(annots=False)
                        page.set_rotation(rotation)
                        _rasterize_redacted_pdf_page(page, scene, [])
                        raster_signatures[page_index] = _raster_page_signature(page)
                    continue

            if _attempt == 0:
                continue  # Only the final, garbage-collected file is published.
            temporary_path.replace(output_path)
            return
        except Exception as exc:
            last_error = exc
        finally:
            temporary_path.unlink(missing_ok=True)

    raise RuntimeError("Unable to create a complete redacted PDF safely.") from last_error


@lru_cache(maxsize=128)
def _docx_font_path(family: str, bold: bool, italic: bool) -> str:
    executable = shutil.which("fc-match")
    if not executable:
        raise RuntimeError("Fontconfig is required to preserve DOCX redaction geometry.")
    suffix = (":weight=bold" if bold else "") + (":slant=italic" if italic else "")
    def match(name):
        result = subprocess.run([executable,"-f","%{family}\n%{file}",name+suffix],
            capture_output=True,text=True,timeout=10,check=True)
        resolved_family, _separator, filename = result.stdout.strip().partition("\n")
        return resolved_family,filename
    resolved_family,filename = match(family)
    # Writer uses these metric-compatible substitutions for common Office
    # families. A generic system alias (e.g. Nimbus Roman) can have a different
    # advance even while LibreOffice uses Liberation Serif for the document.
    compatible = {"times new roman":"Liberation Serif","arial":"Liberation Sans",
                  "courier new":"Liberation Mono","calibri":"Carlito","cambria":"Caladea"}
    exact = family.casefold() in {value.strip().casefold() for value in resolved_family.split(",")}
    if not exact and family.casefold() in compatible:
        replacement = compatible[family.casefold()]
        candidate_family,candidate_file = match(replacement)
        if replacement.casefold() in {value.strip().casefold() for value in candidate_family.split(",")}:
            filename = candidate_file
    if not Path(filename).is_file():
        raise RuntimeError("A required DOCX font cannot be resolved.")
    return filename


def _docx_font_properties(run: Any) -> tuple[str, float, bool, bool]:
    fonts = [run.font]
    try:
        style = run.style
        while style is not None:
            fonts.append(style.font); style = style.base_style
        style = run._parent.style
        while style is not None:
            fonts.append(style.font); style = style.base_style
    except (AttributeError, KeyError):
        pass
    def resolved(key: str, default: Any):
        for font in fonts:
            value = getattr(font, key, None)
            if value is not None:
                return value
        return default
    # The document defaults may supply a font without a Normal style override.
    doc_part = getattr(run.part, "_document_part", run.part)
    defaults = doc_part.document.styles.element.xpath("./w:docDefaults/w:rPrDefault/w:rPr") if hasattr(doc_part, "document") else []
    def family_from_properties(properties):
        if properties is None:
            return None
        rf = properties.find(qn("w:rFonts"))
        if rf is None:
            return None
        explicit = rf.get(qn("w:ascii")) or rf.get(qn("w:hAnsi"))
        if explicit:
            return explicit
        theme = rf.get(qn("w:asciiTheme")) or rf.get(qn("w:hAnsiTheme"))
        if theme:
            for relationship in doc_part.rels.values():
                if relationship.reltype.endswith("/theme") and not relationship.is_external:
                    root = etree.fromstring(relationship.target_part.blob)
                    namespace = {"a": "http://schemas.openxmlformats.org/drawingml/2006/main"}
                    category = "majorFont" if theme.startswith("major") else "minorFont"
                    latin = root.find(".//a:"+category+"/a:latin", namespace)
                    if latin is not None and latin.get("typeface"):
                        return latin.get("typeface")
        return None
    family = None
    for font in fonts:
        family = family_from_properties(getattr(font._element, "rPr", None))
        if family:
            break
    size = resolved("size", None)
    if defaults:
        family = family or family_from_properties(defaults[0])
        sz = defaults[0].find(qn("w:sz"))
        if size is None and sz is not None:
            size = Pt(float(sz.get(qn("w:val"))) / 2)
    return family or "Calibri", float(size.pt if size is not None else 11), bool(resolved("bold", False)), bool(resolved("italic", False))


def _docx_mask_drawing(run: Any, text: str) -> Any:
    family, points, bold, italic = _docx_font_properties(run)
    font = ImageFont.truetype(_docx_font_path(family, bold, italic), size=max(1, round(points * 64)))
    # Word normally disables kerning for ordinary text. Preserve character
    # tracking/scaling where specified rather than replacing it with block glyphs.
    width = font.getlength(text, features=["-kern", "-liga"]) / 64
    rpr = run._r.rPr
    if rpr is not None:
        spacing = rpr.find(qn("w:spacing")); scale = rpr.find(qn("w:w"))
        if spacing is not None:
            width += len(text) * float(spacing.get(qn("w:val"), "0")) / 20
        if scale is not None:
            width *= float(scale.get(qn("w:val"), "100")) / 100
    mask = BytesIO(); Image.new("RGB", (1, 1), "black").save(mask, format="PNG"); mask.seek(0)
    # Inline drawing width is the removed text advance. No original sensitive
    # characters survive in XML, alternative text or a hidden run.
    shape = run.add_picture(mask, width=Pt(max(.05, width)), height=Pt(max(.5, points * .65)))
    for edge in ("distL", "distR", "distT", "distB"):
        shape._inline.set(edge, "0")
    owner_part = getattr(run._parent, "_redaction_owner_part", run.part)
    if owner_part is not run.part:
        from docx.opc.constants import RELATIONSHIP_TYPE
        blip = shape._inline.xpath(".//a:blip")[0]
        image_part = run.part.related_parts[blip.get(qn("r:embed"))]
        blip.set(qn("r:embed"), owner_part.relate_to(image_part, RELATIONSHIP_TYPE.IMAGE))
    drawing = shape._inline.getparent()
    run._r.remove(drawing)
    return drawing


def _docx_render_pdf(path: Path, directory: Path) -> Path:
    executable = os.environ.get("SOFFICE_PATH") or shutil.which("soffice")
    if not executable:
        raise RuntimeError("LibreOffice is required to verify DOCX redaction layout.")
    directory.mkdir(parents=True, exist_ok=True)
    profile = directory / "profile"
    subprocess.run([executable, "-env:UserInstallation=" + profile.resolve().as_uri(),
                    "--headless", "--convert-to", "pdf", "--outdir", str(directory), str(path.resolve())],
                   capture_output=True, timeout=120, check=True)
    output = directory / (path.stem + ".pdf")
    if not output.is_file():
        raise RuntimeError("DOCX redaction layout verification could not render the document.")
    return output


def _pdf_character_records(page: fitz.Page) -> tuple[str, list[dict[str, Any]]]:
    chars = []
    text = []
    for block in page.get_text("rawdict", sort=True)["blocks"]:
        for line in block.get("lines", []):
            for span in line["spans"]:
                for char in span["chars"]:
                    chars.append(dict(char, font=span["font"], size=span["size"], color=span["color"]))
                    text.append(char["c"])
            chars.append({"c": "\n"}); text.append("\n")
    return "".join(text), chars


def _verify_docx_layout(source: Path, output: Path, removed: Sequence[str]) -> None:
    with tempfile.TemporaryDirectory(prefix="redocx-redact-layout-") as temp:
        root = Path(temp)
        before = _docx_render_pdf(source, root / "before")
        after = _docx_render_pdf(output, root / "after")
        with fitz.open(before) as left, fitz.open(after) as right:
            if len(left) != len(right):
                raise RuntimeError("DOCX redaction would change pagination; no output was released.")
            for old, new in zip(left, right):
                if old.rect != new.rect:
                    raise RuntimeError("DOCX redaction would change page geometry.")
                text, chars = _pdf_character_records(old)
                findings = _literal_text_findings(text, custom_redactions=removed)
                expected = [c for i, c in enumerate(chars) if not c["c"].isspace()
                            and not any(f.start <= i < f.end for f in findings)]
                _text, after_chars = _pdf_character_records(new)
                actual = [c for c in after_chars if not c["c"].isspace()]
                if len(expected) != len(actual):
                    raise RuntimeError("DOCX redaction could not verify retained text; no output was released.")
                for a, b in zip(expected, actual):
                    if (a["c"] != b["c"] or a["font"] != b["font"] or abs(a["size"]-b["size"]) > .05
                            or a["color"] != b["color"] or max(abs(x-y) for x,y in zip(a["origin"],b["origin"])) > .8):
                        raise RuntimeError("DOCX redaction would move or restyle retained text; no output was released.")


def _uncached_image_findings(image: Image.Image, *, sdp: GoogleSDPClient, payload: RedactionRequest, custom_redactions=None):
    if not payload.target_data and not _normalize_manual_redactions(custom_redactions):
        return []
    targets = [target for target in payload.target_data if not (
        str(getattr(target, "value", target)) == "signature" and _visual_signature_is_excluded(payload))]
    if not targets and not _normalize_manual_redactions(custom_redactions):
        return []
    source = image.convert("RGB")
    # content.inspect permits a 0.5 MB request. Keep encoded content below
    # 350 KB to leave room for transport encoding and the inspection config.
    limit = 350_000
    def encoded(candidate):
        stream = BytesIO(); candidate.save(stream, format="PNG")
        if len(stream.getvalue()) <= limit:
            return stream.getvalue(), "IMAGE_PNG"
        stream = BytesIO(); candidate.save(stream, format="JPEG", quality=90, subsampling=0)
        return stream.getvalue(), "IMAGE_JPEG"
    data, kind = encoded(source)
    surfaces = []
    if len(data) <= limit:
        surfaces.append((source, (0,0,source.width,source.height), data, kind))
    else:
        # A whole-image preview preserves signature/context evidence. Overlapping
        # full-resolution tiles supplement it without reducing text resolution.
        preview = source.copy()
        while True:
            preview.thumbnail((max(64,round(preview.width*.75)),max(64,round(preview.height*.75))),LANCZOS_RESAMPLING)
            preview_data, preview_kind = encoded(preview)
            if len(preview_data) <= limit:
                break
            if max(preview.size) <= 64:
                raise RuntimeError("An image could not be inspected within the service request limit.")
        surfaces.append((preview,(0,0,source.width,source.height),preview_data,preview_kind))
        def tile(rect):
            crop = source.crop(rect); chunk, chunk_kind = encoded(crop)
            if len(chunk) <= limit:
                surfaces.append((crop,rect,chunk,chunk_kind)); return
            x0,y0,x1,y1 = rect
            if x1-x0 >= y1-y0:
                midpoint=(x0+x1)//2; overlap=min(64,max(1,(x1-x0)//8))
                tile((x0,y0,midpoint+overlap,y1)); tile((midpoint-overlap,y0,x1,y1))
            else:
                midpoint=(y0+y1)//2; overlap=min(64,max(1,(y1-y0)//8))
                tile((x0,y0,x1,midpoint+overlap)); tile((x0,midpoint-overlap,x1,y1))
        tile((0,0,source.width,source.height))
    excluded = {_normalize_text_for_compare(value) for value in payload.review_exclusions}
    findings = []
    for candidate, rect, data, kind in surfaces:
        for finding in inspect_sensitive_image(sdp=sdp, image_bytes=data, image_type=kind,
                targets=targets, review_exclusions=payload.review_exclusions,
                custom_redactions=_normalize_manual_redactions(custom_redactions),
                min_likelihood=getattr(sdp,"min_likelihood",DEFAULT_MIN_LIKELIHOOD)):
            quote = finding.quote or (VISUAL_SIGNATURE_QUOTE if finding.label == "signature" else
                                      "visual " + finding.label.replace("_"," "))
            if _normalize_text_for_compare(quote) in excluded:
                continue
            if finding.label == "custom_redaction":
                _record_manual_match(quote)
            x0,y0,x1,y1=rect; left,top,right,bottom=finding.bbox
            box = _clip_pixel_box((math.floor(x0+left/candidate.width*(x1-x0)),
                                   math.floor(y0+top/candidate.height*(y1-y0)),
                                   math.ceil(x0+right/candidate.width*(x1-x0)),
                                   math.ceil(y0+bottom/candidate.height*(y1-y0))),
                                  width=source.width,height=source.height,padding=0)
            if box is None:
                raise RuntimeError("An image finding could not be mapped to source pixels.")
            duplicate = next((index for index,existing in enumerate(findings) if existing.label==finding.label and
                            _normalize_text_for_compare(existing.quote)==_normalize_text_for_compare(finding.quote) and
                            (fitz.Rect(existing.bbox)&fitz.Rect(box)).get_area() /
                            max(1,min(fitz.Rect(existing.bbox).get_area(),fitz.Rect(box).get_area())) > .5
                            ), None)
            if duplicate is not None:
                previous = findings[duplicate]
                union = fitz.Rect(previous.bbox) | fitz.Rect(box)
                findings[duplicate] = ImageFinding(bbox=tuple(round(value) for value in union),
                    quote=finding.quote,label=finding.label,source=finding.source)
            else:
                findings.append(ImageFinding(bbox=box,quote=finding.quote,label=finding.label,source=finding.source))
    return findings


def _image_findings(image: Image.Image, *, sdp: GoogleSDPClient, payload: RedactionRequest, custom_redactions=None):
    if not analysis_active():
        return _uncached_image_findings(
            image, sdp=sdp, payload=payload, custom_redactions=custom_redactions,
        )
    findings = analysis_cached(
        "image-findings", (image_fingerprint(image), finding_context(sdp, payload, custom_redactions)),
        lambda: _uncached_image_findings(
            image, sdp=sdp, payload=payload, custom_redactions=custom_redactions,
        ),
    )
    for finding in findings:
        if finding.label == "custom_redaction":
            _record_manual_match(finding.quote)
    return list(findings)


def _image_text_passes(image: Image.Image, *, ocr_lang: str, partial: bool):
    for data, processed, config in extract_redaction_ocr_passes(image, ocr_lang=ocr_lang):
        yield _ocr_words_from_image(image, ocr_lang=ocr_lang, include_character_boxes=partial,
                                    ocr_pass=(data, processed, config))


def _ocr_redaction_boxes(image: Image.Image, *, sdp: GoogleSDPClient, payload: RedactionRequest,
                          custom_redactions: Optional[Sequence[str]], ocr_lang: str):
    boxes = []
    for data, processed, config in extract_redaction_ocr_passes(image, ocr_lang=ocr_lang):
        text, words = _ocr_words_from_image(image, ocr_lang=ocr_lang,
                                            ocr_pass=(data, processed, config))
        findings = _redaction_findings(sdp=sdp, text=text, payload=payload, custom_redactions=custom_redactions)
        partial = [word for word in words if any(
            finding.start < word.end and finding.end > word.start and
            not (finding.start <= word.start and finding.end >= word.end)
            for finding in findings)]
        partial = [word for word in partial if len(word.character_boxes) != len(word.text)]
        if partial:
            refined = _attach_ocr_character_boxes(
                image=processed, words=partial,
                ocr_lang=str(data.get("_redaction_ocr_lang") or ocr_lang), ocr_config=config,
            )
            lookup = {word.start: word for word in refined}
            words = [lookup.get(word.start, word) for word in words]
        boxes.extend(_boxes_for_text_spans(findings, words))
    return boxes


def _redaction_boxes_for_image(image: Image.Image, *, sdp: GoogleSDPClient, payload: RedactionRequest,
                               custom_redactions: Optional[Sequence[str]], ocr_lang: str):
    boxes = _ocr_redaction_boxes(image, sdp=sdp, payload=payload,
                                 custom_redactions=custom_redactions, ocr_lang=ocr_lang)
    # Image detections supplement OCR rather than relying on text-adjacent ink.
    boxes.extend(f.bbox for f in _image_findings(image, sdp=sdp, payload=payload, custom_redactions=custom_redactions))
    return list(dict.fromkeys(boxes))


def _manual_image_boxes(image: Image.Image, *, sdp: GoogleSDPClient, payload: RedactionRequest,
                        custom_redactions, ocr_lang: str):
    if not _normalize_manual_redactions(custom_redactions):
        return []
    return _ocr_redaction_boxes(image, sdp=sdp, payload=payload.model_copy(update={"target_data": []}),
                                custom_redactions=custom_redactions, ocr_lang=ocr_lang)


def _stable_image_boxes(image: Image.Image, boxes, *, sdp: GoogleSDPClient,
                        payload: RedactionRequest, custom_redactions, ocr_lang: str):
    """A mask can change OCR segmentation. Remove newly readable manual matches."""
    result = list(boxes)
    if not boxes or not _normalize_manual_redactions(custom_redactions):
        return result
    for _attempt in range(3):
        remaining = _manual_image_boxes(draw_redaction_boxes(image,result), sdp=sdp, payload=payload,
                                         custom_redactions=custom_redactions, ocr_lang=ocr_lang)
        if not remaining:
            return result
        result.extend(remaining)
    raise RuntimeError("Requested text remained readable after image redaction; no output was released.")


def _inverse_exif_image(image: Image.Image, orientation: int) -> Image.Image:
    transpose = getattr(Image, "Transpose", Image)
    operations = {2: transpose.FLIP_LEFT_RIGHT, 3: transpose.ROTATE_180,
                  4: transpose.FLIP_TOP_BOTTOM, 5: transpose.TRANSPOSE,
                  6: transpose.ROTATE_90, 7: transpose.TRANSVERSE, 8: transpose.ROTATE_270}
    return image.transpose(operations[orientation]) if orientation in operations else image


def _display_boxes_to_raw(boxes, size, orientation):
    # A mask plane provides the exact inverse for all eight EXIF transforms.
    plane = Image.new("1", ImageOps.exif_transpose(_orientation_probe(size, orientation)).size)
    draw = ImageDraw.Draw(plane)
    for box in boxes:
        draw.rectangle(box, fill=1)
    raw = _inverse_exif_image(plane, orientation)
    import numpy as np
    pixels = np.asarray(raw)
    # Separate rectangles by connected components, not one page-sized union.
    import cv2
    _count, _labels, stats, _centres = cv2.connectedComponentsWithStats(pixels.astype("uint8"), 8)
    return [(int(x), int(y), int(x+w), int(y+h)) for x,y,w,h,area in stats[1:] if area]


def _orientation_probe(size, orientation):
    image = Image.new("L", size)
    exif = image.getexif(); exif[274] = orientation
    image.info["exif"] = exif.tobytes()
    return image


def _lossless_jpeg_redact(source: Path, output: Path, boxes) -> None:
    """Remove whole affected MCU blocks without recompressing other blocks."""
    worker = Path(__file__).with_name("_jpeg_blocks.py")
    try:
        subprocess.run([sys.executable, str(worker)],
                       input=json.dumps({"source": str(source), "output": str(output), "boxes": boxes}),
                       text=True, capture_output=True, check=True, timeout=90)
        _copy_jpeg_presentation_metadata(source, output)
        with Image.open(output) as verified, Image.open(source) as original:
            verified.load()
            if verified.size != original.size:
                raise RuntimeError("JPEG redaction changed image dimensions.")
    except Exception as exc:
        output.unlink(missing_ok=True)
        raise RuntimeError("The JPEG could not be redacted losslessly; no output was released.") from exc


def _copy_jpeg_presentation_metadata(source: Path, output: Path) -> None:
    with Image.open(source) as original:
        exif = Image.Exif()
        for key in (274, 282, 283, 296):
            if key in original.getexif():
                exif[key] = original.getexif()[key]
        exif_bytes = exif.tobytes() if len(exif) else b""
        icc = original.info.get("icc_profile", b"")
        dpi = original.info.get("dpi")
        jfif_unit = original.info.get("jfif_unit", 0)
        jfif_density = original.info.get("jfif_density", (1, 1))
    data = output.read_bytes()
    if data[:2] != b"\xff\xd8":
        raise RuntimeError("JPEG writer produced invalid data.")
    segments = []
    offset = 2
    while offset < len(data):
        start = offset
        if data[offset] != 255:
            raise RuntimeError("JPEG metadata could not be verified.")
        marker = data[offset+1]; offset += 2
        if marker in (0xDA, 0xD9):
            tail = data[start:]; break
        length = int.from_bytes(data[offset:offset+2], "big")
        segment = data[start:offset+length]; offset += length
        if marker not in (0xE1, 0xE2, 0xED, 0xFE):
            # JFIF density is replaced below when the source reports DPI.
            if marker != 0xE0:
                segments.append(segment)
    else:
        raise RuntimeError("JPEG writer produced a truncated file.")
    def marker(kind, payload):
        return b"\xff" + bytes([kind]) + (len(payload)+2).to_bytes(2,"big") + payload
    presentation = []
    density = dpi if dpi else jfif_density
    unit = 1 if dpi else int(jfif_unit)
    dx,dy = [max(1,min(65535,round(v))) for v in density]
    presentation.append(marker(0xE0,b"JFIF\0\x01\x01"+bytes([unit])+dx.to_bytes(2,"big")+dy.to_bytes(2,"big")+b"\0\0"))
    if exif_bytes:
        presentation.append(marker(0xE1,exif_bytes))
    if icc:
        size=65519; total=math.ceil(len(icc)/size)
        for i in range(total):
            presentation.append(marker(0xE2,b"ICC_PROFILE\0"+bytes([i+1,total])+icc[i*size:(i+1)*size]))
    output.write_bytes(b"\xff\xd8"+b"".join(presentation+segments)+tail)


def _redact_docx(*, source_path: Path, output_path: Path, sdp: GoogleSDPClient,
                 payload: RedactionRequest, custom_redactions: Optional[Sequence[str]] = None,
                 ocr_languages: Optional[Sequence[str]] = None) -> None:
    document = docx.Document(source_path)
    for part in document.part.package.parts:
        if str(part.partname).startswith(("/word/embeddings/", "/word/charts/", "/word/diagrams/")):
            raise ValueError("This Word document contains embedded objects, charts or SmartArt that cannot be inspected safely. "
                             "Replace those objects with ordinary text or PNG/JPEG images before redaction.")
        for relationship in part.rels.values():
            if relationship.is_external and (relationship.reltype.endswith("/image") or relationship.reltype.endswith("/oleObject")):
                raise ValueError("Linked external Word objects cannot be inspected safely; embed their content before redaction.")
    removed = []
    original_images = [part for part in document.part.package.parts if part.content_type.startswith("image/")]
    stories = list(_redaction_docx_parts(document))
    if any(next(root.iter(qn("w:dataBinding")), None) is not None for _part,root in stories):
        raise ValueError("Data-bound Word controls can restore source text; remove their data binding before redaction.")
    for paragraphs,offsets,parts,text in _docx_story_surfaces(document,stories=stories):
        findings = _redaction_findings(sdp=sdp,text=text,payload=payload,
            custom_redactions=custom_redactions) if text.strip() else []
        for paragraph,start,part_text in zip(paragraphs,offsets,parts):
            local=[]
            for finding in findings:
                left=max(0,finding.start-start);right=min(len(part_text),finding.end-start)
                if right>left:
                    quote=part_text[left:right]
                    local.append(TextFinding(start=left,end=right,quote=quote,label=finding.label,source=finding.source))
                    removed.append(quote)
            if local:
                redact_paragraph_runs(paragraph,local)
    # Notes and comments may use generic OPC Parts; serialize their changed roots.
    # Relationship URLs and alternative text are content too, not formatting.
    for part, root in stories:
        for element in root.iter():
            for key, value in list(element.attrib.items()):
                if key.rsplit("}",1)[-1] in {"descr", "title", "author", "initials"}:
                    findings = _redaction_findings(sdp=sdp, text=value, payload=payload, custom_redactions=custom_redactions)
                    if findings:
                        element.set(key, _remove_metadata_findings(value, findings))
        if getattr(part, "_element", None) is None:
            part._blob = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
    for part in document.part.package.parts:
        for relationship in list(part.rels.values()):
            if relationship.is_external and "hyperlink" in relationship.reltype:
                value = str(relationship.target_ref)
                if _redaction_findings(sdp=sdp, text=value, payload=payload, custom_redactions=custom_redactions):
                    relationship._target = "about:blank"
        if part not in original_images:
            continue
        if part.content_type in {"image/png", "image/jpeg"}:
            with tempfile.TemporaryDirectory(prefix="redocx-redact-image-") as temp:
                suffix = ".png" if part.content_type == "image/png" else ".jpg"
                source = Path(temp)/("source"+suffix); target = Path(temp)/("redacted"+suffix)
                source.write_bytes(part.blob)
                _redact_image_file(source_path=source, output_path=target, sdp=sdp,
                                   payload=payload, custom_redactions=custom_redactions, ocr_languages=ocr_languages)
                part._blob = target.read_bytes()
        elif part.content_type.startswith("image/"):
            try:
                with Image.open(BytesIO(part.blob)) as raw:
                    image = ImageOps.exif_transpose(raw)
                    boxes = _redaction_boxes_for_image(image, sdp=sdp, payload=payload,
                        custom_redactions=custom_redactions, ocr_lang=resolve_ocr_lang(ocr_languages))
                    encoded = BytesIO()
                    options = {key: raw.info[key] for key in ("dpi", "icc_profile") if key in raw.info}
                    draw_redaction_boxes(image, boxes).save(encoded,format="PNG",**options)
                    from docx.opc.packuri import PackURI
                    part._blob = encoded.getvalue()
                    part._content_type = "image/png"
                    part._partname = PackURI("/word/media/redacted_"+uuid.uuid4().hex+".png")
            except Exception as exc:
                raise RuntimeError("An embedded Word image cannot be inspected safely; convert that image to PNG or JPEG before redaction.") from exc
    for key in ("author", "last_modified_by", "title", "subject", "keywords", "comments", "identifier", "category"):
        value = getattr(document.core_properties, key, "") or ""
        findings = _redaction_findings(sdp=sdp, text=value, payload=payload, custom_redactions=custom_redactions) if value else []
        if findings:
            setattr(document.core_properties, key, _remove_metadata_findings(value, findings))
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        document.save(output_path)
        if removed:
            try:
                _verify_docx_layout(source_path, output_path, removed)
            except RuntimeError:
                # Inline replacement drawings can move a keep-with-next group
                # across a page boundary. Recover from the original appearance
                # instead of relaxing layout verification or publishing drift.
                output_path.unlink(missing_ok=True)
                _redact_docx_as_page_images(
                    source_path=source_path, output_path=output_path, sdp=sdp,
                    payload=payload, custom_redactions=custom_redactions,
                    ocr_languages=ocr_languages,
                )
    except Exception:
        output_path.unlink(missing_ok=True)
        raise


def _redact_docx_as_page_images(*, source_path, output_path, sdp, payload,
                               custom_redactions=None, ocr_languages=None) -> None:
    """Layout fallback: a fresh DOCX containing only verified redacted pages.

    Original document XML, revisions, fields, hyperlinks, attachments and source
    images are not copied. Affected documents retain their appearance, page sizes
    and DOCX container; page content becomes non-editable raster images.
    """
    from docx.enum.section import WD_SECTION_START
    from zipfile import ZipFile

    with tempfile.TemporaryDirectory(prefix="redocx-redact-word-pages-") as directory:
        root = Path(directory)
        source_pdf = _docx_render_pdf(source_path, root / "source")
        safe_pdf = root / "redacted.pdf"
        _redact_pdf(source_path=source_pdf, output_path=safe_pdf, sdp=sdp,
                    payload=payload, custom_redactions=custom_redactions,
                    ocr_languages=ocr_languages)
        fresh = docx.Document()
        for key in ("author", "last_modified_by", "title", "subject", "keywords", "comments"):
            setattr(fresh.core_properties, key, "")
        expected_images = []
        geometry = []
        with fitz.open(safe_pdf) as pages:
            for index, page in enumerate(pages):
                section = fresh.sections[0] if index == 0 else fresh.add_section(WD_SECTION_START.NEW_PAGE)
                section.page_width, section.page_height = Pt(page.rect.width), Pt(page.rect.height)
                section.top_margin = section.bottom_margin = Pt(0)
                section.left_margin = section.right_margin = Pt(0)
                section.header_distance = section.footer_distance = Pt(0)
                geometry.append((page.rect.width, page.rect.height))
                pix = page.get_pixmap(matrix=fitz.Matrix(4, 4), colorspace=fitz.csRGB, alpha=False)
                image = pix.tobytes("png")
                expected_images.append(hashlib.sha256(image).hexdigest())
                paragraph = fresh.add_paragraph()
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = Pt(1)
                run = paragraph.add_run()
                run.font.size = Pt(1)
                picture = run.add_picture(BytesIO(image), width=Pt(page.rect.width), height=Pt(page.rect.height))
                inline = picture._inline
                anchor = OxmlElement("wp:anchor")
                for key, value in {"distT":"0", "distB":"0", "distL":"0", "distR":"0",
                                   "simplePos":"0", "relativeHeight":"0", "behindDoc":"1",
                                   "locked":"0", "layoutInCell":"1", "allowOverlap":"1"}.items():
                    anchor.set(key, value)
                position = OxmlElement("wp:simplePos")
                position.set("x", "0"); position.set("y", "0"); anchor.append(position)
                for axis in ("H", "V"):
                    position = OxmlElement("wp:position" + axis)
                    position.set("relativeFrom", "page")
                    offset = OxmlElement("wp:posOffset"); offset.text = "0"
                    position.append(offset); anchor.append(position)
                for name in ("wp:extent", "wp:effectExtent"):
                    element = inline.find(qn(name))
                    if element is not None:
                        anchor.append(deepcopy(element))
                anchor.append(OxmlElement("wp:wrapNone"))
                for name in ("wp:docPr", "wp:cNvGraphicFramePr", "a:graphic"):
                    element = inline.find(qn(name))
                    if element is not None:
                        anchor.append(deepcopy(element))
                inline.getparent().replace(inline, anchor)
        fresh.save(output_path)
        # Check stored source data as well as the rendered Word pagination.
        with ZipFile(output_path) as package:
            actual_images = [hashlib.sha256(package.read(name)).hexdigest()
                             for name in package.namelist() if name.startswith("word/media/")]
            if set(actual_images) != set(expected_images):
                raise RuntimeError("Word fallback could not verify stored page images.")
            xml = etree.fromstring(package.read("word/document.xml"))
            if any((node.text or "").strip() for node in xml.iter(qn("w:t"))):
                raise RuntimeError("Word fallback retained unexpected editable text.")
        verified_pdf = _docx_render_pdf(output_path, root / "verified")
        with fitz.open(verified_pdf) as verified:
            if len(verified) != len(geometry) or any(
                    abs(page.rect.width-width) > .1 or abs(page.rect.height-height) > .1
                    for page, (width, height) in zip(verified, geometry)):
                raise RuntimeError("Word fallback could not preserve page geometry.")
            if any(not _page_has_rendered_content(page) for page in verified):
                raise RuntimeError("Word fallback lost a rendered document page.")


def _redact_image_file(*, source_path: Path, output_path: Path, sdp: GoogleSDPClient,
                       payload: RedactionRequest, custom_redactions: Optional[Sequence[str]] = None,
                       ocr_languages: Optional[Sequence[str]] = None) -> None:
    try:
        with Image.open(source_path) as raw:
            raw.load()
            info = dict(raw.info)
            orientation = int(raw.getexif().get(274, 1))
            image = ImageOps.exif_transpose(raw)
            boxes = _redaction_boxes_for_image(image, sdp=sdp, payload=payload,
                                              custom_redactions=custom_redactions,
                                              ocr_lang=resolve_ocr_lang(ocr_languages))
            boxes = _stable_image_boxes(image,boxes,sdp=sdp,payload=payload,
                custom_redactions=custom_redactions,ocr_lang=resolve_ocr_lang(ocr_languages))
            output_path.parent.mkdir(parents=True, exist_ok=True)
            if raw.format == "JPEG":
                raw_boxes = _display_boxes_to_raw(boxes, raw.size, orientation) if boxes else []
                try:
                    _lossless_jpeg_redact(source_path, output_path, raw_boxes)
                except RuntimeError:
                    if raw.mode != "CMYK":
                        raise
                    # Some jpeglib builds cannot verify the fourth component.
                    # Rebuild from decoded, erased RGB pixels rather than
                    # returning a visually black file that still contains K.
                    display = image.convert("RGB")
                    profile = None
                    if info.get("icc_profile"):
                        from PIL import ImageCms
                        source_profile = ImageCms.ImageCmsProfile(BytesIO(info["icc_profile"]))
                        target_profile = ImageCms.ImageCmsProfile(ImageCms.createProfile("sRGB"))
                        display = ImageCms.profileToProfile(image, source_profile, target_profile, outputMode="RGB")
                        profile = target_profile.tobytes()
                    result = _inverse_exif_image(draw_redaction_boxes(display, boxes), orientation)
                    exif = Image.Exif()
                    for key in (274, 282, 283, 296):
                        if key in raw.getexif():
                            exif[key] = raw.getexif()[key]
                    options = {"exif": exif.tobytes(), "quality": 100, "subsampling": 0}
                    if info.get("dpi"):
                        options["dpi"] = info["dpi"]
                    if profile:
                        options["icc_profile"] = profile
                    result.save(output_path, format="JPEG", **options)
                    with Image.open(output_path) as checked:
                        checked.load()
                        if checked.mode != "RGB" or checked.size != raw.size:
                            raise RuntimeError("CMYK recovery could not verify the JPEG geometry.")
            else:
                result = _inverse_exif_image(draw_redaction_boxes(image, boxes), orientation)
                # Only presentation metadata is copied; textual metadata and EXIF
                # thumbnails can contain the original sensitive data.
                options = {key: info[key] for key in ("dpi", "icc_profile") if key in info}
                exif = Image.Exif()
                for key in (274, 282, 283, 296):
                    if key in raw.getexif():
                        exif[key] = raw.getexif()[key]
                if len(exif):
                    options["exif"] = exif.tobytes()
                from PIL.PngImagePlugin import PngInfo
                import struct
                presentation = PngInfo()
                if "gamma" in info:
                    presentation.add(b"gAMA", struct.pack(">I", round(float(info["gamma"])*100000)))
                if "srgb" in info:
                    presentation.add(b"sRGB", bytes([int(info["srgb"])]))
                if "chromaticity" in info:
                    presentation.add(b"cHRM", struct.pack(">8I", *(round(float(value)*100000) for value in info["chromaticity"])))
                result.save(output_path, format="PNG", pnginfo=presentation, **options)
            if boxes and _normalize_manual_redactions(custom_redactions):
                with Image.open(output_path) as encoded:
                    if _manual_image_boxes(ImageOps.exif_transpose(encoded),sdp=sdp,payload=payload,
                            custom_redactions=custom_redactions,ocr_lang=resolve_ocr_lang(ocr_languages)):
                        raise RuntimeError("Requested text remained readable in the encoded image; no output was released.")
    except Exception:
        output_path.unlink(missing_ok=True)
        raise


def _pdf_metadata_findings(
    *,
    document: fitz.Document,
    sdp: GoogleSDPClient,
    payload: RedactionRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[tuple[str, str, list[TextFinding]]]:
    metadata = dict(document.metadata or {})
    target_values = {
        str(getattr(target, "value", target))
        for target in payload.target_data
    }
    exclusions = {
        _normalize_text_for_compare(item)
        for item in payload.review_exclusions
        if item and item.strip()
    }
    results: list[tuple[str, str, list[TextFinding]]] = []

    for key in _PDF_REVIEWABLE_METADATA_KEYS:
        value = metadata.get(key)
        if not isinstance(value, str) or not value.strip():
            continue
        findings = _redaction_findings(
            sdp=sdp,
            text=value,
            payload=payload,
            custom_redactions=custom_redactions,
        )
        normalized_value = _normalize_text_for_compare(value)
        if (
            key == "author"
            and "name" in target_values
            and normalized_value not in exclusions
            and normalized_value not in _GENERIC_METADATA_AUTHORS
        ):
            findings = merge_overlapping_findings(
                [
                    *findings,
                    TextFinding(
                        start=0,
                        end=len(value),
                        quote=value,
                        label="name",
                        source="pdf_metadata",
                    ),
                ],
                original_text=value,
            )
        if findings:
            results.append((key, value, findings))
    return results


def _remove_metadata_findings(value: str, findings: Sequence[TextFinding]) -> str:
    result = value
    for finding in sorted(findings, key=lambda item: item.start, reverse=True):
        result = result[:finding.start] + result[finding.end:]
    return result.strip()


def _redact_pdf_metadata(
    *,
    document: fitz.Document,
    sdp: GoogleSDPClient,
    payload: RedactionRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    metadata = dict(document.metadata or {})
    changed = False
    for key, value, findings in _pdf_metadata_findings(
        document=document,
        sdp=sdp,
        payload=payload,
        custom_redactions=custom_redactions,
    ):
        metadata[key] = _remove_metadata_findings(value, findings)
        changed = True
    if changed:
        document.set_metadata(metadata)


def _verify_pdf_untouched_pixels(page: fitz.Page, before: fitz.Pixmap, masks) -> None:
    import numpy as np
    after = page.get_pixmap(alpha=False, annots=False)
    if (before.width, before.height, before.n) != (after.width, after.height, after.n):
        raise RuntimeError("PDF redaction changed the rendered page geometry.")
    expected = np.frombuffer(before.samples, dtype="uint8").reshape(before.height,before.width,before.n)
    actual = np.frombuffer(after.samples, dtype="uint8").reshape(after.height,after.width,after.n)
    permitted = Image.new("1", (before.width,before.height))
    draw = ImageDraw.Draw(permitted)
    for mask in masks:
        rect = fitz.Rect(mask) * page.rotation_matrix
        # Permit raster edge antialiasing at the redaction boundary only.
        draw.rectangle((math.floor(rect.x0)-2,math.floor(rect.y0)-2,
                        math.ceil(rect.x1)+2,math.ceil(rect.y1)+2),fill=1)
    # Rewriting an image XObject can change subpixel sampling without changing
    # its placement or typography. Require symmetric local agreement within one
    # rendered pixel and eight intensity levels; disappearing artwork fails.
    import cv2
    kernel = np.ones((3,3),dtype="uint8")
    old_min=cv2.erode(expected,kernel).astype("int16");old_max=cv2.dilate(expected,kernel).astype("int16")
    new_min=cv2.erode(actual,kernel).astype("int16");new_max=cv2.dilate(actual,kernel).astype("int16")
    old=expected.astype("int16");new=actual.astype("int16")
    changed = np.any((new < old_min-8)|(new > old_max+8)|(old < new_min-8)|(old > new_max+8),axis=2)
    if np.any(changed & ~np.asarray(permitted,dtype=bool)):
        raise RuntimeError("PDF redaction would change content outside the selected regions; no output was released.")


def _signature_shares_source_graphics(drawings, signature_masks) -> bool:
    """Inspect original artwork only, never the black rectangles we create."""
    for drawing in drawings:
        rect = fitz.Rect(drawing["rect"])
        if drawing.get("type") in {"s", "fs"}:
            padding = max(0.0, float(drawing.get("width") or 0.0)) / 2
            rect = fitz.Rect(rect.x0-padding, rect.y0-padding,
                             rect.x1+padding, rect.y1+padding)
        if not any(not (rect & mask).is_empty for mask in signature_masks):
            continue
        if not any(mask.contains(rect) for mask in signature_masks):
            return True
    return False


def _rasterize_redacted_pdf_page(page, source_display_list, masks) -> None:
    """Replace a difficult page with verified pixels from its original appearance.

    The source display list is captured before any removal. All selected areas
    are erased in the pixels before embedding. Original content streams and
    resources are detached, so no editable text, vector path or original image
    survives behind an opaque overlay. Final garbage collection removes orphan
    objects. This fallback is local to the affected page; its boxes and rotation
    remain unchanged. Other pages keep their native representation.
    """
    area = max(1.0, source_display_list.rect.get_area())
    scale = min(4.0, math.sqrt(24_000_000 / area))
    if scale < 1:
        raise ValueError("This PDF page is too large to redact at a readable resolution.")
    pix = source_display_list.get_pixmap(
        matrix=fitz.Matrix(scale, scale), colorspace=fitz.csRGB, alpha=False,
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    draw = ImageDraw.Draw(image)
    for mask in masks:
        rect = fitz.Rect(mask)
        draw.rectangle((math.floor(rect.x0*scale-pix.x), math.floor(rect.y0*scale-pix.y),
                        math.ceil(rect.x1*scale-pix.x), math.ceil(rect.y1*scale-pix.y)),
                       fill=DEFAULT_BLACK)
    encoded = BytesIO()
    image.save(encoded, format="PNG")

    rotation = page.rotation
    try:
        page.set_rotation(0)
        document = page.parent
        empty_stream = document.get_new_xref()
        document.update_object(empty_stream, "<<>>")
        document.update_stream(empty_stream, b"")
        page.set_contents(empty_stream)
        document.xref_set_key(page.xref, "Resources", "<<>>")
        # A cached thumbnail may carry an unredacted miniature of this page.
        document.xref_set_key(page.xref, "Thumb", "null")
        image_xref = page.insert_image(page.rect, stream=encoded.getvalue(), keep_proportion=False)
        recovered = fitz.Pixmap(document, image_xref)
        if ((recovered.width, recovered.height, recovered.n) != (image.width, image.height, 3)
                or recovered.samples != image.tobytes()
                or page.get_text("text").strip() or page.get_drawings()
                or len(page.get_image_info()) != 1):
            raise RuntimeError("A PDF page could not be flattened and verified safely.")
    finally:
        page.set_rotation(rotation)


def _raster_page_signature(page):
    """Lossless pixel and placement check, independent of renderer ICC sampling."""
    images = page.get_image_info(xrefs=True)
    if (len(images) != 1 or page.get_text("text").strip()
            or page.get_drawings() or page.get_fonts()):
        raise RuntimeError("Verified raster page contains unexpected source objects.")
    info = images[0]
    if not info.get("xref"):
        raise RuntimeError("Verified raster page has no inspectable image object.")
    pix = fitz.Pixmap(page.parent, info["xref"])
    return (pix.width, pix.height, pix.n, hashlib.sha256(pix.samples).hexdigest(),
            tuple(round(value, 3) for value in info["bbox"]),
            tuple(round(value, 3) for value in info["transform"]))


def _redact_pdf(*, source_path: Path, output_path: Path, sdp: GoogleSDPClient,
                payload: RedactionRequest, custom_redactions: Optional[Sequence[str]] = None,
                ocr_languages: Optional[Sequence[str]] = None,
                render_scale: float = DEFAULT_PDF_RENDER_SCALE) -> None:
    temporary_paths: list[Path] = []
    rasterized_pages: set[int] = set()
    ocr_lang = resolve_ocr_lang(ocr_languages)
    try:
        with fitz.open(source_path) as document:
            expected_geometry = [(tuple(p.mediabox), tuple(p.cropbox), p.rotation) for p in document]
            # Bake form/annotation appearances so their visible and editable
            # values pass through the same permanent removal path.
            document.bake(annots=True, widgets=True)
            document.del_xml_metadata()
            _redact_pdf_metadata(document=document, sdp=sdp, payload=payload, custom_redactions=custom_redactions)
            for name in list(document.embfile_names()):
                document.embfile_del(name)
            for page in document:
                # Apply redactions can invalidate MuPDF's cached link chain.
                # Inspect/remove sensitive actions before modifying page streams.
                try:
                    links = list(page.get_links())
                except (IndexError, RuntimeError):
                    # Annotation/widget appearances were baked above. Dropping
                    # an unreadable action array preserves their visible content
                    # without retaining an uninspectable hidden destination.
                    document.xref_set_key(page.xref, "Annots", "[]")
                    links = []
                for link in links:
                    value = " ".join(str(link.get(key) or "") for key in ("uri", "file", "name")).strip()
                    if value and _redaction_findings(sdp=sdp, text=value, payload=payload, custom_redactions=custom_redactions):
                        page.delete_link(link)
                page = document.reload_page(page)
                visual_before = page.get_pixmap(alpha=False, annots=False)
                original_rotation = page.rotation
                page.set_rotation(0)
                original_scene = page.get_displaylist(annots=False)
                source_drawings = page.get_drawings()
                page.set_rotation(original_rotation)
                document.xref_set_key(page.xref, "Thumb", "null")
                bounds = fitz.Rect(page.rect) * page.derotation_matrix
                text, chars = _pdf_character_records(page)
                findings = _redaction_findings(sdp=sdp, text=text, payload=payload, custom_redactions=custom_redactions)
                selected = {i for f in findings for i in range(f.start, min(len(chars), f.end)) if "bbox" in chars[i]}
                masks = []
                signature_masks = []
                for i in selected:
                    rect = fitz.Rect(chars[i]["bbox"]) & bounds
                    if not rect.is_empty and _pdf_rect_has_visible_content(page, rect):
                        masks.append(rect)

                # OCR is additive on mixed pages. Never choose a native-text OR
                # raster branch: a single native heading must not hide a scan.
                if page.get_image_info() or not text.strip() or _signature_target_selected(payload):
                    surfaces = list(redaction_pdf_surfaces(page, original_source=True)) if page.get_image_info() or not text.strip() else []
                    if _signature_target_selected(payload):
                        pix = page.get_pixmap(matrix=fitz.Matrix(render_scale, render_scale),
                                              alpha=False, colorspace=fitz.csRGB, annots=False)
                        surfaces.append((Image.frombytes("RGB", (pix.width, pix.height), pix.samples), fitz.Rect(page.rect)))
                    for surface_index, (image, surface_rect) in enumerate(surfaces):
                        raster_boxes = []
                        # The full-page signature surface is not a third OCR pass.
                        is_signature_surface = bool(_signature_target_selected(payload) and surface_rect == page.rect and surface_index == len(surfaces)-1)
                        if not is_signature_surface:
                            raster_boxes.extend(_ocr_redaction_boxes(image, sdp=sdp, payload=payload,
                                custom_redactions=custom_redactions, ocr_lang=ocr_lang))
                        visual_payload = payload.model_copy(update={"target_data": [target for target in payload.target_data
                            if not is_signature_surface or str(getattr(target,"value",target)) == "signature"]})
                        visual_findings = _image_findings(image, sdp=sdp, payload=visual_payload,
                            custom_redactions=custom_redactions if not is_signature_surface else None)
                        raster_boxes.extend(f.bbox for f in visual_findings)
                        signature_boxes = {f.bbox for f in visual_findings if f.label == "signature"}
                        for box in raster_boxes:
                            rotated = fitz.Rect(surface_rect.x0 + box[0]/image.width*surface_rect.width,
                                                surface_rect.y0 + box[1]/image.height*surface_rect.height,
                                                surface_rect.x0 + box[2]/image.width*surface_rect.width,
                                                surface_rect.y0 + box[3]/image.height*surface_rect.height)
                            rect = (rotated * page.derotation_matrix) & bounds
                            if rect.is_empty:
                                raise RuntimeError("A detected PDF field could not be mapped to page coordinates.")
                            hit = {i for i,c in enumerate(chars) if "bbox" in c and
                                   not (fitz.Rect(c["bbox"]) & rect).is_empty and not c["c"].isspace()}
                            if box in signature_boxes:
                                signature_masks.append(rect)
                            if hit and hit <= selected and box not in signature_boxes:
                                continue
                            selected.update(hit)
                            masks.append(rect)

                saved_rotation = page.rotation
                page.set_rotation(0)
                # Remove only chosen glyph objects. Keep safe neighbouring text
                # in its original font, size, colour and position, even within a word.
                for i in selected:
                    rect = fitz.Rect(chars[i]["bbox"]) & bounds
                    if rect.is_empty:
                        continue
                    centre = (rect.tl + rect.br) / 2
                    probe = fitz.Rect(centre.x-.01, centre.y-.01, centre.x+.01, centre.y+.01)
                    page.add_redact_annot(probe, fill=None, cross_out=False)
                if selected:
                    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_NONE,
                                          graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                                          text=fitz.PDF_REDACT_TEXT_REMOVE)
                # Pixel removal and visible bars run separately so a scan's
                # bounding rectangle cannot erase/restyle neighbouring PDF text.
                for rect in masks:
                    page.add_redact_annot(rect, fill=DEFAULT_BLACK, cross_out=False)
                if masks:
                    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,
                                          graphics=fitz.PDF_REDACT_LINE_ART_NONE,
                                          text=fitz.PDF_REDACT_TEXT_NONE)
                # Confirmed visual signatures may be vector paths. Remove their
                # source graphics, then verify that unrelated artwork survived.
                flatten_page = _signature_shares_source_graphics(source_drawings, signature_masks)
                for rect in signature_masks:
                    page.add_redact_annot(rect, fill=DEFAULT_BLACK, cross_out=False)
                if signature_masks:
                    page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,
                                          graphics=(fitz.PDF_REDACT_LINE_ART_NONE if flatten_page else
                                                    fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED),
                                          text=fitz.PDF_REDACT_TEXT_NONE)
                page.set_rotation(saved_rotation)
                if masks and page.get_image_info() and _normalize_manual_redactions(custom_redactions):
                    for _attempt in range(3):
                        extra_masks=[]
                        for image, surface_rect in redaction_pdf_surfaces(page):
                            boxes = _manual_image_boxes(image,sdp=sdp,payload=payload,
                                custom_redactions=custom_redactions,ocr_lang=ocr_lang)
                            for box in boxes:
                                rotated = fitz.Rect(surface_rect.x0+box[0]/image.width*surface_rect.width,
                                    surface_rect.y0+box[1]/image.height*surface_rect.height,
                                    surface_rect.x0+box[2]/image.width*surface_rect.width,
                                    surface_rect.y0+box[3]/image.height*surface_rect.height)
                                extra_masks.append((rotated*page.derotation_matrix)&bounds)
                        if not extra_masks:
                            break
                        page.set_rotation(0)
                        for rect in extra_masks:
                            if rect.is_empty:
                                raise RuntimeError("A remaining manual redaction could not be mapped safely.")
                            hit={i for i,c in enumerate(chars) if "bbox" in c and not(fitz.Rect(c["bbox"])&rect).is_empty}
                            selected.update(hit)
                            page.add_redact_annot(rect,fill=DEFAULT_BLACK,cross_out=False)
                        page.apply_redactions(images=fitz.PDF_REDACT_IMAGE_PIXELS,
                            graphics=fitz.PDF_REDACT_LINE_ART_NONE,text=fitz.PDF_REDACT_TEXT_REMOVE)
                        masks.extend(extra_masks)
                        page.set_rotation(saved_rotation)
                    else:
                        raise RuntimeError("Requested text remained readable after PDF redaction; no output was released.")
                # A nearby/overlapping glyph bbox must not cause collateral loss.
                expected = [c for i,c in enumerate(chars) if i not in selected and not c["c"].isspace()]
                _new_text, new_chars = _pdf_character_records(page)
                actual = [c for c in new_chars if not c["c"].isspace()]
                native_text_changed = len(expected) != len(actual) or any(
                    a["c"] != b["c"] or a["font"] != b["font"] or a["color"] != b["color"]
                    or abs(a["size"]-b["size"]) > .01
                    or max(abs(x-y) for x,y in zip(a["origin"], b["origin"])) > .05
                    for a,b in zip(expected,actual)
                )
                if not flatten_page and not native_text_changed:
                    try:
                        _verify_pdf_untouched_pixels(page, visual_before, masks)
                    except RuntimeError:
                        # Recover from collateral edits using the original
                        # appearance, not the already modified page.
                        flatten_page = True
                if flatten_page or native_text_changed:
                    _rasterize_redacted_pdf_page(page, original_scene, masks)
                    rasterized_pages.add(page.number)
                    # The fallback verifies every embedded pixel at 288 dpi
                    # against the original scene with only the masks applied.
                    # Comparing its 72-dpi downsample to native font hinting is
                    # not an equivalent test and can reject identical content.
            if expected_geometry != [(tuple(p.mediabox), tuple(p.cropbox), p.rotation) for p in document]:
                raise RuntimeError("PDF redaction changed page geometry; no output was released.")
            _save_pdf_checked(document, output_path, temporary_paths=temporary_paths,
                              rasterized_pages=rasterized_pages)
            with fitz.open(output_path) as verified:
                if expected_geometry != [(tuple(p.mediabox), tuple(p.cropbox), p.rotation) for p in verified]:
                    raise RuntimeError("PDF writer changed page geometry; no output was released.")
    except Exception:
        output_path.unlink(missing_ok=True)
        raise
    finally:
        for path in temporary_paths:
            path.unlink(missing_ok=True)


def apply_redaction(
    request: AnalyzerRequest | Mapping[str, Any],
    *,
    source_path: str | Path,
    output_dir: str | Path,
    sdp: GoogleSDPClient | None = None,
    project_id: str | None = None,
    location: str = DEFAULT_DLP_LOCATION,
    min_likelihood: str = DEFAULT_MIN_LIKELIHOOD,
    client: Any | None = None,
    ocr_languages: Optional[Sequence[str]] = None,
    custom_redactions: Optional[Sequence[str]] = None,
) -> AnalyzerResponse:
    req = validate_analyzer_request(request)
    input_payload, payload = _ensure_redaction_request(req)
    source = Path(source_path)
    if not source.exists():
        raise FileNotFoundError(f"Source file not found: {source}")
    output_path = _output_path(source, Path(output_dir), suffix="_redacted")
    resolved = _resolve_sdp(
        sdp=sdp,
        project_id=project_id,
        location=location,
        min_likelihood=min_likelihood,
        client=client,
    )

    with _manual_redaction_scope(custom_redactions,payload.review_exclusions,output_path):
        fmt = input_payload.metadata.input_format
        if fmt == DocumentInputFormat.docx:
            _redact_docx(
                source_path=source,
                output_path=output_path,
                sdp=resolved,
                payload=payload,
                custom_redactions=custom_redactions,
                ocr_languages=ocr_languages,
            )
        elif fmt == DocumentInputFormat.pdf:
            _redact_pdf(
                source_path=source,
                output_path=output_path,
                sdp=resolved,
                payload=payload,
                custom_redactions=custom_redactions,
                ocr_languages=ocr_languages,
            )
        elif fmt in {DocumentInputFormat.jpg, DocumentInputFormat.jpeg, DocumentInputFormat.png}:
            _redact_image_file(
                source_path=source,
                output_path=output_path,
                sdp=resolved,
                payload=payload,
                custom_redactions=custom_redactions,
                ocr_languages=ocr_languages,
            )
        else:
            raise ValueError("Unsupported redact input format.")
    return _build_redaction_response(request=req, output_path=output_path)


__all__ = [
    "preview_redaction_candidates",
    "redact_text_with_findings",
    "draw_redaction_boxes",
    "redact_paragraph_runs",
    "apply_redaction",
]
