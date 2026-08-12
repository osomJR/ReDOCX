from __future__ import annotations

import re
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
from typing import Any, Mapping, Optional, Sequence
import uuid

import docx
from docx.text.paragraph import Paragraph
import fitz
import pytesseract
from PIL import Image, ImageDraw

from backend.src.extraction import (
    OCR_CONFIG,
    extract_ocr_data,
    preprocess_for_ocr,
    resolve_ocr_lang,
)
from backend.src.processing.data_protection.client import (
    DEFAULT_DLP_LOCATION,
    DEFAULT_MIN_LIKELIHOOD,
    DetectionCandidate,
    GoogleSDPClient,
    TextFinding,
    build_google_sdp_client,
    id_document_field_findings,
    inspect_sensitive_image,
    inspect_sensitive_text,
    is_id_document_payload,
    merge_overlapping_findings,
)
from backend.src.schema import (
    AnalyzerRequest,
    AnalyzerResponse,
    DataMaskingRequest,
    DocumentInputFormat,
    DocumentPayload,
    FeatureType,
    HumanReviewRequirement,
    SensitiveDataType,
)
from backend.src.validation import (
    build_document_file_result,
    validate_analyzer_request,
    validate_analyzer_response,
)

DEFAULT_PDF_RENDER_SCALE = 2.0
MAX_PDF_OCR_RENDER_SCALE = 6.0
MIN_EMBEDDED_IMAGE_AREA_RATIO = 0.001
MAX_EMBEDDED_IMAGE_AREA_RATIO_FOR_SIGNATURE = 0.20
MASK_BOX_PADDING_PX = 3
MAX_INSPECTION_CHARS = 100_000
INSPECTION_OVERLAP_CHARS = 512

_SIGNATURE_ANCHOR_RE = re.compile(
    r"(?i)\b(?:authorized\s+signature|customer\s+signature|applicant\s+signature|"
    r"holder(?:'s)?\s+signature|signature|signed\s+by|signatory|"
    r"signature\s+autoris[eé]e|sign[eé]\s+par)\b"
)
_SIGNATURE_WORDS = {
    "signature",
    "signatory",
    "signed",
    "signé",
    "signée",
}
_DOCX_VISUAL_LOCAL_NAMES = {"drawing", "pict", "object", "AlternateContent"}

_IMAGE_OUTPUT_MAP: dict[DocumentInputFormat, str] = {
    DocumentInputFormat.jpg: "JPEG",
    DocumentInputFormat.jpeg: "JPEG",
    DocumentInputFormat.png: "PNG",
}

_FILE_OUTPUT_MAP = {
    DocumentInputFormat.pdf: "pdf",
    DocumentInputFormat.docx: "docx",
    DocumentInputFormat.jpg: "jpg",
    DocumentInputFormat.jpeg: "jpeg",
    DocumentInputFormat.png: "png",
}

_SUPPORTED_INPUTS = {
    DocumentInputFormat.pdf,
    DocumentInputFormat.docx,
    DocumentInputFormat.jpg,
    DocumentInputFormat.jpeg,
    DocumentInputFormat.png,
}

_CUSTOM_MASK_LABEL = "custom_mask"
VISUAL_SIGNATURE_QUOTE = "visual signature"
PDF_VISIBILITY_RENDER_SCALE = 2.0
PDF_VISIBLE_PIXEL_THRESHOLD = 245
_PDF_REVIEWABLE_METADATA_KEYS = ("author", "title", "subject", "keywords")
_GENERIC_METADATA_AUTHORS = {"admin", "administrator", "anonymous", "unknown", "user"}

_SIGNATURE_ROLE_RE = re.compile(
    r"(?i)\b(?:registrar|registar|istrar|piatrer|vice\s+chancellor|chancellor|authorized\s+officer|"
    r"approving\s+officer|director)\b"
)


@dataclass(frozen=True)
class OCRWord:
    text: str
    bbox: tuple[int, int, int, int]
    start: int
    end: int
    page_num: int = 0
    block_num: int = 0
    line_num: int = 0
    character_boxes: tuple[tuple[int, int, int, int], ...] = ()


@dataclass(frozen=True)
class RunSpan:
    start: int
    end: int
    text: str
    rpr_xml: Any
    run: Any = None


@dataclass(frozen=True)
class PDFWord:
    start: int
    end: int
    text: str
    rect: fitz.Rect
    block_no: int
    line_no: int
    word_no: int


@dataclass(frozen=True)
class PDFOCRRegion:
    rect: fitz.Rect
    image: Image.Image
    transform: fitz.Matrix


def _ocr_words_from_data(data: Mapping[str, Sequence[object]]) -> tuple[str, list[OCRWord]]:
    words: list[OCRWord] = []
    text_parts: list[str] = []
    cursor = 0
    texts = data.get("text", [])
    lefts = data.get("left", [])
    tops = data.get("top", [])
    widths = data.get("width", [])
    heights = data.get("height", [])
    page_numbers = data.get("page_num", [])
    block_numbers = data.get("block_num", [])
    line_numbers = data.get("line_num", [])
    previous_line: tuple[object, object, object] | None = None

    for idx, raw in enumerate(texts):
        token = str(raw or "").strip()
        if not token:
            continue

        def item(values: Sequence[object]) -> object:
            return values[idx] if idx < len(values) else 0

        def integer(values: Sequence[object]) -> int:
            try:
                return int(float(item(values)))
            except (TypeError, ValueError):
                return 0

        current_line = (
            item(page_numbers),
            item(block_numbers),
            item(line_numbers),
        )
        if text_parts:
            separator = "\n" if current_line != previous_line else " "
            text_parts.append(separator)
            cursor += len(separator)
        start = cursor
        text_parts.append(token)
        cursor += len(token)
        end = cursor
        bbox = (
            integer(lefts),
            integer(tops),
            integer(lefts) + max(0, integer(widths)),
            integer(tops) + max(0, integer(heights)),
        )
        words.append(
            OCRWord(
                text=token,
                bbox=bbox,
                start=start,
                end=end,
                page_num=integer(page_numbers),
                block_num=integer(block_numbers),
                line_num=integer(line_numbers),
            )
        )
        previous_line = current_line
    return "".join(text_parts), words


def _ocr_words_from_image(
    image: Image.Image,
    *,
    ocr_lang: Optional[str] = None,
    include_character_boxes: bool = False,
) -> tuple[str, list[OCRWord]]:
    data = extract_ocr_data(
        image.convert("RGB"),
        ocr_lang=ocr_lang or resolve_ocr_lang(),
        config=OCR_CONFIG,
    )
    text, words = _ocr_words_from_data(data)
    if include_character_boxes and words:
        words = _attach_ocr_character_boxes(
            image=image.convert("RGB"),
            words=words,
            ocr_lang=ocr_lang or resolve_ocr_lang(),
            ocr_config=OCR_CONFIG,
        )
        if any(not word.character_boxes for word in words):
            words = _attach_ocr_character_boxes(
                image=preprocess_for_ocr(image.convert("RGB")),
                words=words,
                ocr_lang=ocr_lang or resolve_ocr_lang(),
                ocr_config=OCR_CONFIG,
            )
    return text, words


def _thresholded_ocr_words_from_image(
    image: Image.Image,
    *,
    ocr_lang: Optional[str] = None,
    include_character_boxes: bool = False,
) -> tuple[str, list[OCRWord]]:
    """Return the cleanup OCR pass even when another pass scores higher overall.

    A best-whole-page OCR score can hide a single faint address or identifier.
    This supplemental pass is used only to recover finding types absent from the
    primary pass; it never replaces or duplicates primary detections.
    """
    processed = preprocess_for_ocr(image.convert("RGB"))
    data = pytesseract.image_to_data(
        processed,
        lang=ocr_lang or resolve_ocr_lang(),
        config=OCR_CONFIG,
        output_type=pytesseract.Output.DICT,
    )
    text, words = _ocr_words_from_data(data)
    if include_character_boxes and words:
        words = _attach_ocr_character_boxes(
            image=processed,
            words=words,
            ocr_lang=ocr_lang or resolve_ocr_lang(),
            ocr_config=OCR_CONFIG,
        )
    return text, words


def _attach_ocr_character_boxes(
    *,
    image: Image.Image,
    words: Sequence[OCRWord],
    ocr_lang: str,
    ocr_config: str,
) -> list[OCRWord]:
    """Attach exact symbol boxes to words when Tesseract agrees on the text."""
    try:
        raw_boxes = pytesseract.image_to_boxes(
            image,
            lang=ocr_lang,
            config=ocr_config,
        )
    except Exception:
        return list(words)

    symbols: list[tuple[str, tuple[int, int, int, int]]] = []
    for raw_line in str(raw_boxes or "").splitlines():
        fields = raw_line.rsplit(maxsplit=5)
        if len(fields) < 6:
            continue
        symbol = fields[0]
        try:
            left, bottom, right, top = (int(value) for value in fields[1:5])
        except ValueError:
            continue
        box = (left, image.height - top, right, image.height - bottom)
        if symbol and box[2] > box[0] and box[3] > box[1]:
            symbols.append((symbol, box))

    attached: list[OCRWord] = []
    for word in words:
        x0, y0, x1, y1 = word.bbox
        candidates = [
            (symbol, box)
            for symbol, box in symbols
            if x0 - 3 <= (box[0] + box[2]) / 2 <= x1 + 3
            and y0 - 3 <= (box[1] + box[3]) / 2 <= y1 + 3
        ]
        candidates.sort(key=lambda item: (item[1][0], item[1][1]))
        character_boxes = (
            tuple(box for _symbol, box in candidates)
            if len(candidates) == len(word.text)
            else word.character_boxes
        )
        attached.append(
            OCRWord(
                text=word.text,
                bbox=word.bbox,
                start=word.start,
                end=word.end,
                page_num=word.page_num,
                block_num=word.block_num,
                line_num=word.line_num,
                character_boxes=character_boxes,
            )
        )
    return attached


def _normalize_text_for_compare(value: str) -> str:
    return re.sub(r"\s+", " ", value.strip()).casefold()


def _clean_custom_mask_items(values: Optional[Sequence[str]] = None) -> list[str]:
    cleaned: list[str] = []
    seen: set[str] = set()

    for value in values or ():
        text = str(value).strip()
        if not text:
            continue

        key = _normalize_text_for_compare(text)
        if key in seen:
            continue

        cleaned.append(text)
        seen.add(key)

    return cleaned


def _custom_mask_findings(
    text: str,
    values: Optional[Sequence[str]] = None,
    *,
    review_exclusions: Sequence[str] = (),
) -> list[TextFinding]:
    findings: list[TextFinding] = []
    excluded = {
        _normalize_text_for_compare(item)
        for item in review_exclusions
        if item and item.strip()
    }

    for value in _clean_custom_mask_items(values):
        if _normalize_text_for_compare(value) in excluded:
            continue
        pattern = re.compile(rf"(?=({re.escape(value)}))", flags=re.IGNORECASE)
        for match in pattern.finditer(text):
            start, end = match.span(1)
            quote = text[start:end]
            if not quote.strip():
                continue
            findings.append(
                TextFinding(
                    start=start,
                    end=end,
                    quote=quote,
                    label=_CUSTOM_MASK_LABEL,
                    source="custom_mask",
                )
            )

    return findings


def _mask_findings_for_text(
    *,
    sdp: GoogleSDPClient,
    text: str,
    payload: DataMaskingRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[TextFinding]:
    sensitive_findings = inspect_sensitive_text(
        sdp=sdp,
        text=text,
        targets=payload.target_data,
        review_exclusions=payload.review_exclusions,
        min_likelihood=getattr(sdp, "min_likelihood", DEFAULT_MIN_LIKELIHOOD),
    )
    if is_id_document_payload(payload):
        sensitive_findings = [
            finding
            for finding in sensitive_findings
            if finding.label
            not in {
                SensitiveDataType.name.value,
                SensitiveDataType.contact_address.value,
            }
        ]
        sensitive_findings.extend(id_document_field_findings(text, payload=payload))
    return merge_overlapping_findings(
        [
            *sensitive_findings,
            *_custom_mask_findings(
                text,
                custom_redactions,
                review_exclusions=payload.review_exclusions,
            ),
        ],
        original_text=text,
    )


def _ocr_words_for_finding(
    finding: TextFinding,
    words: Sequence[OCRWord],
) -> list[OCRWord]:
    matched = [
        word
        for word in words
        if not (word.end <= finding.start or word.start >= finding.end)
    ]
    if matched:
        return matched

    needle = _normalize_text_for_compare(finding.quote)
    maximum_window = max(1, len(needle.split()) + 2)
    for start_index in range(len(words)):
        buffered: list[OCRWord] = []
        for end_index in range(start_index, min(len(words), start_index + maximum_window)):
            buffered.append(words[end_index])
            candidate = _normalize_text_for_compare(
                " ".join(item.text for item in buffered),
            )
            if candidate == needle:
                return buffered
            if len(candidate) > len(needle):
                break

    return []


def _ocr_word_overlap_box(
    word: OCRWord,
    finding: TextFinding,
    *,
    use_full_word: bool = False,
) -> tuple[int, int, int, int] | None:
    if use_full_word:
        return word.bbox

    overlap_start = max(word.start, finding.start)
    overlap_end = min(word.end, finding.end)
    if overlap_end <= overlap_start:
        return None

    if overlap_start == word.start and overlap_end == word.end:
        return word.bbox

    x0, y0, x1, y1 = word.bbox
    character_count = max(1, word.end - word.start)
    local_start = overlap_start - word.start
    local_end = overlap_end - word.start
    if len(word.character_boxes) == character_count:
        selected = word.character_boxes[local_start:local_end]
        if selected:
            return (
                min(box[0] for box in selected),
                min(box[1] for box in selected),
                max(box[2] for box in selected),
                max(box[3] for box in selected),
            )
    clipped_x0 = round(
        x0 + (x1 - x0) * (local_start / character_count)
    )
    clipped_x1 = round(
        x0 + (x1 - x0) * (local_end / character_count)
    )
    if clipped_x1 <= clipped_x0:
        clipped_x1 = min(x1, clipped_x0 + 1)
    return clipped_x0, y0, clipped_x1, y1


def _boxes_for_text_spans(findings: Sequence[TextFinding], words: Sequence[OCRWord]) -> list[tuple[int, int, int, int]]:
    boxes: list[tuple[int, int, int, int]] = []
    for finding in findings:
        has_offset_match = any(
            not (word.end <= finding.start or word.start >= finding.end)
            for word in words
        )
        matched = _ocr_words_for_finding(finding, words)
        if not matched:
            continue

        by_line: dict[tuple[int, int, int], list[tuple[int, int, int, int]]] = {}
        for word in matched:
            box = _ocr_word_overlap_box(
                word,
                finding,
                use_full_word=not has_offset_match,
            )
            if box is not None:
                by_line.setdefault(
                    (word.page_num, word.block_num, word.line_num),
                    [],
                ).append(box)

        for line_boxes in by_line.values():
            boxes.append(
                (
                    min(box[0] for box in line_boxes),
                    min(box[1] for box in line_boxes),
                    max(box[2] for box in line_boxes),
                    max(box[3] for box in line_boxes),
                )
            )

    unique: list[tuple[int, int, int, int]] = []
    seen: set[tuple[int, int, int, int]] = set()
    for box in boxes:
        if box not in seen:
            unique.append(box)
            seen.add(box)
    return unique


def _ocr_mask_specs(
    findings: Sequence[TextFinding],
    words: Sequence[OCRWord],
) -> list[tuple[tuple[int, int, int, int], str]]:
    specs: list[tuple[tuple[int, int, int, int], str]] = []
    seen: set[tuple[int, int, int, int]] = set()

    for finding in findings:
        has_offset_match = any(
            not (word.end <= finding.start or word.start >= finding.end)
            for word in words
        )
        matched = _ocr_words_for_finding(finding, words)
        if not matched:
            continue

        by_line: dict[tuple[int, int, int], list[OCRWord]] = {}
        for word in matched:
            by_line.setdefault(
                (word.page_num, word.block_num, word.line_num),
                [],
            ).append(word)

        masked_full = _same_length_mask_value(finding.label, finding.quote)
        for line_words in by_line.values():
            piece_boxes = [
                box
                for word in line_words
                if (
                    box := _ocr_word_overlap_box(
                        word,
                        finding,
                        use_full_word=not has_offset_match,
                    )
                ) is not None
            ]
            if not piece_boxes:
                continue
            box = (
                min(item[0] for item in piece_boxes),
                min(item[1] for item in piece_boxes),
                max(item[2] for item in piece_boxes),
                max(item[3] for item in piece_boxes),
            )
            if box in seen:
                continue

            if has_offset_match:
                line_start = max(finding.start, min(word.start for word in line_words))
                line_end = min(finding.end, max(word.end for word in line_words))
                masked = masked_full[
                    line_start - finding.start:line_end - finding.start
                ]
            else:
                masked = _same_length_mask_value(
                    finding.label,
                    " ".join(word.text for word in line_words),
                )
            if not masked:
                continue

            seen.add(box)
            specs.append((box, masked))

    return specs


def _iter_table_paragraphs(table: Any):
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                yield paragraph
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def _iter_document_paragraphs(document: docx.Document):
    # Retain XML elements rather than only id(element). lxml proxy objects can
    # be released during iteration, allowing Python to reuse an id for a
    # different paragraph and silently skip content.
    seen: set[Any] = set()

    def emit(paragraphs: Any):
        for paragraph in paragraphs:
            element = paragraph._p
            if element in seen:
                continue
            seen.add(element)
            yield paragraph

            # Text boxes are not exposed by python-docx's public paragraph
            # collection. They still need to participate in the same in-place,
            # structure-preserving masking pipeline as body text.
            for nested_element in paragraph._p.xpath(".//w:txbxContent//w:p"):
                if nested_element in seen:
                    continue
                seen.add(nested_element)
                yield Paragraph(nested_element, paragraph._parent)

    yield from emit(document.paragraphs)
    for table in document.tables:
        yield from emit(_iter_table_paragraphs(table))
    for section in document.sections:
        yield from emit(section.header.paragraphs)
        for table in section.header.tables:
            yield from emit(_iter_table_paragraphs(table))
        yield from emit(section.footer.paragraphs)
        for table in section.footer.tables:
            yield from emit(_iter_table_paragraphs(table))


def _output_path(source_path: Path, output_dir: Path, suffix: str = "_masked") -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    return output_dir / f"{source_path.stem}{suffix}{source_path.suffix.lower()}"


def _ensure_mask_request(request: AnalyzerRequest) -> tuple[DocumentPayload, DataMaskingRequest]:
    if request.action != FeatureType.data_mask:
        raise ValueError("data_mask.py only handles action='data_mask'.")
    if not isinstance(request.input, DocumentPayload):
        raise ValueError("data_mask requires DocumentPayload input.")
    if not isinstance(request.payload, DataMaskingRequest):
        raise ValueError("data_mask requires DataMaskingRequest payload.")
    if request.input.metadata.input_format not in _SUPPORTED_INPUTS:
        raise ValueError("data_mask only supports pdf, docx, jpg, jpeg, png.")
    if request.policy.structure_preservation is not True:
        raise ValueError("structure_preservation must be True for data_mask.")
    return request.input, request.payload


def _build_masking_response(*, request: AnalyzerRequest, output_path: Path) -> AnalyzerResponse:
    input_payload, _ = _ensure_mask_request(request)
    result = build_document_file_result(
        filename=output_path.name,
        output_format=_FILE_OUTPUT_MAP[input_payload.metadata.input_format],
        file_size_mb=round(output_path.stat().st_size / (1024 * 1024), 4),
        algorithm_version="google-sdp-data-mask-v2",
    )
    response = AnalyzerResponse(
        action=FeatureType.data_mask,
        input_format=input_payload.metadata.input_format,
        policy=request.policy,
        system_language=request.system_language,
        result=result,
        human_review=HumanReviewRequirement(),
    )
    return validate_analyzer_response(response, request=request)


def _mask_digits(value: str, *, keep_last: int = 4, mask_char: str = "X") -> str:
    digits = [i for i, char in enumerate(value) if char.isdigit()]
    if not digits:
        return value
    to_keep = set(digits[-keep_last:]) if keep_last > 0 else set()
    chars = list(value)
    for idx in digits:
        if idx not in to_keep:
            chars[idx] = mask_char
    return "".join(chars)


def _mask_alpha(value: str, *, keep_first: int = 1, mask_char: str = "X") -> str:
    chars = list(value)
    alpha_indices = [i for i, char in enumerate(chars) if char.isalpha()]
    to_keep = set(alpha_indices[:keep_first]) if keep_first > 0 else set()
    for idx in alpha_indices:
        if idx not in to_keep:
            chars[idx] = mask_char
    return "".join(chars)


def _mask_name(value: str) -> str:
    parts = value.split()
    masked = []
    for part in parts:
        masked.append(part[0] + ("X" * (len(part) - 1)) if len(part) > 1 else "X")
    return " ".join(masked) if masked else "X"


def _mask_email(value: str) -> str:
    if "@" not in value:
        return _mask_alpha(value, keep_first=1)
    local, domain = value.split("@", 1)
    if "." in domain:
        head, _, tld = domain.rpartition(".")
        masked_domain = (head[:1] + "X" * max(1, len(head) - 1)) if head else "X"
        return f"{local[:1]}{'X' * max(1, len(local) - 1)}@{masked_domain}.{tld}"
    return f"{local[:1]}{'X' * max(1, len(local) - 1)}@{'X' * max(1, len(domain))}"


def _mask_preserving_layout(value: str, *, mask_char: str = "X") -> str:
    """Mask visible letters/digits while preserving spacing and punctuation.

    This keeps PDF layout stable while avoiding partially readable results
    such as "NigeriX". For example, "Lagos, Nigeria" becomes
    "XXXXX, XXXXXXX".
    """
    return "".join(mask_char if char.isalnum() else char for char in value)


def mask_value_by_target(label: str, quote: str) -> str:
    label = (label or "").lower()
    if label == _CUSTOM_MASK_LABEL:
        return "".join(char if char.isspace() else "X" for char in quote)
    if label in {SensitiveDataType.name.value, "person_name"}:
        return _mask_name(quote)
    if label in {SensitiveDataType.email_address.value, "email_address"}:
        return _mask_email(quote)
    if label in {SensitiveDataType.phone_number.value, "phone_number"}:
        return _mask_digits(quote, keep_last=2)
    if label in {SensitiveDataType.card_number.value, "credit_card_number", SensitiveDataType.account_number.value}:
        return _mask_digits(quote, keep_last=4)
    if label in {
        SensitiveDataType.national_id.value,
        SensitiveDataType.tax_id.value,
        SensitiveDataType.passport_number.value,
        "passport",
    }:
        return _mask_digits(_mask_alpha(quote, keep_first=1), keep_last=2)
    if label in {SensitiveDataType.contact_address.value, "street_address", "location"}:
        return _mask_preserving_layout(quote)
    if label == SensitiveDataType.date_of_birth.value:
        chars = list(quote)
        digits = [i for i, char in enumerate(chars) if char.isdigit()]
        keep = set(digits[-2:])
        for idx in digits:
            if idx not in keep:
                chars[idx] = "X"
        return "".join(chars)
    if label == SensitiveDataType.age.value:
        return _mask_digits(quote, keep_last=0)
    if label == SensitiveDataType.signature.value:
        return _mask_preserving_layout(quote)
    return _mask_alpha(_mask_digits(quote, keep_last=2), keep_first=1)


def _same_length_mask_value(label: str, quote: str) -> str:
    masked = mask_value_by_target(label, quote)
    if len(masked) == len(quote):
        return masked
    if len(masked) < len(quote):
        return masked + ("X" * (len(quote) - len(masked)))
    return masked[: len(quote)]


def mask_text_with_findings(text: str, findings: Sequence[TextFinding]) -> str:
    result = text
    for finding in sorted(findings, key=lambda item: item.start, reverse=True):
        replacement = _same_length_mask_value(finding.label, finding.quote)
        result = result[:finding.start] + replacement + result[finding.end:]
    return result


def _clip_pixel_box(
    box: tuple[int, int, int, int],
    *,
    width: int,
    height: int,
    padding: int = 0,
) -> tuple[int, int, int, int] | None:
    x0, y0, x1, y1 = box
    clipped = (
        max(0, min(width, x0 - padding)),
        max(0, min(height, y0 - padding)),
        max(0, min(width, x1 + padding)),
        max(0, min(height, y1 + padding)),
    )
    if clipped[2] <= clipped[0] or clipped[3] <= clipped[1]:
        return None
    return clipped


def _dark_content_bbox(
    image: Image.Image,
    search_box: tuple[int, int, int, int],
) -> tuple[int, int, int, int] | None:
    clipped = _clip_pixel_box(
        search_box,
        width=image.width,
        height=image.height,
    )
    if clipped is None:
        return None

    grayscale = image.crop(clipped).convert("L")
    # A permissive threshold retains faint blue/gray pen strokes while the
    # geometry checks below reject an empty underline.
    ink = grayscale.point(lambda value: 255 if value < 232 else 0)
    local = ink.getbbox()
    if local is None:
        return None

    left, top, right, bottom = local
    content_width = right - left
    content_height = bottom - top
    histogram = ink.histogram()
    ink_pixels = sum(histogram[1:])
    if content_width < 5 or content_height < 3 or ink_pixels < 16:
        return None

    return (
        clipped[0] + left,
        clipped[1] + top,
        clipped[0] + right,
        clipped[1] + bottom,
    )


def _ocr_signature_boxes(
    image: Image.Image,
    words: Sequence[OCRWord],
) -> list[tuple[int, int, int, int]]:
    lines: dict[tuple[int, int, int], list[OCRWord]] = {}
    for word in words:
        key = (word.page_num, word.block_num, word.line_num)
        lines.setdefault(key, []).append(word)

    boxes: list[tuple[int, int, int, int]] = []
    for line_words in lines.values():
        ordered = sorted(line_words, key=lambda word: (word.bbox[0], word.start))
        line_text = " ".join(word.text for word in ordered)
        if (
            len(line_text) <= 80
            and _SIGNATURE_ROLE_RE.search(line_text)
            and min(word.bbox[1] for word in ordered) >= image.height * 0.45
        ):
            role_anchor = (
                min(word.bbox[0] for word in ordered),
                min(word.bbox[1] for word in ordered),
                max(word.bbox[2] for word in ordered),
                max(word.bbox[3] for word in ordered),
            )
            role_height = max(8, role_anchor[3] - role_anchor[1])
            role_search = (
                max(0, role_anchor[0] - role_height * 2),
                max(0, role_anchor[1] - max(90, role_height * 8)),
                min(image.width, role_anchor[2] + role_height * 3),
                max(0, role_anchor[1] - 2),
            )
            role_content = _dark_content_bbox(image, role_search)
            if role_content is not None:
                boxes.append(role_content)

        for index, word in enumerate(ordered):
            normalized = re.sub(r"[^\wÀ-ÖØ-öø-ÿ]+", "", word.text.casefold())
            if normalized not in _SIGNATURE_WORDS:
                continue

            anchor = word.bbox
            anchor_height = max(8, anchor[3] - anchor[1])
            right_search = (
                anchor[2] + 2,
                max(0, anchor[1] - anchor_height),
                min(image.width, anchor[2] + max(180, anchor_height * 18)),
                min(image.height, anchor[3] + max(45, anchor_height * 3)),
            )
            below_search = (
                max(0, anchor[0] - anchor_height),
                anchor[3] + 2,
                min(image.width, anchor[0] + max(220, anchor_height * 22)),
                min(image.height, anchor[3] + max(70, anchor_height * 6)),
            )
            above_search = (
                max(0, anchor[0] - anchor_height),
                max(0, anchor[1] - max(70, anchor_height * 6)),
                min(image.width, anchor[0] + max(220, anchor_height * 22)),
                max(0, anchor[1] - 2),
            )

            for search_box in (right_search, below_search, above_search):
                content_box = _dark_content_bbox(image, search_box)
                if content_box is not None:
                    boxes.append(content_box)

            # OCR-readable typed signatures should still be covered if their
            # ink is too light for the visual threshold.
            trailing = [item for item in ordered[index + 1:] if item.bbox[0] >= anchor[2]]
            if trailing:
                boxes.append(
                    (
                        min(item.bbox[0] for item in trailing),
                        min(item.bbox[1] for item in trailing),
                        max(item.bbox[2] for item in trailing),
                        max(item.bbox[3] for item in trailing),
                    )
                )

    unique: list[tuple[int, int, int, int]] = []
    seen: set[tuple[int, int, int, int]] = set()
    for box in boxes:
        clipped = _clip_pixel_box(
            box,
            width=image.width,
            height=image.height,
            padding=MASK_BOX_PADDING_PX,
        )
        if clipped is not None and clipped not in seen:
            unique.append(clipped)
            seen.add(clipped)
    return unique


def _signature_target_selected(payload: DataMaskingRequest) -> bool:
    return any(
        str(getattr(target, "value", target)) == SensitiveDataType.signature.value
        for target in payload.target_data
    )


def _visual_signature_is_excluded(payload: DataMaskingRequest) -> bool:
    excluded = {
        _normalize_text_for_compare(item)
        for item in payload.review_exclusions
        if item and item.strip()
    }
    return _normalize_text_for_compare(VISUAL_SIGNATURE_QUOTE) in excluded


def _google_signature_boxes(
    *,
    image: Image.Image,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
) -> list[tuple[int, int, int, int]]:
    if not _signature_target_selected(payload) or _visual_signature_is_excluded(payload):
        return []

    encoded = BytesIO()
    image.convert("RGB").save(encoded, format="JPEG", quality=90)
    signature_targets = [
        target
        for target in payload.target_data
        if str(getattr(target, "value", target)) == SensitiveDataType.signature.value
    ]
    findings = inspect_sensitive_image(
        sdp=sdp,
        image_bytes=encoded.getvalue(),
        image_type="IMAGE_JPEG",
        targets=signature_targets,
        review_exclusions=payload.review_exclusions,
        min_likelihood=getattr(sdp, "min_likelihood", DEFAULT_MIN_LIKELIHOOD),
    )
    return [
        finding.bbox
        for finding in findings
        if finding.label == SensitiveDataType.signature.value
    ]


def _signature_boxes_for_image(
    *,
    image: Image.Image,
    words: Sequence[OCRWord],
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
) -> list[tuple[int, int, int, int]]:
    if not _signature_target_selected(payload) or _visual_signature_is_excluded(payload):
        return []
    boxes = [
        *_ocr_signature_boxes(image, words),
        *_google_signature_boxes(image=image, sdp=sdp, payload=payload),
    ]
    unique: list[tuple[int, int, int, int]] = []
    seen: set[tuple[int, int, int, int]] = set()
    for box in boxes:
        clipped = _clip_pixel_box(
            box,
            width=image.width,
            height=image.height,
            padding=2,
        )
        if clipped is not None and clipped not in seen:
            unique.append(clipped)
            seen.add(clipped)
    return unique


def _require_ocr_findings_mapped(
    findings: Sequence[TextFinding],
    words: Sequence[OCRWord],
) -> None:
    unmapped = sum(1 for finding in findings if not _ocr_words_for_finding(finding, words))
    if unmapped:
        raise RuntimeError(
            f"{unmapped} sensitive finding(s) could not be mapped to image coordinates; "
            "masking stopped to prevent an unsafe output."
        )


def pixelate_boxes(
    image: Image.Image,
    boxes: Sequence[tuple[int, int, int, int]],
    *,
    blur_radius: int = 3,
) -> Image.Image:
    """Irreversibly cover sensitive image regions.

    The public name is retained for compatibility. Blur/pixelation can preserve
    recognizable strokes and is not appropriate for dependable privacy
    protection, so v2 uses an opaque mask instead.
    """
    del blur_radius
    out = image.convert("RGB").copy()
    draw = ImageDraw.Draw(out)
    for box in boxes:
        clipped = _clip_pixel_box(
            box,
            width=out.width,
            height=out.height,
            padding=MASK_BOX_PADDING_PX,
        )
        if clipped is not None:
            draw.rectangle(clipped, fill=(0, 0, 0))
    return out


def _resolve_sdp(
    *,
    sdp: GoogleSDPClient | None = None,
    project_id: str | None = None,
    location: str = DEFAULT_DLP_LOCATION,
    min_likelihood: str = DEFAULT_MIN_LIKELIHOOD,
    client: Any | None = None,
) -> GoogleSDPClient:
    if sdp is not None:
        return sdp
    if not project_id:
        raise ValueError("Provide either sdp or project_id.")
    return build_google_sdp_client(
        project_id=project_id,
        location=location,
        min_likelihood=min_likelihood,
        client=client,
    )


def preview_data_mask_candidates(
    request: AnalyzerRequest | Mapping[str, Any],
    *,
    sdp: GoogleSDPClient | None = None,
    project_id: str | None = None,
    location: str = DEFAULT_DLP_LOCATION,
    min_likelihood: str = DEFAULT_MIN_LIKELIHOOD,
    client: Any | None = None,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[DetectionCandidate]:
    req = validate_analyzer_request(request)
    input_payload, payload = _ensure_mask_request(req)
    if not input_payload.text:
        raise ValueError(
            "preview_data_mask_candidates requires extracted document text. "
            "Build the request with extraction.py before previewing."
        )
    resolved = _resolve_sdp(
        sdp=sdp,
        project_id=project_id,
        location=location,
        min_likelihood=min_likelihood,
        client=client,
    )
    grouped: dict[tuple[str, str, str], int] = {}
    for finding in _inspect_text_in_chunks(
        sdp=resolved,
        text=input_payload.text,
        payload=payload,
        custom_redactions=custom_redactions,
    ):
        key = (finding.label, finding.quote, finding.source)
        grouped[key] = grouped.get(key, 0) + 1

    if input_payload.filename:
        source = Path(str(input_payload.filename))
        if source.exists() and source.is_file():
            if source.suffix.lower() == ".pdf":
                with fitz.open(source) as document:
                    for _key, _value, metadata_findings in _pdf_metadata_findings(
                        document=document,
                        sdp=resolved,
                        payload=payload,
                        custom_redactions=custom_redactions,
                    ):
                        for finding in metadata_findings:
                            candidate_key = (
                                finding.label,
                                finding.quote,
                                finding.source,
                            )
                            grouped[candidate_key] = grouped.get(candidate_key, 0) + 1
            signature_occurrences = _visual_signature_occurrence_count(
                source_path=source,
                sdp=resolved,
                payload=payload,
                ocr_lang=resolve_ocr_lang(),
            )
            if signature_occurrences:
                grouped[("signature", VISUAL_SIGNATURE_QUOTE, "visual_signature")] = (
                    signature_occurrences
                )
    return [
        DetectionCandidate(
            label=label,
            quote=quote,
            occurrences=occurrences,
            source=source,
        )
        for (label, quote, source), occurrences in sorted(
            grouped.items(),
            key=lambda item: (item[0][0], item[0][1].casefold(), item[0][2]),
        )
    ]


def _paragraph_run_spans(paragraph: Any) -> tuple[str, list[RunSpan]]:
    spans: list[RunSpan] = []
    parts: list[str] = []
    cursor = 0

    for run in paragraph.runs:
        text = run.text or ""
        if not text:
            continue
        start = cursor
        end = start + len(text)
        parts.append(text)
        spans.append(
            RunSpan(
                start=start,
                end=end,
                text=text,
                rpr_xml=deepcopy(run._r.rPr) if run._r.rPr is not None else None,
                run=run,
            )
        )
        cursor = end

    return "".join(parts), spans


def _copy_rpr(dst_run: Any, rpr_xml: Any) -> None:
    if dst_run._r.rPr is not None:
        dst_run._r.remove(dst_run._r.rPr)
    if rpr_xml is not None:
        dst_run._r.insert(0, deepcopy(rpr_xml))


def _rewrite_paragraph_from_fragments(paragraph: Any, fragments: list[tuple[Any, str]]) -> None:
    for run in list(paragraph.runs):
        paragraph._p.remove(run._r)

    for rpr_xml, text in fragments:
        if not text:
            continue
        new_run = paragraph.add_run()
        _copy_rpr(new_run, rpr_xml)
        new_run.text = text


def _run_span_for_offset(spans: Sequence[RunSpan], offset: int) -> RunSpan:
    for span in spans:
        if span.start <= offset < span.end:
            return span
    return spans[-1]


def _xml_local_name(element: Any) -> str:
    tag = str(getattr(element, "tag", ""))
    return tag.rsplit("}", 1)[-1]


def _replace_run_text_preserving_visuals(run: Any, replacement: str) -> None:
    text_nodes = [element for element in run._r.iter() if _xml_local_name(element) == "t"]
    node_text = "".join(element.text or "" for element in text_nodes)
    if node_text == (run.text or "") and len(node_text) == len(replacement):
        cursor = 0
        for element in text_nodes:
            length = len(element.text or "")
            element.text = replacement[cursor:cursor + length]
            cursor += length
        return

    visuals = [
        deepcopy(element)
        for element in run._r
        if _xml_local_name(element) in _DOCX_VISUAL_LOCAL_NAMES
    ]
    run.text = replacement
    for visual in visuals:
        run._r.append(visual)


def mask_paragraph_runs(paragraph: Any, findings: Sequence[TextFinding]) -> None:
    full_text, spans = _paragraph_run_spans(paragraph)
    if not full_text or not spans or not findings:
        return

    merged = merge_overlapping_findings(findings, original_text=full_text)
    for span in spans:
        replacement = list(span.text)
        changed = False
        for finding in merged:
            overlap_start = max(span.start, finding.start)
            overlap_end = min(span.end, finding.end)
            if overlap_end <= overlap_start:
                continue

            masked_full = _same_length_mask_value(finding.label, finding.quote)
            source_start = overlap_start - finding.start
            source_end = overlap_end - finding.start
            target_start = overlap_start - span.start
            target_end = overlap_end - span.start
            replacement[target_start:target_end] = masked_full[source_start:source_end]
            changed = True

        if changed and span.run is not None:
            _replace_run_text_preserving_visuals(span.run, "".join(replacement))


def _page_words_with_offsets(page: fitz.Page) -> tuple[str, list[PDFWord]]:
    raw = page.get_text("words", sort=True)
    words: list[PDFWord] = []
    parts: list[str] = []
    cursor = 0
    prev_line: tuple[int, int] | None = None

    for x0, y0, x1, y1, token, block_no, line_no, word_no in raw:
        token = (token or "").strip()
        if not token:
            continue

        current_line = (block_no, line_no)
        if parts:
            separator = "\n" if prev_line != current_line else " "
            parts.append(separator)
            cursor += len(separator)

        start = cursor
        parts.append(token)
        cursor += len(token)
        end = cursor

        words.append(
            PDFWord(
                start=start,
                end=end,
                text=token,
                rect=fitz.Rect(x0, y0, x1, y1),
                block_no=block_no,
                line_no=line_no,
                word_no=word_no,
            )
        )
        prev_line = current_line

    return "".join(parts), words


def _pdf_word_overlap_rect(word: PDFWord, finding: TextFinding) -> fitz.Rect | None:
    overlap_start = max(word.start, finding.start)
    overlap_end = min(word.end, finding.end)
    if overlap_end <= overlap_start:
        return None

    rect = fitz.Rect(word.rect)
    if overlap_start == word.start and overlap_end == word.end:
        return rect

    character_count = max(1, word.end - word.start)
    original_x0 = rect.x0
    width = rect.width
    rect.x0 = original_x0 + width * (
        (overlap_start - word.start) / character_count
    )
    rect.x1 = original_x0 + width * (
        (overlap_end - word.start) / character_count
    )
    return rect if rect.width > 0 else None


def _pdf_rects_for_findings(
    findings: Sequence[TextFinding],
    words: Sequence[PDFWord],
    *,
    original_text: str,
) -> list[fitz.Rect]:
    rects: list[fitz.Rect] = []
    merged = merge_overlapping_findings(findings, original_text=original_text)

    for finding in merged:
        matched = [w for w in words if not (w.end <= finding.start or w.start >= finding.end)]
        if not matched:
            continue

        by_line: dict[tuple[int, int], list[fitz.Rect]] = {}
        for word in matched:
            rect = _pdf_word_overlap_rect(word, finding)
            if rect is not None:
                by_line.setdefault((word.block_no, word.line_no), []).append(rect)

        for line_rects in by_line.values():
            rect = line_rects[0]
            for word_rect in line_rects[1:]:
                rect = rect | word_rect
            rects.append(rect)

    unique: list[fitz.Rect] = []
    seen: set[tuple[float, float, float, float]] = set()
    for rect in rects:
        key = (round(rect.x0, 3), round(rect.y0, 3), round(rect.x1, 3), round(rect.y1, 3))
        if key not in seen:
            seen.add(key)
            unique.append(rect)
    return unique


def _pdf_rect_key(rect: fitz.Rect) -> tuple[float, float, float, float]:
    return (round(rect.x0, 3), round(rect.y0, 3), round(rect.x1, 3), round(rect.y1, 3))


def _pdf_rect_has_visible_content(page: fitz.Page, rect: fitz.Rect) -> bool:
    candidate = fitz.Rect(rect)
    candidate_area = max(1.0, candidate.get_area())
    try:
        paint_log = list(page.get_bboxlog())
    except Exception:
        paint_log = []
    text_indexes = [
        index
        for index, item in enumerate(paint_log)
        if str(item[0]).endswith("text")
        and not (fitz.Rect(item[1]) & candidate).is_empty
    ]
    if text_indexes:
        text_index = max(text_indexes)
        for item in paint_log[text_index + 1:]:
            if item[0] != "fill-image":
                continue
            covered = fitz.Rect(item[1]) & candidate
            if not covered.is_empty and covered.get_area() / candidate_area >= 0.95:
                return False

    clipped = fitz.Rect(
        candidate.x0 - 0.5,
        candidate.y0 - 0.5,
        candidate.x1 + 0.5,
        candidate.y1 + 0.5,
    ) & page.rect
    if clipped.is_empty:
        return False
    pix = page.get_pixmap(
        matrix=fitz.Matrix(PDF_VISIBILITY_RENDER_SCALE, PDF_VISIBILITY_RENDER_SCALE),
        clip=clipped,
        alpha=False,
        colorspace=fitz.csRGB,
        annots=False,
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples).convert("L")
    histogram = image.histogram()
    visible_pixels = sum(histogram[:PDF_VISIBLE_PIXEL_THRESHOLD])
    minimum_pixels = max(2, round(image.width * image.height * 0.002))
    return visible_pixels >= minimum_pixels


def _tight_pdf_text_rect(rect: fitz.Rect) -> fitz.Rect:
    """Return a safer redaction rectangle for text-only masking.

    PyMuPDF redaction fill rectangles can visually erase nearby separator
    lines when the rectangle is too tall. A slightly tighter vertical rectangle
    still intersects the text glyphs, so the original text is removed, but it
    is less likely to cover horizontal rules or adjacent line art.
    """
    tightened = fitz.Rect(rect)
    if tightened.height <= 2:
        return tightened

    inset_y = min(1.0, max(0.25, tightened.height * 0.12))
    if tightened.y0 + inset_y < tightened.y1 - inset_y:
        tightened.y0 += inset_y
        tightened.y1 -= inset_y

    return tightened


def _fit_pdf_mask_font_size(masked: str, rect: fitz.Rect) -> float:
    """Choose a font size that fits the masked value inside the original span."""
    font_size = max(4.0, min(10.0, rect.height * 0.78))

    if not masked or rect.width <= 0:
        return font_size

    try:
        text_width = fitz.get_text_length(masked, fontname="helv", fontsize=font_size)
    except Exception:
        return font_size

    if text_width > rect.width and text_width > 0:
        font_size = max(3.5, font_size * (rect.width / text_width) * 0.96)

    return font_size


def _pdf_mask_insert_point(rect: fitz.Rect) -> fitz.Point:
    return fitz.Point(rect.x0, rect.y1 - max(0.5, rect.height * 0.15))


def _insert_pdf_mask_text(page: fitz.Page, specs: Sequence[tuple[fitz.Rect, fitz.Rect, str]]) -> None:
    for _redaction_rect, original_rect, masked in specs:
        if not masked:
            continue

        page.insert_text(
            _pdf_mask_insert_point(original_rect),
            masked,
            fontname="helv",
            fontsize=_fit_pdf_mask_font_size(masked, original_rect),
            color=(0, 0, 0),
            overlay=True,
        )


def _masked_rect_specs(
    findings: Sequence[TextFinding],
    words: Sequence[PDFWord],
    *,
    original_text: str,
) -> list[tuple[fitz.Rect, fitz.Rect, str]]:
    specs: list[tuple[fitz.Rect, fitz.Rect, str]] = []
    merged = merge_overlapping_findings(findings, original_text=original_text)

    # PyMuPDF may remove an entire PDF word when any glyph intersects a
    # redaction rectangle. Rewrite each affected word in full so a literal
    # custom match inside that word masks only the requested characters while
    # preserving its prefix and suffix.
    for word in words:
        replacement = list(word.text)
        changed = False
        for finding in merged:
            overlap_start = max(word.start, finding.start)
            overlap_end = min(word.end, finding.end)
            if overlap_end <= overlap_start:
                continue

            masked_full = _same_length_mask_value(finding.label, finding.quote)
            source_start = overlap_start - finding.start
            source_end = overlap_end - finding.start
            target_start = overlap_start - word.start
            target_end = overlap_end - word.start
            replacement[target_start:target_end] = masked_full[source_start:source_end]
            changed = True

        if changed:
            original_rect = fitz.Rect(word.rect)
            specs.append(
                (
                    _tight_pdf_text_rect(original_rect),
                    original_rect,
                    "".join(replacement),
                )
            )

    return specs


def _apply_page_redactions(
    page: fitz.Page,
    *,
    image_mode: str,
    remove_graphics: bool = False,
) -> None:
    kwargs: dict[str, Any] = {}

    if image_mode == "none" and hasattr(fitz, "PDF_REDACT_IMAGE_NONE"):
        kwargs["images"] = fitz.PDF_REDACT_IMAGE_NONE
    elif image_mode == "pixels" and hasattr(fitz, "PDF_REDACT_IMAGE_PIXELS"):
        kwargs["images"] = fitz.PDF_REDACT_IMAGE_PIXELS

    if remove_graphics and hasattr(fitz, "PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED"):
        kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_REMOVE_IF_TOUCHED
    elif hasattr(fitz, "PDF_REDACT_LINE_ART_NONE"):
        kwargs["graphics"] = fitz.PDF_REDACT_LINE_ART_NONE
    if hasattr(fitz, "PDF_REDACT_TEXT_REMOVE"):
        kwargs["text"] = fitz.PDF_REDACT_TEXT_REMOVE

    page.apply_redactions(**kwargs)


def _page_has_rendered_content(page: fitz.Page) -> bool:
    pixmap = page.get_pixmap(
        matrix=fitz.Matrix(0.5, 0.5),
        alpha=False,
        colorspace=fitz.csRGB,
        annots=False,
    )
    return min(pixmap.samples, default=255) < 250


def _save_pdf_checked(
    document: fitz.Document,
    output_path: Path,
    *,
    temporary_paths: list[Path],
) -> None:
    """Atomically publish a readable PDF without accepting blank-page loss."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    expected_content = [
        _page_has_rendered_content(page)
        for page in document
    ]
    last_error: Exception | None = None

    for _attempt in range(2):
        temporary_path = output_path.with_name(
            f".{output_path.stem}.{uuid.uuid4().hex}.tmp{output_path.suffix}"
        )
        temporary_paths.append(temporary_path)
        try:
            document.save(temporary_path, garbage=4, deflate=True, clean=True)
            if not temporary_path.exists() or temporary_path.stat().st_size <= 0:
                raise RuntimeError("PDF writer produced an empty file.")

            with fitz.open(temporary_path) as verification:
                if len(verification) != len(expected_content):
                    raise RuntimeError("PDF writer changed the page count.")
                for page_index, expected in enumerate(expected_content):
                    if expected and not _page_has_rendered_content(verification[page_index]):
                        raise RuntimeError(
                            f"PDF writer lost visible page content on page {page_index + 1}."
                        )

            temporary_path.replace(output_path)
            return
        except Exception as exc:
            last_error = exc
        finally:
            temporary_path.unlink(missing_ok=True)

    raise RuntimeError("Unable to create a complete masked PDF safely.") from last_error


def _paragraph_has_visual(paragraph: Any) -> bool:
    return any(
        _xml_local_name(element) in _DOCX_VISUAL_LOCAL_NAMES
        for element in paragraph._p.iter()
    )


def _image_parts_for_element(paragraph: Any, root: Any) -> set[Any]:
    parts: set[Any] = set()
    relationships = getattr(getattr(paragraph, "part", None), "rels", {})
    for element in root.iter():
        for attribute, relationship_id in element.attrib.items():
            if str(attribute).rsplit("}", 1)[-1] not in {"embed", "id"}:
                continue
            if not isinstance(relationship_id, str) or not relationship_id.startswith("rId"):
                continue
            relationship = relationships.get(relationship_id)
            if relationship is None or getattr(relationship, "is_external", False):
                continue
            target_part = getattr(relationship, "target_part", None)
            if (
                target_part is not None
                and str(getattr(target_part, "content_type", "")).startswith("image/")
            ):
                parts.add(target_part)
    return parts


def _paragraph_image_parts(paragraph: Any) -> set[Any]:
    return _image_parts_for_element(paragraph, paragraph._p)


def _black_picture_graphic(paragraph: Any) -> Any:
    """Create a reusable opaque picture graphic in the paragraph's story part."""
    image_bytes = BytesIO()
    Image.new("RGB", (1, 1), color=(0, 0, 0)).save(image_bytes, format="PNG")
    image_bytes.seek(0)

    placeholder_run = paragraph.add_run()
    placeholder_run.add_picture(image_bytes)
    graphic = next(
        (
            element
            for element in placeholder_run._r.iter()
            if _xml_local_name(element) == "graphic"
        ),
        None,
    )
    if graphic is None:
        paragraph._p.remove(placeholder_run._r)
        raise RuntimeError(
            "Could not create a layout-preserving DOCX signature mask."
        )

    result = deepcopy(graphic)
    paragraph._p.remove(placeholder_run._r)
    return result


def _replace_drawing_graphic_in_place(
    drawing: Any,
    *,
    replacement_graphic: Any,
) -> bool:
    """Replace drawing content while retaining its inline/anchor geometry."""
    graphics = [
        element
        for element in drawing.iter()
        if _xml_local_name(element) == "graphic"
    ]
    if not graphics:
        return False

    outer_extent = next(
        (
            element
            for element in drawing.iter()
            if _xml_local_name(element) == "extent"
            and "wordprocessingDrawing" in str(getattr(element, "tag", ""))
        ),
        None,
    )

    replaced = False
    for graphic in graphics:
        parent = graphic.getparent()
        if parent is None:
            continue
        replacement = deepcopy(replacement_graphic)
        if outer_extent is not None:
            picture_extent = next(
                (
                    element
                    for element in replacement.iter()
                    if _xml_local_name(element) == "ext"
                    and "drawingml/2006/main" in str(getattr(element, "tag", ""))
                ),
                None,
            )
            if picture_extent is not None:
                for attribute in ("cx", "cy"):
                    value = outer_extent.get(attribute)
                    if value is not None:
                        picture_extent.set(attribute, value)
        parent.replace(graphic, replacement)
        replaced = True
    return replaced


def _black_out_vml_visual(pict: Any) -> bool:
    """Turn legacy VML picture/shape containers into opaque rectangles."""
    shapes = [
        element
        for element in pict.iter()
        if _xml_local_name(element) in {"shape", "rect", "roundrect"}
    ]
    for shape in shapes:
        for child in list(shape):
            shape.remove(child)
        shape.set("fillcolor", "#000000")
        shape.set("filled", "t")
        shape.set("stroked", "f")
    return bool(shapes)


def _mask_non_image_signature_visuals(paragraph: Any) -> None:
    """Mask vector/ink signatures without removing their layout containers."""
    replacement_graphic: Any | None = None
    masked_visual = False

    for drawing in [
        element
        for element in paragraph._p.iter()
        if _xml_local_name(element) == "drawing"
    ]:
        if _image_parts_for_element(paragraph, drawing):
            continue
        if replacement_graphic is None:
            replacement_graphic = _black_picture_graphic(paragraph)
        masked_visual = (
            _replace_drawing_graphic_in_place(
                drawing,
                replacement_graphic=replacement_graphic,
            )
            or masked_visual
        )

    for pict in [
        element
        for element in paragraph._p.iter()
        if _xml_local_name(element) == "pict"
    ]:
        if _image_parts_for_element(paragraph, pict):
            continue
        masked_visual = _black_out_vml_visual(pict) or masked_visual

    if not _paragraph_image_parts(paragraph) and not masked_visual:
        raise RuntimeError(
            "A signature visual uses an unsupported DOCX object type; "
            "masking stopped to prevent an unsafe or layout-damaging output."
        )


def _mask_docx_signature_visuals(document: Any) -> set[Any]:
    paragraphs = list(_iter_document_paragraphs(document))
    anchors = {
        index
        for index, paragraph in enumerate(paragraphs)
        if _SIGNATURE_ANCHOR_RE.search(paragraph.text or "")
    }
    targeted_image_parts: set[Any] = set()

    for anchor_index in anchors:
        for candidate_index in range(
            max(0, anchor_index - 1),
            min(len(paragraphs), anchor_index + 3),
        ):
            paragraph = paragraphs[candidate_index]
            if not _paragraph_has_visual(paragraph):
                continue

            paragraph_text = (paragraph.text or "").strip()
            is_anchor = candidate_index == anchor_index
            is_visual_only_neighbor = (
                not paragraph_text
                or len(paragraph_text) <= 40
                or bool(_SIGNATURE_ANCHOR_RE.search(paragraph_text))
            )
            if not is_anchor and not is_visual_only_neighbor:
                continue

            targeted_image_parts.update(_paragraph_image_parts(paragraph))
            _mask_non_image_signature_visuals(paragraph)

    return targeted_image_parts


def _iter_docx_image_parts(document: Any):
    seen: set[int] = set()
    for part in document.part.package.parts:
        content_type = str(getattr(part, "content_type", ""))
        if not content_type.startswith("image/") or id(part) in seen:
            continue
        seen.add(id(part))
        yield part


def _pil_save_format(image: Image.Image, content_type: str) -> str:
    source_format = (image.format or "").upper()
    if source_format in {"JPEG", "PNG", "BMP", "GIF", "TIFF", "WEBP"}:
        return source_format
    return {
        "image/jpeg": "JPEG",
        "image/png": "PNG",
        "image/bmp": "BMP",
        "image/gif": "GIF",
        "image/tiff": "TIFF",
        "image/webp": "WEBP",
    }.get(content_type, "")


def _store_masked_docx_image_part(part: Any, image: Image.Image) -> None:
    save_format = _pil_save_format(image, str(getattr(part, "content_type", "")))
    if not save_format:
        raise RuntimeError(
            "A sensitive DOCX image uses an unsupported image encoding; "
            "masking stopped to prevent an unsafe output."
        )
    output = BytesIO()
    image.convert("RGB").save(output, format=save_format)
    part._blob = output.getvalue()


def _mask_docx_embedded_images(
    *,
    document: Any,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    forced_signature_parts: set[Any],
    ocr_languages: Optional[Sequence[str]] = None,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    ocr_lang = resolve_ocr_lang(ocr_languages)
    for part in _iter_docx_image_parts(document):
        content_type = str(getattr(part, "content_type", "")).lower()
        if part in forced_signature_parts and content_type == "image/svg+xml":
            part._blob = (
                b'<svg xmlns="http://www.w3.org/2000/svg" width="1" height="1" '
                b'viewBox="0 0 1 1"><rect width="1" height="1" fill="#000"/></svg>'
            )
            continue

        try:
            with Image.open(BytesIO(part.blob)) as source_image:
                source_format = source_image.format
                image = source_image.convert("RGB")
                image.format = source_format
        except Exception as exc:
            if part in forced_signature_parts:
                raise RuntimeError(
                    "A signature image in the DOCX could not be decoded; "
                    "masking stopped to prevent an unsafe output."
                ) from exc
            continue

        if part in forced_signature_parts:
            masked = Image.new("RGB", image.size, color=(0, 0, 0))
            masked.format = image.format
            _store_masked_docx_image_part(part, masked)
            continue

        image_text, words = _ocr_words_from_image(
            image,
            ocr_lang=ocr_lang,
            include_character_boxes=bool(_clean_custom_mask_items(custom_redactions)),
        )
        if not image_text.strip():
            continue
        findings = _inspect_text_in_chunks(
            sdp=sdp,
            text=image_text,
            payload=payload,
            custom_redactions=custom_redactions,
        )
        _require_ocr_findings_mapped(findings, words)
        boxes = _boxes_for_text_spans(findings, words)
        boxes.extend(
            _signature_boxes_for_image(
                image=image,
                words=words,
                sdp=sdp,
                payload=payload,
            )
        )
        if not boxes:
            continue

        masked = pixelate_boxes(image, boxes)
        masked.format = image.format
        _store_masked_docx_image_part(part, masked)


def _mask_docx_core_properties(
    *,
    document: Any,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    properties = document.core_properties
    for attribute in (
        "author",
        "category",
        "comments",
        "identifier",
        "keywords",
        "last_modified_by",
        "subject",
        "title",
        "version",
    ):
        value = getattr(properties, attribute, None)
        if not isinstance(value, str) or not value.strip():
            continue
        findings = _mask_findings_for_text(
            sdp=sdp,
            text=value,
            payload=payload,
            custom_redactions=custom_redactions,
        )
        if findings:
            setattr(properties, attribute, mask_text_with_findings(value, findings))


def _inspect_text_in_chunks(
    *,
    sdp: GoogleSDPClient,
    text: str,
    payload: DataMaskingRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[TextFinding]:
    if not text.strip():
        return []

    step = MAX_INSPECTION_CHARS - INSPECTION_OVERLAP_CHARS
    findings: list[TextFinding] = []
    for chunk_start in range(0, len(text), step):
        chunk_end = min(len(text), chunk_start + MAX_INSPECTION_CHARS)
        chunk = text[chunk_start:chunk_end]
        for finding in _mask_findings_for_text(
            sdp=sdp,
            text=chunk,
            payload=payload,
            custom_redactions=custom_redactions,
        ):
            start = chunk_start + finding.start
            end = chunk_start + finding.end
            findings.append(
                TextFinding(
                    start=start,
                    end=end,
                    quote=text[start:end],
                    label=finding.label,
                    source=finding.source,
                )
            )
        if chunk_end >= len(text):
            break

    return merge_overlapping_findings(findings, original_text=text)


def _mask_docx_paragraph_text(
    *,
    document: Any,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    paragraphs: list[Any] = []
    seen: set[Any] = set()
    for paragraph in _iter_document_paragraphs(document):
        element = paragraph._p
        if element in seen:
            continue
        seen.add(element)
        paragraphs.append(paragraph)

    text_parts: list[str] = []
    paragraph_spans: list[tuple[Any, int, int]] = []
    cursor = 0
    for paragraph in paragraphs:
        paragraph_text, _ = _paragraph_run_spans(paragraph)
        if text_parts:
            text_parts.append("\n")
            cursor += 1
        start = cursor
        text_parts.append(paragraph_text)
        cursor += len(paragraph_text)
        paragraph_spans.append((paragraph, start, cursor))

    document_text = "".join(text_parts)
    findings = _inspect_text_in_chunks(
        sdp=sdp,
        text=document_text,
        payload=payload,
        custom_redactions=custom_redactions,
    )
    findings_by_paragraph: dict[Any, list[TextFinding]] = {}
    for finding in findings:
        for paragraph, paragraph_start, paragraph_end in paragraph_spans:
            overlap_start = max(finding.start, paragraph_start)
            overlap_end = min(finding.end, paragraph_end)
            if overlap_end <= overlap_start:
                continue
            local_start = overlap_start - paragraph_start
            local_end = overlap_end - paragraph_start
            paragraph_text, _ = _paragraph_run_spans(paragraph)
            quote = paragraph_text[local_start:local_end]
            if not quote:
                continue
            findings_by_paragraph.setdefault(paragraph._p, []).append(
                TextFinding(
                    start=local_start,
                    end=local_end,
                    quote=quote,
                    label=finding.label,
                    source=finding.source,
                )
            )

    for paragraph in paragraphs:
        paragraph_findings = findings_by_paragraph.get(paragraph._p, [])
        if paragraph_findings:
            mask_paragraph_runs(paragraph, paragraph_findings)


def _mask_docx(
    *,
    source_path: Path,
    output_path: Path,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    ocr_languages: Optional[Sequence[str]] = None,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    document = docx.Document(source_path)
    signature_parts: set[Any] = set()
    if _signature_target_selected(payload) and not _visual_signature_is_excluded(payload):
        signature_parts = _mask_docx_signature_visuals(document)

    _mask_docx_embedded_images(
        document=document,
        sdp=sdp,
        payload=payload,
        forced_signature_parts=signature_parts,
        ocr_languages=ocr_languages,
        custom_redactions=custom_redactions,
    )
    _mask_docx_core_properties(
        document=document,
        sdp=sdp,
        payload=payload,
        custom_redactions=custom_redactions,
    )

    _mask_docx_paragraph_text(
        document=document,
        sdp=sdp,
        payload=payload,
        custom_redactions=custom_redactions,
    )

    document.save(output_path)


def _mask_image_file(
    *,
    source_path: Path,
    output_path: Path,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    ocr_languages: Optional[Sequence[str]] = None,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    image = Image.open(source_path).convert("RGB")
    page_text, words = _ocr_words_from_image(
        image,
        ocr_lang=resolve_ocr_lang(ocr_languages),
        include_character_boxes=bool(_clean_custom_mask_items(custom_redactions)),
    )
    findings = (
        _inspect_text_in_chunks(
            sdp=sdp,
            text=page_text,
            payload=payload,
            custom_redactions=custom_redactions,
        )
        if page_text.strip()
        else []
    )
    _require_ocr_findings_mapped(findings, words)
    boxes = _boxes_for_text_spans(findings, words)
    boxes.extend(
        _signature_boxes_for_image(
            image=image,
            words=words,
            sdp=sdp,
            payload=payload,
        )
    )
    result = pixelate_boxes(image, boxes)
    fmt = DocumentInputFormat(source_path.suffix.lower().lstrip("."))
    result.save(output_path, format=_IMAGE_OUTPUT_MAP[fmt])


def _render_pdf_ocr_region(
    page: fitz.Page,
    *,
    rect: fitz.Rect,
    scale: float,
) -> PDFOCRRegion:
    clipped = fitz.Rect(rect) & page.rect
    pix = page.get_pixmap(
        matrix=fitz.Matrix(scale, scale),
        clip=clipped,
        alpha=False,
        colorspace=fitz.csRGB,
    )
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    transform = fitz.Matrix(
        clipped.width,
        0.0,
        0.0,
        clipped.height,
        clipped.x0,
        clipped.y0,
    )
    return PDFOCRRegion(rect=clipped, image=image, transform=transform)


def _pdf_ocr_regions(
    page: fitz.Page,
    *,
    default_scale: float,
    whole_page_fallback: bool = True,
) -> list[PDFOCRRegion]:
    """Render meaningful image placements exactly as they appear on the page.

    Raw PDF image streams may be rotated, mirrored, or cropped by their
    placement matrix. OCRing the raw stream caused both missed ID fields and
    false masks at unrelated coordinates. Rendering the placement preserves
    visible orientation and gives a direct pixel-to-page mapping.
    """
    page_area = max(1.0, float(page.rect.width * page.rect.height))
    regions: list[PDFOCRRegion] = []
    seen: set[tuple[float, float, float, float]] = set()

    document = page.parent
    for item in page.get_images(full=True):
        xref = int(item[0])
        source_width = int(item[2])
        source_height = int(item[3])
        if (
            document is None
            or xref <= 0
            or source_width < 64
            or source_height < 24
        ):
            continue

        for placement in page.get_image_rects(xref):
            rect = fitz.Rect(placement) & page.rect
            if rect.is_empty or rect.width <= 0 or rect.height <= 0:
                continue
            if float(rect.width * rect.height) / page_area < MIN_EMBEDDED_IMAGE_AREA_RATIO:
                continue

            key = _pdf_rect_key(rect)
            if key in seen:
                continue

            # Do not derive scale from raw image dimensions: rotation/cropping
            # makes those ratios misleading and previously enlarged ID scans
            # enough to degrade OCR into false fields.
            effective_scale = min(
                MAX_PDF_OCR_RENDER_SCALE,
                max(2.0, default_scale),
            )
            regions.append(
                _render_pdf_ocr_region(
                    page,
                    rect=rect,
                    scale=effective_scale,
                )
            )
            seen.add(key)

    if regions:
        return regions

    if not whole_page_fallback:
        return []

    # Vector-only and unusual image encodings still receive a whole-page OCR
    # fallback. Three times PDF resolution is a practical floor for small text.
    fallback_scale = min(
        MAX_PDF_OCR_RENDER_SCALE,
        max(3.0, default_scale),
    )
    return [
        _render_pdf_ocr_region(
            page,
            rect=page.rect,
            scale=fallback_scale,
        )
    ]


def _visual_signature_occurrence_count(
    *,
    source_path: Path,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    ocr_lang: str,
) -> int:
    if not _signature_target_selected(payload) or _visual_signature_is_excluded(payload):
        return 0

    suffix = source_path.suffix.lower()
    if suffix in {".jpg", ".jpeg", ".png"}:
        image = Image.open(source_path).convert("RGB")
        _text, words = _ocr_words_from_image(image, ocr_lang=ocr_lang)
        return len(
            _signature_boxes_for_image(
                image=image,
                words=words,
                sdp=sdp,
                payload=payload,
            )
        )

    if suffix != ".pdf":
        return 0

    count = 0
    with fitz.open(source_path) as document:
        for page in document:
            scale = min(4.0, max(2.5, 2000 / max(1.0, page.rect.width)))
            region = _render_pdf_ocr_region(page, rect=page.rect, scale=scale)
            _text, words = _ocr_words_from_image(region.image, ocr_lang=ocr_lang)
            count += len(
                _signature_boxes_for_image(
                    image=region.image,
                    words=words,
                    sdp=sdp,
                    payload=payload,
                )
            )
    return count


def _pdf_rect_from_ocr_box(
    region: PDFOCRRegion,
    box: tuple[int, int, int, int],
) -> fitz.Rect:
    x0, y0, x1, y1 = box
    width = max(1, region.image.width)
    height = max(1, region.image.height)
    corners = (
        fitz.Point(x0 / width, y0 / height) * region.transform,
        fitz.Point(x1 / width, y0 / height) * region.transform,
        fitz.Point(x0 / width, y1 / height) * region.transform,
        fitz.Point(x1 / width, y1 / height) * region.transform,
    )
    return fitz.Rect(
        min(point.x for point in corners),
        min(point.y for point in corners),
        max(point.x for point in corners),
        max(point.y for point in corners),
    ) & region.rect


def _require_pdf_findings_mapped(
    findings: Sequence[TextFinding],
    words: Sequence[PDFWord],
) -> None:
    unmapped = sum(
        1
        for finding in findings
        if not any(
            not (word.end <= finding.start or word.start >= finding.end)
            for word in words
        )
    )
    if unmapped:
        raise RuntimeError(
            f"{unmapped} sensitive finding(s) could not be mapped to PDF coordinates; "
            "masking stopped to prevent an unsafe output."
        )


def _rect_overlaps_native_text(
    rect: fitz.Rect,
    words: Sequence[PDFWord],
) -> bool:
    candidate = fitz.Rect(rect)
    candidate_area = max(1.0, candidate.get_area())
    for word in words:
        intersection = candidate & word.rect
        if not intersection.is_empty and intersection.get_area() / candidate_area >= 0.20:
            return True
    return False


def _pdf_signature_anchors(
    page: fitz.Page,
    words: Sequence[PDFWord],
) -> list[fitz.Rect]:
    anchors: list[fitz.Rect] = []
    by_line: dict[tuple[int, int], list[PDFWord]] = {}
    for word in words:
        by_line.setdefault((word.block_no, word.line_no), []).append(word)
        normalized = re.sub(r"[^\wÀ-ÖØ-öø-ÿ]+", "", word.text.casefold())
        if normalized in _SIGNATURE_WORDS and _pdf_rect_has_visible_content(page, word.rect):
            anchors.append(fitz.Rect(word.rect))
    for line_words in by_line.values():
        ordered = sorted(line_words, key=lambda item: item.word_no)
        line_text = " ".join(word.text for word in ordered)
        if (
            len(line_text) > 80
            or not _SIGNATURE_ROLE_RE.search(line_text)
            or min(word.rect.y0 for word in ordered)
            < page.rect.y0 + page.rect.height * 0.45
        ):
            continue
        rect = fitz.Rect(ordered[0].rect)
        for word in ordered[1:]:
            rect |= word.rect
        if _pdf_rect_has_visible_content(page, rect):
            anchors.append(rect)
    return anchors


def _pdf_signature_zones(
    page: fitz.Page,
    anchor: fitz.Rect,
) -> tuple[fitz.Rect, fitz.Rect, fitz.Rect]:
    anchor_height = max(8.0, anchor.height)
    right = fitz.Rect(
        anchor.x1 + 1,
        max(page.rect.y0, anchor.y0 - anchor_height),
        min(page.rect.x1, anchor.x1 + max(140.0, anchor_height * 16)),
        min(page.rect.y1, anchor.y1 + max(36.0, anchor_height * 3)),
    )
    below = fitz.Rect(
        max(page.rect.x0, anchor.x0 - anchor_height),
        anchor.y1 + 1,
        min(page.rect.x1, anchor.x0 + max(180.0, anchor_height * 19)),
        min(page.rect.y1, anchor.y1 + max(55.0, anchor_height * 5)),
    )
    above = fitz.Rect(
        max(page.rect.x0, anchor.x0 - anchor_height),
        max(page.rect.y0, anchor.y0 - max(55.0, anchor_height * 5)),
        min(page.rect.x1, anchor.x0 + max(180.0, anchor_height * 19)),
        max(page.rect.y0, anchor.y0 - 1),
    )
    return right, below, above


def _rect_intersects(rect: fitz.Rect, zone: fitz.Rect) -> bool:
    return not (fitz.Rect(rect) & fitz.Rect(zone)).is_empty


def _pdf_signature_regions(
    page: fitz.Page,
    words: Sequence[PDFWord],
) -> list[fitz.Rect]:
    regions: list[fitz.Rect] = []
    widgets_to_delete: list[Any] = []
    annotations_to_delete: list[Any] = []

    try:
        widgets = list(page.widgets() or [])
    except Exception:
        widgets = []
    signature_widget_type = getattr(fitz, "PDF_WIDGET_TYPE_SIGNATURE", None)
    for widget in widgets:
        field_name = str(getattr(widget, "field_name", "") or "")
        field_label = str(getattr(widget, "field_label", "") or "")
        if (
            getattr(widget, "field_type", None) == signature_widget_type
            or _SIGNATURE_ANCHOR_RE.search(f"{field_name} {field_label}")
        ):
            regions.append(fitz.Rect(widget.rect))
            widgets_to_delete.append(widget)

    try:
        annotations = list(page.annots() or [])
    except Exception:
        annotations = []
    for annotation in annotations:
        annotation_type = str((getattr(annotation, "type", ("", "")) or ("", ""))[1]).casefold()
        info = getattr(annotation, "info", {}) or {}
        descriptor = " ".join(str(value) for value in info.values() if value)
        if annotation_type == "ink" or (
            annotation_type == "stamp" and _SIGNATURE_ANCHOR_RE.search(descriptor)
        ):
            regions.append(fitz.Rect(annotation.rect))
            annotations_to_delete.append(annotation)

    anchors = _pdf_signature_anchors(page, words)
    zones = [zone for anchor in anchors for zone in _pdf_signature_zones(page, anchor)]
    page_area = max(1.0, float(page.rect.width * page.rect.height))

    if zones:
        for item in page.get_images(full=True):
            xref = int(item[0])
            try:
                placements = page.get_image_rects(xref)
            except Exception:
                continue
            for placement in placements:
                rect = fitz.Rect(placement) & page.rect
                if rect.is_empty:
                    continue
                if float(rect.width * rect.height) / page_area > MAX_EMBEDDED_IMAGE_AREA_RATIO_FOR_SIGNATURE:
                    continue
                if any(_rect_intersects(rect, zone) for zone in zones):
                    regions.append(rect)

        try:
            drawings = page.get_drawings()
        except Exception:
            drawings = []
        for zone in zones:
            drawing_rects: list[fitz.Rect] = []
            for drawing in drawings:
                rect = fitz.Rect(drawing.get("rect", fitz.Rect()))
                if (
                    rect.is_empty
                    or rect.width < 4
                    or rect.height < 2.5
                    or not _rect_intersects(rect, zone)
                ):
                    continue
                drawing_rects.append(rect)
            if drawing_rects:
                combined = drawing_rects[0]
                for rect in drawing_rects[1:]:
                    combined |= rect
                regions.append(combined)

    for widget in widgets_to_delete:
        try:
            page.delete_widget(widget)
        except Exception:
            pass
    for annotation in annotations_to_delete:
        try:
            page.delete_annot(annotation)
        except Exception:
            pass

    unique: list[fitz.Rect] = []
    seen: set[tuple[float, float, float, float]] = set()
    for region in regions:
        clipped = fitz.Rect(region) & page.rect
        if clipped.is_empty:
            continue
        clipped.x0 = max(page.rect.x0, clipped.x0 - 1.5)
        clipped.y0 = max(page.rect.y0, clipped.y0 - 1.5)
        clipped.x1 = min(page.rect.x1, clipped.x1 + 1.5)
        clipped.y1 = min(page.rect.y1, clipped.y1 + 1.5)
        key = _pdf_rect_key(clipped)
        if key not in seen:
            unique.append(clipped)
            seen.add(key)
    return unique


def _pdf_metadata_findings(
    *,
    document: fitz.Document,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> list[tuple[str, str, list[TextFinding]]]:
    metadata = dict(document.metadata or {})
    target_values = {
        str(getattr(target, "value", target))
        for target in payload.target_data
    }
    exclusions = {
        _normalize_text_for_compare(item)
        for item in payload.review_exclusions
        if item and item.strip()
    }
    results: list[tuple[str, str, list[TextFinding]]] = []

    for key in _PDF_REVIEWABLE_METADATA_KEYS:
        value = metadata.get(key)
        if not isinstance(value, str) or not value.strip():
            continue
        findings = _mask_findings_for_text(
            sdp=sdp,
            text=value,
            payload=payload,
            custom_redactions=custom_redactions,
        )
        normalized_value = _normalize_text_for_compare(value)
        if (
            key == "author"
            and SensitiveDataType.name.value in target_values
            and normalized_value not in exclusions
            and normalized_value not in _GENERIC_METADATA_AUTHORS
        ):
            findings = merge_overlapping_findings(
                [
                    *findings,
                    TextFinding(
                        start=0,
                        end=len(value),
                        quote=value,
                        label=SensitiveDataType.name.value,
                        source="pdf_metadata",
                    ),
                ],
                original_text=value,
            )
        if findings:
            results.append((key, value, findings))
    return results


def _mask_pdf_metadata(
    *,
    document: fitz.Document,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    metadata = dict(document.metadata or {})
    changed = False
    for key, value, findings in _pdf_metadata_findings(
        document=document,
        sdp=sdp,
        payload=payload,
        custom_redactions=custom_redactions,
    ):
        metadata[key] = mask_text_with_findings(value, findings)
        changed = True
    if changed:
        document.set_metadata(metadata)


def _mask_pdf(
    *,
    source_path: Path,
    output_path: Path,
    sdp: GoogleSDPClient,
    payload: DataMaskingRequest,
    ocr_languages: Optional[Sequence[str]] = None,
    render_scale: float = DEFAULT_PDF_RENDER_SCALE,
    custom_redactions: Optional[Sequence[str]] = None,
) -> None:
    doc = fitz.open(source_path)
    temporary_paths: list[Path] = []
    detected_findings = 0
    applied_regions = 0
    ocr_lang = resolve_ocr_lang(ocr_languages)
    try:
        _mask_pdf_metadata(
            document=doc,
            sdp=sdp,
            payload=payload,
            custom_redactions=custom_redactions,
        )
        for embedded_name in list(doc.embfile_names()):
            doc.embfile_del(embedded_name)

        for page in doc:
            page_text, words = _page_words_with_offsets(page)
            signature_regions = (
                _pdf_signature_regions(page, words)
                if _signature_target_selected(payload) and not _visual_signature_is_excluded(payload)
                else []
            )
            regular_specs: list[tuple[fitz.Rect, fitz.Rect, str]] = []
            regular_rect_keys: set[tuple[float, float, float, float]] = set()
            regular_region_count = 0
            touches_images = False

            if words:
                findings = _inspect_text_in_chunks(
                    sdp=sdp,
                    text=page_text,
                    payload=payload,
                    custom_redactions=custom_redactions,
                )
                detected_findings += len(findings)
                _require_pdf_findings_mapped(findings, words)

                specs = _masked_rect_specs(findings, words, original_text=page_text)
                for redaction_rect, original_rect, masked in specs:
                    if not _pdf_rect_has_visible_content(page, original_rect):
                        continue
                    key = _pdf_rect_key(redaction_rect)
                    if key in regular_rect_keys:
                        continue
                    page.add_redact_annot(
                        redaction_rect,
                        fill=(1, 1, 1),
                        cross_out=False,
                    )
                    regular_specs.append((redaction_rect, original_rect, masked))
                    regular_rect_keys.add(key)
                    regular_region_count += 1
                    applied_regions += 1

            for region in _pdf_ocr_regions(
                page,
                default_scale=render_scale,
                whole_page_fallback=not bool(words),
            ):
                region_text, ocr_words = _ocr_words_from_image(
                    region.image,
                    ocr_lang=ocr_lang,
                    include_character_boxes=bool(
                        _clean_custom_mask_items(custom_redactions)
                    ),
                )
                ocr_passes = [(region, region_text, ocr_words)]
                if not is_id_document_payload(payload):
                    primary_scale = region.image.width / max(1.0, region.rect.width)
                    supplemental_region = _render_pdf_ocr_region(
                        page,
                        rect=region.rect,
                        scale=max(3.0, primary_scale),
                    )
                    supplemental_text, supplemental_words = _thresholded_ocr_words_from_image(
                        supplemental_region.image,
                        ocr_lang=ocr_lang,
                        include_character_boxes=bool(
                            _clean_custom_mask_items(custom_redactions)
                        ),
                    )
                    if (
                        supplemental_text.strip()
                        and _normalize_text_for_compare(supplemental_text)
                        != _normalize_text_for_compare(region_text)
                    ):
                        ocr_passes.append(
                            (supplemental_region, supplemental_text, supplemental_words)
                        )

                primary_finding_keys: set[tuple[str, str]] = set()
                for pass_index, (ocr_region, pass_text, pass_words) in enumerate(ocr_passes):
                    if not pass_text.strip():
                        continue
                    findings = _inspect_text_in_chunks(
                        sdp=sdp,
                        text=pass_text,
                        payload=payload,
                        custom_redactions=custom_redactions,
                    )
                    if pass_index == 0:
                        primary_finding_keys = {
                            (finding.label, _normalize_text_for_compare(finding.quote))
                            for finding in findings
                        }
                    else:
                        findings = [
                            finding
                            for finding in findings
                            if (
                                finding.label,
                                _normalize_text_for_compare(finding.quote),
                            )
                            not in primary_finding_keys
                        ]
                    detected_findings += len(findings)
                    _require_ocr_findings_mapped(findings, pass_words)

                    for box, masked in _ocr_mask_specs(findings, pass_words):
                        original_rect = _pdf_rect_from_ocr_box(ocr_region, box)
                        if original_rect.is_empty:
                            continue
                        if words and _rect_overlaps_native_text(original_rect, words):
                            continue
                        redaction_rect = _tight_pdf_text_rect(original_rect)
                        key = _pdf_rect_key(redaction_rect)
                        if key in regular_rect_keys:
                            continue
                        page.add_redact_annot(
                            redaction_rect,
                            fill=(1, 1, 1),
                            cross_out=False,
                        )
                        regular_specs.append((redaction_rect, original_rect, masked))
                        regular_rect_keys.add(key)
                        regular_region_count += 1
                        applied_regions += 1
                        touches_images = True

                if _signature_target_selected(payload) and not _visual_signature_is_excluded(payload):
                    for box in _signature_boxes_for_image(
                        image=region.image,
                        words=ocr_words,
                        sdp=sdp,
                        payload=payload,
                    ):
                        signature_rect = _pdf_rect_from_ocr_box(region, box)
                        if not signature_rect.is_empty:
                            signature_regions.append(signature_rect)

            if regular_region_count:
                _apply_page_redactions(
                    page,
                    image_mode="pixels" if touches_images else "none",
                )
                _insert_pdf_mask_text(page, regular_specs)

            signature_keys: set[tuple[float, float, float, float]] = set()
            signature_region_count = 0
            for rect in signature_regions:
                clipped = fitz.Rect(rect) & page.rect
                if clipped.is_empty:
                    continue
                key = _pdf_rect_key(clipped)
                if key in signature_keys:
                    continue
                page.add_redact_annot(
                    clipped,
                    fill=(0, 0, 0),
                    cross_out=False,
                )
                signature_keys.add(key)
                signature_region_count += 1
                applied_regions += 1

            if signature_region_count:
                _apply_page_redactions(
                    page,
                    image_mode="pixels",
                    remove_graphics=True,
                )

        if detected_findings > 0 and applied_regions == 0:
            raise RuntimeError(
                "Sensitive data was detected, but no PDF regions could be mapped for masking."
            )

        _save_pdf_checked(
            doc,
            output_path,
            temporary_paths=temporary_paths,
        )
    finally:
        doc.close()
        for temporary_path in temporary_paths:
            temporary_path.unlink(missing_ok=True)


def apply_data_mask(
    request: AnalyzerRequest | Mapping[str, Any],
    *,
    source_path: str | Path,
    output_dir: str | Path,
    sdp: GoogleSDPClient | None = None,
    project_id: str | None = None,
    location: str = DEFAULT_DLP_LOCATION,
    min_likelihood: str = DEFAULT_MIN_LIKELIHOOD,
    client: Any | None = None,
    ocr_languages: Optional[Sequence[str]] = None,
    custom_redactions: Optional[Sequence[str]] = None,
) -> AnalyzerResponse:
    req = validate_analyzer_request(request)
    input_payload, payload = _ensure_mask_request(req)
    source = Path(source_path)
    if not source.exists():
        raise FileNotFoundError(f"Source file not found: {source}")
    output_path = _output_path(source, Path(output_dir), suffix="_masked")
    resolved = _resolve_sdp(
        sdp=sdp,
        project_id=project_id,
        location=location,
        min_likelihood=min_likelihood,
        client=client,
    )

    fmt = input_payload.metadata.input_format
    if fmt == DocumentInputFormat.docx:
        _mask_docx(
            source_path=source,
            output_path=output_path,
            sdp=resolved,
            payload=payload,
            ocr_languages=ocr_languages,
            custom_redactions=custom_redactions,
        )
    elif fmt == DocumentInputFormat.pdf:
        _mask_pdf(
            source_path=source,
            output_path=output_path,
            sdp=resolved,
            payload=payload,
            ocr_languages=ocr_languages,
            custom_redactions=custom_redactions,
        )
    elif fmt in {DocumentInputFormat.jpg, DocumentInputFormat.jpeg, DocumentInputFormat.png}:
        _mask_image_file(
            source_path=source,
            output_path=output_path,
            sdp=resolved,
            payload=payload,
            ocr_languages=ocr_languages,
            custom_redactions=custom_redactions,
        )
    else:
        raise ValueError("Unsupported data_mask input format.")
    return _build_masking_response(request=req, output_path=output_path)


__all__ = [
    "preview_data_mask_candidates",
    "mask_value_by_target",
    "mask_text_with_findings",
    "pixelate_boxes",
    "mask_paragraph_runs",
    "apply_data_mask",
]