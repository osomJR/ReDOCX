from __future__ import annotations

"""
V1 document-conversion processing.

Purpose:
- hold conversion-specific processing logic outside analyzer.py
- keep schema/validation/extraction unchanged
- keep analyzer responsible only for orchestration, routing, and response building

Design notes:
- stateless at the processor layer
- schema-agnostic: this module returns conversion artifact metadata
- provider-backed by default via concrete conversion backends
- performs real file conversion work for the contract-allowed pairs
- persists produced artifacts through a separate storage layer
- analyzer.py remains responsible for mapping ConversionArtifact into FileResult
"""

from collections import Counter
from copy import deepcopy
from dataclasses import dataclass
from io import BytesIO
from html.parser import HTMLParser
from pathlib import Path
from tempfile import TemporaryDirectory, NamedTemporaryFile
from typing import Any, Iterable, Mapping, Optional, Protocol, Sequence
import json
import hashlib
import csv
import sys
import threading
import logging
import math
import os
import re
import signal
import unicodedata
import shutil
import subprocess
import zipfile
from urllib.parse import unquote, urlsplit
from xml.etree import ElementTree

import fitz  # PyMuPDF
from PIL import Image, ImageChops, ImageFilter, ImageOps, ImageStat, UnidentifiedImageError
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from backend.src.storage.artifacts import (
    StorageBackend,
    LocalArtifactStorage,
)

try:
    from pdf2docx import Converter as PDFToDOCXConverter
except ImportError:  # pragma: no cover
    PDFToDOCXConverter = None

try:
    from openpyxl import Workbook, load_workbook
    from openpyxl.drawing.image import Image as OpenPyXLImage
except ImportError:  # pragma: no cover
    Workbook = None
    load_workbook = None
    OpenPyXLImage = None

try:
    from pptx import Presentation as PptxPresentation
except ImportError:  # pragma: no cover
    PptxPresentation = None

try:
    from weasyprint import HTML as WeasyHTML, default_url_fetcher as weasy_default_url_fetcher
except ImportError:  # pragma: no cover
    WeasyHTML = None
    weasy_default_url_fetcher = None


from . import native, quality

_MAX_CONVERSION_WORKERS = max(1, min(16, int(os.getenv("REDOCX_CONVERSION_MAX_WORKERS", "2"))))
_CONVERSION_SLOTS = threading.BoundedSemaphore(_MAX_CONVERSION_WORKERS)

logger = logging.getLogger(__name__)


CONVERSION_RULES = """
TASK: DOCUMENT CONVERSION
RULES:
- Convert only between contract-allowed format pairs
- Preserve source content fidelity as much as the target format permits
- Do not add, remove, summarize, translate, or reinterpret content
- Output must be a downloadable file artifact, never inline text
- Conversion is format transformation only, not semantic editing
""".strip()

ALLOWED_CONVERSION_PAIRS: dict[str, set[str]] = {
    "pdf": {"docx", "jpg", "pptx", "xlsx", "pdfa"},
    "docx": {"pdf"},
    "xlsx": {"pdf"},
    "html": {"pdf"},
    "htm": {"pdf"},
    "pptx": {"pdf"},
    "jpg": {"pdf", "docx"},
    "jpeg": {"pdf", "docx"},
    "png": {"jpg", "jpeg"},
}

OUTPUT_EXTENSION_ALIASES: dict[str, str] = {
    "jpg": "jpg",
    "jpeg": "jpeg",
    "pdf": "pdf",
    "pdfa": "pdfa",
    "docx": "docx",
    "png": "png",
    "xlsx": "xlsx",
    "html": "html",
    "htm": "htm",
    "pptx": "pptx",
    "zip": "zip",
}

PHYSICAL_OUTPUT_FORMAT_BY_TARGET: dict[str, str] = {
    "pdf": "pdf",
    "pdfa": "pdf",
    "docx": "docx",
    "jpg": "jpg",
    "jpeg": "jpeg",
    "pptx": "pptx",
    "xlsx": "xlsx",
}

PDF_TO_DOCX_MODES = {"auto", "editable", "fidelity", "native"}
DEFAULT_PDF_TO_DOCX_MODE = os.getenv("REDOCX_PDF_TO_DOCX_MODE", "auto").strip().lower()
DEFAULT_PDF_TO_DOCX_OCR_DPI = int(os.getenv("REDOCX_PDF_TO_DOCX_OCR_DPI", "200"))
DEFAULT_PDF_TO_DOCX_OCR_LANGUAGE = os.getenv("REDOCX_PDF_TO_DOCX_OCR_LANGUAGE", "eng").strip() or "eng"
DEFAULT_PDF_TO_DOCX_FIDELITY_DPI = int(os.getenv("REDOCX_PDF_TO_DOCX_FIDELITY_DPI", "180"))
DEFAULT_PDF_TO_DOCX_VISUAL_MIN_SCORE = float(
    os.getenv("REDOCX_PDF_TO_DOCX_VISUAL_MIN_SCORE", "0.80")
)
DEFAULT_PDF_TO_DOCX_PAGE_INSET_PT = float(
    os.getenv("REDOCX_PDF_TO_DOCX_PAGE_INSET_PT", "0.5")
)
DEFAULT_DOCX_TO_PDF_TIMEOUT_SECONDS = int(os.getenv("REDOCX_DOCX_TO_PDF_TIMEOUT_SECONDS", "90"))
DEFAULT_OFFICE_TO_PDF_TIMEOUT_SECONDS = int(os.getenv("REDOCX_OFFICE_TO_PDF_TIMEOUT_SECONDS", str(DEFAULT_DOCX_TO_PDF_TIMEOUT_SECONDS)))
DEFAULT_PDFA_TIMEOUT_SECONDS = int(os.getenv("REDOCX_PDFA_TIMEOUT_SECONDS", "120"))
DEFAULT_PDFA_TEXT_TOKEN_COVERAGE = float(
    os.getenv("REDOCX_PDFA_TEXT_TOKEN_COVERAGE", "0.80")
)
DEFAULT_PDFA_VISUAL_MIN_SCORE = float(
    os.getenv("REDOCX_PDFA_VISUAL_MIN_SCORE", "0.98")
)
DEFAULT_PDF_RASTER_DPI = int(os.getenv("REDOCX_PDF_RASTER_DPI", "200"))
DEFAULT_IMAGE_PDF_DPI = float(os.getenv("REDOCX_IMAGE_PDF_DPI", "150"))
DEFAULT_IMAGE_JPEG_QUALITY = int(os.getenv("REDOCX_IMAGE_JPEG_QUALITY", "95"))
MAX_PDF_TO_DOCX_PAGES = int(os.getenv("REDOCX_PDF_TO_DOCX_MAX_PAGES", "250"))
MIN_EDITABLE_DOCX_TEXT_CHARS = int(os.getenv("REDOCX_MIN_EDITABLE_DOCX_TEXT_CHARS", "20"))
MIN_PDF_TEXT_RETENTION_RATIO = float(os.getenv("REDOCX_MIN_PDF_TEXT_RETENTION_RATIO", "0.70"))

# Conversion reliability / fidelity gates. These reject structurally valid but
# materially incomplete outputs instead of silently returning untrustworthy files.
DEFAULT_TEXT_TOKEN_COVERAGE = float(os.getenv("REDOCX_CONVERSION_TEXT_TOKEN_COVERAGE", "0.95"))
DEFAULT_XLSX_PDF_TOKEN_COVERAGE = float(os.getenv("REDOCX_XLSX_PDF_TOKEN_COVERAGE", "0.985"))
DEFAULT_XLSX_PDF_CELL_COVERAGE = float(os.getenv("REDOCX_XLSX_PDF_CELL_COVERAGE", "0.995"))
DEFAULT_XLSX_PDF_CHARACTER_COVERAGE = float(os.getenv("REDOCX_XLSX_PDF_CHARACTER_COVERAGE", "0.9975"))
DEFAULT_XLSX_LAYOUT_MODE = os.getenv("REDOCX_XLSX_PDF_LAYOUT_MODE", "readable").strip().lower()
DEFAULT_XLSX_MIN_COLUMN_WIDTH = float(os.getenv("REDOCX_XLSX_MIN_COLUMN_WIDTH", "10"))
DEFAULT_XLSX_MAX_COLUMN_WIDTH = float(os.getenv("REDOCX_XLSX_MAX_COLUMN_WIDTH", "42"))
DEFAULT_XLSX_TOTAL_COLUMN_WIDTH = float(os.getenv("REDOCX_XLSX_TOTAL_COLUMN_WIDTH", "108"))
DEFAULT_XLSX_LANDSCAPE_THRESHOLD = float(os.getenv("REDOCX_XLSX_LANDSCAPE_THRESHOLD", "72"))
DEFAULT_XLSX_PAPER_SIZE = os.getenv("REDOCX_XLSX_PAPER_SIZE", "A4").strip().upper() or "A4"
DEFAULT_PDF_TO_XLSX_RASTER_DPI = int(os.getenv("REDOCX_PDF_TO_XLSX_RASTER_DPI", "144"))
DEFAULT_PDF_TO_XLSX_MAX_IMAGE_WIDTH = int(os.getenv("REDOCX_PDF_TO_XLSX_MAX_IMAGE_WIDTH", "1600"))
DEFAULT_PDFA_VALIDATOR_MODE = os.getenv("REDOCX_PDFA_VALIDATOR_MODE", "required").strip().lower()
DEFAULT_LIBREOFFICE_UNO_PYTHON = os.getenv(
    "REDOCX_LIBREOFFICE_UNO_PYTHON", "/usr/bin/python3"
).strip() or "/usr/bin/python3"
LIBREOFFICE_CALC_PDF_WORKER = Path(__file__).with_name("libreoffice_calc_pdf_worker.py")

CONTENT_TYPES_BY_FORMAT: dict[str, str] = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
    "png": "image/png",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "zip": "application/zip",
}

WORDPROCESSINGML_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
XML_NS = "http://www.w3.org/XML/1998/namespace"
REDOCX_FIDELITY_DOCX_MARKER = "redocx-pdf-fidelity-v1"


@dataclass(frozen=True)
class ConversionArtifact:
    """
    Schema-agnostic conversion artifact descriptor.

    file_path is the persisted artifact location, not the temporary conversion path.
    storage_key and download_url are exposed for future orchestration layers.
    """

    file_name: str
    file_extension: str
    file_size_mb: float
    file_path: str
    storage_key: Optional[str] = None
    download_url: Optional[str] = None


@dataclass(frozen=True)
class PdfHyperlink:
    """External link annotation recovered from a source PDF."""

    page_number: int
    label: str
    target: str


@dataclass(frozen=True)
class PdfTextLine:
    """One visual PDF line and its native text spans."""

    bbox: tuple[float, float, float, float]
    spans: tuple[dict[str, Any], ...]

    @property
    def text(self) -> str:
        return "".join(str(span.get("text", "")) for span in self.spans)


@dataclass(frozen=True)
class PdfTextGroup:
    """Consecutive visual lines that should become one editable Word paragraph."""

    bbox: tuple[float, float, float, float]
    lines: tuple[PdfTextLine, ...]

    @property
    def text(self) -> str:
        return " ".join(line.text.strip() for line in self.lines if line.text.strip())


class ConversionBackend(Protocol):
    """Provider interface for document-conversion runtime."""

    def convert(
        self,
        *,
        source_reference: str,
        input_format: str,
        output_format: str,
        planned_output_name: str,
    ) -> ConversionArtifact:
        ...


class RealConversionBackend:
    """Allowed conversions with native Office content, fail-closed validation,
    and one isolated interpreter per job. Artifact ownership remains in the caller.
    """

    def __init__(self, storage_backend: Optional[StorageBackend] = None, *, _in_worker: bool = False) -> None:
        self.storage_backend = storage_backend or LocalArtifactStorage()
        self._in_worker = _in_worker

    def convert(
        self,
        *,
        source_reference: str,
        input_format: str,
        output_format: str,
        planned_output_name: str,
    ) -> ConversionArtifact:
        if not self._in_worker:
            return self._convert_isolated(source_reference=source_reference, input_format=input_format,
                                          output_format=output_format, planned_output_name=planned_output_name)
        ensure_allowed_conversion_pair(input_format=input_format, output_format=output_format)
        normalized_input = _normalize_format(input_format)
        normalized_output = _normalize_format(output_format)
        source_path = Path(_normalize_source_reference(source_reference)).resolve()

        if not source_path.exists():
            raise FileNotFoundError(f"Source file not found: {source_path}")

        _validate_source_format(source_path, normalized_input)
        planned_name = _normalize_file_name(planned_output_name)
        if Path(planned_name).name != planned_name or "\\" in planned_name:
            raise ValueError("Output name must be a file name without a directory.")

        with TemporaryDirectory(prefix="convert-work-") as workdir:
            output_path = (Path(workdir) / planned_name).resolve()
            physical_output_format = PHYSICAL_OUTPUT_FORMAT_BY_TARGET[normalized_output]

            if normalized_input == "pdf" and normalized_output == "docx":
                self._convert_pdf_to_docx(source_path, output_path)
            elif normalized_input in {"docx", "xlsx", "pptx"} and normalized_output == "pdf":
                self._convert_office_to_pdf(source_path, output_path, normalized_input)
            elif normalized_input in {"html", "htm"} and normalized_output == "pdf":
                self._convert_html_to_pdf(source_path, output_path)
            elif normalized_input == "pdf" and normalized_output == "jpg":
                physical_output_format = self._convert_pdf_to_jpg(source_path, output_path)
                if physical_output_format == "zip":
                    output_path = output_path.with_suffix(".zip")
            elif normalized_input == "pdf" and normalized_output == "pptx":
                self._convert_pdf_to_pptx(source_path, output_path)
            elif normalized_input == "pdf" and normalized_output == "xlsx":
                self._convert_pdf_to_xlsx(source_path, output_path)
            elif normalized_input == "pdf" and normalized_output == "pdfa":
                self._convert_pdf_to_pdfa(source_path, output_path)
            elif normalized_input in {"jpg", "jpeg"} and normalized_output == "pdf":
                self._convert_image_to_pdf(source_path, output_path)
            elif normalized_input in {"jpg", "jpeg"} and normalized_output == "docx":
                self._convert_image_to_docx(source_path, output_path)
            elif normalized_input == "png" and normalized_output in {"jpg", "jpeg"}:
                self._convert_png_to_jpeg_family(source_path, output_path)
            else:
                raise ValueError(
                    f"Unsupported conversion pair: {normalized_input} -> {normalized_output}."
                )

            if not output_path.exists() or output_path.stat().st_size <= 0:
                raise RuntimeError("Conversion completed without producing a valid output file.")

            _validate_converted_output(
                output_path,
                physical_output_format,
                source_path=source_path,
                input_format=normalized_input,
            )
            _validate_conversion_fidelity(
                source_path=source_path,
                input_format=normalized_input,
                requested_output_format=normalized_output,
                output_path=output_path,
                physical_output_format=physical_output_format,
            )
            if normalized_output == "pdfa":
                _validate_pdfa_file(output_path)

            stored = self.storage_backend.persist(
                source_file_path=str(output_path),
                artifact_name=output_path.name,
                content_type=CONTENT_TYPES_BY_FORMAT[physical_output_format],
            )

            stored_path = Path(stored.stored_path)
            _validate_converted_output(
                stored_path,
                physical_output_format,
                source_path=source_path,
                input_format=normalized_input,
            )
            if normalized_output == "pdfa":
                _validate_pdfa_file(stored_path)

            return ConversionArtifact(
                file_name=output_path.name,
                file_extension=physical_output_format,
                file_size_mb=_get_file_size_mb(stored_path),
                file_path=stored.stored_path,
                storage_key=stored.storage_key,
                download_url=stored.download_url,
            )

    def _convert_isolated(self, **request) -> ConversionArtifact:
        ensure_allowed_conversion_pair(input_format=request["input_format"], output_format=request["output_format"])
        timeout = max(10, min(1800, int(os.getenv("REDOCX_CONVERSION_TIMEOUT_SECONDS", "300"))))
        if not _CONVERSION_SLOTS.acquire(timeout=timeout):
            raise RuntimeError("Conversion capacity is busy; retry the request.")
        try:
            return self._convert_in_slot(request, timeout)
        finally:
            _CONVERSION_SLOTS.release()

    def _convert_in_slot(self, request, timeout) -> ConversionArtifact:
        worker = Path(__file__).with_name("conversion_worker.py")
        with TemporaryDirectory(prefix="redocx-conversion-") as work:
            root = Path(work).resolve()
            request_file, response_file = root/"request.json", root/"response.json"
            request_file.write_text(json.dumps(request), encoding="utf-8")
            process = subprocess.Popen([sys.executable, "-B", str(worker), str(request_file), str(response_file)],
                stdin=subprocess.DEVNULL, stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                text=True, start_new_session=(os.name != "nt"))
            try:
                stdout, stderr = process.communicate(timeout=timeout)
            except subprocess.TimeoutExpired as exc:
                if os.name != "nt":
                    try:
                        os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError:
                        pass
                else:
                    process.kill()
                process.communicate()
                raise RuntimeError("Conversion exceeded the configured time limit.") from exc
            if not response_file.is_file():
                raise RuntimeError("Conversion worker exited without a result.")
            result = json.loads(response_file.read_text(encoding="utf-8"))
            if not result.get("ok"):
                error_type = {"ValueError": ValueError, "FileNotFoundError": FileNotFoundError}.get(result.get("type"), RuntimeError)
                raise error_type(result.get("error", "Conversion worker failed."))
            if process.returncode != 0:
                raise RuntimeError("Conversion worker exited abnormally.")
            staged = Path(result["path"]).resolve()
            if staged.parent != root or not staged.is_file():
                raise RuntimeError("Conversion worker returned an invalid artifact location.")
            digest = hashlib.sha256(staged.read_bytes()).hexdigest()
            if digest != result["sha256"]:
                raise RuntimeError("Conversion artifact failed its integrity check.")
            stored = self.storage_backend.persist(source_file_path=str(staged), artifact_name=staged.name,
                content_type=CONTENT_TYPES_BY_FORMAT[result["format"]])
            stored_path = Path(stored.stored_path)
            if hashlib.sha256(stored_path.read_bytes()).hexdigest() != digest:
                raise RuntimeError("Stored conversion artifact failed its integrity check.")
            return ConversionArtifact(file_name=staged.name, file_extension=result["format"],
                file_size_mb=_get_file_size_mb(stored_path), file_path=stored.stored_path,
                storage_key=stored.storage_key, download_url=stored.download_url)

    def _convert_pdf_to_docx(self, source_path: Path, output_path: Path) -> None:
        mode = _pdf_to_docx_mode()
        links = _extract_pdf_hyperlinks(source_path)
        with fitz.open(source_path) as pdf:
            has_page_background = any(any(fitz.Rect(info['bbox']).get_area() >= page.rect.get_area()*.75
                                         for info in page.get_image_info()) for page in pdf)
        if mode in {"auto", "editable"} and PDFToDOCXConverter is not None and not has_page_background:
            try:
                converter = PDFToDOCXConverter(str(source_path))
                try:
                    converter.convert(str(output_path))
                finally:
                    converter.close()
                _validate_docx_file(output_path)
                _validate_pdf_to_docx_editability(source_path, output_path)
                with fitz.open(source_path) as pdf:
                    expected_tables = sum(len(native.extract_tables(page)) for page in pdf)
                if len(Document(output_path).tables) < expected_tables:
                    raise RuntimeError("PDF-to-Word did not reconstruct the detected tables.")
                _apply_pdf_hyperlinks(output_path, links)
                _assert_pdf_hyperlinks_preserved(output_path, links)
                self._assert_pdf_to_docx_visual_fidelity(source_path, output_path)
                return
            except Exception as exc:
                logger.info("PDF-to-Word native engine requires positioned reconstruction: %s", exc)
                output_path.unlink(missing_ok=True)
        # 'fidelity' remains an accepted configuration alias, but creates visible
        # editable text and native tables instead of hidden text and page images.
        native.convert_docx(source_path, output_path)
        _validate_docx_file(output_path)
        _validate_pdf_to_docx_editability(source_path, output_path)
        _apply_pdf_hyperlinks(output_path, links)
        _assert_pdf_hyperlinks_preserved(output_path, links)
        self._assert_pdf_to_docx_visual_fidelity(source_path, output_path)

    def _convert_pdf_to_fidelity_docx(self, source_path: Path, output_path: Path) -> None:
        native.convert_docx(source_path, output_path)

    def _assert_pdf_to_docx_visual_fidelity(
        self,
        source_path: Path,
        docx_path: Path,
    ) -> None:
        """Round-trip the DOCX to PDF and reject material layout drift."""

        with TemporaryDirectory(prefix="pdf-docx-visual-audit-") as audit_dir:
            roundtrip_pdf = Path(audit_dir) / "roundtrip.pdf"
            self._convert_office_to_pdf(docx_path, roundtrip_pdf, "docx")
            _assert_pdf_visual_layout_fidelity(
                source_path,
                roundtrip_pdf,
                minimum_score=max(0.0, min(1.0, DEFAULT_PDF_TO_DOCX_VISUAL_MIN_SCORE)),
                label="PDF-to-Word conversion",
            )

    def _convert_pdf_to_native_editable_docx(self, source_path: Path, output_path: Path) -> None:
        native.convert_docx(source_path, output_path)

    def _convert_docx_to_pdf(self, source_path: Path, output_path: Path) -> None:
        self._convert_office_to_pdf(source_path, output_path, "docx")

    def _convert_office_to_pdf(
        self,
        source_path: Path,
        output_path: Path,
        input_format: str,
    ) -> None:
        """Convert Office documents through an isolated LibreOffice process.

        XLSX conversion is environment-hardened. ReDOCX first uses the balanced
        OOXML print-layout normalizer and verifies the resulting PDF with a
        cell-aware retention audit that is tolerant of PDF extractor line wraps.
        If the first export is materially incomplete, ReDOCX automatically retries
        from the untouched source with a conservative ``safe`` layout that wraps
        every populated visible cell and releases row heights for auto-sizing.

        This retry is intentionally scoped to XLSX: DOCX and PPTX already carry
        their own page/flow semantics and are validated by the generic fidelity
        gates after export.
        """
        soffice_bin = (
            os.getenv("SOFFICE_PATH")
            or shutil.which("soffice")
            or shutil.which("libreoffice")
            or shutil.which("soffice.exe")
        )
        if not soffice_bin:
            raise RuntimeError(
                "LibreOffice not found. Set SOFFICE_PATH or add soffice/libreoffice to PATH."
            )

        source_path = source_path.resolve()
        output_path = output_path.resolve()
        output_dir = output_path.parent.resolve()
        if not source_path.exists():
            raise FileNotFoundError(f"Source Office file not found: {source_path}")

        export_filters = {
            "docx": "writer_pdf_Export",
            "xlsx": "calc_pdf_Export",
            "pptx": "impress_pdf_Export",
        }
        export_filter = export_filters.get(input_format)
        if not export_filter:
            raise ValueError(f"Unsupported Office-to-PDF input format: {input_format}.")

        filter_options: dict[str, dict[str, str]] = {
            "UseLosslessCompression": {"type": "boolean", "value": "true"},
            "ReduceImageResolution": {"type": "boolean", "value": "false"},
            "UseTaggedPDF": {"type": "boolean", "value": "true"},
            "ExportBookmarks": {"type": "boolean", "value": "true"},
        }
        if input_format == "pptx":
            # Impress can emit invalid structure trees for independent text boxes.
            # A regular PDF export retains selectable text without claiming tags.
            filter_options["UseTaggedPDF"] = {"type": "boolean", "value": "false"}
            filter_options["ExportHiddenSlides"] = {"type": "boolean", "value": "false"}
        filter_spec = f"pdf:{export_filter}:{json.dumps(filter_options, separators=(',', ':'))}"

        layout_profiles: tuple[str, ...]
        if input_format != "xlsx" or DEFAULT_XLSX_LAYOUT_MODE == "preserve":
            layout_profiles = ("preserve",)
        else:
            # A second export is performed only when the first artifact does not
            # satisfy the XLSX-specific retention audit. This keeps normal jobs fast
            # while making production output resilient to font substitution and
            # LibreOffice build differences.
            layout_profiles = ("balanced", "safe")

        last_fidelity_error: Optional[RuntimeError] = None
        for attempt_index, layout_profile in enumerate(layout_profiles, start=1):
            output_path.unlink(missing_ok=True)
            default_output = output_dir / f"{source_path.stem}.pdf"
            if default_output != output_path:
                default_output.unlink(missing_ok=True)

            with TemporaryDirectory(prefix="office-pdf-") as office_dir:
                office_root = Path(office_dir).resolve()
                staged_source = office_root / source_path.name
                shutil.copy2(source_path, staged_source)

                if input_format == "xlsx" and layout_profile != "preserve":
                    _prepare_xlsx_for_pdf(staged_source, profile=layout_profile)

                use_calc_uno = input_format == "xlsx"
                if use_calc_uno:
                    # XLSX is rendered through UNO so Calc itself recalculates live
                    # row heights with the fonts installed in this exact container.
                    # This removes the environment-dependent clipping that cannot be
                    # solved reliably with OOXML wrap/height hints alone.
                    self._convert_xlsx_to_pdf_with_uno(
                        staged_source,
                        output_path,
                        soffice_bin=soffice_bin,
                        layout_profile=layout_profile,
                    )
                else:
                    profile_dir = office_root / "profile"
                    profile_dir.mkdir(parents=True, exist_ok=True)
                    profile_uri = profile_dir.as_uri()
                    cmd = [
                        soffice_bin,
                        "--headless",
                        "--nologo",
                        "--nodefault",
                        "--nolockcheck",
                        "--nofirststartwizard",
                        f"-env:UserInstallation={profile_uri}",
                        "--convert-to",
                        filter_spec,
                        "--outdir",
                        str(output_dir),
                        str(staged_source),
                    ]
                    env = {
                        **os.environ,
                        "HOME": str(profile_dir),
                        "TMPDIR": str(office_root),
                        "SAL_USE_VCLPLUGIN": os.getenv("SAL_USE_VCLPLUGIN", "svp"),
                        "LANG": os.getenv("LANG", "C.UTF-8"),
                        "LC_ALL": os.getenv("LC_ALL", "C.UTF-8"),
                    }
                    try:
                        result = subprocess.run(
                            cmd,
                            check=False,
                            capture_output=True,
                            text=True,
                            timeout=DEFAULT_OFFICE_TO_PDF_TIMEOUT_SECONDS,
                            env=env,
                        )
                    except subprocess.TimeoutExpired as exc:
                        raise RuntimeError(
                            f"LibreOffice timed out while converting {input_format} to pdf."
                        ) from exc

            if not use_calc_uno:
                if result.returncode != 0:
                    detail = (result.stderr or result.stdout or "").strip()[-2000:]
                    suffix = f" Details: {detail}" if detail else ""
                    raise RuntimeError(
                        f"LibreOffice failed while converting {input_format} to pdf. "
                        f"Return code: {result.returncode}.{suffix}"
                    )

                default_output = output_dir / f"{source_path.stem}.pdf"
                if not default_output.exists() or default_output.stat().st_size <= 0:
                    raise RuntimeError(
                        f"LibreOffice completed without producing a PDF for the {input_format} source."
                    )
                if default_output != output_path:
                    default_output.replace(output_path)
            elif not output_path.exists() or output_path.stat().st_size <= 0:
                raise RuntimeError(
                    "LibreOffice Calc UNO completed without producing a PDF for the XLSX source."
                )

            if input_format != "xlsx":
                return
            if layout_profile == "preserve":
                # Explicit preserve mode retains the caller's spreadsheet print
                # semantics. It still passes the same fail-closed fidelity audit.
                _assert_xlsx_pdf_fidelity(source_path, output_path)
                return

            try:
                _assert_xlsx_pdf_fidelity(source_path, output_path)
                if attempt_index > 1:
                    logger.info(
                        "XLSX-to-PDF conversion passed fidelity validation after safe-layout retry."
                    )
                return
            except RuntimeError as exc:
                last_fidelity_error = exc
                if attempt_index < len(layout_profiles):
                    logger.warning(
                        "XLSX-to-PDF balanced layout did not satisfy fidelity validation; "
                        "retrying with safe wrapping and automatic row sizing. Audit: %s",
                        exc,
                    )
                    continue
                raise

        if last_fidelity_error is not None:  # pragma: no cover - defensive guard.
            raise last_fidelity_error

    def _convert_xlsx_to_pdf_with_uno(self, source_path: Path, output_path: Path,
                                     *, soffice_bin: str, layout_profile: str) -> None:
        engine = os.getenv("REDOCX_XLSX_PDF_ENGINE", "auto").strip().lower()
        if engine not in {"auto", "uno", "cli"}:
            raise ValueError("REDOCX_XLSX_PDF_ENGINE must be auto, uno, or cli.")
        if engine != "cli":
            try:
                self._convert_xlsx_to_pdf_using_uno(source_path, output_path,
                    soffice_bin=soffice_bin, layout_profile=layout_profile)
                return
            except RuntimeError:
                if engine == "uno":
                    raise
                logger.warning("Calc UNO export was unavailable or incomplete; retrying with audited CLI export.")
                output_path.unlink(missing_ok=True)
        self._convert_xlsx_to_pdf_using_cli(source_path, output_path, soffice_bin=soffice_bin)

    def _convert_xlsx_to_pdf_using_cli(self, source_path: Path, output_path: Path, *, soffice_bin: str) -> None:
        """Portable Calc export, audited against Calc's own formatted CSV values.

        CSV token 9 preserves displayed values; token 12 exports all sheets. These
        are audit files only. No CSV is returned to users or used as a PDF source.
        """
        if load_workbook is None:
            raise RuntimeError("openpyxl is required for spreadsheet display validation.")
        with TemporaryDirectory(prefix="calc-cli-") as work:
            root = Path(work)
            profile = root / "profile"
            profile.mkdir()
            # The isolated profile trusts no document macro sources.
            user = profile / "user"
            user.mkdir()
            (user/"registrymodifications.xcu").write_text(
                '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
                '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
                '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
                '</item></oor:items>', encoding="utf-8")
            common = [soffice_bin, "--headless", "--nologo", "--nodefault", "--nolockcheck",
                      "--nofirststartwizard", f"-env:UserInstallation={profile.as_uri()}"]
            environment = {**os.environ, "HOME": str(profile), "TMPDIR": str(root),
                           "SAL_USE_VCLPLUGIN": "svp", "LANG": "C.UTF-8", "LC_ALL": "C.UTF-8"}
            pdf_options = {"UseLosslessCompression": {"type": "boolean", "value": "true"},
                           "ReduceImageResolution": {"type": "boolean", "value": "false"}}
            for filter_spec in ["pdf:calc_pdf_Export:"+json.dumps(pdf_options,separators=(",", ":")),
                                "csv:Text - txt - csv (StarCalc):44,34,76,1,,0,false,true,true,false,false,-1"]:
                try:
                    result = subprocess.run(common+["--convert-to",filter_spec,"--outdir",str(root),str(source_path)],
                        stdin=subprocess.DEVNULL,capture_output=True,text=True,env=environment,
                        timeout=DEFAULT_OFFICE_TO_PDF_TIMEOUT_SECONDS)
                except subprocess.TimeoutExpired as exc:
                    raise RuntimeError("Calc CLI export timed out.") from exc
                if result.returncode != 0:
                    raise RuntimeError("Calc CLI could not complete the spreadsheet export.")
            rendered = root/(source_path.stem+".pdf")
            if not rendered.is_file() or not rendered.stat().st_size:
                raise RuntimeError("Calc CLI did not produce a PDF.")
            book = load_workbook(source_path, read_only=False, data_only=False, keep_links=False)
            displayed = []
            limit = max(1, int(os.getenv("REDOCX_XLSX_DISPLAY_CELL_LIMIT", "200000")))
            count = 0
            try:
                for sheet in book.worksheets:
                    if sheet.sheet_state != "visible":
                        continue
                    audit = root/(source_path.stem+"-"+sheet.title+".csv")
                    if audit.resolve().parent != root.resolve() or not audit.is_file():
                        raise RuntimeError("Calc did not produce every visible sheet's display audit.")
                    hidden_columns = set()
                    for dim in sheet.column_dimensions.values():
                        if dim.hidden:
                            hidden_columns.update(range(dim.min or 1,(dim.max or dim.min or 1)+1))
                    with audit.open(encoding="utf-8-sig",newline="") as handle:
                        for row_index,row in enumerate(csv.reader(handle),1):
                            if sheet.row_dimensions[row_index].hidden:
                                continue
                            for column_index,value in enumerate(row,1):
                                if not value or column_index in hidden_columns:
                                    continue
                                count += 1
                                if count > limit:
                                    raise RuntimeError("Workbook exceeds the configured display audit cell limit.")
                                displayed.append(value)
            finally:
                book.close()
            quality.assert_text("\n".join(displayed), _pdf_native_text(rendered), minimum=.995,
                                label="Excel-to-PDF displayed cells")
            shutil.copy2(rendered,output_path)

    def _convert_xlsx_to_pdf_using_uno(
        self,
        source_path: Path,
        output_path: Path,
        *,
        soffice_bin: str,
        layout_profile: str,
    ) -> None:
        """Render XLSX through Calc's live UNO layout engine.

        XML-only row-height hints are not authoritative in headless Calc. Different
        LibreOffice/font environments may retain default row heights even when
        wrapText is present, which clips long cell values. This worker loads the
        temporary workbook in the exact Calc process that exports the PDF, applies
        live row OptimalHeight/page-scaling properties, and exports without closing
        and reopening the document.

        The UNO bridge is intentionally executed with Debian's system Python because
        python3-uno is ABI-coupled to the LibreOffice packages. The application
        virtualenv must not attempt to pip-install or import pyuno directly.
        """
        uno_python = Path(DEFAULT_LIBREOFFICE_UNO_PYTHON).expanduser()
        worker = LIBREOFFICE_CALC_PDF_WORKER.resolve()
        if not uno_python.is_file():
            raise RuntimeError(
                "LibreOffice UNO Python runtime is unavailable. Install python3-uno "
                "and set REDOCX_LIBREOFFICE_UNO_PYTHON to the system Python executable."
            )
        if not worker.is_file():
            raise RuntimeError(
                "ReDOCX LibreOffice Calc PDF worker is missing from the deployment image."
            )

        command = [
            str(uno_python),
            str(worker),
            "--audit-manifest",
            str(output_path.with_suffix(".display.json")),
            "--source",
            str(source_path),
            "--output",
            str(output_path),
            "--soffice",
            str(soffice_bin),
            "--profile",
            layout_profile if layout_profile in {"balanced", "safe", "preserve"} else "balanced",
            "--timeout",
            str(DEFAULT_OFFICE_TO_PDF_TIMEOUT_SECONDS),
        ]
        process = subprocess.Popen(
            command,
            stdin=subprocess.DEVNULL,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            start_new_session=(os.name != "nt"),
            env={
                **os.environ,
                "LANG": os.getenv("LANG", "C.UTF-8"),
                "LC_ALL": os.getenv("LC_ALL", "C.UTF-8"),
                "SAL_USE_VCLPLUGIN": os.getenv("SAL_USE_VCLPLUGIN", "svp"),
            },
        )
        try:
            stdout, stderr = process.communicate(
                timeout=DEFAULT_OFFICE_TO_PDF_TIMEOUT_SECONDS + 30
            )
        except subprocess.TimeoutExpired as exc:
            if os.name != "nt":
                try:
                    os.killpg(process.pid, signal.SIGTERM)
                    process.wait(timeout=5)
                except (ProcessLookupError, subprocess.TimeoutExpired):
                    try:
                        os.killpg(process.pid, signal.SIGKILL)
                    except ProcessLookupError:
                        pass
            else:  # pragma: no cover - production image is Linux.
                process.kill()
            output_path.unlink(missing_ok=True)
            raise RuntimeError(
                "LibreOffice Calc UNO layout/export timed out while converting XLSX to PDF."
            ) from exc

        if process.returncode != 0:
            output_path.unlink(missing_ok=True)
            detail = (stderr or stdout or "").strip()[-2000:]
            # Temporary customer paths are not useful diagnostics outside this
            # process. Redact them before the message can reach API/log surfaces.
            for sensitive in (str(source_path), str(output_path), str(worker)):
                if sensitive:
                    detail = detail.replace(sensitive, "<runtime-path>")
            raise RuntimeError(
                "LibreOffice Calc UNO layout/export failed for XLSX-to-PDF."
                + (f" Details: {detail}" if detail else "")
            )
        if not output_path.is_file() or output_path.stat().st_size <= 0:
            raise RuntimeError(
                "LibreOffice Calc UNO completed without producing a non-empty PDF."
            )

        manifest = output_path.with_suffix(".display.json")
        try:
            displayed = json.loads(manifest.read_text(encoding="utf-8"))
            if not isinstance(displayed, list) or not all(isinstance(x, str) for x in displayed):
                raise RuntimeError("Calc produced an invalid display audit manifest.")
            quality.assert_text("\n".join(displayed), _pdf_native_text(output_path), minimum=.995,
                                label="Excel-to-PDF displayed cells")
        finally:
            manifest.unlink(missing_ok=True)

    def _convert_html_to_pdf(self, source_path: Path, output_path: Path) -> None:
        from .html_render import convert
        convert(source_path,output_path)

    def _convert_pdf_to_jpg(self, source_path: Path, output_path: Path) -> str:
        scale = max(1.0, float(DEFAULT_PDF_RASTER_DPI) / 72.0)
        matrix = fitz.Matrix(scale, scale)
        with fitz.open(source_path) as pdf:
            if pdf.page_count < 1:
                raise ValueError("PDF has no pages to convert.")
            if pdf.page_count == 1:
                pix = pdf.load_page(0).get_pixmap(matrix=matrix, alpha=False)
                image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                image.save(
                    output_path,
                    "JPEG",
                    quality=DEFAULT_IMAGE_JPEG_QUALITY,
                    optimize=True,
                    dpi=(DEFAULT_PDF_RASTER_DPI, DEFAULT_PDF_RASTER_DPI),
                )
                return "jpg"

            zip_path = output_path.with_suffix(".zip")
            with TemporaryDirectory(prefix="pdf-jpg-pages-") as pages_dir:
                page_paths: list[Path] = []
                digits = max(4, len(str(pdf.page_count)))
                for page_index in range(pdf.page_count):
                    pix = pdf.load_page(page_index).get_pixmap(matrix=matrix, alpha=False)
                    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
                    page_path = Path(pages_dir) / f"page-{page_index + 1:0{digits}d}.jpg"
                    image.save(
                        page_path,
                        "JPEG",
                        quality=DEFAULT_IMAGE_JPEG_QUALITY,
                        optimize=True,
                        dpi=(DEFAULT_PDF_RASTER_DPI, DEFAULT_PDF_RASTER_DPI),
                    )
                    page_paths.append(page_path)

                with zipfile.ZipFile(zip_path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
                    for page_path in page_paths:
                        archive.write(page_path, arcname=page_path.name)
            return "zip"

    def _convert_pdf_to_pptx(self, source_path: Path, output_path: Path) -> None:
        from .page_geometry import slide_reference
        with slide_reference(source_path) as reference:
            native.convert_pptx(reference, output_path)
            _validate_pptx_file(output_path)
            quality.assert_text(_pdf_native_text(source_path), _pptx_visible_slide_text(output_path),
                                minimum=.98, label="PDF-to-PowerPoint")
            with TemporaryDirectory(prefix="pptx-visual-") as work:
                rendered=Path(work)/"rendered.pdf"
                self._convert_office_to_pdf(output_path,rendered,"pptx")
                _assert_pdf_visual_layout_fidelity(reference,rendered,minimum_score=.80,
                                                   label="PDF-to-PowerPoint")

    def _convert_pdf_to_xlsx(self, source_path: Path, output_path: Path) -> None:
        native.convert_xlsx(source_path, output_path)
        native.validate_xlsx(source_path, output_path)

    def _convert_pdf_to_pdfa(self, source_path: Path, output_path: Path) -> None:
        gs_bin = (
            os.getenv("GHOSTSCRIPT_PATH")
            or shutil.which("gs")
            or shutil.which("gswin64c")
            or shutil.which("gswin32c")
        )
        if not gs_bin:
            raise RuntimeError(
                "Ghostscript is required for PDF/A conversion. Set GHOSTSCRIPT_PATH or add Ghostscript to PATH."
            )

        icc_profile = _resolve_pdfa_icc_profile()
        pdfa_def = _resolve_pdfa_definition_file(gs_bin)
        with TemporaryDirectory(prefix="pdfa-config-") as config_dir:
            local_def = Path(config_dir) / "PDFA_def.ps"
            definition = pdfa_def.read_text(encoding="latin-1")
            escaped_icc = str(icc_profile).replace("\\", "/").replace("(", "\\(").replace(")", "\\)")
            definition = re.sub(
                r"/ICCProfile\s*\([^\r\n]*\)",
                f"/ICCProfile ({escaped_icc})",
                definition,
                count=1,
            )
            local_def.write_text(definition, encoding="latin-1")

            cmd = [
                gs_bin,
                "-dPDFA=2",
                "-dBATCH",
                "-dNOPAUSE",
                "-dSAFER",
                f"--permit-file-read={icc_profile}",
                f"--permit-file-read={local_def}",
                f"--permit-file-read={source_path}",
                "-sDEVICE=pdfwrite",
                "-sColorConversionStrategy=RGB",
                "-dAutoRotatePages=/None",
                "-dPreserveAnnots=true",
                "-dEmbedAllFonts=true",
                "-dSubsetFonts=true",
                "-dDownsampleColorImages=false",
                "-dDownsampleGrayImages=false",
                "-dDownsampleMonoImages=false",
                "-dPDFACompatibilityPolicy=1",
                f"-sOutputFile={output_path}",
                str(local_def),
                str(source_path),
            ]
            try:
                result = subprocess.run(
                    cmd,
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=DEFAULT_PDFA_TIMEOUT_SECONDS,
                )
            except subprocess.TimeoutExpired as exc:
                raise RuntimeError("PDF/A conversion timed out.") from exc

        if result.returncode != 0:
            output_path.unlink(missing_ok=True)
            raise RuntimeError(
                f"Ghostscript failed while creating PDF/A-2b. Return code: {result.returncode}."
            )

        quality.restore_pdf_links(source_path, output_path)
        quality.assert_pdfa_content(source_path, output_path)

    def _convert_image_to_pdf(self, source_path: Path, output_path: Path) -> None:
        with Image.open(source_path) as image:
            orientation = image.getexif().get(274, 1)
            original_format = image.format
            image = ImageOps.exif_transpose(image)
            if orientation == 1 and original_format == "JPEG":
                data = source_path.read_bytes()
            else:
                stream = BytesIO()
                quality.rgb(image).save(stream, "PNG")
                data = stream.getvalue()
            pdf = fitz.open()
            try:
                page = pdf.new_page(width=image.width*72/DEFAULT_IMAGE_PDF_DPI,
                                    height=image.height*72/DEFAULT_IMAGE_PDF_DPI)
                page.insert_image(page.rect, stream=data)
                pdf.save(output_path)
            finally:
                pdf.close()

    def _convert_image_to_docx(self, source_path: Path, output_path: Path) -> None:
        mode = os.getenv("REDOCX_IMAGE_TO_DOCX_MODE", "auto").strip().lower()
        if mode not in {"auto", "editable", "image"}:
            raise ValueError("REDOCX_IMAGE_TO_DOCX_MODE must be auto, editable, or image.")
        if mode != "image":
            with TemporaryDirectory(prefix="image-ocr-") as work:
                editable = Path(work) / "editable.pdf"
                if native.image_to_editable_pdf(source_path, editable, dpi=DEFAULT_IMAGE_PDF_DPI):
                    self._convert_pdf_to_docx(editable, output_path)
                    rendered = Path(work) / "rendered.pdf"
                    self._convert_office_to_pdf(output_path, rendered, "docx")
                    quality.assert_image_pdf(source_path, rendered)
                    native.validate_image_word_text(source_path, rendered)
                    return
                if mode == "editable":
                    raise RuntimeError("Image has insufficient readable text for editable Word output.")
        source_path = Path(source_path)
        output_path = Path(output_path)

        if not source_path.exists():
            raise FileNotFoundError(f"Source file not found: {source_path}")
        if source_path.stat().st_size == 0:
            raise ValueError(f"Source image is empty: {source_path}")

        try:
            with Image.open(source_path) as image:
                image = ImageOps.exif_transpose(image)
                rgb = _flatten_image_to_rgb(image)
                with NamedTemporaryFile(suffix=".jpg", delete=False) as tmp:
                    normalized_path = Path(tmp.name)
                try:
                    rgb.save(normalized_path, format="JPEG", quality=DEFAULT_IMAGE_JPEG_QUALITY, optimize=True)
                    document = Document()
                    section = document.sections[0]
                    _configure_section_for_image(section, rgb.width, rgb.height)
                    paragraph = document.paragraphs[0] if document.paragraphs else document.add_paragraph()
                    _configure_full_page_image_paragraph(paragraph)
                    picture_width, picture_height = _fit_image_to_section(
                        section,
                        width_px=rgb.width,
                        height_px=rgb.height,
                    )
                    vertical_gap = max(
                        0.0,
                        (float(section.page_height.pt) - float(picture_height.pt)) / 2.0,
                    )
                    paragraph.paragraph_format.space_before = Pt(vertical_gap)
                    paragraph.add_run().add_picture(
                        str(normalized_path),
                        width=picture_width,
                        height=picture_height,
                    )
                    document.save(output_path)
                finally:
                    normalized_path.unlink(missing_ok=True)
        except UnidentifiedImageError as e:
            with open(source_path, "rb") as f:
                header = f.read(32)
            raise ValueError(
                f"Uploaded file is not a readable image. path={source_path}, header={header!r}"
            ) from e

    def _convert_png_to_jpeg_family(self, source_path: Path, output_path: Path) -> None:
        try:
            with Image.open(source_path) as image:
                image = ImageOps.exif_transpose(image)
                rgb = _flatten_image_to_rgb(image)
                rgb.save(output_path, "JPEG", quality=DEFAULT_IMAGE_JPEG_QUALITY, optimize=True)
        except UnidentifiedImageError as exc:
            raise ValueError(f"Uploaded file is not a readable PNG image: {source_path}") from exc


@dataclass(frozen=True)
class ConvertConfig:
    """Optional knobs for future provider-backed conversion."""

    algorithm_version: Optional[str] = None


class ConvertProcessor:
    """
    Stateless document-conversion processor.

    Responsibilities:
    - validate local conversion preconditions
    - enforce contract-allowed conversion pairs
    - plan deterministic output naming
    - delegate actual conversion to a backend
    - return schema-agnostic artifact metadata only

    Non-responsibilities:
    - request validation
    - response/result model construction
    - analyzer response language-field orchestration
    - document text extraction
    """

    def __init__(
        self,
        backend: Optional[ConversionBackend] = None,
        config: Optional[ConvertConfig] = None,
    ) -> None:
        self.backend = backend or RealConversionBackend()
        self.config = config or ConvertConfig()

    def convert(
        self,
        *,
        input_format: str,
        output_format: str,
        source_reference: str,
        source_name_hint: Optional[str] = None,
    ) -> ConversionArtifact:
        normalized_input = _normalize_format(input_format)
        normalized_output = _normalize_format(output_format)
        normalized_reference = _normalize_source_reference(source_reference)

        ensure_allowed_conversion_pair(
            input_format=normalized_input,
            output_format=normalized_output,
        )

        planned_name = plan_output_file_name(
            input_format=normalized_input,
            output_format=normalized_output,
            source_name_hint=source_name_hint,
        )

        return self.backend.convert(
            source_reference=normalized_reference,
            input_format=normalized_input,
            output_format=normalized_output,
            planned_output_name=planned_name,
        )


def ensure_allowed_conversion_pair(*, input_format: str, output_format: str) -> None:
    normalized_input = _normalize_format(input_format)
    normalized_output = _normalize_format(output_format)

    allowed_targets = ALLOWED_CONVERSION_PAIRS.get(normalized_input, set())
    if normalized_output not in allowed_targets:
        allowed = ", ".join(sorted(allowed_targets)) if allowed_targets else "(none)"
        raise ValueError(
            f"Unsupported conversion pair: {normalized_input} -> {normalized_output}. "
            f"Allowed outputs for {normalized_input}: {allowed}."
        )


def plan_output_file_name(
    *,
    input_format: str,
    output_format: str,
    source_name_hint: Optional[str] = None,
) -> str:
    normalized_input = _normalize_format(input_format)
    normalized_output = _normalize_format(output_format)
    ensure_allowed_conversion_pair(
        input_format=normalized_input,
        output_format=normalized_output,
    )

    base_source = source_name_hint or normalized_input
    base = _safe_basename(Path(str(base_source)).stem)
    physical_extension = PHYSICAL_OUTPUT_FORMAT_BY_TARGET[normalized_output]
    return f"{base}.converted.{physical_extension}"


def build_conversion_instructions(*, input_format: str, output_format: str) -> str:
    normalized_input = _normalize_format(input_format)
    normalized_output = _normalize_format(output_format)
    ensure_allowed_conversion_pair(
        input_format=normalized_input,
        output_format=normalized_output,
    )

    return (
        f"{CONVERSION_RULES}\n\n"
        f"SOURCE FORMAT:\n{normalized_input}\n\n"
        f"TARGET FORMAT:\n{normalized_output}"
    )


def convert_document(
    *,
    input_format: str,
    output_format: str,
    source_reference: str,
    source_name_hint: Optional[str] = None,
    backend: Optional[ConversionBackend] = None,
    config: Optional[ConvertConfig] = None,
) -> ConversionArtifact:
    """Functional convenience wrapper for analyzer integration."""
    processor = ConvertProcessor(backend=backend, config=config)
    return processor.convert(
        input_format=input_format,
        output_format=output_format,
        source_reference=source_reference,
        source_name_hint=source_name_hint,
    )


def _normalize_format(value: str) -> str:
    if not isinstance(value, str):
        raise TypeError("format value must be a string.")
    normalized = value.strip().lower()
    if normalized not in OUTPUT_EXTENSION_ALIASES:
        raise ValueError(
            "format must be one of: pdf, pdfa, docx, jpg, jpeg, png, xlsx, html, htm, pptx, zip."
        )
    return OUTPUT_EXTENSION_ALIASES[normalized]


def _normalize_source_reference(value: str) -> str:
    if not isinstance(value, str):
        raise TypeError("source_reference must be a string.")
    normalized = value.strip()
    if not normalized:
        raise ValueError("source_reference cannot be empty.")
    return normalized


def _normalize_file_name(value: str) -> str:
    if not isinstance(value, str):
        raise TypeError("planned_output_name must be a string.")
    normalized = value.strip()
    if not normalized:
        raise ValueError("planned_output_name cannot be empty.")
    return normalized


def _safe_basename(value: str) -> str:
    raw = str(value).strip().lower()
    if not raw:
        return "document"
    filtered = "".join(ch if ch.isalnum() or ch in {"-", "_"} else "-" for ch in raw)
    compact = "-".join(part for part in filtered.split("-") if part)
    return compact or "document"


def _get_file_size_mb(path: Path) -> float:
    return round(path.stat().st_size / (1024 * 1024), 4)


SPREADSHEETML_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_XLSX_CELL_REF_RE = re.compile(r"^([A-Z]+)([1-9][0-9]*)$")


def _xlsx_q(tag: str) -> str:
    return f"{{{SPREADSHEETML_NS}}}{tag}"


def _xlsx_column_index(cell_reference: str) -> Optional[int]:
    match = _XLSX_CELL_REF_RE.match(str(cell_reference or "").upper())
    if not match:
        return None
    value = 0
    for character in match.group(1):
        value = value * 26 + (ord(character) - 64)
    return value


def _xlsx_shared_strings(package: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in package.namelist():
        return []
    root = ElementTree.fromstring(package.read("xl/sharedStrings.xml"))
    return [
        "".join(node.text or "" for node in item.iter(_xlsx_q("t")))
        for item in root.findall(_xlsx_q("si"))
    ]


def _xlsx_cell_text(cell: Any, shared_strings: Sequence[str]) -> str:
    cell_type = str(cell.attrib.get("t") or "")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.iter(_xlsx_q("t")))

    value = cell.find(_xlsx_q("v"))
    if value is None:
        return ""
    raw = value.text or ""
    if cell_type == "s":
        try:
            return shared_strings[int(raw)]
        except (ValueError, IndexError):
            return ""
    if cell_type == "b":
        return "TRUE" if raw == "1" else "FALSE"
    return raw


def _xlsx_effective_column_attributes(columns: Sequence[Any], index: int) -> dict[str, str]:
    attributes: dict[str, str] = {}
    for column in columns:
        try:
            minimum = int(column.attrib.get("min", "0"))
            maximum = int(column.attrib.get("max", "0"))
        except ValueError:
            continue
        if minimum <= index <= maximum:
            attributes = dict(column.attrib)
    return attributes


def _xlsx_add_wrapped_styles(
    styles_xml: bytes,
    style_ids: set[int],
) -> tuple[bytes, dict[int, int]]:
    if not style_ids:
        return styles_xml, {}

    ElementTree.register_namespace("", SPREADSHEETML_NS)
    root = ElementTree.fromstring(styles_xml)
    cell_xfs = root.find(_xlsx_q("cellXfs"))
    if cell_xfs is None:
        raise RuntimeError("XLSX styles.xml is missing cellXfs.")

    original_xfs = list(cell_xfs)
    if not original_xfs:
        raise RuntimeError("XLSX styles.xml contains no cell formats.")

    mapping: dict[int, int] = {}
    for requested_style_id in sorted(style_ids):
        style_id = requested_style_id if 0 <= requested_style_id < len(original_xfs) else 0
        copied = deepcopy(original_xfs[style_id])
        alignment = copied.find(_xlsx_q("alignment"))
        if alignment is None:
            alignment = ElementTree.SubElement(copied, _xlsx_q("alignment"))
        alignment.set("wrapText", "1")
        if "vertical" not in alignment.attrib:
            alignment.set("vertical", "top")
        copied.set("applyAlignment", "1")
        cell_xfs.append(copied)
        mapping[requested_style_id] = len(cell_xfs) - 1

    cell_xfs.set("count", str(len(cell_xfs)))
    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True), mapping


def _xlsx_insert_page_setup(root: Any) -> Any:
    existing = root.find(_xlsx_q("pageSetup"))
    if existing is not None:
        return existing

    page_setup = ElementTree.Element(_xlsx_q("pageSetup"))
    trailing_tags = {
        _xlsx_q("headerFooter"),
        _xlsx_q("rowBreaks"),
        _xlsx_q("colBreaks"),
        _xlsx_q("customProperties"),
        _xlsx_q("cellWatches"),
        _xlsx_q("ignoredErrors"),
        _xlsx_q("smartTags"),
        _xlsx_q("drawing"),
        _xlsx_q("legacyDrawing"),
        _xlsx_q("legacyDrawingHF"),
        _xlsx_q("picture"),
        _xlsx_q("oleObjects"),
        _xlsx_q("controls"),
        _xlsx_q("webPublishItems"),
        _xlsx_q("tableParts"),
        _xlsx_q("extLst"),
    }
    children = list(root)
    for position, child in enumerate(children):
        if child.tag in trailing_tags:
            root.insert(position, page_setup)
            return page_setup
    root.append(page_setup)
    return page_setup


def _xlsx_patch_sheet_for_readable_pdf(
    sheet_xml: bytes,
    shared_strings: Sequence[str],
    *,
    style_mapping: Optional[dict[int, int]] = None,
    scan_only: bool = False,
    profile: str = "balanced",
) -> tuple[bytes, set[int]]:
    normalized_profile = str(profile or "balanced").strip().lower()
    safe_profile = normalized_profile == "safe"

    ElementTree.register_namespace("", SPREADSHEETML_NS)
    root = ElementTree.fromstring(sheet_xml)
    sheet_data = root.find(_xlsx_q("sheetData"))
    if sheet_data is None:
        return sheet_xml, set()

    columns_element = root.find(_xlsx_q("cols"))
    original_columns = list(columns_element) if columns_element is not None else []

    cells: list[tuple[Any, Any, int, str, int]] = []
    lengths_by_column: dict[int, list[int]] = {}
    maximum_column = 0

    for row in sheet_data.findall(_xlsx_q("row")):
        if str(row.attrib.get("hidden", "0")).lower() in {"1", "true"}:
            continue
        for cell in row.findall(_xlsx_q("c")):
            column_index = _xlsx_column_index(cell.attrib.get("r", ""))
            if not column_index:
                continue
            maximum_column = max(maximum_column, column_index)
            text = _xlsx_cell_text(cell, shared_strings)
            if not text:
                continue
            longest_line = max((len(part) for part in text.splitlines()), default=0)
            lengths_by_column.setdefault(column_index, []).append(longest_line)
            cells.append((row, cell, column_index, text, longest_line))

    if maximum_column < 1:
        return sheet_xml, set()

    hidden: dict[int, bool] = {}
    current_widths: dict[int, float] = {}
    desired_widths: dict[int, float] = {}

    for index in range(1, maximum_column + 1):
        attributes = _xlsx_effective_column_attributes(original_columns, index)
        hidden[index] = str(attributes.get("hidden", "0")).lower() in {"1", "true"}
        try:
            current_widths[index] = float(attributes.get("width", "8.43"))
        except ValueError:
            current_widths[index] = 8.43

        values = lengths_by_column.get(index, [])
        if hidden[index]:
            desired_widths[index] = current_widths[index]
            continue
        if not values:
            desired_widths[index] = max(
                DEFAULT_XLSX_MIN_COLUMN_WIDTH,
                min(current_widths[index], DEFAULT_XLSX_MAX_COLUMN_WIDTH),
            )
            continue

        sorted_values = sorted(values)
        percentile = 1.0 if safe_profile else 0.95
        percentile_index = max(0, math.ceil(len(sorted_values) * percentile) - 1)
        representative_length = sorted_values[percentile_index]
        max_column_width = max(DEFAULT_XLSX_MAX_COLUMN_WIDTH, 48.0) if safe_profile else DEFAULT_XLSX_MAX_COLUMN_WIDTH
        width_multiplier = 1.18 if safe_profile else 1.10
        width_padding = 3.0 if safe_profile else 2.0
        target = min(
            max_column_width,
            max(DEFAULT_XLSX_MIN_COLUMN_WIDTH, representative_length * width_multiplier + width_padding),
        )
        desired_widths[index] = max(current_widths[index], target)

    visible_columns = [index for index in range(1, maximum_column + 1) if not hidden[index]]
    desired_total = sum(desired_widths[index] for index in visible_columns)
    total_column_width = (
        max(DEFAULT_XLSX_TOTAL_COLUMN_WIDTH, 128.0)
        if safe_profile
        else DEFAULT_XLSX_TOTAL_COLUMN_WIDTH
    )
    if desired_total > total_column_width and visible_columns:
        floor_total = DEFAULT_XLSX_MIN_COLUMN_WIDTH * len(visible_columns)
        distributable = max(0.0, total_column_width - floor_total)
        wants = {
            index: max(0.0, desired_widths[index] - DEFAULT_XLSX_MIN_COLUMN_WIDTH)
            for index in visible_columns
        }
        want_total = sum(wants.values())
        for index in visible_columns:
            desired_widths[index] = DEFAULT_XLSX_MIN_COLUMN_WIDTH + (
                distributable * wants[index] / want_total if want_total else 0.0
            )

    wrap_candidates: list[tuple[Any, Any, int]] = []
    requested_style_ids: set[int] = set()
    for row, cell, column_index, text, longest_line in cells:
        if hidden[column_index]:
            continue
        if safe_profile or "\n" in text or longest_line > desired_widths[column_index] * 0.92:
            try:
                style_id = int(cell.attrib.get("s", "0") or "0")
            except ValueError:
                style_id = 0
            requested_style_ids.add(style_id)
            wrap_candidates.append((row, cell, style_id))

    if scan_only:
        return sheet_xml, requested_style_ids

    style_mapping = style_mapping or {}
    for row, cell, style_id in wrap_candidates:
        cell.set("s", str(style_mapping.get(style_id, style_id)))
        # Explicit fixed row heights are a frequent second source of clipping once
        # wrapping is enabled. On the temporary conversion copy, let Calc auto-size.
        row.attrib.pop("ht", None)
        row.attrib.pop("customHeight", None)

    if safe_profile:
        # Font substitution can change glyph widths enough that a row which looked
        # safe under development fonts clips in a production container. Release the
        # height of every populated visible row in the fallback profile so Calc can
        # recompute it using the fonts that actually exist in that runtime.
        rows_with_visible_content = {id(row): row for row, _, column_index, _, _ in cells if not hidden[column_index]}
        for row in rows_with_visible_content.values():
            row.attrib.pop("ht", None)
            row.attrib.pop("customHeight", None)

    if columns_element is None:
        columns_element = ElementTree.Element(_xlsx_q("cols"))
        children = list(root)
        root.insert(children.index(sheet_data), columns_element)
    else:
        for child in list(columns_element):
            columns_element.remove(child)

    for index in range(1, maximum_column + 1):
        attributes = _xlsx_effective_column_attributes(original_columns, index)
        attributes["min"] = str(index)
        attributes["max"] = str(index)
        if not hidden[index]:
            attributes["width"] = f"{desired_widths[index]:.3f}".rstrip("0").rstrip(".")
            attributes["customWidth"] = "1"
            attributes.pop("bestFit", None)
        ElementTree.SubElement(columns_element, _xlsx_q("col"), attributes)

    # Preserve any explicit column definitions that start beyond the populated area.
    for column in original_columns:
        try:
            if int(column.attrib.get("min", "0")) > maximum_column:
                columns_element.append(deepcopy(column))
        except ValueError:
            continue

    sheet_properties = root.find(_xlsx_q("sheetPr"))
    if sheet_properties is None:
        sheet_properties = ElementTree.Element(_xlsx_q("sheetPr"))
        root.insert(0, sheet_properties)
    page_setup_properties = sheet_properties.find(_xlsx_q("pageSetUpPr"))
    if page_setup_properties is None:
        page_setup_properties = ElementTree.SubElement(sheet_properties, _xlsx_q("pageSetUpPr"))
    page_setup_properties.set("fitToPage", "1")

    page_setup = _xlsx_insert_page_setup(root)
    page_setup.set("fitToWidth", "1")
    page_setup.set("fitToHeight", "0")
    page_setup.attrib.pop("scale", None)
    if "orientation" not in page_setup.attrib:
        page_setup.set(
            "orientation",
            "landscape"
            if sum(desired_widths[index] for index in visible_columns) >= DEFAULT_XLSX_LANDSCAPE_THRESHOLD
            else "portrait",
        )
    if "paperSize" not in page_setup.attrib:
        paper_sizes = {"LETTER": "1", "LEGAL": "5", "A4": "9"}
        page_setup.set("paperSize", paper_sizes.get(DEFAULT_XLSX_PAPER_SIZE, "9"))

    return ElementTree.tostring(root, encoding="utf-8", xml_declaration=True), requested_style_ids


def _prepare_xlsx_for_pdf(path: Path, *, profile: str = "balanced") -> None:
    """Normalize a temporary XLSX copy for readable, complete PDF output.

    The OOXML package is rewritten in-place without loading/saving it through an
    office library, so charts, drawings, formulas, relationships, macros/extensions
    and unknown package parts remain byte-for-byte untouched unless they are one of
    the worksheet/style parts required for print readability.
    """
    path = Path(path).resolve()
    if DEFAULT_XLSX_LAYOUT_MODE == "preserve":
        return
    if not zipfile.is_zipfile(path):
        raise ValueError(f"File content is not an OOXML XLSX package: {path}")

    temp_path = path.with_name(f".{path.name}.redocx-normalizing")
    temp_path.unlink(missing_ok=True)
    try:
        with zipfile.ZipFile(path, "r") as source_package:
            names = set(source_package.namelist())
            sheet_paths = sorted(
                name
                for name in names
                if name.startswith("xl/worksheets/") and name.endswith(".xml")
            )
            if not sheet_paths or "xl/styles.xml" not in names:
                raise ValueError("XLSX package is missing worksheet/style parts.")

            shared_strings = _xlsx_shared_strings(source_package)
            requested_style_ids: set[int] = set()
            for sheet_path in sheet_paths:
                _, style_ids = _xlsx_patch_sheet_for_readable_pdf(
                    source_package.read(sheet_path),
                    shared_strings,
                    scan_only=True,
                    profile=profile,
                )
                requested_style_ids.update(style_ids)

            styles_xml, style_mapping = _xlsx_add_wrapped_styles(
                source_package.read("xl/styles.xml"),
                requested_style_ids,
            )

            with zipfile.ZipFile(temp_path, "w") as target_package:
                for info in source_package.infolist():
                    payload = source_package.read(info.filename)
                    if info.filename == "xl/styles.xml":
                        payload = styles_xml
                    elif info.filename in sheet_paths:
                        payload, _ = _xlsx_patch_sheet_for_readable_pdf(
                            payload,
                            shared_strings,
                            style_mapping=style_mapping,
                            profile=profile,
                        )
                    target_package.writestr(info, payload)
        temp_path.replace(path)
    except Exception:
        temp_path.unlink(missing_ok=True)
        raise


def _pdf_to_docx_mode() -> str:
    mode = DEFAULT_PDF_TO_DOCX_MODE
    if mode not in PDF_TO_DOCX_MODES:
        raise ValueError("Invalid REDOCX_PDF_TO_DOCX_MODE.")
    return mode

def _normalized_tokens(value: str) -> list[str]:
    return quality.tokens(value)

def _token_coverage(source_text: str, output_text: str) -> tuple[float, int, int]:
    source = Counter(_normalized_tokens(source_text))
    if not source:
        return 1.0, 0, 0
    output = Counter(_normalized_tokens(output_text))
    matched = sum(min(count, output.get(token, 0)) for token, count in source.items())
    total = sum(source.values())
    return matched / total, matched, total


def _assert_text_token_coverage(*, source_text: str, output_text: str, minimum: float, label: str) -> None:
    quality.assert_text(source_text, output_text, minimum=minimum, label=label)

def _docx_package_text(path: Path) -> str:
    return quality.visible_word_text(path)

def _pptx_visible_slide_text(path: Path) -> str:
    drawing_text_tag = "{http://schemas.openxmlformats.org/drawingml/2006/main}t"
    chunks: list[str] = []
    with zipfile.ZipFile(path) as package:
        slide_paths = sorted(
            name
            for name in package.namelist()
            if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
        )
        for slide_path in slide_paths:
            try:
                root = ElementTree.fromstring(package.read(slide_path))
            except ElementTree.ParseError:
                continue
            if str(root.attrib.get("show", "1")).lower() in {"0", "false"}:
                continue
            chunks.extend(node.text or "" for node in root.iter(drawing_text_tag))
    return "\n".join(chunks)


def _xlsx_nonliteral_number_format(format_code: str) -> str:
    value = re.sub(r'"(?:[^"\\]|\\.)*"', "", str(format_code or ""))
    value = re.sub(r"\\.", "", value)
    return value.casefold()


def _xlsx_numeric_styles_to_skip_for_text_gate(package: zipfile.ZipFile) -> set[int]:
    """Return style IDs whose raw stored numeric is not comparable to PDF text.

    Dates/times and scientific formats transform the stored scalar substantially at
    display time. They are intentionally excluded from raw-token comparison rather
    than producing false conversion failures. Ordinary number/currency/accounting
    formats remain eligible and are normalized for common thousands separators.
    """

    if "xl/styles.xml" not in package.namelist():
        return set()
    try:
        root = ElementTree.fromstring(package.read("xl/styles.xml"))
    except ElementTree.ParseError:
        return set()

    custom_formats: dict[int, str] = {}
    num_fmts = root.find(_xlsx_q("numFmts"))
    if num_fmts is not None:
        for item in num_fmts.findall(_xlsx_q("numFmt")):
            try:
                num_fmt_id = int(item.attrib.get("numFmtId", ""))
            except ValueError:
                continue
            custom_formats[num_fmt_id] = item.attrib.get("formatCode", "")

    # ECMA/Excel built-in date/time styles plus built-in scientific formats.
    built_in_date_ids = set(range(14, 23)) | set(range(27, 37)) | set(range(45, 48)) | set(range(50, 59))
    built_in_scientific_ids = {11, 48}
    skipped: set[int] = set()
    cell_xfs = root.find(_xlsx_q("cellXfs"))
    if cell_xfs is None:
        return skipped

    for style_id, xf in enumerate(cell_xfs.findall(_xlsx_q("xf"))):
        try:
            num_fmt_id = int(xf.attrib.get("numFmtId", "0"))
        except ValueError:
            num_fmt_id = 0
        if num_fmt_id in built_in_date_ids or num_fmt_id in built_in_scientific_ids:
            skipped.add(style_id)
            continue
        custom = _xlsx_nonliteral_number_format(custom_formats.get(num_fmt_id, ""))
        if custom and (
            re.search(r"[ydhs]", custom)
            or ("m" in custom and any(marker in custom for marker in ("/", "-", ":")))
            or "e+" in custom
            or "e-" in custom
        ):
            skipped.add(style_id)
    return skipped


def _xlsx_visible_cell_values(path: Path) -> list[str]:
    """Return user-visible XLSX cell values suitable for fidelity auditing.

    Raw spreadsheet storage numbers whose display depends on number-format rules
    are intentionally excluded, matching the previous text gate. Textual values,
    booleans, string formula results, and numeric identifiers without display-only
    formatting remain auditable. Hidden rows and columns are excluded because they
    are not expected to render in a normal spreadsheet PDF export.
    """
    values: list[str] = []
    with zipfile.ZipFile(path) as package:
        shared_strings = _xlsx_shared_strings(package)
        numeric_styles_to_skip = _xlsx_numeric_styles_to_skip_for_text_gate(package)
        hidden_paths = set()
        rels = ElementTree.fromstring(package.read("xl/_rels/workbook.xml.rels"))
        targets = {r.get("Id"): r.get("Target", "") for r in rels}
        workbook = ElementTree.fromstring(package.read("xl/workbook.xml"))
        for item in workbook.iter(_xlsx_q("sheet")):
            if item.get("state", "visible") != "visible":
                target = targets.get(item.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"), "")
                hidden_paths.add(target.lstrip("/") if target.startswith("/") else "xl/"+target)
        for sheet_path in sorted(
            name
            for name in package.namelist()
            if name.startswith("xl/worksheets/") and name.endswith(".xml")
        ):
            if sheet_path in hidden_paths:
                continue
            try:
                root = ElementTree.fromstring(package.read(sheet_path))
            except ElementTree.ParseError:
                continue
            columns_element = root.find(_xlsx_q("cols"))
            columns = list(columns_element) if columns_element is not None else []
            sheet_data = root.find(_xlsx_q("sheetData"))
            if sheet_data is None:
                continue
            for row in sheet_data.findall(_xlsx_q("row")):
                if str(row.attrib.get("hidden", "0")).lower() in {"1", "true"}:
                    continue
                for cell in row.findall(_xlsx_q("c")):
                    column_index = _xlsx_column_index(cell.attrib.get("r", ""))
                    if not column_index:
                        continue
                    attributes = _xlsx_effective_column_attributes(columns, column_index)
                    if str(attributes.get("hidden", "0")).lower() in {"1", "true"}:
                        continue
                    cell_type = str(cell.attrib.get("t") or "")
                    text = _xlsx_cell_text(cell, shared_strings)
                    if not text:
                        continue

                    if cell_type in {"s", "inlineStr", "str", "b"}:
                        values.append(text)
                    else:
                        try:
                            style_id = int(cell.attrib.get("s", "0") or "0")
                        except ValueError:
                            style_id = 0
                        if style_id not in numeric_styles_to_skip:
                            values.append(text)
    return values


def _xlsx_visible_cell_text(path: Path) -> str:
    return "\n".join(_xlsx_visible_cell_values(path))


def _xlsx_retention_key(value: str, *, preserve_punctuation: bool) -> str:
    """Canonicalize text for cell-aware PDF retention checks.

    PDF extractors may insert whitespace/newlines inside a visually complete cell
    when wrapping occurs. Removing whitespace makes the audit invariant to those
    extractor artifacts while still requiring the actual characters to survive.
    A punctuation-preserving form is used first; an alphanumeric form provides a
    fallback for benign punctuation-normalization differences.
    """
    normalized = unicodedata.normalize("NFKC", str(value or "")).casefold()
    normalized = (
        normalized.replace("\u2010", "-")
        .replace("\u2011", "-")
        .replace("\u2012", "-")
        .replace("\u2013", "-")
        .replace("\u2014", "-")
        .replace("\u2212", "-")
        .replace("\u00a0", " ")
        .replace("\u202f", " ")
    )
    if preserve_punctuation:
        allowed_punctuation = set("@._+-/:#%&()[]{}")
        return "".join(
            character
            for character in normalized
            if character.isalnum() or character in allowed_punctuation
        )
    return "".join(character for character in normalized if character.isalnum())


def _count_nonoverlapping_occurrences(haystack: str, needle: str) -> int:
    if not haystack or not needle:
        return 0
    return haystack.count(needle)


def _xlsx_pdf_cell_retention_metrics(
    source_path: Path,
    output_path: Path,
) -> dict[str, float | int]:
    """Measure XLSX-to-PDF fidelity without depending on word token boundaries.

    The previous gate compared whole tokens. That is useful diagnostically, but a
    PDF extractor is allowed to emit ``exam\nple`` for a wrapped visual string
    ``example``. In a different LibreOffice/font environment this made a complete
    conversion appear incomplete. The cell-aware audit removes extractor-inserted
    whitespace and verifies source cell values against the full PDF text stream.

    Multiplicity is retained: repeated source cell values must occur at least the
    same number of times in the PDF stream to receive full credit. Coverage is
    reported both by cell count and by source-character weight so one missing long
    identifier cannot be hidden by many short successfully rendered cells.
    """
    source_values = _xlsx_visible_cell_values(source_path)
    output_text = _pdf_native_text(output_path)

    strict_output = _xlsx_retention_key(output_text, preserve_punctuation=True)
    relaxed_output = _xlsx_retention_key(output_text, preserve_punctuation=False)

    strict_source = Counter(
        key
        for value in source_values
        if len(key := _xlsx_retention_key(value, preserve_punctuation=True)) >= 4
    )
    relaxed_source = Counter(
        key
        for value in source_values
        if len(key := _xlsx_retention_key(value, preserve_punctuation=False)) >= 4
    )

    def score(source_counter: Counter[str], output_compact: str) -> tuple[int, int, int, int]:
        matched_cells = 0
        total_cells = 0
        matched_chars = 0
        total_chars = 0
        for key, required_count in source_counter.items():
            available_count = _count_nonoverlapping_occurrences(output_compact, key)
            matched_count = min(required_count, available_count)
            matched_cells += matched_count
            total_cells += required_count
            matched_chars += matched_count * len(key)
            total_chars += required_count * len(key)
        return matched_cells, total_cells, matched_chars, total_chars

    strict = score(strict_source, strict_output)
    relaxed = score(relaxed_source, relaxed_output)

    # Use the stronger punctuation-preserving match when it works. The relaxed
    # match may only improve coverage where PDF Unicode normalization changes
    # punctuation, never where alphanumeric source content disappeared.
    matched_cells = max(strict[0], relaxed[0])
    total_cells = max(strict[1], relaxed[1])
    matched_chars = max(strict[2], relaxed[2])
    total_chars = max(strict[3], relaxed[3])

    cell_ratio = matched_cells / total_cells if total_cells else 1.0
    character_ratio = matched_chars / total_chars if total_chars else 1.0
    token_ratio, matched_tokens, total_tokens = _token_coverage(
        _xlsx_visible_cell_text(source_path),
        output_text,
    )

    return {
        "cell_ratio": cell_ratio,
        "character_ratio": character_ratio,
        "matched_cells": matched_cells,
        "total_cells": total_cells,
        "matched_characters": matched_chars,
        "total_characters": total_chars,
        "token_ratio": token_ratio,
        "matched_tokens": matched_tokens,
        "total_tokens": total_tokens,
    }


def _xlsx_pdf_fidelity_passes(metrics: Mapping[str, float | int]) -> bool:
    # Token coverage remains diagnostic only. PDF extractors may split a visually
    # complete wrapped value at arbitrary glyph positions, so token boundaries must
    # never be the acceptance authority for XLSX -> PDF. The cell-aware metrics are
    # stricter about actual source characters while being invariant to layout-only
    # whitespace inserted by the extractor.
    cell_ratio = float(metrics.get("cell_ratio", 0.0))
    character_ratio = float(metrics.get("character_ratio", 0.0))
    return (
        cell_ratio + 1e-9 >= max(0.0, min(1.0, DEFAULT_XLSX_PDF_CELL_COVERAGE))
        and character_ratio + 1e-9
        >= max(0.0, min(1.0, DEFAULT_XLSX_PDF_CHARACTER_COVERAGE))
    )


def _assert_xlsx_pdf_fidelity(source_path: Path, output_path: Path) -> None:
    metrics = _xlsx_pdf_cell_retention_metrics(source_path, output_path)
    if _xlsx_pdf_fidelity_passes(metrics):
        return

    raise RuntimeError(
        "Excel-to-PDF conversion failed ReDOCX content-retention validation after "
        "environment-tolerant cell auditing: "
        f"token coverage {int(metrics['matched_tokens'])}/{int(metrics['total_tokens'])} "
        f"({float(metrics['token_ratio']):.1%}; token target {DEFAULT_XLSX_PDF_TOKEN_COVERAGE:.1%}); "
        f"cell coverage {int(metrics['matched_cells'])}/{int(metrics['total_cells'])} "
        f"({float(metrics['cell_ratio']):.1%}; target {DEFAULT_XLSX_PDF_CELL_COVERAGE:.1%}); "
        f"character-weighted coverage {float(metrics['character_ratio']):.1%} "
        f"(target {DEFAULT_XLSX_PDF_CHARACTER_COVERAGE:.1%}). "
        "ReDOCX refused to return a structurally valid but materially incomplete conversion."
    )


class _VisibleHtmlTextParser(HTMLParser):
    _IGNORED_TAGS = {"script", "style", "head", "template", "noscript"}
    _VOID_TAGS = {
        "area", "base", "br", "col", "embed", "hr", "img", "input",
        "link", "meta", "param", "source", "track", "wbr",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._hidden_depth = 0
        self._stack: list[tuple[str, bool]] = []
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        normalized_tag = tag.casefold()
        attributes = {str(key).casefold(): str(value or "") for key, value in attrs}
        style = attributes.get("style", "").replace(" ", "").casefold()
        hidden = (
            normalized_tag in self._IGNORED_TAGS
            or "hidden" in attributes
            or "display:none" in style
            or "visibility:hidden" in style
        )
        if normalized_tag in self._VOID_TAGS:
            return
        self._stack.append((normalized_tag, hidden))
        if hidden:
            self._hidden_depth += 1

    def handle_startendtag(self, tag: str, attrs: list[tuple[str, Optional[str]]]) -> None:
        # Self-closing tags cannot contribute visible text directly.
        return

    def handle_endtag(self, tag: str) -> None:
        normalized_tag = tag.casefold()
        # HTML can be malformed. Unwind to the matching open tag instead of
        # blindly popping one entry and corrupting hidden-depth accounting.
        for index in range(len(self._stack) - 1, -1, -1):
            if self._stack[index][0] != normalized_tag:
                continue
            closing = self._stack[index:]
            del self._stack[index:]
            self._hidden_depth = max(
                0, self._hidden_depth - sum(1 for _, hidden in closing if hidden)
            )
            break

    def handle_data(self, data: str) -> None:
        if self._hidden_depth == 0 and data.strip():
            self.parts.append(data)


def _html_visible_text(path: Path) -> str:
    parser = _VisibleHtmlTextParser()
    parser.feed(path.read_text(encoding="utf-8"))
    parser.close()
    return "\n".join(parser.parts)


def _pdf_page_count(path: Path) -> int:
    with fitz.open(path) as document:
        return int(document.page_count)


def _validate_pdf_to_jpg_page_count(source_path: Path, output_path: Path, physical_format: str) -> None:
    expected = _pdf_page_count(source_path)
    if physical_format == "jpg":
        if expected != 1:
            raise RuntimeError("Multi-page PDF-to-JPG conversion must return a page-image archive.")
        return
    if physical_format != "zip":
        raise RuntimeError("PDF-to-JPG produced an unexpected physical output format.")
    with zipfile.ZipFile(output_path) as archive:
        actual = len([info for info in archive.infolist() if not info.is_dir()])
    if actual != expected:
        raise RuntimeError(
            f"PDF-to-JPG page-count mismatch: source has {expected} pages, output has {actual} images."
        )


def _validate_pdf_to_pptx_page_count(source_path: Path, output_path: Path) -> None:
    if PptxPresentation is None:
        return
    expected = _pdf_page_count(source_path)
    actual = len(PptxPresentation(str(output_path)).slides)
    if actual != expected:
        raise RuntimeError(
            f"PDF-to-PowerPoint page-count mismatch: source has {expected} pages, output has {actual} slides."
        )


def _validate_pdf_to_xlsx_page_count(source_path: Path, output_path: Path) -> None:
    native.validate_xlsx(source_path, output_path)


def _validate_image_dimensions_preserved(source_path: Path, output_path: Path) -> None:
    with Image.open(source_path) as source_image, Image.open(output_path) as output_image:
        source_size = ImageOps.exif_transpose(source_image).size
        output_size = output_image.size
    if tuple(source_size) != tuple(output_size):
        raise RuntimeError(
            f"Image conversion changed pixel dimensions from {source_size} to {output_size}."
        )


def _pdf_page_comparison_image(
    page: fitz.Page,
    *,
    width_px: int = 720,
    target_height_px: Optional[int] = None,
) -> Image.Image:
    """Render a PDF page to a deterministic low-resolution grayscale QA image."""

    width_px = max(240, min(1600, int(width_px)))
    page_width = max(1.0, float(page.rect.width))
    page_height = max(1.0, float(page.rect.height))
    scale = max(1.0, width_px / page_width)
    pix = page.get_pixmap(
        matrix=fitz.Matrix(scale, scale),
        colorspace=fitz.csGRAY,
        alpha=False,
    )
    image = Image.frombytes("L", (pix.width, pix.height), pix.samples)
    target_height = (
        max(1, int(target_height_px))
        if target_height_px is not None
        else max(1, round(width_px * page_height / page_width))
    )
    if image.size != (width_px, target_height):
        image = image.resize((width_px, target_height), Image.Resampling.LANCZOS)
    return image


def _binary_ink_mask(image: Image.Image, *, threshold: int = 238) -> Image.Image:
    return image.point(lambda value: 255 if value < threshold else 0, mode="L")


def _mask_pixel_count(mask: Image.Image) -> int:
    histogram = mask.histogram()
    return int(histogram[255]) if len(histogram) > 255 else 0


def _page_ink_layout_score(source: Image.Image, output: Image.Image) -> float:
    """Return a font-tolerant visual-layout score in [0, 1].

    A small dilation radius tolerates renderer/font antialiasing differences while
    still penalizing blocks that moved, merged, reflowed, or changed orientation.
    """

    source_mask = _binary_ink_mask(source)
    output_mask = _binary_ink_mask(output)
    source_count = _mask_pixel_count(source_mask)
    output_count = _mask_pixel_count(output_mask)
    if source_count == 0 and output_count == 0:
        return 1.0
    if source_count == 0 or output_count == 0:
        return 0.0

    source_dilated = source_mask.filter(ImageFilter.MaxFilter(5))
    output_dilated = output_mask.filter(ImageFilter.MaxFilter(5))
    matched_source = _mask_pixel_count(ImageChops.multiply(source_mask, output_dilated))
    matched_output = _mask_pixel_count(ImageChops.multiply(output_mask, source_dilated))
    recall = matched_source / source_count
    precision = matched_output / output_count
    if precision + recall <= 0:
        return 0.0
    return 2.0 * precision * recall / (precision + recall)


def _page_grayscale_similarity(source: Image.Image, output: Image.Image) -> float:
    if source.size != output.size:
        output = output.resize(source.size, Image.Resampling.LANCZOS)
    difference = ImageChops.difference(source, output)
    mean_difference = float(ImageStat.Stat(difference).mean[0])
    return max(0.0, min(1.0, 1.0 - mean_difference / 255.0))


def _assert_pdf_visual_layout_fidelity(
    source_path: Path,
    output_path: Path,
    *,
    minimum_score: float,
    label: str,
) -> None:
    """Fail closed when a PDF round-trip materially changes page layout."""

    minimum_score = max(0.0, min(1.0, float(minimum_score)))
    with fitz.open(source_path) as source_pdf, fitz.open(output_path) as output_pdf:
        if source_pdf.page_count != output_pdf.page_count:
            raise RuntimeError(
                f"{label} changed page count from {source_pdf.page_count} to {output_pdf.page_count}."
            )

        for page_index in range(source_pdf.page_count):
            source_page = source_pdf.load_page(page_index)
            output_page = output_pdf.load_page(page_index)
            source_width = max(1.0, float(source_page.rect.width))
            source_height = max(1.0, float(source_page.rect.height))
            output_width = max(1.0, float(output_page.rect.width))
            output_height = max(1.0, float(output_page.rect.height))

            source_landscape = source_width > source_height
            output_landscape = output_width > output_height
            if source_landscape != output_landscape:
                raise RuntimeError(
                    f"{label} changed page {page_index + 1} orientation."
                )
            if (
                abs(source_width - output_width) / source_width > 0.02
                or abs(source_height - output_height) / source_height > 0.02
            ):
                raise RuntimeError(
                    f"{label} changed page {page_index + 1} dimensions from "
                    f"{source_width:.1f}x{source_height:.1f} pt to "
                    f"{output_width:.1f}x{output_height:.1f} pt."
                )

            source_image = _pdf_page_comparison_image(source_page)
            output_image = _pdf_page_comparison_image(
                output_page,
                width_px=source_image.width,
                target_height_px=source_image.height,
            )
            source_text_chars = len(
                _normalized_alphanumeric_text(source_page.get_text("text"))
            )
            score = (
                _page_ink_layout_score(source_image, output_image)
                if source_text_chars >= 20
                else _page_grayscale_similarity(source_image, output_image)
            )
            if score + 1e-9 < minimum_score:
                raise RuntimeError(
                    f"{label} failed visual-layout validation on page {page_index + 1}: "
                    f"score {score:.3f}, required {minimum_score:.3f}. ReDOCX refused "
                    "to return a materially reflowed or rearranged artifact."
                )


def _validate_conversion_fidelity(
    *,
    source_path: Path,
    input_format: str,
    requested_output_format: str,
    output_path: Path,
    physical_output_format: str,
) -> None:
    pair = (input_format, requested_output_format)

    if pair == ("xlsx", "pdf"):
        _assert_xlsx_pdf_fidelity(source_path, output_path)
        return

    if pair == ("docx", "pdf"):
        if not _docx_package_text(source_path).strip():
            if _docx_is_redocx_fidelity_document(source_path) and len(Document(source_path).sections) != _pdf_page_count(output_path):
                raise RuntimeError("Legacy Word-to-PDF changed the page count.")
            quality.assert_legacy_word_images(source_path, output_path)
            return
        _assert_text_token_coverage(
            source_text=_docx_package_text(source_path),
            output_text=_pdf_native_text(output_path),
            minimum=DEFAULT_TEXT_TOKEN_COVERAGE,
            label="Word-to-PDF conversion",
        )
        return

    if pair == ("pptx", "pdf"):
        _assert_text_token_coverage(
            source_text=_pptx_visible_slide_text(source_path),
            output_text=_pdf_native_text(output_path),
            minimum=DEFAULT_TEXT_TOKEN_COVERAGE,
            label="PowerPoint-to-PDF conversion",
        )
        return

    if pair in {("html", "pdf"), ("htm", "pdf")}:
        _assert_text_token_coverage(
            source_text=_html_visible_text(source_path),
            output_text=_pdf_native_text(output_path),
            minimum=DEFAULT_TEXT_TOKEN_COVERAGE,
            label="HTML-to-PDF conversion",
        )
        return

    if pair == ("pdf", "docx"):
        _assert_text_token_coverage(
            source_text=_pdf_native_text(source_path),
            output_text=_docx_package_text(output_path),
            minimum=.98,
            label="PDF-to-Word conversion",
        )
        return

    if pair == ("pdf", "jpg"):
        _validate_pdf_to_jpg_page_count(source_path, output_path, physical_output_format)
        quality.assert_pdf_jpegs(source_path, output_path, physical_output_format)
        return

    if pair == ("pdf", "pptx"):
        _validate_pdf_to_pptx_page_count(source_path, output_path)
        quality.assert_text(_pdf_native_text(source_path), _pptx_visible_slide_text(output_path),
                            minimum=.98, label="PDF-to-PowerPoint")
        return

    if pair == ("pdf", "xlsx"):
        native.validate_xlsx(source_path, output_path)
        return

    if pair == ("pdf", "pdfa"):
        quality.assert_pdfa_content(source_path, output_path)
        _assert_pdf_visual_layout_fidelity(source_path, output_path,
            minimum_score=max(.98, DEFAULT_PDFA_VISUAL_MIN_SCORE), label="PDF-to-PDF/A")
        return

    if pair in {("jpg", "pdf"), ("jpeg", "pdf")}:
        quality.assert_image_pdf(source_path, output_path)
        return

    if pair in {("jpg", "docx"), ("jpeg", "docx")}:
        with zipfile.ZipFile(output_path) as package:
            media = [name for name in package.namelist() if name.startswith("word/media/")]
        if not media and not quality.visible_word_text(output_path).strip():
            raise RuntimeError("Image-to-Word conversion produced no content.")
        return

    if input_format == "png" and requested_output_format in {"jpg", "jpeg"}:
        _validate_image_dimensions_preserved(source_path, output_path)
        with Image.open(source_path) as src, Image.open(output_path) as dst:
            quality.assert_pixels(src, dst, label="PNG-to-JPEG")
        return


def _validate_source_format(path: Path, expected_format: str) -> None:
    """Reject extension/MIME spoofing before any conversion engine is invoked."""

    if not path.is_file() or path.stat().st_size <= 0:
        raise ValueError(f"Source file is missing or empty: {path}")

    if expected_format == "pdf":
        _validate_pdf_file(path)
    elif expected_format == "docx":
        _validate_docx_file(path)
    elif expected_format == "xlsx":
        _validate_xlsx_file(path)
    elif expected_format == "pptx":
        _validate_pptx_file(path)
    elif expected_format in {"html", "htm"}:
        _validate_html_file(path)
    elif expected_format in {"jpg", "jpeg", "png"}:
        _validate_image_file(path, expected_format)
    else:  # pragma: no cover - guarded by _normalize_format
        raise ValueError(f"Unsupported source format validation: {expected_format}")


def _validate_converted_output(
    path: Path,
    expected_format: str,
    *,
    source_path: Optional[Path] = None,
    input_format: Optional[str] = None,
) -> None:
    """Verify output bytes, not merely the filename or declared MIME type."""

    if not path.is_file() or path.stat().st_size <= 0:
        raise RuntimeError(f"Conversion output is missing or empty: {path}")

    expected_suffixes = {"jpg": {".jpg"}, "jpeg": {".jpeg"}}.get(
        expected_format,
        {f".{expected_format}"},
    )
    if path.suffix.lower() not in expected_suffixes:
        raise RuntimeError(
            f"Conversion output extension {path.suffix or '(none)'} does not match {expected_format}."
        )

    if expected_format == "pdf":
        _validate_pdf_file(path)
    elif expected_format == "docx":
        _validate_docx_file(path)
        if input_format == "pdf" and source_path is not None:
            _validate_pdf_to_docx_editability(source_path, path)
            _assert_pdf_hyperlinks_preserved(path, _extract_pdf_hyperlinks(source_path))
    elif expected_format in {"jpg", "jpeg"}:
        _validate_image_file(path, expected_format)
    elif expected_format == "pptx":
        _validate_pptx_file(path)
    elif expected_format == "xlsx":
        _validate_xlsx_file(path)
    elif expected_format == "zip":
        _validate_jpg_archive(path)
    else:
        raise RuntimeError(f"Unsupported conversion output validation: {expected_format}")


def _validate_pdf_file(path: Path) -> None:
    try:
        with path.open("rb") as stream:
            header = stream.read(8)
        if not header.startswith(b"%PDF-"):
            raise ValueError(f"File content is not a PDF: {path}")
        with fitz.open(path) as pdf:
            if pdf.page_count < 1:
                raise ValueError(f"PDF contains no pages: {path}")
            if pdf.is_encrypted or pdf.needs_pass:
                raise ValueError(f"Password-protected PDF is not supported: {path}")
            # Force page-tree parsing so truncated/corrupt files fail now, before storage.
            for page_number in range(pdf.page_count):
                _ = pdf.load_page(page_number).rect
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"File is not a readable PDF: {path}") from exc



def _validate_html_file(path: Path) -> None:
    try:
        data = path.read_bytes()
        if b"\x00" in data:
            raise ValueError("HTML contains null bytes.")
        text = data.decode("utf-8")
    except UnicodeDecodeError as exc:
        raise ValueError("HTML input must be valid UTF-8.") from exc
    if not text.strip():
        raise ValueError("HTML input is empty.")


def _validate_xlsx_file(path: Path) -> None:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required for XLSX conversion.")
    try:
        workbook = load_workbook(path, read_only=True, data_only=False)
        try:
            if not workbook.sheetnames:
                raise ValueError("XLSX workbook contains no worksheets.")
        finally:
            workbook.close()
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"File is not a readable XLSX workbook: {path}") from exc


def _validate_pptx_file(path: Path) -> None:
    if PptxPresentation is None:
        raise RuntimeError("python-pptx is required for PPTX conversion.")
    try:
        presentation = PptxPresentation(str(path))
        if len(presentation.slides) < 1:
            raise ValueError("PPTX presentation contains no slides.")
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"File is not a readable PPTX presentation: {path}") from exc


def _validate_jpg_archive(path: Path) -> None:
    try:
        with zipfile.ZipFile(path) as archive:
            infos = [info for info in archive.infolist() if not info.is_dir()]
            if not infos:
                raise RuntimeError("PDF to JPG archive contains no page images.")
            for info in infos:
                name = info.filename.replace("\\", "/")
                if name.startswith("/") or ".." in Path(name).parts:
                    raise RuntimeError("PDF to JPG archive contains an unsafe entry path.")
                if not name.lower().endswith(".jpg"):
                    raise RuntimeError("PDF to JPG archive contains a non-JPG entry.")
                if info.file_size <= 0:
                    raise RuntimeError("PDF to JPG archive contains an empty page image.")
                payload = archive.read(info)
                if not payload.startswith(b"\xff\xd8\xff"):
                    raise RuntimeError("PDF to JPG archive contains invalid JPEG data.")
    except zipfile.BadZipFile as exc:
        raise RuntimeError("PDF to JPG output is not a valid ZIP archive.") from exc


def _validate_pdfa_file(path: Path) -> None:
    _validate_pdf_file(path)
    with fitz.open(path) as pdf:
        xmp_xref = pdf.xref_xml_metadata()
        if not xmp_xref:
            raise RuntimeError("PDF/A output is missing XMP conformance metadata.")
        root = ElementTree.fromstring(pdf.xref_stream(xmp_xref))
        ns = "{http://www.aiim.org/pdfa/ns/id/}"
        def metadata(name):
            return [n.text for n in root.iter(ns+name)] + [n.attrib[ns+name] for n in root.iter() if ns+name in n.attrib]
        if "2" not in metadata("part") or "B" not in [str(x).upper() for x in metadata("conformance")]:
            raise RuntimeError("PDF/A output does not declare PDF/A-2b conformance.")
        catalog = pdf.xref_object(pdf.pdf_catalog(), compressed=False).casefold()
        if "/outputintents" not in catalog:
            raise RuntimeError("PDF/A output is missing an OutputIntent color profile.")

    _validate_pdfa_with_verapdf(path)


def _validate_pdfa_with_verapdf(path: Path) -> None:
    """PDF/A certification always requires an independent validator verdict."""
    mode = "required"
    if DEFAULT_PDFA_VALIDATOR_MODE != "required":
        raise RuntimeError("PDF/A validation cannot be disabled; use REDOCX_PDFA_VALIDATOR_MODE=required.")
    validator = (
        os.getenv("VERAPDF_PATH")
        or shutil.which("verapdf")
        or shutil.which("verapdf.bat")
    )
    if not validator:
        if mode == "required":
            raise RuntimeError(
                "veraPDF is required for production PDF/A validation but was not found. "
                "Install veraPDF or set VERAPDF_PATH."
            )
        return

    try:
        result = subprocess.run(
            [validator, "--loglevel", "1", "--format", "raw", "-f", "2b", str(path)],
            check=False,
            capture_output=True,
            text=True,
            timeout=DEFAULT_PDFA_TIMEOUT_SECONDS,
        )
    except subprocess.TimeoutExpired as exc:
        raise RuntimeError("veraPDF timed out while validating PDF/A-2b output.") from exc
    except OSError as exc:
        if mode == "required":
            raise RuntimeError("veraPDF could not be executed.") from exc
        return

    report = result.stdout or ""
    compliant = False
    if result.returncode == 0 and report.strip():
        try:
            root = ElementTree.fromstring(report)
            reports = [element for element in root.iter()
                       if element.tag.split("}")[-1] == "validationResult"]
            compliant = (len(reports) == 1 and reports[0].get("isCompliant") == "true"
                         and reports[0].get("flavour") == "PDFA_2_B"
                         and reports[0].get("jobEndStatus") == "NORMAL")
        except ElementTree.ParseError:
            compliant = False

    if not compliant:
        raise RuntimeError(
            "PDF/A conversion failed independent veraPDF PDF/A-2b validation. "
            "ReDOCX refused to return a non-conformant archival PDF."
        )


def _resolve_pdfa_icc_profile() -> Path:
    configured = os.getenv("REDOCX_PDFA_ICC_PROFILE", "").strip()
    candidates = [
        Path(configured).expanduser() if configured else None,
        Path("/usr/share/color/icc/ghostscript/srgb.icc"),
        Path("/usr/share/color/icc/sRGB.icc"),
        Path("/usr/share/color/icc/colord/sRGB.icc"),
    ]
    for candidate in candidates:
        if candidate and candidate.is_file() and candidate.stat().st_size > 0:
            return candidate.resolve()
    raise RuntimeError(
        "A valid sRGB ICC profile is required for PDF/A conversion. "
        "Set REDOCX_PDFA_ICC_PROFILE."
    )


def _resolve_pdfa_definition_file(gs_bin: str) -> Path:
    configured = os.getenv("REDOCX_PDFA_DEF_PS", "").strip()
    candidates: list[Path] = []
    if configured:
        candidates.append(Path(configured).expanduser())
    try:
        version = subprocess.run(
            [gs_bin, "--version"],
            check=False,
            capture_output=True,
            text=True,
            timeout=10,
        ).stdout.strip()
    except Exception:
        version = ""
    if version:
        candidates.append(Path(f"/usr/share/ghostscript/{version}/lib/PDFA_def.ps"))
    candidates.extend(sorted(Path("/usr/share/ghostscript").glob("*/lib/PDFA_def.ps"), reverse=True))
    for candidate in candidates:
        if candidate.is_file() and candidate.stat().st_size > 0:
            return candidate.resolve()
    raise RuntimeError(
        "Ghostscript PDFA_def.ps was not found. Set REDOCX_PDFA_DEF_PS to its path."
    )

def _validate_docx_file(path: Path) -> None:
    required_parts = {
        "[Content_Types].xml",
        "_rels/.rels",
        "word/document.xml",
    }
    expected_main_content_type = (
        "application/vnd.openxmlformats-officedocument."
        "wordprocessingml.document.main+xml"
    )

    try:
        if not zipfile.is_zipfile(path):
            raise ValueError(f"File content is not an OOXML DOCX package: {path}")
        with zipfile.ZipFile(path) as package:
            names = set(package.namelist())
            missing = sorted(required_parts - names)
            if missing:
                raise ValueError(f"DOCX package is missing required parts: {', '.join(missing)}")

            content_types = ElementTree.fromstring(package.read("[Content_Types].xml"))
            main_types = {
                node.attrib.get("ContentType")
                for node in content_types
                if node.attrib.get("PartName") == "/word/document.xml"
            }
            if expected_main_content_type not in main_types:
                raise ValueError("OOXML package is not a standard .docx Word document.")

            ElementTree.fromstring(package.read("word/document.xml"))

        # python-docx exercises the OPC relationships and catches packages that
        # contain XML parts but cannot actually be opened as Word documents.
        Document(str(path))
    except (ValueError, zipfile.BadZipFile):
        raise
    except Exception as exc:
        raise ValueError(f"File is not a readable DOCX document: {path}") from exc


def _validate_image_file(path: Path, expected_format: str) -> None:
    expected_pillow_format = "JPEG" if expected_format in {"jpg", "jpeg"} else "PNG"
    try:
        with Image.open(path) as image:
            actual_format = (image.format or "").upper()
            width, height = image.size
            image.verify()
        if actual_format != expected_pillow_format:
            raise ValueError(
                f"Image content is {actual_format or 'unknown'}, not {expected_pillow_format}."
            )
        if width < 1 or height < 1:
            raise ValueError("Image has invalid dimensions.")
    except ValueError:
        raise
    except (UnidentifiedImageError, OSError) as exc:
        raise ValueError(f"File is not a readable {expected_pillow_format} image: {path}") from exc


def _docx_text(path: Path) -> str:
    with zipfile.ZipFile(path) as package:
        root = ElementTree.fromstring(package.read("word/document.xml"))
    text_tag = f"{{{WORDPROCESSINGML_NS}}}t"
    return "".join(node.text or "" for node in root.iter(text_tag))


def _normalized_alphanumeric_text(value: str) -> str:
    return "".join(character for character in value if character.isalnum())


def _pdf_native_text(path: Path) -> str:
    with fitz.open(path) as pdf:
        return "\n".join(page.get_text("text") for page in pdf)


def _pdf_page_searchable_text(page: fitz.Page) -> str:
    """Return native page text, with best-effort OCR for scanned pages.

    OCR is an enhancement for the fidelity DOCX text layer, not a prerequisite for
    visual conversion. A scanned PDF can still be converted faithfully when the OCR
    runtime is unavailable; in that case the page image remains authoritative.
    """

    native = page.get_text("text", sort=True).strip()
    if native:
        return native

    try:
        text_page = page.get_textpage_ocr(
            language=DEFAULT_PDF_TO_DOCX_OCR_LANGUAGE,
            dpi=max(96, DEFAULT_PDF_TO_DOCX_OCR_DPI),
            full=True,
        )
        return page.get_text("text", textpage=text_page, sort=True).strip()
    except Exception:
        return ""


def _docx_is_redocx_fidelity_document(path: Path) -> bool:
    try:
        with zipfile.ZipFile(path) as package:
            if "docProps/core.xml" not in package.namelist():
                return False
            core = package.read("docProps/core.xml").decode("utf-8", errors="ignore")
    except (OSError, zipfile.BadZipFile):
        return False
    return REDOCX_FIDELITY_DOCX_MARKER in core


def _validate_pdf_to_docx_editability(source_path: Path, docx_path: Path) -> None:
    with fitz.open(source_path) as pdf:
        source = "\n".join(native.page_text(page) for page in pdf)
    quality.assert_text(source, quality.visible_word_text(docx_path), minimum=.98,
                        label="PDF-to-Word native editability")

def _extract_pdf_hyperlinks(path: Path) -> list[PdfHyperlink]:
    links: list[PdfHyperlink] = []
    with fitz.open(path) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            for link in page.get_links():
                target = str(link.get("uri") or "").strip()
                if not target:
                    # External web/mail links are portable to DOCX. PDF-only
                    # actions and page-coordinate jumps have no reliable Word
                    # equivalent without rewriting document navigation.
                    continue

                label = ""
                source_rect = link.get("from")
                if source_rect:
                    rect = fitz.Rect(source_rect)
                    words = [
                        str(word[4]).strip()
                        for word in page.get_text("words", clip=rect, sort=True)
                        if len(word) > 4 and str(word[4]).strip()
                    ]
                    words = [word for word in words if not _is_link_separator(word)]
                    label = " ".join(words).strip()
                    if not label:
                        label = " ".join(page.get_textbox(rect).split()).strip()

                if not label:
                    label = _display_text_from_link_target(target)

                links.append(
                    PdfHyperlink(
                        page_number=page_number,
                        label=label,
                        target=target,
                    )
                )
    return links


def _is_link_separator(value: str) -> bool:
    return not any(character.isalnum() for character in value) and "@" not in value


def _display_text_from_link_target(target: str) -> str:
    decoded = unquote(target).strip()
    if decoded.lower().startswith("mailto:"):
        decoded = decoded[7:].split("?", 1)[0]
        decoded = decoded.split("|", 1)[0].strip()
    return decoded or target


def _link_search_labels(link: PdfHyperlink) -> list[str]:
    candidates = [link.label, _display_text_from_link_target(link.target)]
    if link.target.lower().startswith("mailto:"):
        candidates.append(unquote(link.target[7:]).split("?", 1)[0].split("|", 1)[0].strip())

    unique: list[str] = []
    for candidate in candidates:
        normalized = " ".join(str(candidate or "").split()).strip(" :|\t\r\n")
        if normalized and normalized not in unique:
            unique.append(normalized)
    return unique


def _docx_external_hyperlink_targets(path: Path) -> Counter[str]:
    with zipfile.ZipFile(path) as package:
        document_root = ElementTree.fromstring(package.read("word/document.xml"))
        relationships_path = "word/_rels/document.xml.rels"
        if relationships_path not in package.namelist():
            return Counter()
        relationships_root = ElementTree.fromstring(package.read(relationships_path))

    rel_id_attr = "Id"
    rel_type_attr = "Type"
    rel_target_attr = "Target"
    rel_mode_attr = "TargetMode"
    relationship_by_id: dict[str, str] = {}
    for relationship in relationships_root:
        if not relationship.attrib.get(rel_type_attr, "").endswith("/hyperlink"):
            continue
        if relationship.attrib.get(rel_mode_attr) != "External":
            continue
        relationship_by_id[relationship.attrib.get(rel_id_attr, "")] = relationship.attrib.get(
            rel_target_attr,
            "",
        )

    hyperlink_tag = f"{{{WORDPROCESSINGML_NS}}}hyperlink"
    relationship_id_attr = f"{{{OFFICE_REL_NS}}}id"
    targets: Counter[str] = Counter()
    for hyperlink in document_root.iter(hyperlink_tag):
        target = relationship_by_id.get(hyperlink.attrib.get(relationship_id_attr, ""))
        if target:
            targets[target] += 1
    return targets


def _assert_pdf_hyperlinks_preserved(
    docx_path: Path,
    source_links: Sequence[PdfHyperlink],
) -> None:
    expected = Counter(link.target for link in source_links if link.target)
    if not expected:
        return
    actual = _docx_external_hyperlink_targets(docx_path)
    missing: list[str] = []
    for target, expected_count in expected.items():
        missing_count = expected_count - actual.get(target, 0)
        missing.extend([target] * max(0, missing_count))
    if missing:
        raise RuntimeError(
            "PDF-to-Word conversion could not preserve every external hyperlink: "
            + ", ".join(missing)
        )


def _apply_pdf_hyperlinks(docx_path: Path, source_links: Sequence[PdfHyperlink]) -> None:
    if not source_links:
        return

    document = Document(str(docx_path))
    existing = _docx_external_hyperlink_targets(docx_path)
    consumed_existing: Counter[str] = Counter()
    paragraph_tag = qn("w:p")
    paragraphs = list(document.element.body.iter(paragraph_tag))

    changed = False
    for link in source_links:
        if consumed_existing[link.target] < existing.get(link.target, 0):
            consumed_existing[link.target] += 1
            continue

        linked = False
        for label in _link_search_labels(link):
            for paragraph_element in paragraphs:
                if _wrap_text_range_with_hyperlink(
                    document,
                    paragraph_element,
                    label=label,
                    target=link.target,
                ):
                    linked = True
                    changed = True
                    break
            if linked:
                break

        if not linked:
            raise RuntimeError(
                f"Could not map PDF hyperlink label {link.label!r} to editable Word text."
            )

    if changed:
        document.save(docx_path)


def _run_visible_text(run_element: Any) -> str:
    text_tag = qn("w:t")
    tab_tag = qn("w:tab")
    break_tag = qn("w:br")
    chunks: list[str] = []
    for child in run_element:
        if child.tag == text_tag:
            chunks.append(child.text or "")
        elif child.tag == tab_tag:
            chunks.append("\t")
        elif child.tag == break_tag:
            chunks.append("\n")
    return "".join(chunks)


def _set_run_visible_text(run_element: Any, text: str) -> None:
    run_properties_tag = qn("w:rPr")
    for child in list(run_element):
        if child.tag != run_properties_tag:
            run_element.remove(child)

    text_element = OxmlElement("w:t")
    if text[:1].isspace() or text[-1:].isspace():
        text_element.set(f"{{{XML_NS}}}space", "preserve")
    text_element.text = text
    run_element.append(text_element)


def _clone_run_with_text(run_element: Any, text: str, *, hyperlink_style: bool = False) -> Any:
    cloned = deepcopy(run_element)
    _set_run_visible_text(cloned, text)
    if hyperlink_style:
        _style_hyperlink_run(cloned)
    return cloned


def _style_hyperlink_run(run_element: Any) -> None:
    run_properties = run_element.find(qn("w:rPr"))
    if run_properties is None:
        run_properties = OxmlElement("w:rPr")
        run_element.insert(0, run_properties)

    for tag_name in ("w:rStyle", "w:color", "w:u"):
        for existing in list(run_properties.findall(qn(tag_name))):
            run_properties.remove(existing)

    run_style = OxmlElement("w:rStyle")
    run_style.set(qn("w:val"), "Hyperlink")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_properties.extend((run_style, color, underline))


def _wrap_text_range_with_hyperlink(
    document: Any,
    paragraph_element: Any,
    *,
    label: str,
    target: str,
) -> bool:
    # Conversion engines used by ReDOCX emit ordinary body runs. Existing
    # hyperlink runs are excluded so repeated labels map to distinct occurrences.
    run_tag = qn("w:r")
    hyperlink_tag = qn("w:hyperlink")
    runs = [child for child in paragraph_element if child.tag == run_tag]
    if not runs:
        return False

    run_texts = [_run_visible_text(run) for run in runs]
    combined = "".join(run_texts)
    tokens = re.findall(r"\S+", label)
    if not tokens:
        return False
    pattern = re.compile(r"\s+".join(re.escape(token) for token in tokens), re.IGNORECASE)
    match = pattern.search(combined)
    if match is None:
        return False

    start_offset, end_offset = match.span()
    ranges: list[tuple[Any, str, int, int]] = []
    cursor = 0
    for run, run_text in zip(runs, run_texts):
        next_cursor = cursor + len(run_text)
        if next_cursor > start_offset and cursor < end_offset:
            ranges.append((run, run_text, cursor, next_cursor))
        cursor = next_cursor
    if not ranges:
        return False

    first_run, first_text, first_start, _ = ranges[0]
    last_run, last_text, last_start, _ = ranges[-1]
    insertion_index = paragraph_element.index(first_run)
    replacement_nodes: list[Any] = []

    before = first_text[: max(0, start_offset - first_start)]
    if before:
        replacement_nodes.append(_clone_run_with_text(first_run, before))

    relationship_id = document.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), relationship_id)
    hyperlink.set(qn("w:history"), "1")

    for run, run_text, run_start, run_end in ranges:
        selected_start = max(start_offset, run_start) - run_start
        selected_end = min(end_offset, run_end) - run_start
        selected = run_text[selected_start:selected_end]
        if selected:
            hyperlink.append(_clone_run_with_text(run, selected, hyperlink_style=True))
    replacement_nodes.append(hyperlink)

    after = last_text[max(0, end_offset - last_start) :]
    if after:
        replacement_nodes.append(_clone_run_with_text(last_run, after))

    for run, _run_text, _run_start, _run_end in ranges:
        paragraph_element.remove(run)
    for offset, node in enumerate(replacement_nodes):
        paragraph_element.insert(insertion_index + offset, node)

    return any(child.tag == hyperlink_tag for child in replacement_nodes)


def _pdf_dict_text_char_count(page_dict: dict[str, Any]) -> int:
    return sum(
        len(_normalized_alphanumeric_text(str(span.get("text", ""))))
        for block in page_dict.get("blocks", [])
        if block.get("type") == 0
        for line in block.get("lines", [])
        for span in line.get("spans", [])
    )


def _rect_tuple(value: Sequence[float]) -> tuple[float, float, float, float]:
    return (float(value[0]), float(value[1]), float(value[2]), float(value[3]))


def _rect_union(rectangles: Iterable[Sequence[float]]) -> tuple[float, float, float, float]:
    values = [_rect_tuple(rectangle) for rectangle in rectangles]
    if not values:
        return (0.0, 0.0, 0.0, 0.0)
    return (
        min(rectangle[0] for rectangle in values),
        min(rectangle[1] for rectangle in values),
        max(rectangle[2] for rectangle in values),
        max(rectangle[3] for rectangle in values),
    )


def _rect_intersection_area(first: Sequence[float], second: Sequence[float]) -> float:
    x0 = max(float(first[0]), float(second[0]))
    y0 = max(float(first[1]), float(second[1]))
    x1 = min(float(first[2]), float(second[2]))
    y1 = min(float(first[3]), float(second[3]))
    return max(0.0, x1 - x0) * max(0.0, y1 - y0)


def _rect_area(rectangle: Sequence[float]) -> float:
    return max(0.0, float(rectangle[2]) - float(rectangle[0])) * max(
        0.0,
        float(rectangle[3]) - float(rectangle[1]),
    )


def _mostly_inside_any(rectangle: Sequence[float], containers: Sequence[Sequence[float]]) -> bool:
    area = _rect_area(rectangle)
    if area <= 0:
        return False
    return any(_rect_intersection_area(rectangle, container) / area >= 0.55 for container in containers)


def _safe_find_pdf_tables(page: fitz.Page) -> list[Any]:
    try:
        finder = page.find_tables()
        return list(getattr(finder, "tables", []) or [])
    except Exception:
        # Table detection is an enhancement. Native text remains available if a
        # particular PDF drawing pattern is unsupported by PyMuPDF's detector.
        return []


def _build_pdf_page_elements(
    page: fitz.Page,
    *,
    page_dict: dict[str, Any],
    native_page_dict: dict[str, Any],
) -> tuple[list[tuple[str, tuple[float, float, float, float], Any]], tuple[float, float, float, float]]:
    tables = _safe_find_pdf_tables(page)
    table_rectangles = [_rect_tuple(table.bbox) for table in tables]

    raw_lines: list[PdfTextLine] = []
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            bbox = _rect_tuple(line.get("bbox") or block.get("bbox") or (0, 0, 0, 0))
            if _mostly_inside_any(bbox, table_rectangles):
                continue
            spans = tuple(
                dict(span)
                for span in line.get("spans", [])
                if str(span.get("text", ""))
            )
            combined_text = "".join(str(span.get("text", "")) for span in spans)
            has_text_or_marker = bool(
                _normalized_alphanumeric_text(combined_text)
                or any(marker in combined_text for marker in ("\uf0b7", "\uf0a7", "•", "●", "▪"))
            )
            if spans and has_text_or_marker:
                raw_lines.append(PdfTextLine(bbox=bbox, spans=spans))

    visual_lines = _merge_parallel_pdf_lines(raw_lines)
    text_groups = _group_pdf_lines_into_paragraphs(visual_lines)

    elements: list[tuple[str, tuple[float, float, float, float], Any]] = [
        ("text", group.bbox, group) for group in text_groups
    ]
    elements.extend(("table", _rect_tuple(table.bbox), table) for table in tables)

    # Keep native embedded images (logos, signatures, photos) as individual Word
    # images. This is materially different from the removed full-page screenshot
    # fallback: all source text remains real editable Word text.
    for block in native_page_dict.get("blocks", []):
        if block.get("type") != 1 or not block.get("image"):
            continue
        bbox = _rect_tuple(block.get("bbox") or (0, 0, 0, 0))
        if _mostly_inside_any(bbox, table_rectangles):
            continue
        elements.append(("image", bbox, block))

    elements.sort(key=lambda item: (round(item[1][1], 1), item[1][0], item[0] != "image"))
    content_rectangles = [bbox for _kind, bbox, _payload in elements if _rect_area(bbox) > 0]
    content_bbox = _rect_union(content_rectangles) if content_rectangles else _rect_tuple(page.rect)
    return elements, content_bbox


def _merge_parallel_pdf_lines(lines: Sequence[PdfTextLine]) -> list[PdfTextLine]:
    merged: list[PdfTextLine] = []
    for line in sorted(lines, key=lambda item: (item.bbox[1], item.bbox[0])):
        if not merged:
            merged.append(line)
            continue

        previous = merged[-1]
        previous_center = (previous.bbox[1] + previous.bbox[3]) / 2.0
        current_center = (line.bbox[1] + line.bbox[3]) / 2.0
        height = max(1.0, min(previous.bbox[3] - previous.bbox[1], line.bbox[3] - line.bbox[1]))
        same_visual_row = abs(previous_center - current_center) <= max(2.5, height * 0.35)

        if same_visual_row:
            spans = tuple(
                sorted(
                    (*previous.spans, *line.spans),
                    key=lambda span: float((span.get("bbox") or (0, 0, 0, 0))[0]),
                )
            )
            merged[-1] = PdfTextLine(
                bbox=_rect_union((previous.bbox, line.bbox)),
                spans=spans,
            )
        else:
            merged.append(line)
    return merged


def _line_dominant_size(line: PdfTextLine) -> float:
    sizes = [float(span.get("size") or 10.0) for span in line.spans]
    return max(sizes) if sizes else 10.0


def _line_is_bold(line: PdfTextLine) -> bool:
    weighted_total = 0
    weighted_bold = 0
    for span in line.spans:
        weight = max(1, len(str(span.get("text", ""))))
        weighted_total += weight
        if int(span.get("flags") or 0) & 16 or "bold" in str(span.get("font", "")).lower():
            weighted_bold += weight
    return bool(weighted_total and weighted_bold / weighted_total >= 0.5)


def _looks_like_heading(value: str) -> bool:
    letters = [character for character in value if character.isalpha()]
    return bool(letters and len(value.strip()) <= 90 and all(character.isupper() for character in letters))


def _starts_with_bullet(value: str) -> bool:
    return value.lstrip().startswith(("\uf0b7", "•", "●", "▪", "- "))


def _starts_new_pdf_paragraph(group: Sequence[PdfTextLine], current: PdfTextLine) -> bool:
    previous = group[-1]
    gap = current.bbox[1] - previous.bbox[3]
    previous_size = _line_dominant_size(previous)
    current_size = _line_dominant_size(current)
    current_text = current.text.strip()
    previous_text = previous.text.strip()

    if gap > max(4.0, min(previous_size, current_size) * 0.45):
        return True
    if abs(current_size - previous_size) > 0.9:
        return True
    if _line_is_bold(current) != _line_is_bold(previous) and (
        _looks_like_heading(current_text) or _looks_like_heading(previous_text)
    ):
        return True
    if _looks_like_heading(current_text) or _looks_like_heading(previous_text):
        return True

    group_left = group[0].bbox[0]
    left_delta = current.bbox[0] - previous.bbox[0]
    if left_delta < -8.0:
        return True
    if left_delta > 14.0 and not _starts_with_bullet(group[0].text):
        return True
    if _starts_with_bullet(current_text) and not _starts_with_bullet(previous_text):
        return True
    if abs(current.bbox[0] - group_left) > 24.0 and not _starts_with_bullet(group[0].text):
        return True
    return False


def _group_pdf_lines_into_paragraphs(lines: Sequence[PdfTextLine]) -> list[PdfTextGroup]:
    groups: list[list[PdfTextLine]] = []
    for line in lines:
        if not groups or _starts_new_pdf_paragraph(groups[-1], line):
            groups.append([line])
        else:
            groups[-1].append(line)
    return [
        PdfTextGroup(
            bbox=_rect_union(line.bbox for line in group),
            lines=tuple(group),
        )
        for group in groups
    ]


def _configure_section_for_editable_pdf_page(
    section: Any,
    page_rect: fitz.Rect,
    content_bbox: Sequence[float],
) -> None:
    page_width = float(page_rect.width)
    page_height = float(page_rect.height)
    section.page_width = Pt(page_width)
    section.page_height = Pt(page_height)

    left = min(90.0, max(18.0, float(content_bbox[0])))
    top = min(90.0, max(18.0, float(content_bbox[1])))
    right = min(90.0, max(18.0, page_width - float(content_bbox[2])))
    # The last glyph's y-position is not a semantic bottom margin. Keeping a
    # compact editing margin prevents small Word font-metric differences from
    # spilling a source page just before the explicit section break.
    bottom = 18.0
    section.left_margin = Pt(left)
    section.right_margin = Pt(right)
    section.top_margin = Pt(top)
    section.bottom_margin = Pt(bottom)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def _section_margin_points(section: Any, name: str) -> float:
    value = getattr(section, name, None)
    return float(getattr(value, "pt", 0.0) or 0.0)


def _should_keep_pdf_line_break(group: PdfTextGroup, line_index: int) -> bool:
    if line_index <= 0:
        return False
    previous_text = group.lines[line_index - 1].text.strip()
    current_text = group.lines[line_index].text.strip()
    if not previous_text or not current_text:
        return True
    # Addresses/signatures are intentionally line-oriented; ordinary wrapped prose
    # should reflow naturally when edited in Word.
    compact_line_block = (
        len(group.lines) >= 2
        and max(len(line.text.strip()) for line in group.lines) <= 60
    )
    return (
        compact_line_block
        or (
            len(previous_text) <= 45
            and len(current_text) <= 45
            and previous_text.endswith((",", ".", ":"))
        )
    )


def _normalized_pdf_span_text(value: str) -> str:
    return (
        str(value)
        .replace("\uf0b7", "•")
        .replace("\uf0a7", "▪")
        .replace("\u00a0", " ")
    )


def _append_pdf_span(paragraph: Any, span: dict[str, Any], *, prefix: str = "") -> None:
    source_text = str(span.get("text", ""))
    text = prefix + _normalized_pdf_span_text(source_text)
    if not text:
        return
    run = paragraph.add_run(text)
    font_name = str(span.get("font") or "").split("+")[-1].strip()
    if any(marker in source_text for marker in ("\uf0b7", "\uf0a7", "•", "●", "▪")):
        # Private-use bullet codepoints commonly come from Symbol/Wingdings PDF
        # fonts. After mapping them to Unicode, use a Unicode-capable Word font.
        font_name = "Arial"
    if font_name:
        run.font.name = font_name
    size = min(72.0, max(5.0, float(span.get("size") or 10.0)))
    run.font.size = Pt(size)
    flags = int(span.get("flags") or 0)
    font_name_lower = font_name.lower()
    run.bold = bool(flags & 16 or "bold" in font_name_lower)
    run.italic = bool(flags & 2 or "italic" in font_name_lower or "oblique" in font_name_lower)

    color_value = int(span.get("color") or 0)
    run.font.color.rgb = RGBColor(
        (color_value >> 16) & 0xFF,
        (color_value >> 8) & 0xFF,
        color_value & 0xFF,
    )


def _populate_pdf_text_paragraph(
    paragraph: Any,
    group: PdfTextGroup,
    *,
    page_rect: fitz.Rect,
    section: Any,
    previous_bottom: Optional[float],
) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_after = Pt(0)
    fmt.line_spacing = 1.0
    left_margin = _section_margin_points(section, "left_margin")
    fmt.left_indent = Pt(max(0.0, group.bbox[0] - left_margin))
    # PDF x1 is the end of the rendered glyphs, not a semantic paragraph edge.
    # Turning it into a Word right indent makes short address lines and headings
    # wrap unnecessarily and can add pages. Let normal paragraphs reflow across
    # the available text width.
    fmt.right_indent = Pt(0)

    if previous_bottom is None:
        fmt.space_before = Pt(max(0.0, group.bbox[1] - _section_margin_points(section, "top_margin")))
    else:
        fmt.space_before = Pt(min(18.0, max(0.0, group.bbox[1] - previous_bottom)))

    group_width = group.bbox[2] - group.bbox[0]
    group_center = (group.bbox[0] + group.bbox[2]) / 2.0
    group_text = group.text.strip()
    if (
        group_width < page_rect.width * 0.55
        and len(group_text) <= 90
        and abs(group_center - page_rect.width / 2.0) < 18.0
    ):
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT

    previous_span_right: Optional[float] = None
    for line_index, line in enumerate(group.lines):
        if line_index:
            separator = "\n" if _should_keep_pdf_line_break(group, line_index) else " "
            paragraph.add_run(separator)
        for span in line.spans:
            span_bbox = span.get("bbox") or (0, 0, 0, 0)
            prefix = ""
            if previous_span_right is not None and float(span_bbox[0]) - previous_span_right > 18.0:
                prefix = "\t"
            _append_pdf_span(paragraph, span, prefix=prefix)
            previous_span_right = float(span_bbox[2])
        previous_span_right = None


def _add_pdf_table_to_docx(document: Any, pdf_table: Any, *, previous_bottom: Optional[float]) -> None:
    data = list(pdf_table.extract() or [])
    if not data:
        return
    column_count = max(len(row or []) for row in data)
    if column_count < 1:
        return

    if previous_bottom is not None:
        spacer = document.add_paragraph()
        spacer.paragraph_format.space_before = Pt(
            min(12.0, max(0.0, float(pdf_table.bbox[1]) - previous_bottom))
        )
        spacer.paragraph_format.space_after = Pt(0)

    table = document.add_table(rows=len(data), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for row_index, source_row in enumerate(data):
        source_row = list(source_row or [])
        for column_index in range(column_count):
            value = source_row[column_index] if column_index < len(source_row) else ""
            cell = table.cell(row_index, column_index)
            cell.text = str(value or "").strip()
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(9.0)
                    if row_index == 0:
                        run.bold = True


def _add_pdf_image_to_docx(
    document: Any,
    image_block: dict[str, Any],
    *,
    bbox: Sequence[float],
    section: Any,
    previous_bottom: Optional[float],
) -> None:
    image_bytes = image_block.get("image")
    if not image_bytes:
        return
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    if previous_bottom is not None:
        paragraph.paragraph_format.space_before = Pt(
            min(12.0, max(0.0, float(bbox[1]) - previous_bottom))
        )

    available_width = (
        float(section.page_width.pt)
        - _section_margin_points(section, "left_margin")
        - _section_margin_points(section, "right_margin")
    )
    width = min(max(1.0, float(bbox[2]) - float(bbox[0])), max(1.0, available_width))
    aspect = max(0.01, (float(bbox[3]) - float(bbox[1])) / max(1.0, float(bbox[2]) - float(bbox[0])))
    paragraph.add_run().add_picture(BytesIO(image_bytes), width=Pt(width), height=Pt(width * aspect))


def _configure_section_for_fixed_pdf_page(section, page_rect: fitz.Rect) -> None:
    """Mirror one PDF page's physical size and orientation in a Word section."""

    width = max(1.0, float(page_rect.width))
    height = max(1.0, float(page_rect.height))
    section.page_width = Pt(width)
    section.page_height = Pt(height)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)
    try:
        section.gutter = Pt(0)
    except (AttributeError, ValueError):
        pass


def _fit_image_to_section(
    section: Any,
    *,
    width_px: int,
    height_px: int,
) -> tuple[Pt, Pt]:
    """Fit an image inside a Word page without altering its aspect ratio."""

    width_px = max(1, int(width_px))
    height_px = max(1, int(height_px))
    available_width = max(1.0, float(section.page_width.pt))
    available_height = max(1.0, float(section.page_height.pt))
    for attr in ("left_margin", "right_margin"):
        value = getattr(section, attr, None)
        if value is not None:
            available_width -= max(0.0, float(value.pt))
    for attr in ("top_margin", "bottom_margin"):
        value = getattr(section, attr, None)
        if value is not None:
            available_height -= max(0.0, float(value.pt))
    available_width = max(1.0, available_width)
    available_height = max(1.0, available_height)

    aspect = width_px / height_px
    fitted_width = available_width
    fitted_height = fitted_width / aspect
    if fitted_height > available_height:
        fitted_height = available_height
        fitted_width = fitted_height * aspect
    return Pt(max(1.0, fitted_width)), Pt(max(1.0, fitted_height))


def _configure_section_for_image(section, width_px: int, height_px: int) -> None:
    # Fit common image documents onto an A4-equivalent portrait/landscape page while
    # preserving image aspect ratio and avoiding the old fixed 6-inch insertion.
    if width_px >= height_px:
        section.page_width = Inches(11.69)
        section.page_height = Inches(8.27)
    else:
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
    section.top_margin = Pt(0)
    section.bottom_margin = Pt(0)
    section.left_margin = Pt(0)
    section.right_margin = Pt(0)
    section.header_distance = Pt(0)
    section.footer_distance = Pt(0)


def _configure_full_page_image_paragraph(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(0)
    fmt.space_after = Pt(0)
    fmt.left_indent = Pt(0)
    fmt.right_indent = Pt(0)
    fmt.first_line_indent = Pt(0)
    fmt.line_spacing = 1


def _flatten_image_to_rgb(image: Image.Image) -> Image.Image:
    if image.mode in {"RGBA", "LA"} or (image.mode == "P" and "transparency" in image.info):
        rgba = image.convert("RGBA")
        background = Image.new("RGBA", rgba.size, "WHITE")
        background.alpha_composite(rgba)
        return background.convert("RGB")
    return image.convert("RGB")


__all__ = [
    "CONVERSION_RULES",
    "ALLOWED_CONVERSION_PAIRS",
    "ConversionArtifact",
    "ConversionBackend",
    "RealConversionBackend",
    "ConvertConfig",
    "ConvertProcessor",
    "ensure_allowed_conversion_pair",
    "plan_output_file_name",
    "build_conversion_instructions",
    "convert_document",
]