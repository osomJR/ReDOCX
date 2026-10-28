"""Flowing Word paragraphs, semantic headings, columns and native tables.

Fixed geometry is kept inside diagram groups and footers. Ordinary body text is
not converted into hundreds of page-positioned text boxes.
"""
from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import Path
import statistics
import unicodedata

import cv2
import fitz
import numpy as np
from PIL import Image
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT,WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt,RGBColor
from lxml import etree

from . import native


@dataclass
class Paragraph:
    bbox: tuple
    lines: list

    @property
    def spans(self):return [span for line in self.lines for span in line['spans']]
    @property
    def text(self):return ' '.join(''.join(s['text'] for s in line['spans']).strip() for line in self.lines)
    @property
    def size(self):return statistics.median(s['size'] for s in self.spans)


def _rtl(text):
    directions=[unicodedata.bidirectional(c) for c in text if c.isalpha()]
    return bool(directions and sum(d in ('R','AL') for d in directions)>len(directions)/2)


def paragraphs(lines):
    clusters=[]
    for line in sorted(lines,key=lambda l:(l['bbox'][0],l['bbox'][1])):
        for cluster in clusters:
            if abs(cluster[0]['bbox'][0]-line['bbox'][0])<5:
                cluster.append(line);break
        else:clusters.append([line])
    result=[]
    for cluster in clusters:
        current=None
        for line in sorted(cluster,key=lambda l:l['bbox'][1]):
            rect=fitz.Rect(line['bbox'])
            size=statistics.median(s['size'] for s in line['spans'])
            if (current is not None and abs(current.size-size)<.6 and
                    0<=rect.y0-current.bbox[3]<size*.9 and
                    bool(current.spans[-1]['flags']&16)==bool(line['spans'][0]['flags']&16)):
                current.lines.append(line);current.bbox=tuple(fitz.Rect(current.bbox)|rect)
            else:
                current=Paragraph(tuple(rect),[line]);result.append(current)
    return sorted(result,key=lambda p:(p.bbox[1],p.bbox[0]))


def _raster_diagrams(page,blocks):
    kind=page.parent.xref_get_key(page.xref,'ReDOCXOCR')[1]
    if kind!='true':return []
    regions=[]
    for block in blocks:
        if block['type']!=1 or fitz.Rect(block['bbox']).get_area()<page.rect.get_area()*.7:continue
        with Image.open(BytesIO(block['image'])) as image:
            array=np.array(image.convert('RGB'))
        hsv=cv2.cvtColor(array,cv2.COLOR_RGB2HSV)
        mask=((hsv[:,:,1]>55)&(hsv[:,:,2]<250)).astype(np.uint8)*255
        count,labels,stats,centers=cv2.connectedComponentsWithStats(mask,8)
        boxes=[]
        for x,y,w,h,area in stats[1:]:
            if area<array.shape[0]*array.shape[1]*.0004 or w<8 or h<8:continue
            if w>array.shape[1]*.85 or area>array.shape[0]*array.shape[1]*.2:continue
            boxes.append(fitz.Rect(int(x),int(y),int(x+w),int(y+h)))
        groups=[]
        for rect in boxes:
            for index,old in enumerate(groups):
                if min(old.y1,rect.y1)-max(old.y0,rect.y0)>=0:
                    groups[index]=old|rect;break
            else:groups.append(rect)
        image_rect=fitz.Rect(block['bbox'])
        sx,sy=image_rect.width/array.shape[1],image_rect.height/array.shape[0]
        for rect in groups:
            r=fitz.Rect(image_rect.x0+(rect.x0-35)*sx,image_rect.y0+(rect.y0-35)*sy,
                        image_rect.x0+(rect.x1+35)*sx,image_rect.y0+(rect.y1+40)*sy)&page.rect
            regions.append(tuple(r))
    return regions


def _diagram(document,page,rect,paragraph,left_margin):
    rect=fitz.Rect(rect)
    available=document.sections[-1].page_width.pt-document.sections[-1].left_margin.pt-document.sections[-1].right_margin.pt
    scale=min(1,available/rect.width)
    paragraph.paragraph_format.left_indent=Pt(max(0,rect.x0-left_margin)*scale)
    paragraph.paragraph_format.space_after=Pt(0)
    pict=OxmlElement('w:pict');paragraph.add_run()._r.append(pict)
    group=native._vml(pict,'group',{'style':f'width:{rect.width*scale}pt;height:{rect.height*scale}pt',
        'coordorigin':'0,0','coordsize':f'{round(rect.width*100)},{round(rect.height*100)}'})

    def position(r):
        r=fitz.Rect(r)
        return (f'position:absolute;left:{(r.x0-rect.x0)*100};top:{(r.y0-rect.y0)*100};'
                f'width:{max(.1,r.width)*100};height:{max(.1,r.height)*100}')

    # OCR backgrounds have already had their recognized glyphs removed. Crop
    # just the figure artwork; the labels below are editable native Word text.
    for block in page.get_text('dict')['blocks']:
        if block['type']!=1 or not fitz.Rect(block['bbox']).intersects(rect):continue
        bounds=fitz.Rect(block['bbox']);intersection=bounds&rect
        with Image.open(BytesIO(block['image'])) as image:
            crop=image.crop((round((intersection.x0-bounds.x0)/bounds.width*image.width),
                             round((intersection.y0-bounds.y0)/bounds.height*image.height),
                             round((intersection.x1-bounds.x0)/bounds.width*image.width),
                             round((intersection.y1-bounds.y0)/bounds.height*image.height)))
            stream=BytesIO();crop.save(stream,'PNG')
        rid,_=document.part.get_or_add_image(BytesIO(stream.getvalue()))
        shape=native._vml(group,'rect',{'style':position(intersection),'stroked':'f'})
        data=native._vml(shape,'imagedata',{});data.set(qn('r:id'),rid)
    for drawing in page.get_drawings():
        if not rect.contains(drawing['rect']):continue
        for item in drawing['items']:
            if item[0]=='re':
                native._vml(group,'rect',{'style':position(item[1]),'filled':'t' if drawing.get('fill') else 'f',
                    'stroked':'t' if drawing.get('color') else 'f','fillcolor':'#'+native.color(drawing.get('fill'),'FFFFFF'),
                    'strokecolor':'#'+native.color(drawing.get('color')),'strokeweight':str(drawing.get('width') or .5)+'pt'})
            elif item[0]=='l':
                a,b=item[1:3]
                native._vml(group,'line',{'style':'position:absolute',
                    'from':f'{(a.x-rect.x0)*100},{(a.y-rect.y0)*100}',
                    'to':f'{(b.x-rect.x0)*100},{(b.y-rect.y0)*100}',
                    'strokecolor':'#'+native.color(drawing.get('color')),'strokeweight':str(drawing.get('width') or .5)+'pt'})
    for block in page.get_text('dict')['blocks']:
        if block['type']!=0:continue
        for line in block['lines']:
            r=fitz.Rect(line['bbox'])
            if not rect.contains(fitz.Point((r.x0+r.x1)/2,(r.y0+r.y1)/2)):continue
            r=fitz.Rect(r.x0,r.y0-1,r.x1+10,r.y1+10)
            shape=native._vml(group,'rect',{'style':position(r),'stroked':'f','filled':'f'})
            box=native._vml(shape,'textbox',{'inset':'0,0,0,0'})
            content=OxmlElement('w:txbxContent');box.append(content)
            node=OxmlElement('w:p');content.append(node)
            for span in line['spans']:
                run=OxmlElement('w:r');props=OxmlElement('w:rPr');run.append(props)
                fonts=OxmlElement('w:rFonts')
                for key in ('ascii','hAnsi','cs','eastAsia'):fonts.set(qn('w:'+key),native.family(span))
                props.append(fonts)
                size=OxmlElement('w:sz');size.set(qn('w:val'),str(round(span['size']*scale*2)));props.append(size)
                color=OxmlElement('w:color');color.set(qn('w:val'),native.color(span['color']));props.append(color)
                if span['flags']&16:props.append(OxmlElement('w:b'))
                text=OxmlElement('w:t');text.text=span['text'];text.set('{http://www.w3.org/XML/1998/namespace}space','preserve')
                run.append(text);node.append(run)


def _paragraph(document,model,body_size,gap=0,page=None):
    style='Normal'
    bold=bool(model.spans and all(s['flags']&16 for s in model.spans))
    if len(model.text)<150:
        if model.size>=body_size*1.6:style='Heading 1'
        elif model.size>=body_size*1.15 and bold:style='Heading 2'
    paragraph=document.add_paragraph(style=style)
    fmt=paragraph.paragraph_format
    fmt.space_before=Pt(max(0,min(gap,40)))
    fmt.space_after=Pt(0);fmt.line_spacing=1.0
    fmt.keep_with_next=style.startswith('Heading')
    if _rtl(model.text):
        bidi=OxmlElement('w:bidi');paragraph._p.get_or_add_pPr().append(bidi)
    for index,line in enumerate(model.lines):
        if index:paragraph.add_run(' ')
        for span in line['spans']:
            run=paragraph.add_run(span['text'])
            run.font.name=native.family(span);run.font.size=Pt(span['size'])
            run.font.bold=bool(span['flags']&16);run.font.italic=bool(span['flags']&2)
            run.font.color.rgb=RGBColor.from_string(native.color(span['color']))
            if _rtl(span['text']):run._r.get_or_add_rPr().append(OxmlElement('w:rtl'))
    # A dark banner must not turn white native text invisible when its raster
    # background is decomposed into flowing body content.
    if page and any(s.get('color',0)>0xDDDDDD for s in model.spans):
        background=None
        for block in page.get_text('dict')['blocks']:
            if block['type']!=1:continue
            bounds=fitz.Rect(block['bbox']);r=fitz.Rect(model.bbox)
            if not bounds.contains(r):continue
            with Image.open(BytesIO(block['image'])) as image:
                point=(max(0,min(image.width-1,round((r.x0-bounds.x0)/bounds.width*image.width))),
                       max(0,min(image.height-1,round((r.y0-bounds.y0)/bounds.height*image.height))))
                background=image.convert('RGB').getpixel(point)
            break
        if background and sum(background)<600:
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),''.join(f'{c:02X}' for c in background));paragraph._p.get_or_add_pPr().append(shd)
    return paragraph


def _table(document,source,left_margin):
    target=document.add_table(rows=len(source.rows),cols=len(source.rows[0]))
    target.autofit=False
    props=target._tbl.tblPr
    indent=OxmlElement('w:tblInd');indent.set(qn('w:w'),str(round(max(0,source.bbox[0]-left_margin)*20)));indent.set(qn('w:type'),'dxa');props.append(indent)
    borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        node=OxmlElement('w:'+side);node.set(qn('w:val'),'nil');borders.append(node)
    props.append(borders)
    available=document.sections[-1].page_width.pt-document.sections[-1].left_margin.pt-document.sections[-1].right_margin.pt
    factor=min(1,available/(source.bbox[2]-source.bbox[0]))
    for i,column in enumerate(target.columns):column.width=Pt((source.edges[i+1]-source.edges[i])*factor)
    for i,values in enumerate(source.rows):
        row=target.rows[i]
        row.height=Pt(source.row_edges[i+1]-source.row_edges[i]);row.height_rule=WD_ROW_HEIGHT_RULE.AT_LEAST
        if i==0:row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j,value in enumerate(values):
            cell=row.cells[j];cell.width=Pt((source.edges[j+1]-source.edges[j])*factor)
            cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER;cell.text=value
            shd=OxmlElement('w:shd');shd.set(qn('w:fill'),source.fills[i]);cell._tc.get_or_add_tcPr().append(shd)
            paragraph=cell.paragraphs[0]
            paragraph.paragraph_format.space_after=Pt(0);paragraph.paragraph_format.line_spacing=1.0
            if _rtl(value):paragraph._p.get_or_add_pPr().append(OxmlElement('w:bidi'))
            for run in paragraph.runs:
                run.font.name='DejaVu Sans';run.font.size=Pt(source.font_size)
                run.font.bold=source.bold[i];run.font.color.rgb=RGBColor.from_string(source.colors[i])
    for r0,c0,r1,c1 in source.merges:target.cell(r0,c0).merge(target.cell(r1,c1))
    return target


def _set_columns(section,count):
    cols=section._sectPr.find(qn('w:cols'))
    if cols is None:cols=OxmlElement('w:cols');section._sectPr.append(cols)
    cols.set(qn('w:num'),str(count));cols.set(qn('w:space'),'400')


def _flush_text(document,items,body_size,left,width,page,previous):
    if not items:return previous
    # Infer parallel columns only from multiple, overlapping narrow paragraphs.
    narrow=[p for p in items if p.bbox[2]-p.bbox[0]<width*.58]
    lhs=[p for p in narrow if p.bbox[0]<left+width*.3]
    rhs=[p for p in narrow if p.bbox[0]>left+width*.45]
    common=(min(max((p.bbox[3] for p in lhs),default=0),max((p.bbox[3] for p in rhs),default=0))-
            max(min((p.bbox[1] for p in lhs),default=0),min((p.bbox[1] for p in rhs),default=0)))
    if len(lhs)>=3 and len(rhs)>=3 and common>body_size*3:
        top=min(p.bbox[1] for p in lhs+rhs);bottom=max(p.bbox[3] for p in lhs+rhs)
        before=[p for p in items if p not in lhs+rhs and p.bbox[3]<=top]
        after=[p for p in items if p not in lhs+rhs and p not in before]
        for model in before:
            _paragraph(document,model,body_size,model.bbox[1]-previous,page);previous=model.bbox[3]
        section=document.add_section(WD_SECTION.CONTINUOUS);_set_columns(section,2)
        for index,column in enumerate((lhs,rhs)):
            if index:document.add_paragraph().add_run().add_break(WD_BREAK.COLUMN)
            prior=top
            for model in sorted(column,key=lambda p:p.bbox[1]):
                _paragraph(document,model,body_size,model.bbox[1]-prior,page);prior=model.bbox[3]
        section=document.add_section(WD_SECTION.CONTINUOUS);_set_columns(section,1)
        previous=bottom
        for model in sorted(after,key=lambda p:p.bbox[1]):
            _paragraph(document,model,body_size,model.bbox[1]-previous,page);previous=model.bbox[3]
    else:
        for model in sorted(items,key=lambda p:(p.bbox[1],p.bbox[0])):
            _paragraph(document,model,body_size,model.bbox[1]-previous,page);previous=model.bbox[3]
    return previous


def convert(source: Path,output: Path):
    document=Document()
    normal=document.styles['Normal'];normal.font.name='DejaVu Sans';normal.font.size=Pt(10)
    normal.paragraph_format.space_after=Pt(0)
    with fitz.open(source) as pdf:
        for index,page in enumerate(pdf):
            section=document.sections[0] if index==0 else document.add_section(WD_SECTION.NEW_PAGE)
            section.page_width,section.page_height=Pt(page.rect.width),Pt(page.rect.height)
            _set_columns(section,1)
            data=page.get_text('dict')['blocks'];tables=native.extract_tables(page)
            lines=[line for block in data if block['type']==0 for line in block['lines'] if line['spans']]
            sizes=[s['size'] for line in lines for s in line['spans'] if s['text'].strip()]
            body=statistics.median(sizes) if sizes else 10
            diagrams=[tuple(a['bbox']) for a in native.sheet_assets(page,tables) if a.get('diagram')]
            diagrams.extend(_raster_diagrams(page,data))
            diagrams=[r for i,r in enumerate(diagrams) if not any(fitz.Rect(old).contains(r) for old in diagrams[:i])]
            footer=[line for line in lines if line['bbox'][1]>page.rect.height*.92]
            body_lines=[line for line in lines if line not in footer and not native.inside(line['bbox'],tables)
                        and not any(fitz.Rect(r).contains(fitz.Point((line['bbox'][0]+line['bbox'][2])/2,
                              (line['bbox'][1]+line['bbox'][3])/2)) for r in diagrams)]
            left=max(15,min([line['bbox'][0] for line in body_lines] or [40]))
            section.left_margin=Pt(left);section.right_margin=Pt(max(20,min(50,page.rect.width*.08)))
            top=max(12,min([line['bbox'][1] for line in body_lines] or [40])-3)
            section.top_margin=Pt(top);section.bottom_margin=Pt(30)
            events=[(model.bbox[1],'text',model) for model in paragraphs(body_lines)]
            events.extend((table.bbox[1],'table',table) for table in tables)
            events.extend((rect[1],'diagram',rect) for rect in diagrams)
            is_ocr=page.parent.xref_get_key(page.xref,'ReDOCXOCR')[1]=='true'
            for block in data:
                if block['type']!=1:continue
                if is_ocr and fitz.Rect(block['bbox']).get_area()>page.rect.get_area()*.7:continue
                if any(fitz.Rect(r).contains(block['bbox']) for r in diagrams):continue
                events.append((block['bbox'][1],'picture',block))
            width=section.page_width.pt-section.left_margin.pt-section.right_margin.pt
            prior=top;pending=[]
            for _,kind,item in sorted(events,key=lambda e:e[0]):
                if kind=='text':pending.append(item);continue
                width=section.page_width.pt-section.left_margin.pt-section.right_margin.pt
                prior=_flush_text(document,pending,body,left,width,page,prior);pending=[]
                if kind=='table':
                    gap=document.add_paragraph();gap.paragraph_format.space_before=Pt(max(0,min(25,item.bbox[1]-prior)))
                    gap.paragraph_format.line_spacing=Pt(1);gap.paragraph_format.space_after=Pt(0)
                    _table(document,item,left);prior=item.bbox[3]
                elif kind=='diagram':
                    paragraph=document.add_paragraph();paragraph.paragraph_format.space_before=Pt(max(0,min(25,item[1]-prior)))
                    _diagram(document,page,item,paragraph,left);prior=item[3]
                else:
                    r=fitz.Rect(item['bbox']);paragraph=document.add_paragraph()
                    paragraph.paragraph_format.space_before=Pt(max(0,min(25,r.y0-prior)))
                    paragraph.add_run().add_picture(BytesIO(item['image']),width=Pt(min(width,r.width)))
                    prior=r.y1
            _flush_text(document,pending,body,left,width if events else 400,page,prior)
            if footer:
                section.footer.is_linked_to_previous=False
                paragraph=section.footer.paragraphs[0]
                paragraph.paragraph_format.space_after=Pt(0)
                for li,line in enumerate(footer):
                    if li:paragraph.add_run('\t')
                    for span in line['spans']:
                        run=paragraph.add_run(span['text']);run.font.name=native.family(span)
                        run.font.size=Pt(span['size']);run.font.color.rgb=RGBColor.from_string(native.color(span['color']))
                section.footer_distance=Pt(20)
    document.core_properties.keywords='redocx-native-flow-v2'
    document.save(output)
