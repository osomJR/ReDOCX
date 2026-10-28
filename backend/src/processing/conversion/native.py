"""Editable document reconstruction; page screenshots are never Office text."""

from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path
import re
import statistics
from functools import lru_cache

import fitz
from PIL import Image, ImageDraw
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.drawing.image import Image as SheetImage
from pptx import Presentation
from pptx.util import Pt as SlidePt
from pptx.dml.color import RGBColor as SlideColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR
from lxml import etree
from fontTools.ttLib import TTFont

from .quality import assert_text, rgb


@dataclass
class Table:
    bbox: tuple
    rows: list[list[str]]
    edges: list[float]
    row_edges: list[float]
    font_size: float
    fills: list[str]
    colors: list[str]
    bold: list[bool]
    merges: list[tuple[int,int,int,int]] = field(default_factory=list)


def color(value, default="000000"):
    if value is None:
        return default
    if isinstance(value, int):
        return f"{value:06X}"
    return "".join(f"{round(max(0,min(1,float(c)))*255):02X}" for c in value[:3])


def family(span):
    font = re.sub(r"^[A-Z]{6}\+", "", span.get("font", "DejaVu Sans"))
    font = re.sub(
        r"[-,]?(BoldItalic|BoldOblique|Bold|Italic|Oblique|Regular)$", "", font
    )
    return {
        "DejaVuSans": "DejaVu Sans",
        "DejaVuSerif": "DejaVu Serif",
        "Times-Roman": "Times New Roman",
        "Helvetica": "Arial",
    }.get(font, font)


def word_rows(words):
    rows = []
    for word in sorted(words, key=lambda w: ((w[1] + w[3]) / 2, w[0])):
        center = (word[1] + word[3]) / 2
        if (
            rows
            and abs(center - statistics.mean((w[1] + w[3]) / 2 for w in rows[-1])) < 3
        ):
            rows[-1].append(word)
        else:
            rows.append([word])
    return [sorted(row, key=lambda w: w[0]) for row in rows]


def _spans(page):
    return [
        s
        for b in page.get_text("dict")["blocks"]
        if b["type"] == 0
        for line in b["lines"]
        for s in line["spans"]
        if s.get("text", "").strip()
    ]


def extract_tables(page) -> list[Table]:
    spans = _spans(page)
    drawings = page.get_drawings()

    def make(bbox, rows, edges, row_edges, merges=None):
        fills, colors, bold, sizes = [], [], [], []
        for y0, y1 in zip(row_edges, row_edges[1:]):
            yy = (y0 + y1) / 2
            region = [
                s
                for s in spans
                if y0 <= (s["bbox"][1] + s["bbox"][3]) / 2 <= y1
                and bbox[0] <= s["bbox"][0] < bbox[2]
            ]
            fill = "FFFFFF"
            for d in drawings:
                r = d["rect"]
                if (
                    d.get("fill")
                    and r.y0 <= yy <= r.y1
                    and r.width >= (bbox[2] - bbox[0]) * 0.8
                ):
                    fill = color(d["fill"])
            fills.append(fill)
            colors.append(color(region[0]["color"]) if region else "000000")
            bold.append(bool(region and region[0].get("flags", 0) & 16))
            sizes.extend(s["size"] for s in region)
        return Table(
            tuple(bbox),
            rows,
            edges,
            row_edges,
            statistics.median(sizes) if sizes else 10,
            fills,
            colors,
            bold,
            merges or [],
        )

    tables = []
    try:
        for table in page.find_tables().tables:
            cells = sorted({tuple(round(v,2) for v in cell) for cell in table.cells if cell})
            edges = sorted({r[0] for r in cells} | {r[2] for r in cells})
            row_edges = sorted({r[1] for r in cells} | {r[3] for r in cells})
            if len(edges)<3 or len(row_edges)<3:
                continue
            rows = [['']*(len(edges)-1) for _ in range(len(row_edges)-1)]
            merges=[]
            for rect in cells:
                c0,c1=edges.index(rect[0]),edges.index(rect[2])
                r0,r1=row_edges.index(rect[1]),row_edges.index(rect[3])
                rows[r0][c0]=page.get_textbox(fitz.Rect(rect)).strip()
                if r1-r0>1 or c1-c0>1:
                    merges.append((r0,c0,r1-1,c1-1))
            tables.append(make(table.bbox,rows,edges,row_edges,merges))
    except (AttributeError, ValueError, RuntimeError):
        pass

    candidates = []
    for row in word_rows(page.get_text("words")):
        if any(
            fitz.Rect(t.bbox).contains(fitz.Point(row[0][0], row[0][1])) for t in tables
        ):
            continue
        cells = []
        for word in row:
            if cells and word[0] - cells[-1][-1][2] < max(
                18, (word[3] - word[1]) * 1.6
            ):
                cells[-1].append(word)
            else:
                cells.append([word])
        candidates.append(cells)

    i = 0
    while i < len(candidates):
        row = candidates[i]
        if len(row) < 2:
            i += 1
            continue
        group = [row]
        j = i + 1
        while j < len(candidates):
            nxt, prev = candidates[j], group[-1]
            gap = nxt[0][0][1] - prev[0][0][1]
            if (
                len(nxt) != len(row)
                or not 0 < gap < 65
                or any(abs(a[0][0] - b[0][0]) > 9 for a, b in zip(row, nxt))
            ):
                break
            group.append(nxt)
            j += 1
        if len(group) >= 3:
            left = min(c[0][0] for r in group for c in r) - 6
            right = max(c[-1][2] for r in group for c in r) + 6
            centers = [
                statistics.mean((w[1] + w[3]) / 2 for c in r for w in c) for r in group
            ]
            boundaries = [
                (centers[k] + centers[k + 1]) / 2 for k in range(len(centers) - 1)
            ]
            row_edges = [
                centers[0] - (boundaries[0] - centers[0]),
                *boundaries,
                centers[-1] + (centers[-1] - boundaries[-1]),
            ]
            for d in drawings:
                r = d["rect"]
                if (
                    d.get("fill")
                    and r.y0 < centers[0] < r.y1
                    and r.width > right - left
                ):
                    left, right = r.x0, r.x1
                    row_edges[0] = r.y0
                    break
            for d in drawings:
                r = d["rect"]
                if (
                    d.get("fill")
                    and r.y0 < centers[-1] < r.y1
                    and r.width >= (right - left) * 0.9
                ):
                    row_edges[-1] = r.y1
            edges = [left] + [c[0][0] - 6 for c in group[0][1:]] + [right]
            data = [[" ".join(w[4] for w in c) for c in r] for r in group]
            tables.append(
                make((left, row_edges[0], right, row_edges[-1]), data, edges, row_edges)
            )
            i = j
        else:
            i += 1
    return sorted(tables, key=lambda t: t.bbox[1])


def inside(rect, tables):
    rect = fitz.Rect(rect)
    return any(
        fitz.Rect(t.bbox).contains(
            fitz.Point((rect.x0 + rect.x1) / 2, (rect.y0 + rect.y1) / 2)
        )
        for t in tables
    )


def page_text(page):
    text = page.get_text()
    if text.strip():
        return text
    pix = page.get_pixmap(dpi=100, alpha=False)
    image = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    if image.convert("L").getextrema()[0] > 245:
        return ""
    try:
        text = page.get_text(
            textpage=page.get_textpage_ocr(language="eng", dpi=200, full=True)
        )
    except Exception as exc:
        raise RuntimeError("OCR is required for this scanned PDF.") from exc
    if not text.strip():
        raise RuntimeError(
            "Scanned PDF contains no reliably recoverable editable text."
        )
    return text


def cell_value(text):
    # Keep identifiers, dates, currencies and ambiguous localized numbers as text.
    if re.fullmatch(r"-?(?:0|[1-9]\d{0,14})(?:\.\d+)?", text.strip()):
        value = text.strip()
        if len(value.replace("-", "").replace(".", "")) <= 15:
            return float(value) if "." in value else int(value)
    return text


def sheet_assets(page, tables):
    """Keep source pictures and isolated diagrams, without a page screenshot."""
    assets = [b for b in page.get_text("dict")["blocks"] if b["type"] == 1]
    groups = []
    for drawing in page.get_drawings():
        rect = fitz.Rect(drawing["rect"])
        if rect.width < 2 or rect.height < 2 or inside(rect, tables):
            continue
        for group in groups:
            old = group[0]
            if min(old.y1, rect.y1) - max(old.y0, rect.y0) >= 0:
                group[0] = old | rect
                break
        else:
            groups.append([rect])
    spans = _spans(page)
    for group in groups:
        rect = group[0]
        if rect.get_area() > page.rect.get_area() * 0.6:
            continue
        nearby = fitz.Rect(rect.x0 - 14, rect.y0 - 25, rect.x1 + 14, rect.y1 + 25)
        for span in spans:
            r = fitz.Rect(span["bbox"])
            if nearby.contains(r) and not inside(r, tables):
                rect |= r
        rect = fitz.Rect(rect.x0 - 3, rect.y0 - 3, rect.x1 + 3, rect.y1 + 3) & page.rect
        if any(fitz.Rect(t.bbox).intersects(rect) for t in tables):
            continue
        assets.append(
            {
                "bbox": tuple(rect),
                "diagram": True,
                "image": page.get_pixmap(clip=rect, dpi=144, alpha=False).tobytes(
                    "png"
                ),
            }
        )
    return sorted(assets, key=lambda x: x["bbox"][1])


def convert_xlsx(source: Path, output: Path):
    from . import spreadsheet_semantics as semantics
    book=Workbook();book.remove(book.active)
    caches={}
    with fitz.open(source) as pdf:
        for index,page in enumerate(pdf,1):
            sheet=book.create_sheet(f'Page {index}')
            tables,data=extract_tables(page),page.get_text('dict')
            charts=semantics.chart_evidence(page,tables)
            events=[(t.bbox[1],'table',(ti,t)) for ti,t in enumerate(tables)]
            for asset in sheet_assets(page,tables):
                if asset.get('diagram') and any(fitz.Rect(asset['bbox']).intersects(ev.bbox) for ev in charts):
                    continue
                events.append((asset['bbox'][1],'image',asset))
            events.extend((ev.bbox[1],'chart',ev) for ev in charts)
            for block in data['blocks']:
                if block['type']==0:
                    for line in block['lines']:
                        if not inside(line['bbox'],tables):events.append((line['bbox'][1],'text',line))
            row=1;starts={};pending=[];sheet_formulas={}
            for _,kind,item in sorted(events,key=lambda e:e[0]):
                if kind=='table':
                    ti,table=item;starts[ti]=row
                    for ri,values in enumerate(table.rows):
                        for ci,value in enumerate(values,1):
                            cell=sheet.cell(row,ci,cell_value(value))
                            if isinstance(cell.value,str):cell.data_type='s'
                            if isinstance(cell.value,(int,float)) and '.' in value:
                                cell.number_format='0.'+'0'*len(value.split('.')[-1])
                            cell.font=Font(name='DejaVu Sans',size=table.font_size,bold=table.bold[ri],color=table.colors[ri])
                            cell.fill=PatternFill('solid',fgColor=table.fills[ri])
                            cell.alignment=Alignment(vertical='center',wrap_text=True)
                        sheet.row_dimensions[row].height=max(18,table.row_edges[ri+1]-table.row_edges[ri])
                        row+=1
                    for r0,c0,r1,c1 in table.merges:
                        sheet.merge_cells(start_row=starts[ti]+r0,start_column=c0+1,
                                          end_row=starts[ti]+r1,end_column=c1+1)
                    formulas=semantics.formulas_for(table,starts[ti])
                    semantics.apply_formulas(sheet,formulas);sheet_formulas.update(formulas)
                    row+=1
                elif kind=='image':
                    picture=SheetImage(BytesIO(item['image']));factor=min(1,480/picture.width)
                    picture.width,picture.height=picture.width*factor,picture.height*factor
                    sheet.add_image(picture,f'A{row}');row+=int(picture.height/18)+2
                elif kind=='chart':
                    pending.append((item,f'A{row}'));row+=16
                else:
                    value=''.join(span['text'] for span in item['spans'])
                    cell=sheet.cell(row,1,value);cell.data_type='s'
                    if item['spans']:
                        span=item['spans'][0]
                        cell.font=Font(name=family(span),size=span['size'],color=color(span['color']),bold=bool(span['flags']&16))
                    row+=1
            for chart,anchor in pending:
                semantics.add_chart(sheet,chart,starts[chart.table_index],anchor)
            caches[index]=sheet_formulas
            for ci in range(1,max(4,sheet.max_column)+1):
                sheet.column_dimensions[sheet.cell(1,ci).column_letter].width=24
            sheet.freeze_panes='A2';sheet.sheet_view.showGridLines=False
            sheet.page_setup.orientation='landscape' if page.rect.width>page.rect.height else 'portrait'
            sheet.page_setup.fitToWidth,sheet.page_setup.fitToHeight=1,0
            sheet.sheet_properties.pageSetUpPr.fitToPage=True
    book.calculation.fullCalcOnLoad=True
    book.save(output)
    semantics.write_formula_caches(output,caches)


def validate_xlsx(source,output):
    from . import spreadsheet_semantics as semantics
    book=load_workbook(output,data_only=False)
    cached=load_workbook(output,data_only=True)
    try:
        with fitz.open(source) as pdf:
            if len(book.worksheets)!=len(pdf):raise RuntimeError('PDF-to-XLSX changed page/sheet count.')
            for page,sheet,values_sheet in zip(pdf,book.worksheets,cached.worksheets):
                tables=extract_tables(page);charts=semantics.chart_evidence(page,tables)
                assets=[a for a in sheet_assets(page,tables) if not(a.get('diagram') and
                    any(fitz.Rect(a['bbox']).intersects(ev.bbox) for ev in charts))]
                if len(assets)!=len(sheet._images):raise RuntimeError('PDF-to-XLSX lost source artwork.')
                from .quality import assert_pixels
                for expected,actual in zip(assets,sheet._images):
                    with Image.open(BytesIO(expected['image'])) as a,Image.open(BytesIO(actual._data())) as b:
                        assert_pixels(a,b,label='PDF-to-XLSX image')
                text=[]
                for row in sheet:
                    for cell in row:
                        value=values_sheet[cell.coordinate].value if cell.data_type=='f' else cell.value
                        if value is None:continue
                        if isinstance(value,(int,float)) and re.fullmatch(r'0\.0+',cell.number_format):
                            value=f'{value:.{len(cell.number_format)-2}f}'
                        text.append(str(value))
                assert_text(page_text(page),'\n'.join(text),label='PDF-to-XLSX')
                starts={}
                for ti,table in enumerate(tables):
                    for row in range(1,sheet.max_row-len(table.rows)+2):
                        valid=True;formulas=semantics.formulas_for(table,row)
                        for i,values in enumerate(table.rows):
                            for j,value in enumerate(values):
                                cell=sheet.cell(row+i,j+1);actual=values_sheet.cell(row+i,j+1).value if cell.data_type=='f' else cell.value
                                expected=cell_value(value)
                                if str(actual if actual is not None else '')!=str(expected):valid=False
                                if isinstance(expected,(int,float)) and cell.data_type not in ('n','f'):valid=False
                                if cell.data_type=='f' and (cell.coordinate not in formulas or cell.value!=formulas[cell.coordinate][0]):valid=False
                        if valid:
                            starts[ti]=row;break
                    else:raise RuntimeError('PDF-to-XLSX did not preserve an editable table and numeric values.')
                    for r0,c0,r1,c1 in table.merges:
                        from openpyxl.utils import get_column_letter
                        address=f'{get_column_letter(c0+1)}{starts[ti]+r0}:{get_column_letter(c1+1)}{starts[ti]+r1}'
                        if address not in {str(x) for x in sheet.merged_cells.ranges}:
                            raise RuntimeError('PDF-to-XLSX lost a merged table cell.')
                if len(charts)!=len(sheet._charts):raise RuntimeError('PDF-to-XLSX lost an editable chart.')
                from openpyxl.chart import Reference
                for evidence,chart in zip(charts,sheet._charts):
                    start=starts[evidence.table_index]
                    expected=str(Reference(sheet,min_col=evidence.value_column+1,min_row=start+1,max_row=start+len(evidence.values)))
                    if chart.series[0].val.numRef.f!=expected:
                        raise RuntimeError('Editable chart is not linked to the reconstructed numeric cells.')
    finally:
        book.close();cached.close()


def _ppt_style(shape, drawing):
    shape.shadow.inherit = False
    for effect in shape._element.xpath(".//a:effectRef"):
        effect.set("idx", "0")
    if hasattr(shape, "fill"):
        if drawing.get("fill") is None:
            shape.fill.background()
        else:
            shape.fill.solid()
            shape.fill.fore_color.rgb = SlideColor.from_string(color(drawing["fill"]))
    if drawing.get("color") is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = SlideColor.from_string(color(drawing["color"]))
        shape.line.width = SlidePt(max(0.1, drawing.get("width") or 1))


def convert_pptx(source, output):
    from .page_geometry import slide_reference
    with slide_reference(Path(source)) as reference:
        _convert_pptx_uniform(reference,output)


def _convert_pptx_uniform(source, output):
    deck = Presentation()
    with fitz.open(source) as pdf:
        deck.slide_width, deck.slide_height = SlidePt(pdf[0].rect.width), SlidePt(
            pdf[0].rect.height
        )
        for page in pdf:
            if (
                abs(page.rect.width - pdf[0].rect.width) > 1
                or abs(page.rect.height - pdf[0].rect.height) > 1
            ):
                raise RuntimeError(
                    "PowerPoint requires one slide size; mixed PDF page sizes need normalization."
                )
            slide = deck.slides.add_slide(deck.slide_layouts[6])
            tables = extract_tables(page)
            for d in page.get_drawings():
                if inside(d["rect"], tables):
                    continue
                for item in d["items"]:
                    if item[0] == "re":
                        r = item[1]
                        shape = slide.shapes.add_shape(
                            MSO_SHAPE.RECTANGLE,
                            SlidePt(r.x0),
                            SlidePt(r.y0),
                            SlidePt(r.width),
                            SlidePt(r.height),
                        )
                    elif item[0] == "l":
                        a, b = item[1:3]
                        shape = slide.shapes.add_connector(
                            MSO_CONNECTOR.STRAIGHT,
                            SlidePt(a.x),
                            SlidePt(a.y),
                            SlidePt(b.x),
                            SlidePt(b.y),
                        )
                    else:
                        # Preserve editable paths, approximating cubic segments densely.
                        if item[0] != "c":
                            continue
                        a, b, c, e = item[1:5]
                        points = []
                        for n in range(17):
                            t = n / 16
                            u = 1 - t
                            points.append(
                                (
                                    u**3 * a.x
                                    + 3 * u * u * t * b.x
                                    + 3 * u * t * t * c.x
                                    + t**3 * e.x,
                                    u**3 * a.y
                                    + 3 * u * u * t * b.y
                                    + 3 * u * t * t * c.y
                                    + t**3 * e.y,
                                )
                            )
                        builder = slide.shapes.build_freeform(
                            round(points[0][0] * 12700), round(points[0][1] * 12700)
                        )
                        builder.add_line_segments(
                            [
                                (round(x * 12700), round(y * 12700))
                                for x, y in points[1:]
                            ],
                            close=False,
                        )
                        shape = builder.convert_to_shape()
                    _ppt_style(shape, d)
            for block in page.get_text("dict")["blocks"]:
                if block["type"] == 1:
                    r = fitz.Rect(block["bbox"])
                    slide.shapes.add_picture(
                        BytesIO(block["image"]),
                        SlidePt(r.x0),
                        SlidePt(r.y0),
                        SlidePt(r.width),
                        SlidePt(r.height),
                    )
                elif block["type"] == 0:
                    for line in block["lines"]:
                        if inside(line["bbox"], tables):
                            continue
                        r = fitz.Rect(line["bbox"])
                        shape = slide.shapes.add_textbox(
                            SlidePt(r.x0),
                            SlidePt(r.y0 - 1),
                            SlidePt(min(page.rect.width - r.x0, r.width + 12)),
                            SlidePt(r.height + 10),
                        )
                        frame = shape.text_frame
                        frame.word_wrap = False
                        frame.margin_left = frame.margin_right = frame.margin_top = (
                            frame.margin_bottom
                        ) = 0
                        for span in line["spans"]:
                            run = frame.paragraphs[0].add_run()
                            run.text = span["text"]
                            run.font.name = family(span)
                            run.font.size = SlidePt(span["size"])
                            run.font.bold = bool(span["flags"] & 16)
                            run.font.italic = bool(span["flags"] & 2)
                            run.font.color.rgb = SlideColor.from_string(
                                color(span["color"])
                            )
                            for link in page.get_links():
                                if link.get("uri") and fitz.Rect(
                                    link["from"]
                                ).intersects(r):
                                    run.hyperlink.address = link["uri"]
            for table in tables:
                x0, y0, x1, y1 = table.bbox
                shape = slide.shapes.add_table(
                    len(table.rows),
                    len(table.rows[0]),
                    SlidePt(x0),
                    SlidePt(y0),
                    SlidePt(x1 - x0),
                    SlidePt(y1 - y0),
                )
                target = shape.table
                for j in range(len(table.edges) - 1):
                    target.columns[j].width = SlidePt(
                        table.edges[j + 1] - table.edges[j]
                    )
                for i, values in enumerate(table.rows):
                    target.rows[i].height = SlidePt(
                        table.row_edges[i + 1] - table.row_edges[i]
                    )
                    for j, value in enumerate(values):
                        cell = target.cell(i, j)
                        cell.text = value
                        cell.margin_left = SlidePt(6)
                        cell.margin_right = SlidePt(2)
                        cell.margin_top = cell.margin_bottom = 0
                        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                        cell.fill.solid()
                        cell.fill.fore_color.rgb = SlideColor.from_string(
                            table.fills[i]
                        )
                        for run in cell.text_frame.paragraphs[0].runs:
                            run.font.name = "DejaVu Sans"
                            run.font.size = SlidePt(table.font_size)
                            run.font.bold = table.bold[i]
                            run.font.color.rgb = SlideColor.from_string(table.colors[i])
                        tcpr = cell._tc.get_or_add_tcPr()
                        for side in ("L", "R", "T", "B"):
                            ln = etree.SubElement(
                                tcpr,
                                "{http://schemas.openxmlformats.org/drawingml/2006/main}ln"
                                + side,
                            )
                            etree.SubElement(
                                ln,
                                "{http://schemas.openxmlformats.org/drawingml/2006/main}noFill",
                            )
                for r0,c0,r1,c1 in table.merges:
                    target.cell(r0,c0).merge(target.cell(r1,c1))
    deck.save(output)


def _vml(parent, name, attrs):
    element = etree.SubElement(parent, "{urn:schemas-microsoft-com:vml}" + name)
    for key, value in attrs.items():
        element.set(key, str(value))
    return element


def _position(rect):
    x0, y0, x1, y1 = rect
    return (
        f"position:absolute;margin-left:{x0:.3f}pt;margin-top:{y0:.3f}pt;"
        f"width:{max(.1,x1-x0):.3f}pt;height:{max(.1,y1-y0):.3f}pt;"
        "mso-position-horizontal-relative:page;mso-position-vertical-relative:page"
    )


def _word_picture(paragraph, data, rect):
    x0, y0, x1, y1 = rect
    inline = (
        paragraph.add_run()
        .add_picture(BytesIO(data), width=Pt(x1 - x0), height=Pt(y1 - y0))
        ._inline
    )
    ns = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}"
    anchor = etree.Element(ns + "anchor", nsmap=inline.nsmap)
    for k, v in {
        "distT": "0",
        "distB": "0",
        "distL": "0",
        "distR": "0",
        "simplePos": "0",
        "relativeHeight": "0",
        "behindDoc": "0",
        "locked": "0",
        "layoutInCell": "1",
        "allowOverlap": "1",
    }.items():
        anchor.set(k, v)
    etree.SubElement(anchor, ns + "simplePos", x="0", y="0")
    for axis, offset in [("H", x0), ("V", y0)]:
        pos = etree.SubElement(anchor, ns + "position" + axis, relativeFrom="page")
        etree.SubElement(pos, ns + "posOffset").text = str(round(offset * 12700))
    for child in list(inline):
        if child.tag == ns + "extent":
            anchor.append(child)
    etree.SubElement(anchor, ns + "wrapNone")
    for child in list(inline):
        anchor.append(child)
    inline.getparent().replace(inline, anchor)


def convert_docx(source, output):
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "DejaVu Sans"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.space_before = Pt(0)
    with fitz.open(source) as pdf:
        for index, page in enumerate(pdf):
            if index:
                section = document.add_section(WD_SECTION.NEW_PAGE)
            else:
                section = document.sections[0]
            section.page_width, section.page_height = Pt(page.rect.width), Pt(
                page.rect.height
            )
            section.top_margin = section.bottom_margin = section.left_margin = (
                section.right_margin
            ) = Pt(0)
            section.header_distance = section.footer_distance = Pt(0)
            paragraph = document.add_paragraph()
            paragraph.paragraph_format.line_spacing = Pt(1)
            paragraph.paragraph_format.space_after = Pt(0)
            tables = extract_tables(page)
            for d in page.get_drawings():
                if inside(d["rect"], tables):
                    continue
                for item in d["items"]:
                    pict = OxmlElement("w:pict")
                    paragraph.add_run()._r.append(pict)
                    attrs = {
                        "style": _position(d["rect"]),
                        "stroked": "t" if d.get("color") else "f",
                        "filled": "t" if d.get("fill") else "f",
                        "fillcolor": "#" + color(d.get("fill"), "FFFFFF"),
                        "strokecolor": "#" + color(d.get("color")),
                        "strokeweight": f'{d.get("width") or .5}pt',
                    }
                    if item[0] == "re":
                        attrs["style"] = _position(item[1])
                        _vml(pict, "rect", attrs)
                    elif item[0] == "l":
                        a, b = item[1:3]
                        attrs.update(
                            {
                                "style": "position:absolute;mso-position-horizontal-relative:page;mso-position-vertical-relative:page",
                                "from": f"{a.x}pt,{a.y}pt",
                                "to": f"{b.x}pt,{b.y}pt",
                            }
                        )
                        _vml(pict, "line", attrs)
                    elif item[0] == "c":
                        r = d["rect"]
                        a, b, c, e = item[1:5]
                        attrs.update({"coordorigin": "0,0", "coordsize": "10000,10000"})
                        pt = (
                            lambda p: f"{round((p.x-r.x0)/max(.1,r.width)*10000)},{round((p.y-r.y0)/max(.1,r.height)*10000)}"
                        )
                        attrs["path"] = f"m {pt(a)} c {pt(b)},{pt(c)},{pt(e)} e"
                        _vml(pict, "shape", attrs)
            for block in page.get_text("dict")["blocks"]:
                if block["type"] == 1:
                    _word_picture(paragraph, block["image"], block["bbox"])
                elif block["type"] == 0:
                    for line in block["lines"]:
                        if inside(line["bbox"], tables):
                            continue
                        x0, y0, x1, y1 = line["bbox"]
                        pict = OxmlElement("w:pict")
                        paragraph.add_run()._r.append(pict)
                        shape = _vml(
                            pict,
                            "rect",
                            {
                                "style": _position(
                                    (x0, y0 - 1, min(page.rect.width, x1 + 14), y1 + 14)
                                ),
                                "stroked": "f",
                                "filled": "f",
                            },
                        )
                        box = _vml(shape, "textbox", {"inset": "0,0,0,0"})
                        content = OxmlElement("w:txbxContent")
                        box.append(content)
                        p = OxmlElement("w:p")
                        content.append(p)
                        ppr = OxmlElement("w:pPr")
                        p.append(ppr)
                        spacing = OxmlElement("w:spacing")
                        spacing.set(qn("w:before"), "0")
                        spacing.set(qn("w:after"), "0")
                        ppr.append(spacing)
                        for span in line["spans"]:
                            run = OxmlElement("w:r")
                            props = OxmlElement("w:rPr")
                            run.append(props)
                            fonts = OxmlElement("w:rFonts")
                            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                                fonts.set(qn("w:" + key), family(span))
                            props.append(fonts)
                            sz = OxmlElement("w:sz")
                            sz.set(qn("w:val"), str(round(span["size"] * 2)))
                            props.append(sz)
                            col = OxmlElement("w:color")
                            col.set(qn("w:val"), color(span["color"]))
                            props.append(col)
                            if span["flags"] & 16:
                                props.append(OxmlElement("w:b"))
                            if span["flags"] & 2:
                                props.append(OxmlElement("w:i"))
                            text = OxmlElement("w:t")
                            text.set(
                                "{http://www.w3.org/XML/1998/namespace}space",
                                "preserve",
                            )
                            text.text = span["text"]
                            run.append(text)
                            p.append(run)
            for table in tables:
                target = document.add_table(
                    rows=len(table.rows), cols=len(table.rows[0])
                )
                target.autofit = False
                x0, y0, x1, y1 = table.bbox
                props = target._tbl.tblPr
                position = OxmlElement("w:tblpPr")
                for k, v in {
                    "vertAnchor": "page",
                    "horzAnchor": "page",
                    "tblpX": str(round(x0 * 20)),
                    "tblpY": str(round(y0 * 20)),
                    "leftFromText": "0",
                    "rightFromText": "0",
                    "topFromText": "0",
                    "bottomFromText": "0",
                }.items():
                    position.set(qn("w:" + k), v)
                props.append(position)
                margins = OxmlElement("w:tblCellMar")
                for side, value in [
                    ("top", 0),
                    ("bottom", 0),
                    ("left", 120),
                    ("right", 40),
                ]:
                    elem = OxmlElement("w:" + side)
                    elem.set(qn("w:w"), str(value))
                    elem.set(qn("w:type"), "dxa")
                    margins.append(elem)
                props.append(margins)
                borders = OxmlElement("w:tblBorders")
                for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
                    elem = OxmlElement("w:" + side)
                    elem.set(qn("w:val"), "nil")
                    borders.append(elem)
                props.append(borders)
                for j, column in enumerate(target.columns):
                    column.width = Pt(table.edges[j + 1] - table.edges[j])
                for i, values in enumerate(table.rows):
                    row = target.rows[i]
                    row.height = Pt(table.row_edges[i + 1] - table.row_edges[i])
                    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
                    for j, value in enumerate(values):
                        cell = row.cells[j]
                        cell.width = Pt(table.edges[j + 1] - table.edges[j])
                        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                        cell.text = value
                        shd = OxmlElement("w:shd")
                        shd.set(qn("w:fill"), table.fills[i])
                        cell._tc.get_or_add_tcPr().append(shd)
                        p = cell.paragraphs[0]
                        p.paragraph_format.line_spacing = 1
                        for run in p.runs:
                            run.font.name = "DejaVu Sans"
                            run.font.size = Pt(table.font_size)
                            run.font.bold = table.bold[i]
                            run.font.color.rgb = RGBColor.from_string(table.colors[i])
                for r0,c0,r1,c1 in table.merges:
                    target.cell(r0,c0).merge(target.cell(r1,c1))
            # An anchor paragraph after floating tables keeps each page's objects together.
            document.add_paragraph().paragraph_format.line_spacing = Pt(1)
    document.save(output)


def image_to_editable_pdf(source: Path, output: Path, *, dpi=150) -> bool:
    """OCR creates visible editable glyphs and erases their raster counterparts."""
    import pytesseract

    with Image.open(source) as original:
        image = rgb(original)
    upscale = min(2, max(1, 2000 / max(image.size)))
    scan = image.resize(
        (round(image.width * upscale), round(image.height * upscale)),
        Image.Resampling.LANCZOS,
    )
    data = pytesseract.image_to_data(
        scan, config="--psm 3", output_type=pytesseract.Output.DICT
    )
    words = []
    for i, text in enumerate(data["text"]):
        confidence = float(data["conf"][i])
        if text.strip() and confidence >= 45:
            words.append((i, text, confidence))
    if len(words) < 3:
        return False
    if statistics.mean(w[2] for w in words) < 70:
        raise RuntimeError("OCR confidence is too low for editable Word conversion.")
    background = image.copy()
    draw = ImageDraw.Draw(background)
    glyphs = []
    scale = 72 / dpi
    for i, text, _ in words:
        x, y, w, h = [
            round(int(data[k][i]) / upscale) for k in ("left", "top", "width", "height")
        ]
        patch = image.crop(
            (
                max(0, x - 2),
                max(0, y - 2),
                min(image.width, x + w + 2),
                min(image.height, y + h + 2),
            )
        )
        dominant = patch.getcolors(patch.width * patch.height)
        bg = max(dominant, key=lambda p: p[0])[1]
        pixels = [
            p for p in patch.getdata() if sum(abs(p[j] - bg[j]) for j in range(3)) > 90
        ]
        pixels = sorted(
            pixels, key=lambda p: sum(abs(p[j] - bg[j]) for j in range(3)), reverse=True
        )[: max(1, len(pixels) // 4)]
        fg = (
            tuple(int(statistics.median(p[j] for p in pixels)) for j in range(3))
            if pixels
            else (0, 0, 0)
        )
        draw.rectangle((x - 1, y - 1, x + w + 1, y + h + 1), fill=bg)
        glyphs.append((text, x, y, w, h, fg))
    stream = BytesIO()
    background.save(stream, "PNG")
    pdf = fitz.open()
    page = pdf.new_page(width=image.width * scale, height=image.height * scale)
    page.insert_image(page.rect, stream=stream.getvalue())
    fontfile = "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"
    boldfile = "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf"
    for text, x, y, w, h, fg in glyphs:
        bold = h >= 24 or (text.isupper() and h >= 18)
        path = boldfile if bold else fontfile
        font = fitz.Font(fontfile=path)
        xmin, ymin, xmax, ymax = _glyph_metrics(path, text)
        width_size = w * scale / max(0.1, xmax - xmin)
        height_size = h * scale / max(0.1, ymax - ymin)
        size = min(width_size, height_size * 1.08)
        point = fitz.Point(x * scale - xmin * size, y * scale + ymax * size)
        page.insert_text(
            point,
            text,
            fontname="ocrbold" if bold else "ocrregular",
            fontfile=path,
            fontsize=size,
            color=tuple(c / 255 for c in fg),
        )
    pdf.save(output)
    pdf.close()
    return True


@lru_cache(maxsize=2)
def _font_metrics(path):
    font = TTFont(path)
    return font, font["head"].unitsPerEm, font.getBestCmap()


def _glyph_metrics(path, text):
    font, units, cmap = _font_metrics(path)
    advance = 0
    bounds = []
    for char in text:
        name = cmap.get(ord(char), ".notdef")
        glyph = font["glyf"][name]
        if hasattr(glyph, "xMin"):
            bounds.append(
                (advance + glyph.xMin, glyph.yMin, advance + glyph.xMax, glyph.yMax)
            )
        advance += font["hmtx"][name][0]
    if not bounds:
        return 0, 0, max(0.1, advance / units), 1
    return (
        min(b[0] for b in bounds) / units,
        min(b[1] for b in bounds) / units,
        max(b[2] for b in bounds) / units,
        max(b[3] for b in bounds) / units,
    )


def validate_image_word_text(source, output):
    import pytesseract
    from .quality import assert_text_regions, page_image

    with Image.open(source) as original:
        image = rgb(original)
    upscale = min(2, max(1, 2000 / max(image.size)))
    scan = image.resize(
        (round(image.width * upscale), round(image.height * upscale)),
        Image.Resampling.LANCZOS,
    )
    data = pytesseract.image_to_data(
        scan, config="--psm 3", output_type=pytesseract.Output.DICT
    )
    regions = []
    for i, text in enumerate(data["text"]):
        if any(c.isalnum() for c in text) and float(data["conf"][i]) >= 70:
            regions.append(
                tuple(
                    round(data[k][i] / upscale)
                    for k in ("left", "top", "width", "height")
                )
            )
    with fitz.open(output) as pdf:
        if len(pdf) != 1:
            raise RuntimeError("Image-to-Word changed page count.")
        assert_text_regions(image, page_image(pdf[0], dpi=150), regions)
