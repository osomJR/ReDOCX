"""Evidence-backed spreadsheet formulas and editable column charts.

PDFs do not retain original workbook formulas. Only explicit quantity/price/amount
relationships and labelled totals that agree with the displayed values qualify.
Ambiguous charts remain graphics; no business logic is invented.
"""
from __future__ import annotations

from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
import re
import unicodedata
import zipfile
from xml.etree import ElementTree as ET

import fitz
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.marker import DataPoint
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter


@dataclass(frozen=True)
class BarEvidence:
    table_index: int
    value_column: int
    categories: tuple[str,...]
    values: tuple[float,...]
    colors: tuple[str,...]
    bbox: tuple[float,float,float,float]
    title: str


def key(value):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(value)).casefold()
                   if c.isalnum() and not unicodedata.combining(c))


def _numeric(value):
    try:
        value=str(value).strip()
        if not re.fullmatch(r'-?(?:0|[1-9]\d*)(?:\.\d+)?',value):
            return None
        return Decimal(value)
    except InvalidOperation:
        return None


def chart_evidence(page,tables):
    from .native import inside,color
    bars=[]
    for drawing in page.get_drawings():
        if not drawing.get('fill') or inside(drawing['rect'],tables):
            continue
        for item in drawing['items']:
            if item[0]=='re':
                r=fitz.Rect(item[1])
                if r.width>=12 and r.height>=5 and max(drawing['fill'])-min(drawing['fill'])>.1:
                    bars.append((r,color(drawing['fill'])))
    groups=[]
    for rect,fill in sorted(bars,key=lambda b:b[0].x0):
        for group in groups:
            if abs(group[0][0].y1-rect.y1)<=2 and abs(group[0][0].width-rect.width)/rect.width<.2:
                group.append((rect,fill));break
        else:
            groups.append([(rect,fill)])
    words=page.get_text('words')
    evidence=[]
    for group in groups:
        if len(group)<2 or len(group)>30:
            continue
        group.sort(key=lambda b:b[0].x0)
        if any(a[0].x1>=b[0].x0 for a,b in zip(group,group[1:])):
            continue
        values=[]
        for rect,_ in group:
            candidates=[w for w in words if _numeric(w[4]) is not None and
                        rect.y0-24<=w[1]<rect.y0 and
                        abs((w[0]+w[2])/2-(rect.x0+rect.x1)/2)<=rect.width*.5]
            if not candidates:
                break
            word=min(candidates,key=lambda w:abs(w[3]-rect.y0))
            values.append(float(_numeric(word[4])))
        if len(values)!=len(group) or any(v<=0 for v in values):
            continue
        slopes=[bar[0].height/value for bar,value in zip(group,values)]
        if (max(slopes)-min(slopes))/max(slopes)>.04:
            continue
        matches=[]
        for ti,table in enumerate(tables):
            data=table.rows[1:]
            if len(data)!=len(group):
                continue
            for column in range(1,len(table.rows[0])):
                numbers=[_numeric(row[column]) for row in data]
                if any(x is None for x in numbers) or any(abs(float(a)-b)>1e-7 for a,b in zip(numbers,values)):
                    continue
                labels=[row[0] for row in data]
                # Every bar must also have its matching category label nearby.
                if all(any(key(w[4])==key(label) and rect.y1<=w[1]<=rect.y1+28
                           and rect.x0-10<=w[0]<=rect.x1+10 for w in words)
                       for (rect,_),label in zip(group,labels)):
                    matches.append((ti,column,labels))
        if len(matches)!=1:
            continue
        ti,column,labels=matches[0]
        bbox=fitz.Rect(group[0][0])
        for rect,_ in group[1:]:bbox|=rect
        title=''
        headings=[s for b in page.get_text('dict')['blocks'] if b['type']==0
                  for l in b['lines'] for s in l['spans']
                  if s['bbox'][3]<bbox.y0 and bbox.y0-s['bbox'][3]<110 and s['size']>=11]
        if headings:title=max(headings,key=lambda s:s['bbox'][3])['text']
        evidence.append(BarEvidence(ti,column,tuple(labels),tuple(values),
            tuple(fill for _,fill in group),tuple(bbox),title))
    return evidence


def add_chart(sheet,evidence,table_start,anchor):
    chart=BarChart()
    chart.type='col';chart.style=10
    chart.title=evidence.title or None
    chart.legend=None
    chart.width=15;chart.height=6.5
    chart.y_axis.delete=True
    chart.x_axis.title=None
    start=table_start+1
    end=start+len(evidence.categories)-1
    chart.add_data(Reference(sheet,min_col=evidence.value_column+1,
                             min_row=table_start,max_row=end),titles_from_data=True)
    chart.set_categories(Reference(sheet,min_col=1,min_row=start,max_row=end))
    chart.dLbls=DataLabelList();chart.dLbls.showVal=True
    for index,fill in enumerate(evidence.colors):
        chart.series[0].data_points.append(DataPoint(idx=index,
            spPr=GraphicalProperties(solidFill=fill)))
    sheet.add_chart(chart,anchor)
    return chart


def formulas_for(table,start_row):
    """Return coordinate -> (safe generated formula, original displayed value)."""
    headers=[key(x) for x in table.rows[0]]
    aliases={'quantity':{'quantity','qty','quantite'},
             'price':{'unitprice','price','prixunitaire','prix'},
             'amount':{'amount','linetotal','montant'}}
    columns={role:[i for i,h in enumerate(headers) if h in choices] for role,choices in aliases.items()}
    products=[]
    formulas={}
    if all(len(columns[role])==1 for role in aliases):
        quantity,price,amount=(columns[role][0] for role in aliases)
        for index,row in enumerate(table.rows[1:],1):
            q,p,a=(_numeric(row[i]) for i in (quantity,price,amount))
            if None in (q,p,a):continue
            decimals=len(row[amount].split('.')[-1]) if '.' in row[amount] else 0
            precision=Decimal(1).scaleb(-decimals)
            if (q*p).quantize(precision)!=a:continue
            rr=start_row+index
            coordinate=f'{get_column_letter(amount+1)}{rr}'
            formula=f'=ROUND({get_column_letter(quantity+1)}{rr}*{get_column_letter(price+1)}{rr},{decimals})'
            formulas[coordinate]=(formula,a)
    totals={'total','subtotal','grandtotal','soustotal'}
    for index,row in enumerate(table.rows[2:],2):
        if key(row[0]) not in totals:continue
        prior=table.rows[1:index]
        if len(prior)<2:continue
        for column in range(1,len(row)):
            displayed=_numeric(row[column]);values=[_numeric(r[column]) for r in prior]
            if displayed is None or any(v is None for v in values) or sum(values)!=displayed:
                continue
            letter=get_column_letter(column+1);rr=start_row+index
            formulas[f'{letter}{rr}']=(f'=SUM({letter}{start_row+1}:{letter}{rr-1})',displayed)
    return formulas


def apply_formulas(sheet,formulas):
    for coordinate,(formula,value) in formulas.items():
        cell=sheet[coordinate]
        cell.value=formula
        cell.comment=Comment('Formula reconstructed from the visible numeric relationship. '
                             'The source PDF does not preserve the original workbook formula.','ReDOCX')


def write_formula_caches(path: Path,caches):
    if not caches:return
    namespace='{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
    with zipfile.ZipFile(path) as source:
        entries={name:source.read(name) for name in source.namelist()}
    for sheet_index,formulas in caches.items():
        name=f'xl/worksheets/sheet{sheet_index}.xml'
        root=ET.fromstring(entries[name])
        for cell in root.iter(namespace+'c'):
            if cell.get('r') not in formulas:continue
            formula,value=formulas[cell.get('r')]
            node=cell.find(namespace+'f')
            if node is None or node.text!=formula[1:]:
                raise RuntimeError('Generated spreadsheet formula did not survive serialization.')
            cached=cell.find(namespace+'v')
            if cached is None:cached=ET.SubElement(cell,namespace+'v')
            cached.text=format(value,'f')
        entries[name]=ET.tostring(root,encoding='utf-8',xml_declaration=True)
    with zipfile.ZipFile(path,'w',compression=zipfile.ZIP_DEFLATED) as target:
        for name,data in entries.items():target.writestr(name,data)
